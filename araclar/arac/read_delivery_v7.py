"""Read-only XLSX source-chain and lossless archive decoder; no database imports."""
import hashlib
import json
from collections import defaultdict
from pathlib import Path
import sys
import openpyxl

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE if (HERE / 'core').is_dir() else HERE.parents[2]))
from core.content_delivery.delivery_values import parse_value

SCHEMA = json.loads((HERE / 'base-schema.json').read_text())
FIELDS = json.loads((HERE / 'provenance-fields.json').read_text())

def archive_values(workbook):
    if 'UzunAlanlar' not in workbook.sheetnames:
        return {}
    groups = defaultdict(list)
    for row in workbook['UzunAlanlar'].iter_rows(min_row=2, values_only=True):
        if row[0]:
            groups[row[0]].append(row[:7])
    output = {}
    for key, rows in groups.items():
        rows.sort(key=lambda r: r[5])
        if [r[5] for r in rows] != list(range(len(rows))):
            raise ValueError(f'{key}: missing or duplicate archive chunks')
        if len({tuple(r[:5]) for r in rows}) != 1:
            raise ValueError(f'{key}: conflicting archive metadata')
        value = ''.join(r[6] for r in rows)
        sha = hashlib.sha256(value.encode()).hexdigest()
        if sha != rows[0][4]:
            raise ValueError(f'{key}: archive checksum mismatch')
        output[key] = {'value': value, 'sha256': sha, 'table': rows[0][1], 'row_id': rows[0][2], 'field': rows[0][3]}
    return output

def restore(value, archives, table, row_id, field):
    candidate = value
    if isinstance(value, str) and value.startswith('{'):
        try:
            candidate = json.loads(value)
        except json.JSONDecodeError:
            return value
    if not isinstance(candidate, dict) or 'corteqs_long_value' not in candidate:
        return value
    archived = archives[candidate['corteqs_long_value']]
    if (archived['table'], archived['row_id'], archived['field'], archived['sha256']) != (table, row_id, field, candidate['sha256']):
        raise ValueError('Archive reference does not match its row and field')
    config = next(f for f in SCHEMA['tables'][table]['fields'] if f['key'] == field)
    return json.loads(archived['value']) if config['type'] == 'object' else archived['value']

def read(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        archives = archive_values(wb)
        data = {}
        for key, config in SCHEMA['tables'].items():
            fields = config['fields']; sheet = wb[config['sheet']]
            expected = [f['key'] for f in fields]
            header = list(next(sheet.iter_rows(min_row=7, max_row=7, values_only=True)))
            if header[:len(expected)] != expected or any(v is not None for v in header[len(expected):]):
                raise ValueError(f'{sheet.title}: headers changed')
            records = []
            for row in sheet.iter_rows(min_row=8):
                if not any(c.value not in (None, '') for c, f in zip(row, fields) if not f.get('readonly')):
                    continue
                record = {}
                for cell, field in zip(row, fields):
                    if field.get('readonly'):
                        continue
                    if cell.data_type == 'f':
                        raise ValueError(f'{sheet.title}!{cell.coordinate}: unexpected formula')
                    record[field['key']] = parse_value(cell.value, field)
                first = fields[0]['key']
                for field in fields:
                    if not field.get('readonly'):
                        name = field['key']
                        record[name] = restore(record[name], archives, key, record[first], name)
                records.append(record)
            data[key] = records
        sheet = wb['FotoğrafKaynakları']
        if list(next(sheet.iter_rows(min_row=7, max_row=7, values_only=True)))[:len(FIELDS)] != [f[0] for f in FIELDS]:
            raise ValueError('Source-chain headers changed')
        provenance = [{f[0]: (v if v is not None else '') for f, v in zip(FIELDS, row)} for row in sheet.iter_rows(min_row=8, values_only=True) if row[0]]
        return data, provenance, archives
    finally:
        wb.close()

if __name__ == '__main__':
    result, provenance, archives = read(sys.argv[1])
    target = Path(sys.argv[2])
    target.write_text(json.dumps({'tables': result, 'photo_provenance': provenance, 'archive_count': len(archives)}, ensure_ascii=False, indent=2))
    print(json.dumps({'tables': {k: len(v) for k, v in result.items()}, 'photo_provenance': len(provenance), 'archives': len(archives)}))
