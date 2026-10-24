"""xlsx_fidelity.py: sablon ile cikti arasindaki yapisal fark denetimi."""
import json
import shutil
import zipfile
from pathlib import Path

import openpyxl
import pytest

import xlsx_fidelity


def resave(src: Path, dst: Path, mutate=None) -> Path:
    wb = openpyxl.load_workbook(src)
    if mutate:
        mutate(wb)
    wb.save(dst)
    return dst


def test_roundtrip_copy_passes(tmp_path, template):
    out = resave(template, tmp_path / "out.xlsx", lambda wb: wb["Yanıtlar"].cell(8, 1, "veri"))
    assert xlsx_fidelity.compare(template, out) == []
    assert xlsx_fidelity.compare(template, shutil.copy(template, tmp_path / "copy.xlsx")) == []


def drop_dv(wb):
    ws = wb["Yanıtlar"]
    ws.data_validations.dataValidation = ws.data_validations.dataValidation[:1]


def shrink_dv(wb):
    wb["Kanıtlar"].data_validations.dataValidation[0].sqref = openpyxl.worksheet.cell_range.MultiCellRange("I8:I20")


def add_name(wb):
    wb.defined_names["Alan"] = openpyxl.workbook.defined_name.DefinedName("Alan", attr_text="'Yanıtlar'!$A$8:$A$37")


@pytest.mark.parametrize("mutate,needle", [
    (drop_dv, "data validation"),
    (shrink_dv, "data validation"),
    (lambda wb: wb["Yanıtlar"].cell(7, 2, "sehir_id"), "row 7"),
    (lambda wb: wb["Yanıtlar"].cell(6, 1, "X"), "row 6"),
    (lambda wb: setattr(wb["Yanıtlar"]["A5"], "value", None), "row 5"),
    (lambda wb: setattr(wb["Yanıtlar"], "freeze_panes", "B2"), "freeze"),
    (lambda wb: wb["Yanıtlar"].merge_cells("A1:C1"), "merged"),
    (lambda wb: setattr(wb["Yanıtlar"].column_dimensions["A"], "width", 42), "column width"),
    (lambda wb: setattr(wb["Yanıtlar"].row_dimensions[6], "height", 40), "row height"),
    (lambda wb: wb["Konular"].cell(9, 8, 0), "formula"),
    (lambda wb: wb.move_sheet("Kanıtlar", offset=-2), "sheet"),
    (lambda wb: wb.create_sheet("Fazla"), "sheet"),
    (add_name, "defined name"),
    (lambda wb: setattr(wb["Özet"]["A1"], "value", "Değişti"), "static cells"),
    (lambda wb: wb["Yanıtlar"].add_table(openpyxl.worksheet.table.Table(displayName="T1", ref="A7:L37")), "tables"),
])
def test_damaged_output_detected(tmp_path, template, mutate, needle):
    out = resave(template, tmp_path / "bad.xlsx", mutate)
    diffs = xlsx_fidelity.compare(template, out)
    assert diffs and needle in " ".join(diffs).lower(), diffs


def test_zip_part_difference_detected(tmp_path, template):
    out = shutil.copy(template, tmp_path / "parts.xlsx")
    with zipfile.ZipFile(out, "a") as archive:
        archive.writestr("xl/media/image1.png", b"png")
    diffs = xlsx_fidelity.compare(template, out)
    assert any("zip part" in d and "image1.png" in d for d in diffs)


def test_known_openpyxl_parts_only_accepted_on_request_and_in_known_direction():
    tpl = ["xl/sharedStrings.xml", "xl/workbook.xml"]
    out = ["docProps/app.xml", "docProps/core.xml", "xl/workbook.xml"]
    assert xlsx_fidelity.split_known_parts(tpl, out, False) == (tpl, out, [])
    a, b, accepted = xlsx_fidelity.split_known_parts(tpl, out, True)
    assert a == b == ["xl/workbook.xml"] and len(accepted) == 3
    # ters yon (sablonda docProps var, ciktida yok) kabul edilmez
    a, b, accepted = xlsx_fidelity.split_known_parts(out, tpl, True)
    assert accepted == [] and a == out


def test_cli_exit_codes_and_report(tmp_path, template, capsys):
    good = resave(template, tmp_path / "good.xlsx")
    report = tmp_path / "r.json"
    assert xlsx_fidelity.main([str(template), str(good), "--report", str(report)]) == 0
    assert json.loads(capsys.readouterr().out)["fidelity"] == "ok"
    assert json.loads(report.read_text(encoding="utf-8"))["diffs"] == []
    bad = resave(template, tmp_path / "bad.xlsx", drop_dv)
    assert xlsx_fidelity.main([str(template), str(bad)]) == 2
    assert json.loads(capsys.readouterr().out)["fidelity"]
    assert xlsx_fidelity.main([str(template), str(tmp_path / "yok.xlsx")]) == 2
    assert "error" in json.loads(capsys.readouterr().out)
