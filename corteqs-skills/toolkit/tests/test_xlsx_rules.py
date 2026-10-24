"""xlsx_rules: profil turetme ve hucre donusum kurallari (birim testleri)."""
import datetime as dt
import sys

import openpyxl
import pytest

import xlsx_rules


def test_load_field_types_leaves_tools_folder_untouched(tmp_path):
    """Resmi araclar klasorune bayt kodu (__pycache__) birakmamali; sys.dont_write_bytecode geri yuklenmeli."""
    schema_dir = tmp_path / "arac" / "core" / "content_delivery"
    schema_dir.mkdir(parents=True)
    (schema_dir / "schema_definition.py").write_text(
        "TABLES = {'a': {'sheet': 'Yanıtlar', 'fields': [{'key': 'source_ids', 'type': 'id_list'}]}}\n", encoding="utf-8")
    before = sys.dont_write_bytecode
    sys.dont_write_bytecode = False
    try:
        types = xlsx_rules.load_field_types(tmp_path / "arac")
        assert sys.dont_write_bytecode is False  # eski deger geri yuklendi
    finally:
        sys.dont_write_bytecode = before
    assert types.kind("Yanıtlar", "source_ids") == "id_list"
    assert not list((tmp_path / "arac").rglob("*.pyc")) and not list((tmp_path / "arac").rglob("__pycache__"))


def profile(name="SayfaAlanları", keys=("field_id", "value_type", "value"), **overrides):
    base = dict(name=name, keys={k: i for i, k in enumerate(keys, start=1)}, formulas=frozenset(), enums={},
                kinds={k: "string" for k in keys}, capacity=None, merge=False, identity={})
    return xlsx_rules.SheetProfile(**{**base, **overrides})


def test_json_value_type_requires_valid_json():
    p = profile()
    ok = xlsx_rules.check_record(p, {"value_type": "JSON", "value": {"a": 1}}, False)
    assert ok.values["value"] == '{"a":1}' and not ok.errors
    bad = xlsx_rules.check_record(p, {"value_type": "JSON", "value": "a: 1"}, False)
    assert bad.errors and "json" in bad.errors[0][1].lower()
    plain = xlsx_rules.check_record(p, {"value_type": "Metin", "value": "a: 1"}, False)
    assert plain.values["value"] == "a: 1"


def test_dates_and_numbers():
    assert xlsx_rules.to_date("2026-10-20T08:30") == dt.datetime(2026, 10, 20, 8, 30)
    with pytest.raises(xlsx_rules.RuleError):
        xlsx_rules.to_date("2026-02-30")
    assert xlsx_rules.to_number(" 51.43 ") == 51.43 and xlsx_rules.to_number("7") == 7
    for bad in ("nan", "inf", float("nan"), [1]):
        with pytest.raises(xlsx_rules.RuleError):
            xlsx_rules.to_number(bad)


def test_lists_and_literal_dv():
    assert xlsx_rules.from_list([{"a": 1}], "id_list") == ('[{"a":1}]', 0)
    assert xlsx_rules.from_list(["a", "a", ""], "id_list") == ("a", 1)
    assert xlsx_rules.literal_list("=Listeler!$A$1:$A$3") is None
    assert xlsx_rules.literal_list('"a,b"') == ("a", "b")


def test_capacity_prefers_table_ref_over_dv():
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table
    wb = openpyxl.Workbook()
    ws = wb.active
    for col, key in enumerate(("a", "b", "c"), start=1):
        ws.cell(7, col, key)
    dv = DataValidation(type="list", formula1='"x,y"')
    dv.add("B8:B37")
    ws.add_data_validation(dv)
    assert xlsx_rules.sheet_capacity(ws, merge=False) == 30
    ws.add_table(Table(displayName="CT_x", ref="A7:C12"))
    assert xlsx_rules.sheet_capacity(ws, merge=False) == 5


def test_capacity_unknown_for_unformatted_sheet_and_identity_numbers():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(7, 1, "x")
    assert xlsx_rules.sheet_capacity(ws, merge=False) is None
    p = profile(name="Şehir", keys=("latitude",), identity={"latitude": 51.4}, merge=True)
    assert not xlsx_rules.check_record(p, {"latitude": 51.4000000001}, False).errors
    assert xlsx_rules.check_record(p, {"latitude": 52}, False).errors
