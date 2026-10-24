"""make_template testlerinin ortak yardimcilari (test_make_template): sahte sozlesme (araclar/arac), sentetik mini
Berlin ornegi (openpyxl + elle eklenen x14 extLst, Excel'in sekme adi listesi, paylasimli metin tablosu) ve zip/XML
inceleme (sayfa parcalari, [Content_Types].xml, OPC kurallari, ornek belirteci taramasi). Test dosyasi 800 satir
sinirinin altinda kalsin diye ayrildi."""
from __future__ import annotations

import datetime as dt
import json
import posixpath
import re
import unicodedata
import zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from PIL import Image

import common
import make_template as mt

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
FORBIDDEN = ("berlin", "city-013", "city-000")
STATUS = ["Taslak", "İnceleme", "Düzeltme", "Onaylandı"]
DATA_FILL = PatternFill("solid", fgColor="FFFFF4D7")
SPACER_FILL = PatternFill("solid", fgColor="FFE7EDDF")
NOTE_FILL = PatternFill("solid", fgColor="FFFFF0E9")


# --- sentetik sozlesme (sahte araclar/arac) ve mini Berlin ornegi --------------------------------------

def fld(key: str, required: bool = False, kind: str = "string", **extra) -> dict:
    return {"key": key, "label": key.replace("_", " ").capitalize(), "type": kind, "required": required, **extra}


SCHEMA = {
    "schema_version": "1.0",
    "tables": {
        "metadata": {"sheet": "Teslimat", "capacity": 1, "fields": [
            fld("schema_version", True), fld("delivery_id", True), fld("city_id", True), fld("editor", True), fld("reviewer"),
            fld("scope", True, choices=["Tam teslim", "Kısmi örnek"]), fld("updated_on", True, "date"),
            fld("editorial_status", True, choices=STATUS), fld("notes")]},
        "city": {"sheet": "Şehir", "capacity": 1, "fields": [
            fld("city_id", True), fld("name", True), fld("country", True), fld("country_code", True),
            fld("region", True, choices=["EU", "AS", "AF", "NA", "SA", "OC"]), fld("kind", True, choices=["city", "region", "island", "country"]),
            fld("latitude", True, "number"), fld("longitude", True, "number"), fld("source_position", True, "integer", minimum=0),
            fld("priority", True, "integer", minimum=0), fld("aliases"), fld("editorial_status", True, choices=STATUS)]},
        "topics": {"sheet": "Konular", "capacity": 3, "fields": [
            fld("topic_id", True), fld("city_id", True), fld("title", True), fld("used_count", kind="integer", readonly=True)]},
        "answers": {"sheet": "Yanıtlar", "capacity": 4, "fields": [
            fld("entry_id", True), fld("city_id", True), fld("title", True), fld("featured", True, "boolean"), fld("topic_id", True)]},
    },
}
PROVENANCE = [["photo_id", "Fotoğraf kimliği", 32], ["author_url", "Fotoğrafçı profili", 48]]
EXT_LST = (
    '<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}" xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
    '<x14:dataValidations count="1" xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
    '<x14:dataValidation type="list" allowBlank="1"><x14:formula1><xm:f>Konular!$A$8:$A$10</xm:f></x14:formula1>'
    "<xm:sqref>E8:E14</xm:sqref></x14:dataValidation></x14:dataValidations></ext>"
    '<ext uri="{78C0D931-6437-407d-A8EE-F0AAD7539E65}" xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
    '<x14:conditionalFormattings><x14:conditionalFormatting xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
    '<x14:cfRule type="dataBar" id="{6A1E2B3C-0000-4000-8000-000000000001}"><x14:dataBar minLength="0" maxLength="100">'
    '<x14:cfvo type="autoMin"/><x14:cfvo type="autoMax"/></x14:dataBar></x14:cfRule><xm:sqref>C8:C14</xm:sqref>'
    "</x14:conditionalFormatting></x14:conditionalFormattings></ext></extLst>")
MC_DECL = ('xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" mc:Ignorable="x14ac" '
           'xmlns:x14ac="http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac" ')
VT = "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"
EXCEL_APP = ('<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" xmlns:vt="' + VT + '">'
             "<Application>Microsoft Excel</Application><HeadingPairs><vt:vector size=\"2\" baseType=\"variant\"><vt:variant>"
             "<vt:lpstr>Worksheets</vt:lpstr></vt:variant><vt:variant><vt:i4>7</vt:i4></vt:variant></vt:vector></HeadingPairs>"
             "<TitlesOfParts><vt:vector size=\"7\" baseType=\"lpstr\">{titles}</vt:vector></TitlesOfParts></Properties>")


def make_tools(root: Path) -> Path:
    root.mkdir(parents=True, exist_ok=True)
    (root / "base-schema.json").write_text(json.dumps(SCHEMA, ensure_ascii=False), encoding="utf-8")
    (root / "provenance-fields.json").write_text(json.dumps(PROVENANCE, ensure_ascii=False), encoding="utf-8")
    return root


def keyed_sheet(wb, name: str, keys: list[str], rows: list[list], table_end: int):
    ws = wb.create_sheet(name)
    ws["A2"] = f"Corteqs · {name}"
    for col, key in enumerate(keys, start=1):
        ws.cell(5, col, "Zorunlu" if col == 1 else "İsteğe bağlı")
        ws.cell(6, col, key.replace("_", " ").capitalize())
        ws.cell(7, col, key)
    for row in range(common.FIRST_DATA_ROW, table_end + 1):
        for col in range(1, len(keys) + 1):
            ws.cell(row, col).fill = DATA_FILL
        ws.row_dimensions[row].height = 54
    for offset, values in enumerate(rows):
        for col, value in enumerate(values, start=1):
            if value is not None:
                ws.cell(common.FIRST_DATA_ROW + offset, col, value)
        ws.row_dimensions[common.FIRST_DATA_ROW + offset].height = 68  # icerige gore buyumus ornek satiri
    ws.freeze_panes = "C8"
    ws.column_dimensions["A"].width = 30
    last = openpyxl.utils.get_column_letter(len(keys))
    ws.add_table(Table(displayName=f"CT_{len(wb.sheetnames)}", ref=f"A7:{last}{table_end}",
                       tableStyleInfo=TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)))
    return ws


def build_ozet(wb, variant: str) -> None:
    ws = wb.create_sheet("Özet")
    ws["A2"] = "Corteqs · Şehir içerik teslimi"
    ws["A3"] = "Berlin · Tam veritabanı örneği · CITY-013"
    for col, label in enumerate(["İçerik türü", "Dosyada", "Asgari", "Hedef", "Asgariye kalan", "Hedefe kalan"], start=1):
        ws.cell(5, col, label)
    ws["A6"], ws["B6"], ws["C6"], ws["D6"] = "Pratik yaşam yanıtı", "=COUNTA('Yanıtlar'!$A$8:$A$14)", 3, 4
    ws["A7"], ws["B7"], ws["C7"], ws["D7"] = "Yanıt içeren konu", "=SUM('Konular'!$D$8:$D$10)", 2, 2
    for row in (6, 7):
        ws[f"E{row}"], ws[f"F{row}"] = f"=MAX(C{row}-B{row},0)", f"=MAX(D{row}-B{row},0)"
    ws["A9"] = "Dosyayı doldurma sırası"
    ws.merge_cells("A9:F9")
    for row in (10, 11, 14):  # bos ara satirlar
        for col in range(1, 7):
            ws.cell(row, col).fill = SPACER_FILL
        ws.row_dimensions[row].height = 30
    ws["A12"] = ("BERLİN TAM VERİTABANI ÖRNEĞİ: DB KAYITLARI YALNIZ BU DOSYADADIR." if variant == "upper_case" else
                 "Berlin tam veritabanı örneği: DB kayıtları ve albüm yalnız bu dosyadadır.")  # Turkce buyuk I: casefold tuzagi
    ws.merge_cells("A12:F13")
    ws["A15"] = "İkinci örnek notu: fotoğraf albümü sekmesi yalnız örnekte vardır."  # ornek belirteci yok
    ws.merge_cells("A15:F17")
    for row in (12, 13, 15, 16, 17):
        for col in range(1, 7):
            ws.cell(row, col).fill = NOTE_FILL
        ws.row_dimensions[row].height = 28


def build_base_workbook(tmp: Path, variant: str):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    build_ozet(wb, variant)
    keys = {name: [f["key"] for f in t["fields"]] for name, t in ((t["sheet"], t) for t in SCHEMA["tables"].values())}
    keyed_sheet(wb, "Teslimat", keys["Teslimat"], [["1.0", "CITY-013-v001", "berlin", "Berlin ekibi", None, "Tam teslim",
                                                    dt.date(2026, 1, 5), "İnceleme", "Berlin örnek dosyası"]], 8)
    keyed_sheet(wb, "Şehir", keys["Şehir"], [["berlin", "Berlin", "Almanya", "DE", "EU", "city", 52.52, 13.405, 13, 1,
                                             "Berlin; Berlino", "İnceleme"]], 8)
    konular = keyed_sheet(wb, "Konular", keys["Konular"], [["berlin:planlama", "berlin", "Planlama", None],
                                                          ["berlin:oturum", "berlin", "Oturum", None]], 10)
    for row in range(8, 11):
        konular[f"D{row}"] = f"=IF(A{row}=\"\",0,IF(COUNTIFS('Yanıtlar'!$E$8:$E$30,A{row})>0,1,0))"
    answers = [[f"berlin:soru-{i}", "berlin", f"Berlin sorusu {i}", "Evet", "berlin:planlama"] for i in range(1, 6)]
    yan = keyed_sheet(wb, "Yanıtlar", keys["Yanıtlar"], answers, 14)  # tablo 7 satir; sozlesme kapasitesi 4
    dv = DataValidation(type="list", formula1='"Evet,Hayır"', allow_blank=True)
    dv.add("D8:D14")
    yan.add_data_validation(dv)
    for rng in ("A8:A14", "E8:E14"):
        col = rng[0]
        yan.conditional_formatting.add(rng, FormulaRule(formula=[f'AND($A8<>"",ISBLANK({col}8))'], fill=NOTE_FILL))
    if variant == "hyperlink":
        yan["C8"].hyperlink = "https://example.org/berlin"
    if variant == "header_token":
        konular.cell(6, 3, "Berlin başlığı")
    raw = wb.create_sheet("DBKayıtları")
    raw["A6"], raw["B6"], raw["C6"] = "id", "city", "text"
    for row in range(7, 10):
        raw.cell(row, 1, f"berlin:db-{row}")
        raw.cell(row, 2, "berlin")
        raw.cell(row, 3, "Berlin kaydı")
    raw.add_table(Table(displayName="CT_BerlinRaw", ref="A6:C9"))
    png = tmp / "foto.png"
    Image.new("RGB", (8, 8), (200, 30, 30)).save(png)
    raw.add_image(XLImage(str(png)), "E2")
    if variant != "no_provenance":
        keyed_sheet(wb, mt.PROVENANCE_SHEET, [k for k, _, _ in PROVENANCE], [["berlin:photo:1", "https://example.org/berlin"]], 12)
    wb.defined_names["BerlinAlan"] = DefinedName("BerlinAlan", attr_text="'DBKayıtları'!$A$7:$A$9")
    wb.defined_names["KonuListesi"] = DefinedName("KonuListesi", attr_text="'Konular'!$A$8:$A$10")
    wb.active = wb.sheetnames.index("DBKayıtları")
    return wb


def sheet_parts(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        book = ET.fromstring(archive.read("xl/workbook.xml"))
        rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {r.get("Id"): r.get("Target") for r in rels}
    parts = {}
    for sheet in book.iter(f"{{{MAIN}}}sheet"):
        target = targets[sheet.get(f"{{{REL}}}id")]
        parts[sheet.get("name")] = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
    return parts


def patch_part(path: Path, part: str, edit, binary: bool = False) -> None:
    with zipfile.ZipFile(path) as archive:
        items = [(info, archive.read(info.filename)) for info in archive.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for info, data in items:
            new = (edit(data) if binary else edit(data.decode("utf-8")).encode("utf-8")) if info.filename == part else data
            archive.writestr(info, new)


PKG_RELS = "http://schemas.openxmlformats.org/package/2006/relationships"
DOCX_PARTS = {  # Excel olmayan Office paketi (Word): officeDocument var, sekme listesi yok
    "[Content_Types].xml": '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                           '<Default Extension="xml" ContentType="application/xml"/></Types>',
    "_rels/.rels": f'<Relationships xmlns="{PKG_RELS}"><Relationship Id="rId1" Target="word/document.xml" '
                   f'Type="{REL}/officeDocument"/></Relationships>',
    "word/document.xml": '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body/></w:document>',
    "word/_rels/document.xml.rels": f'<Relationships xmlns="{PKG_RELS}"/>',
}


def content_types(path: Path) -> list[tuple[str, str, str]]:
    root = ET.fromstring(read_part(path, "[Content_Types].xml").encode("utf-8"))
    return [(e.tag.rsplit("}", 1)[-1], e.get("PartName") or e.get("Extension"), e.get("ContentType")) for e in root]


def opc_problems(path: Path) -> list[str]:
    """Excel'in ve .NET System.IO.Packaging'in reddettigi icerik turu durumlari (ad karsilastirmasi buyuk/kucuk harf duyarsiz)."""
    entries = content_types(path)
    with zipfile.ZipFile(path) as archive:
        names = {n.casefold() for n in archive.namelist()} - {"[content_types].xml"}
    overrides = [key.casefold() for tag, key, _ in entries if tag == "Override"]
    defaults = [key.casefold() for tag, key, _ in entries if tag == "Default"]
    problems = [f"yinelenen kayit: {k}" for k, n in (Counter(overrides) + Counter(defaults)).items() if n > 1]
    problems += [f"parcasi olmayan kayit: {k}" for k in overrides if k.lstrip("/") not in names]
    extension = lambda name: name.rsplit("/", 1)[-1].rpartition(".")[2] if "." in name.rsplit("/", 1)[-1] else ""  # ".rels" -> rels
    untyped = sorted(names - {k.lstrip("/") for k in overrides})
    return problems + [f"icerik turu yok: {n}" for n in untyped if extension(n) not in defaults]


def fold(text: str) -> str:
    """Turkce I/İ/ı ve aksan duyarsiz kucuk harf: 'BERLİN' -> 'berlin' (casefold 'berli̇n' verir, belirteci kacirir)."""
    text = unicodedata.normalize("NFKD", text.replace("ı", "i").replace("İ", "I"))
    return "".join(ch for ch in text if not unicodedata.combining(ch)).casefold()


INLINE_RE = re.compile(r'<c ([^>]*?)t="inlineStr"([^>]*)><is><t(?: [^>]*)?>(.*?)</t></is></c>', re.S)


def to_shared_strings(path: Path) -> None:
    """openpyxl 3.1 metni satir ici yazar; Excel'in kaydettigi gibi paylasimli metin tablosuna (t="s") cevirir."""
    strings: list[str] = []

    def shared(match: re.Match) -> str:
        strings.append(match.group(3))
        return f'<c {match.group(1)}t="s"{match.group(2)}><v>{len(strings) - 1}</v></c>'

    with zipfile.ZipFile(path) as archive:
        items = [(info, archive.read(info.filename)) for info in archive.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for info, data in items:
            text = data.decode("utf-8") if info.filename.endswith((".xml", ".rels")) else None
            if info.filename.startswith("xl/worksheets/sheet"):
                text = INLINE_RE.sub(shared, text)
            elif info.filename == "xl/_rels/workbook.xml.rels":
                text = text.replace("</Relationships>", '<Relationship Type="http://schemas.openxmlformats.org/officeDocument/'
                                    '2006/relationships/sharedStrings" Target="sharedStrings.xml" Id="rIdSst"/></Relationships>')
            elif info.filename == "[Content_Types].xml":
                text = text.replace("</Types>", '<Override PartName="/xl/sharedStrings.xml" ContentType="application/'
                                    'vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')
            archive.writestr(info, text.encode("utf-8") if text is not None else data)
        body = "".join(f"<si><t>{s}</t></si>" for s in strings)
        archive.writestr("xl/sharedStrings.xml", f'<sst xmlns="{MAIN}" count="{len(strings)}" uniqueCount="{len(strings)}">{body}</sst>')


def make_base(tmp: Path, variant: str = "") -> Path:
    path = tmp / "ekler" / "CITY-000_berlin_ornek_v001.xlsx"  # ornek kendi (salt okunur) klasorunde
    path.parent.mkdir(parents=True, exist_ok=True)
    build_base_workbook(tmp, variant).save(path)
    parts = sheet_parts(path)
    patch_part(path, parts["Yanıtlar"], lambda xml: xml.replace("<worksheet ", "<worksheet " + MC_DECL, 1).replace("</worksheet>", EXT_LST + "</worksheet>"))
    titles = "".join(f"<vt:lpstr>{name}</vt:lpstr>" for name in parts)  # Excel'in kaydettigi gibi sekme adi listesi
    patch_part(path, "docProps/app.xml", lambda xml: EXCEL_APP.format(titles=titles))
    if variant == "upper_case":  # OPC parca adlari harf duyarsiz: buyuk harfli kayit da Teslimat'in (sheet2) kaydidir
        patch_part(path, "[Content_Types].xml", lambda xml: xml.replace("/xl/worksheets/sheet2.xml", "/XL/Worksheets/Sheet2.xml"))
    to_shared_strings(path)
    return path


def read_part(path: Path, part: str) -> str:
    with zipfile.ZipFile(path) as archive:
        return archive.read(part).decode("utf-8-sig")


def sheet_root(path: Path, sheet: str) -> ET.Element:
    return ET.fromstring(read_part(path, sheet_parts(path)[sheet]).encode("utf-8"))


def cells_of(root: ET.Element) -> dict[str, ET.Element]:
    return {c.get("r"): c for c in root.iter(f"{{{MAIN}}}c")}


def leaked_tokens(path: Path, skip: set[tuple[str, str]]) -> list[str]:
    """Ornek belirtecleri: hucrelerde (skip haric) ve sayfa olmayan tum XML parcalarinda."""
    hits = []
    parts = sheet_parts(path)
    sheet_of = {part: name for name, part in parts.items()}
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            data = archive.read(name)
            if not name.endswith((".xml", ".rels")):
                continue
            if name in sheet_of:
                root = ET.fromstring(data)
                for cell in root.iter(f"{{{MAIN}}}c"):
                    if (sheet_of[name], cell.get("r")) in skip:
                        cell.clear()
                text = ET.tostring(root, encoding="unicode")
            else:
                text = data.decode("utf-8", "ignore")
            hits += [f"{name}:{t}" for t in FORBIDDEN if t in fold(text)]
    return hits
