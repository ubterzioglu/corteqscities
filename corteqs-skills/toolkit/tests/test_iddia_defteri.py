"""iddia_defteri: yazar iddia defteri + bagimsiz denetim sonucu teslime (Sevilla dis denetimi: A3/A4 kanit pakette yoktu)."""
from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import iddia_defteri as idf  # noqa: E402

SHEETS = {
    "Yanıtlar": ["entry_id", "title"],
    "Mekânlar": ["entry_id", "title"],
    "Sözlük": ["term_id", "term"],
    "Kaynaklar": ["source_id", "url", "snapshot_file", "snapshot_sha256"],
}
ROWS = {
    "Yanıtlar": [{"entry_id": "t:y1", "title": "Bir"}, {"entry_id": "t:y2", "title": "İki"},
                 {"entry_id": "t:y3", "title": "Üç"}],
    "Mekânlar": [{"entry_id": "t:m1", "title": "Müze"}],
    "Sözlük": [{"term_id": "t:s1", "term": "NIE"}],
    "Kaynaklar": [{"source_id": "t:kaynak:001", "url": "https://a.example/1", "snapshot_file": "Lisans kanıtları/a.html",
                   "snapshot_sha256": "ab" * 32},
                  {"source_id": "t:kaynak:002", "url": "https://a.example/2"}],  # kopyasi yok
}


def _xlsx(path: Path) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, keys in SHEETS.items():
        ws = wb.create_sheet(sheet)
        for col, key in enumerate(keys, start=1):
            ws.cell(6, col, key.upper())
            ws.cell(7, col, key)
        for row, record in enumerate(ROWS[sheet], start=8):
            for key, value in record.items():
                ws.cell(row, keys.index(key) + 1, value)
    wb.save(path)
    return path


def _jsonl(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("".join(json.dumps(r, ensure_ascii=False) + "\n" for r in rows), encoding="utf-8")


def _claim(cid: str, entry: str, source: str = "t:kaynak:001") -> dict:
    return {"claim_id": cid, "entry_id": entry, "field": "summary", "claim": f"iddia {cid}", "value": "6", "unit": "€",
            "period": "2026", "source_id": source, "url": "https://a.example/1", "quote": "seis euros"}


def _audit(cid: str, entry: str, result: str) -> dict:
    return {"claim_id": cid, "entry_id": entry, "iddia": f"denetim {cid}", "resmî_URL": "https://a.example/1",
            "kontrol_tarihi": "2026-10-25", "sonuç": result, "ağırlık": "-" if result == "doğru" else "orta",
            "alıntı": "seis euros", "doğrusu": ""}


@pytest.fixture()
def work(tmp_path: Path) -> dict:
    claims, audit, out = tmp_path / "claims", tmp_path / "denetim", tmp_path / "PKG" / "Teslimler"
    _jsonl(claims / "konu.jsonl", [_claim("C-1", "t:y1"), _claim("C-2", "t:y1"), _claim("C-3", "t:y2", "t:kaynak:002"),
                                   _claim("C-9", "t:silindi")])  # silinen kartin iddiasi deftere girmez
    _jsonl(claims / "mekan.jsonl", [_claim("C-4", "t:m1")])
    _jsonl(audit / "konu_iddia.jsonl", [_audit("D-1", "t:y1", "yanlış"), _audit("D-2", "t:y2", "doğru")])
    _jsonl(audit / "konu_karar.jsonl", [{"entry_id": "t:y1", "karar": "düzelt", "en_agir": "orta", "ozet": "x"},
                                        {"entry_id": "t:y2", "karar": "temiz", "en_agir": "-", "ozet": ""}])
    _jsonl(audit / "konu-d2_iddia.jsonl", [_audit("D2-1", "t:y1", "doğru")])
    _jsonl(audit / "konu-d2_karar.jsonl", [{"entry_id": "t:y1", "karar": "temiz", "en_agir": "-", "ozet": ""}])
    _jsonl(audit / "konu_degisiklik.jsonl", [{"entry_id": "t:y1", "alan": "summary"}])
    _jsonl(audit / "mekan_karar.jsonl", [{"entry_id": "t:m1", "karar": "düzelt", "en_agir": "hafif", "ozet": "x"}])
    _jsonl(audit / "mekan_degisiklik.jsonl", [{"entry_id": "t:m1", "alan": "note"}])
    _jsonl(audit / "mekan_dogrulama.jsonl", [{"entry_id": "t:m1", "sonuc": "hatali_duzeltme"}])
    _jsonl(audit / "acik_mi_iddia_notu.jsonl", [{"entry_id": "t:s1"}])  # iddia/karar dosyasi degil: okunmaz
    return {"xlsx": _xlsx(tmp_path / "CITY-1_t_v001.xlsx"), "claims": claims, "audit": audit, "out": out}


def _run(work: dict, capsys) -> tuple[int, dict]:
    code = idf.main(["--xlsx", str(work["xlsx"]), "--claims-dir", str(work["claims"]), "--audit-dir", str(work["audit"]),
                     "--out-dir", str(work["out"])])
    return code, json.loads(capsys.readouterr().out.strip().splitlines()[-1])


def _csv(path: Path) -> list[dict]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def test_writes_ledger_audit_and_summary_for_records_in_the_final_excel(work, capsys):
    code, s = _run(work, capsys)
    ledger = _csv(work["out"] / idf.LEDGER_NAME)
    assert [r["claim_id"] for r in ledger] == ["C-1", "C-2", "C-3", "C-4"]  # C-9 silinen kart
    first = ledger[0]
    assert (first["sayfa"], first["satir"], first["alinti"]) == ("Yanıtlar", "8", "seis euros")
    assert first["kaynak_kopyasi"] == "Lisans kanıtları/a.html" and first["kaynak_kopyasi_sha256"] == "ab" * 32
    audit = _csv(work["out"] / idf.AUDIT_NAME)
    assert [(r["tur"], r["claim_id"], r["sonuc"]) for r in audit] == [
        ("konu", "D-1", "yanlış"), ("konu", "D-2", "doğru"), ("konu-d2", "D2-1", "doğru")]
    summary = json.loads((work["out"] / idf.SUMMARY_NAME).read_text(encoding="utf-8"))
    assert summary["iddia"] == 4 and summary["kayit_iddiali"] == 3 and summary["silinen_kayit_iddiasi"] == 1
    assert summary["denetim_karari"] == {"t:y1": "temiz", "t:y2": "temiz", "t:m1": "düzelt"}  # d2 turu ilk turu ezer
    assert summary["son_durum"] == {"t:y1": "temiz", "t:y2": "temiz", "t:m1": "düzeltildi"}  # degisiklik kaydi var
    assert summary["dogrulama_sorunlu"] == ["t:m1"]
    assert summary["iddiasiz_kayit"] == {"Yanıtlar": ["t:y3"], "Sözlük": ["t:s1"]}
    assert summary["denetlenmemis_kayit"] == {"Yanıtlar": ["t:y3"], "Sözlük": ["t:s1"]}
    assert summary["kopyasiz_iddia"] == ["C-3"] and summary["duzelt_kayitsiz"] == []
    assert code == 1 and s["kopyasiz_iddia"] == 1  # kopyasi olmayan kaynaga dayanan iddia bulgudur


def test_term_and_route_claims_are_owned_by_their_own_ids(work, capsys):
    """Sozluk defteri term_id, rota defteri route_id tasir (entry_id yok); ikisi de deftere girer."""
    term = {**_claim("C-T1", ""), "term_id": "t:s1", "field": "definition"}
    term.pop("entry_id")
    _jsonl(work["claims"] / "sozluk.jsonl", [term, {**_claim("C-R1", ""), "entry_id": None, "route_id": "t:route:yok"}])
    code, s = _run(work, capsys)
    ledger = {r["claim_id"]: r for r in _csv(work["out"] / idf.LEDGER_NAME)}
    assert ledger["C-T1"]["entry_id"] == "t:s1" and ledger["C-T1"]["sayfa"] == "Sözlük"
    assert "C-R1" not in ledger  # Excel'de olmayan rota: silinen kayit gibi sayilir
    summary = json.loads((work["out"] / idf.SUMMARY_NAME).read_text(encoding="utf-8"))
    assert "Sözlük" not in summary["iddiasiz_kayit"] and summary["silinen_kayit_iddiasi"] == 2


def test_later_audit_round_supersedes_old_verifier_flag(work, capsys):
    """Ilk turun dogrulayici isareti, kaydin sonraki bagimsiz denetimi (d5) varsa sayilmaz; ayni/yeni turda sayilir."""
    _jsonl(work["audit"] / "mekan-d5_karar.jsonl", [{"entry_id": "t:m1", "karar": "temiz", "en_agir": "-", "ozet": ""}])
    code, s = _run(work, capsys)
    summary = json.loads((work["out"] / idf.SUMMARY_NAME).read_text(encoding="utf-8"))
    assert summary["dogrulama_sorunlu"] == [] and summary["son_durum"]["t:m1"] == "temiz"
    _jsonl(work["audit"] / "mekan-d5_dogrulama.jsonl", [{"entry_id": "t:m1", "sonuc": "eksik_kalan"}])
    code, s = _run(work, capsys)
    summary = json.loads((work["out"] / idf.SUMMARY_NAME).read_text(encoding="utf-8"))
    assert summary["dogrulama_sorunlu"] == ["t:m1"]


def test_clean_ledger_exits_zero(work, capsys):
    _jsonl(work["claims"] / "konu.jsonl", [_claim("C-1", "t:y1")])
    code, s = _run(work, capsys)
    assert code == 0 and s["iddia"] == 2 and s["kopyasiz_iddia"] == 0 and s["duzelt_kayitsiz"] == 0


def test_later_round_of_another_part_overrides_first_round(work, capsys):
    """Inceleme bulgusu: 'mekan-d2' turu 'mekan1' ilk turundan sonra okunmali (parca adina gore degil)."""
    _jsonl(work["audit"] / "mekan1_karar.jsonl", [{"entry_id": "t:m1", "karar": "düzelt"}])
    _jsonl(work["audit"] / "mekan-d2_karar.jsonl", [{"entry_id": "t:m1", "karar": "temiz"}])
    _run(work, capsys)
    summary = json.loads((work["out"] / idf.SUMMARY_NAME).read_text(encoding="utf-8"))
    assert summary["denetim_karari"]["t:m1"] == "temiz"


def test_flagged_card_without_change_record_is_a_finding(work, capsys):
    _jsonl(work["claims"] / "konu.jsonl", [_claim("C-1", "t:y1")])
    _jsonl(work["audit"] / "konu-d2_karar.jsonl", [{"entry_id": "t:y1", "karar": "temiz"},
                                                   {"entry_id": "t:y2", "karar": "düzelt"}])
    code, s = _run(work, capsys)
    summary = json.loads((work["out"] / idf.SUMMARY_NAME).read_text(encoding="utf-8"))
    assert summary["son_durum"]["t:y2"] == "düzelt_kayitsiz" and summary["duzelt_kayitsiz"] == ["t:y2"]
    assert code == 1 and s["duzelt_kayitsiz"] == 1


@pytest.mark.parametrize("stem, rank", [("konu", (0, 0)), ("konu-ek", (1, 0)), ("konu-ek2", (1, 2)), ("konu-d2", (2, 2)),
                                        ("konu-d10", (2, 10))])
def test_round_order(stem, rank):
    assert idf.round_rank(stem) == rank


def test_country_pack_audit_dir_is_read_first_and_linked_by_slug(work, capsys, tmp_path):
    pack = tmp_path / "ulke" / "denetim"
    _jsonl(pack / "oturum_iddia.jsonl", [_audit("P-1", "es:y3", "doğru")])  # paket kimligi: es:<slug>
    _jsonl(pack / "oturum_karar.jsonl", [{"entry_id": "es:y3", "karar": "temiz"}, {"entry_id": "es:y1", "karar": "düzelt"}])
    code = idf.main(["--xlsx", str(work["xlsx"]), "--claims-dir", str(work["claims"]), "--audit-dir", str(pack),
                     "--audit-dir", str(work["audit"]), "--out-dir", str(work["out"])])
    capsys.readouterr()
    summary = json.loads((work["out"] / idf.SUMMARY_NAME).read_text(encoding="utf-8"))
    assert summary["son_durum"]["t:y3"] == "temiz" and "Yanıtlar" not in summary["denetlenmemis_kayit"]
    assert summary["denetim_karari"]["t:y1"] == "temiz"  # sehir denetimi (son klasor) paketinkini ezer
    audit = _csv(work["out"] / idf.AUDIT_NAME)
    assert audit[0]["claim_id"] == "P-1" and audit[0]["entry_id"] == "t:y3"
    assert code == 1  # C-3 kopyasiz


def test_usage_errors(work, capsys, tmp_path):
    for args in (["--xlsx", str(tmp_path / "yok.xlsx")], ["--claims-dir", str(tmp_path / "yok")]):
        base = {"--xlsx": str(work["xlsx"]), "--claims-dir": str(work["claims"]), "--audit-dir": str(work["audit"]),
                "--out-dir": str(work["out"])}
        base[args[0]] = args[1]
        code = idf.main([x for pair in base.items() for x in pair])
        assert code == 2 and "error" in capsys.readouterr().out
