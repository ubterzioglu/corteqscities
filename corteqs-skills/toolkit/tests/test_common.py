import datetime as dt

import openpyxl
import pytest

import common


def test_sha256_file(tmp_path):
    p = tmp_path / "a.bin"
    p.write_bytes(b"abc")
    assert common.sha256_file(p) == "ba7816bf8f01cfea414140de5dae2223b00361a396177a9cb410ff61f20015ad"


@pytest.mark.parametrize("bad", ["../x", "/abs/x", "C:/x", "a/../../x", "", "  "])
def test_safe_relpath_rejects(tmp_path, bad):
    with pytest.raises(ValueError):
        common.safe_relpath(tmp_path, bad)


def test_safe_relpath_ok(tmp_path):
    assert common.safe_relpath(tmp_path, "Fotoğraflar/a.webp") == (tmp_path / "Fotoğraflar" / "a.webp").resolve()
    assert common.safe_relpath(tmp_path, "Lisans kanıtları\\x.jpg").name == "x.jpg"


def test_sheet_keys_and_records(template):
    wb = openpyxl.load_workbook(template)
    ws = wb["Yanıtlar"]
    keys = common.sheet_keys(ws)
    assert keys["entry_id"] == 1 and "visual_json" in keys
    ws.cell(8, 1, "t:a")
    ws.cell(8, 7, dt.datetime(2026, 10, 20))
    ws.cell(8, 8, 3.0)
    recs = list(common.iter_records(ws))
    assert recs == [{"entry_id": "t:a", "reviewed_on": "2026-10-20", "order": 3, "_row": 8}]


def test_formula_cells_excluded_and_listed(template):
    wb = openpyxl.load_workbook(template)
    ws = wb["Konular"]
    assert common.formula_columns(ws) == {"used_count"}
    assert list(common.iter_records(ws)) == []  # yalniz formul var -> kayit sayilmaz
    assert list(common.iter_records(ws, include_formulas=True))[0]["used_count"].startswith("=")


def test_duplicate_key_rejected(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(7, 1, "a")
    ws.cell(7, 2, "a")
    with pytest.raises(ValueError):
        common.sheet_keys(ws)


def test_canonical_source_url_merges_commons_forms():
    a = common.canonical_source_url("https://commons.wikimedia.org/wiki/File:Foo_bar.jpg")
    b = common.canonical_source_url("https://commons.wikimedia.org/w/index.php?title=File:Foo bar.jpg&uselang=nl")
    c = common.canonical_source_url("https://www.commons.wikimedia.org/wiki/File:Foo_bar.jpg#file")
    assert a == b == c
    assert a != common.canonical_source_url("https://commons.wikimedia.org/wiki/File:Other.jpg")


def test_group_unique_photos_is_transitive():
    rows = [
        {"photo_id": "A", "sha256": "s1", "source_url": "https://x.org/a"},
        {"photo_id": "B", "sha256": "s1", "source_url": "https://x.org/b"},   # A ile ayni sha
        {"photo_id": "C", "sha256": "s2", "source_url": "https://x.org/b"},   # B ile ayni kaynak -> A,B,C tek kume
        {"photo_id": "D", "sha256": "s3", "source_url": "https://x.org/d"},
    ]
    groups = sorted(sorted(g) for g in common.group_unique_photos(rows))
    assert groups == [["A", "B", "C"], ["D"]]


def test_setup_utf8_and_emit(capsys):
    common.setup_utf8()
    common.emit({"a": "ş", "n": 1})
    assert capsys.readouterr().out.strip() == '{"a":"ş","n":1}'
