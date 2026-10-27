"""gorsel_kalite testleri: kirpma kesirleri, netlik/cozunurluk/kapak isaretleri, JSON bicimi, paftalar, cikis kodlari."""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest
from PIL import Image, ImageFilter

import gorsel_kalite as gk

KEYS = ["photo_id", "entry_id", "role", "file", "web_file"]


def _checker(path: Path, size: tuple[int, int], cell: int = 4, blur: float = 0) -> None:
    """Keskin dama tahtasi (cok yuksek Laplace varyansi); blur>0 ise gaussian bulaniklastirilir."""
    img = Image.new("L", size)
    px = img.load()
    for y in range(size[1]):
        for x in range(size[0]):
            px[x, y] = 255 if (x // cell + y // cell) % 2 else 0
    if blur:
        img = img.filter(ImageFilter.GaussianBlur(blur))
    path.parent.mkdir(parents=True, exist_ok=True)
    img.convert("RGB").save(path, "PNG")


def _xlsx(path: Path, rows: list[dict]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fotoğraflar"
    for col, key in enumerate(KEYS, start=1):
        ws.cell(6, col, key.upper())
        ws.cell(7, col, key)
    for r, record in enumerate(rows, start=8):
        for key, value in record.items():
            ws.cell(r, KEYS.index(key) + 1, value)
    wb.save(path)
    return path


def _pkg(tmp_path: Path) -> Path:
    pkg = tmp_path / "paket"
    (pkg / "Fotoğraflar").mkdir(parents=True)
    _checker(pkg / "Fotoğraflar/keskin.png", (1200, 900))  # 4:3
    _checker(pkg / "Fotoğraflar/bulanik.png", (1200, 900), blur=12)
    _checker(pkg / "Fotoğraflar/kapak_dikey.png", (800, 1200))
    _checker(pkg / "Fotoğraflar/kucuk.png", (300, 200))
    _checker(pkg / "Fotoğraflar/kapak_iyi.png", (1500, 1000))  # 3:2
    return pkg


def _row(pid: str, role: str, file: str, **extra) -> dict:
    return {"photo_id": f"t:photo:{pid}", "entry_id": f"t:{pid}", "role": role, "file": f"Fotoğraflar/{file}", **extra}


ROWS = [
    _row("kapak-iyi", "Kapak", "kapak_iyi.png"),
    _row("kapak-dikey", "Kapak", "kapak_dikey.png"),
    _row("keskin", "Mekân", "keskin.png"),
    _row("bulanik", "Mekân", "bulanik.png"),
    _row("kucuk", "Mekân", "kucuk.png"),
    _row("yok", "Mekân", "yok.png"),
]


def _run(capsys, tmp_path, rows=None, pkg=None, extra=()):
    pkg = pkg or _pkg(tmp_path)
    xlsx = _xlsx(tmp_path / "t.xlsx", ROWS if rows is None else rows)
    out = tmp_path / "OUT"
    code = gk.main(["--xlsx", str(xlsx), "--pkg", str(pkg), "--out-dir", str(out), *extra])
    text = capsys.readouterr().out.strip()
    return code, (json.loads(text) if text else {}), out, pkg


def _report(out: Path) -> dict:
    return json.loads((out / gk.REPORT_NAME).read_text(encoding="utf-8"))


def _by_id(report: dict) -> dict:
    return {p["photo_id"].split(":")[-1]: p for p in report["photos"]}


def test_center_crop_fractions_and_boxes():
    crop = gk.center_crop(1200, 900, 16, 9)
    assert crop.kept == 0.75 and crop.box == (0, 112, 1200, 787)
    assert gk.center_crop(1200, 900, 4, 3).kept == 1.0
    assert gk.center_crop(800, 1200, 1, 1).box == (0, 200, 800, 1000)
    assert gk.center_crop(800, 1200, 16, 9).kept == pytest.approx(0.375, abs=0.001)
    assert gk.center_crop(1280, 574, 16, 9).kept == pytest.approx(0.797, abs=0.001)


def test_parse_ratios_and_percentiles():
    assert gk.parse_ratios("16:9, 1:1") == [("16:9", 16, 9), ("1:1", 1, 1)]
    for bad in ("16x9", "0:1", "", "a:b"):
        with pytest.raises(ValueError):
            gk.parse_ratios(bad)
    assert gk.percentiles([1, 2, 3, 4, 5])["p50"] == 3 and gk.percentiles([]) == {}


def test_flags_crop_fractions_and_json_shape(capsys, tmp_path):
    code, summary, out, _ = _run(capsys, tmp_path)
    assert code == 1
    report = _report(out)
    rows = _by_id(report)
    assert rows["keskin"]["flags"] == [] and rows["keskin"]["crops"]["16:9"]["kept"] == 0.75
    assert rows["keskin"]["width"] == 1200 and rows["keskin"]["short_side"] == 900 and rows["keskin"]["aspect"] == 1.3333
    assert rows["bulanik"]["flags"] == [gk.R_BLUR] and rows["bulanik"]["sharpness"] < gk.BLUR_THRESHOLD < rows["keskin"]["sharpness"]
    assert rows["kucuk"]["flags"] == [gk.R_LOW_RES] and rows["kucuk"]["short_side"] == 200
    assert rows["yok"]["flags"] == [gk.R_UNREADABLE] and "error" in rows["yok"]
    assert rows["kapak-iyi"]["flags"] == [] and rows["kapak-iyi"]["crops"]["16:9"]["kept"] == pytest.approx(0.844, abs=0.001)
    assert set(rows["kapak-dikey"]["flags"]) == {gk.R_COVER_PORTRAIT, gk.R_COVER_CROP, gk.R_STRONG}
    assert rows["kapak-dikey"]["strong_crop_ratios"] == ["16:9", "4:3"]
    assert {f["photo_id"] for f in report["flagged"]} == {f"t:photo:{n}" for n in ("kapak-dikey", "bulanik", "kucuk", "yok")}
    assert summary["photos"] == 6 and summary["covers"] == 2 and summary["flagged"] == 4
    assert summary["by_reason"] == {gk.R_UNREADABLE: 1, gk.R_LOW_RES: 1, gk.R_STRONG: 1, gk.R_BLUR: 1,
                                    gk.R_COVER_PORTRAIT: 1, gk.R_COVER_CROP: 1}
    assert report["summary"] == {k: v for k, v in summary.items() if k not in ("tool", "out_dir")}
    assert set(report) == {"tool", "version", "params", "summary", "photos", "flagged", "sheets"}
    assert "_path" not in rows["keskin"] and report["params"]["ratios"] == ["16:9", "4:3", "1:1"]
    assert {"p5", "p50", "min", "max"} <= set(summary["sharpness_percentiles"])


def test_sheets_covers_first_and_size_limits(capsys, tmp_path):
    _, summary, out, _ = _run(capsys, tmp_path)
    sheets = _report(out)["sheets"]
    assert [s["kind"] for s in sheets] == ["kapak", "mekan"] and summary["sheets"] == 2
    assert sheets[0]["photo_ids"] == ["t:photo:kapak-iyi", "t:photo:kapak-dikey"]
    for s in sheets:
        path = out / s["sheet"]
        assert path.is_file() and path.suffix == ".jpg" and s["sheet"].startswith("kirpma/kirpma_")
        with Image.open(path) as img:
            assert img.format == "JPEG" and img.width <= gk.MAX_SHEET_WIDTH and img.size == (s["width"], s["height"])


def test_more_than_24_photos_split_into_sheets(capsys, tmp_path):
    rows = [_row(f"k{i:02d}", "Mekân", "keskin.png") for i in range(25)] + [_row("kapak", "Kapak", "kapak_iyi.png")]
    code, summary, out, _ = _run(capsys, tmp_path, rows=rows)
    sheets = _report(out)["sheets"]
    assert [len(s["photo_ids"]) for s in sheets] == [1, 24, 1] and [s["kind"] for s in sheets] == ["kapak", "mekan", "mekan"]
    assert sorted(p.name for p in (out / "kirpma").iterdir()) == ["kirpma_01.jpg", "kirpma_02.jpg", "kirpma_03.jpg"]
    with Image.open(out / sheets[1]["sheet"]) as img:
        assert img.width <= gk.MAX_SHEET_WIDTH


def test_clean_set_exits_zero_and_stale_sheets_removed(capsys, tmp_path):
    stale = tmp_path / "OUT" / "kirpma" / "kirpma_09.jpg"
    stale.parent.mkdir(parents=True)
    stale.write_bytes(b"x")
    rows = [ROWS[0], ROWS[2]]
    code, summary, out, _ = _run(capsys, tmp_path, rows=rows)
    assert code == 0 and summary["flagged"] == 0 and _report(out)["flagged"] == [] and not stale.exists()


def test_web_file_preferred_over_file(capsys, tmp_path):
    rows = [_row("w", "Mekân", "keskin.png", web_file="Fotoğraflar/kucuk.png")]
    code, _, out, _ = _run(capsys, tmp_path, rows=rows)
    row = _by_id(_report(out))["w"]
    assert row["source_field"] == "web_file" and row["width"] == 300 and row["flags"] == [gk.R_LOW_RES] and code == 1


def test_custom_ratios_min_short_and_blur_threshold(capsys, tmp_path):
    rows = [ROWS[2], ROWS[3]]
    code, _, out, _ = _run(capsys, tmp_path, rows=rows, extra=("--ratios", "1:1", "--min-short", "1000", "--blur-threshold", "0"))
    report = _report(out)
    rows_by = _by_id(report)
    assert list(rows_by["keskin"]["crops"]) == ["1:1"] and rows_by["keskin"]["crops"]["1:1"]["kept"] == 0.75
    assert rows_by["keskin"]["flags"] == [gk.R_LOW_RES] and rows_by["bulanik"]["flags"] == [gk.R_LOW_RES]  # esik 0: bulanik yok
    assert report["params"]["min_short"] == 1000 and report["params"]["blur_threshold"] == 0


def test_square_crop_measured_but_not_flagged_by_default(capsys, tmp_path):
    """16:9 kare 1:1'de %56 kalir: olculur ve raporlanir, ama varsayilan isaret oranlari (16:9, 4:3) disinda."""
    pkg = _pkg(tmp_path)
    _checker(pkg / "Fotoğraflar/genis.png", (1600, 900))
    rows = [_row("genis", "Mekân", "genis.png")]
    code, summary, out, _ = _run(capsys, tmp_path, rows=rows, pkg=pkg)
    report = _report(out)
    row = _by_id(report)["genis"]
    assert code == 0 and row["flags"] == [] and summary["flagged"] == 0
    assert row["crops"]["1:1"]["kept"] == pytest.approx(0.5625, abs=0.001) and row["min_kept"] < gk.STRONG_CROP_KEPT
    assert report["params"]["flag_ratios"] == ["16:9", "4:3"] and report["params"]["ratios"] == ["16:9", "4:3", "1:1"]


def test_flag_ratios_option_restores_square_flag(capsys, tmp_path):
    pkg = _pkg(tmp_path)
    _checker(pkg / "Fotoğraflar/genis.png", (1600, 900))
    rows = [_row("genis", "Mekân", "genis.png")]
    code, _, out, _ = _run(capsys, tmp_path, rows=rows, pkg=pkg, extra=("--flag-ratios", "16:9,4:3,1:1"))
    row = _by_id(_report(out))["genis"]
    assert code == 1 and row["flags"] == [gk.R_STRONG] and row["strong_crop_ratios"] == ["1:1"]


def test_flag_ratios_default_intersects_and_explicit_must_be_measured(capsys, tmp_path):
    ratios = gk.parse_ratios("1:1")
    assert gk.resolve_flag_ratios(None, ratios) == []
    assert gk.resolve_flag_ratios("1:1", ratios) == ["1:1"]
    with pytest.raises(ValueError):
        gk.resolve_flag_ratios("16:9", ratios)
    pkg = _pkg(tmp_path)
    xlsx = _xlsx(tmp_path / "t.xlsx", ROWS)
    out = tmp_path / "OUT"
    assert gk.main(["--xlsx", str(xlsx), "--pkg", str(pkg), "--out-dir", str(out), "--ratios", "16:9",
                    "--flag-ratios", "1:1"]) == 2
    assert "olculmeyen" in capsys.readouterr().err and not out.exists()


def test_path_escape_is_unreadable_not_read(capsys, tmp_path):
    outside = tmp_path / "disari.png"
    _checker(outside, (1200, 900))
    rows = [_row("kacak", "Mekân", "keskin.png", web_file="../disari.png")]
    _, _, out, _ = _run(capsys, tmp_path, rows=rows)
    assert _by_id(_report(out))["kacak"]["flags"] == [gk.R_UNREADABLE]


def test_never_writes_into_package_and_refuses_out_inside(capsys, tmp_path):
    pkg = _pkg(tmp_path)
    before = sorted(p.relative_to(pkg).as_posix() for p in pkg.rglob("*"))
    xlsx = _xlsx(tmp_path / "t.xlsx", ROWS)
    for inside in (pkg / "OUT", pkg / "Fotoğraflar" / "x", pkg):
        assert gk.main(["--xlsx", str(xlsx), "--pkg", str(pkg), "--out-dir", str(inside)]) == 2
    assert "paket icinde" in capsys.readouterr().err
    # xlsx paketin icinde (Fotoğraflar klasoru yaninda) ise Excel'in klasoru de paket sayilir
    inner = _xlsx(pkg / "p.xlsx", ROWS)
    assert gk.main(["--xlsx", str(inner), "--pkg", str(pkg), "--out-dir", str(pkg / "Teslimler" / "o")]) == 2
    (pkg / "p.xlsx").unlink()
    assert sorted(p.relative_to(pkg).as_posix() for p in pkg.rglob("*")) == before
    code, _, out, _ = _run(capsys, tmp_path, pkg=pkg)
    assert code == 1 and sorted(p.relative_to(pkg).as_posix() for p in pkg.rglob("*")) == before


def test_usage_errors_exit_two(capsys, tmp_path):
    pkg = _pkg(tmp_path)
    xlsx = _xlsx(tmp_path / "t.xlsx", ROWS)
    out = str(tmp_path / "OUT")
    assert gk.main(["--xlsx", str(tmp_path / "yok.xlsx"), "--pkg", str(pkg), "--out-dir", out]) == 2
    assert gk.main(["--xlsx", str(xlsx), "--pkg", str(tmp_path / "yok"), "--out-dir", out]) == 2
    assert gk.main(["--xlsx", str(xlsx), "--pkg", str(pkg), "--out-dir", out, "--ratios", "16x9"]) == 2
    assert gk.main(["--xlsx", str(xlsx), "--pkg", str(pkg), "--out-dir", out, "--min-short", "0"]) == 2
    wb = openpyxl.Workbook()
    wb.save(tmp_path / "bos.xlsx")
    assert gk.main(["--xlsx", str(tmp_path / "bos.xlsx"), "--pkg", str(pkg), "--out-dir", out]) == 2
    assert not (tmp_path / "OUT").exists()
    with pytest.raises(SystemExit) as exc:
        gk.main(["--pkg", str(pkg)])
    assert exc.value.code == 2
