"""kart_tutarlilik: ayni terimi/form kodunu tasiyan kartlari gruplar (Sevilla dis denetimi S05/S12: kartlar birbiriyle celisti)."""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import kart_tutarlilik as kt  # noqa: E402

SHEETS = {"Yanıtlar": ["entry_id", "title", "summary", "steps", "note"], "Sözlük": ["term_id", "term", "definition"],
          "Mekânlar": ["entry_id", "title", "summary", "note"]}
ROWS = {
    "Yanıtlar": [
        {"entry_id": "t:dogum", "title": "Doğum kaydı", "summary": "Kayıttan sonra Libro de Familia verilir.",
         "steps": "1. Hastane bildirir. 2. Libro de Familia için 10 gün içinde başvur."},
        {"entry_id": "t:aile", "title": "Aile cüzdanı", "summary": "Libro de Familia 30 Nisan 2021'den beri verilmiyor.",
         "note": "Libro de Familia kaydı için 30 gün süre vardır. Ofis 9:00'da açılır."},  # terimsiz cumle sayilmaz
        {"entry_id": "t:tie", "title": "TIE", "summary": "TIE için 790-012 harcı 16,08 € ödenir."},
        {"entry_id": "t:izin", "title": "İzin yenileme", "summary": "Yenilemede 790-012 harcı 12 € ödenir."},
        {"entry_id": "t:yalniz", "title": "Tek", "summary": "Adres kaydı belediyede yapılır."},  # terim yalniz sozlukte
    ],
    "Sözlük": [{"term_id": "t:s1", "term": "Libro de Familia", "definition": "Aile cüzdanı; yeni kayıtlarda verilir."},
               {"term_id": "t:s2", "term": "Empadronamiento", "definition": "Adres kaydı."}],
    "Mekânlar": [{"entry_id": "t:m1", "title": "Müze", "summary": "Giriş 6 €."}],
}


def _xlsx(path: Path) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, keys in SHEETS.items():
        ws = wb.create_sheet(sheet)
        for col, key in enumerate(keys, start=1):
            ws.cell(6, col, key.upper())
            ws.cell(7, col, key)
        for row, record in enumerate(ROWS[sheet], start=8):
            for key, value in record.items():
                ws.cell(row, keys.index(key) + 1, value)
    wb.save(path)
    return path


def _run(tmp_path: Path, capsys) -> tuple[int, dict, dict]:
    out = tmp_path / "tutarlilik.json"
    code = kt.main([str(_xlsx(tmp_path / "t.xlsx")), "--out", str(out)])
    summary = json.loads(capsys.readouterr().out.strip().splitlines()[-1])
    return code, summary, json.loads(out.read_text(encoding="utf-8"))


def test_groups_cards_sharing_a_glossary_term_or_form_code(tmp_path, capsys):
    code, s, d = _run(tmp_path, capsys)
    groups = {g["anchor"]: g for g in d["groups"]}
    assert set(groups) == {"libro de familia", "790-012"}  # Empadronamiento tek kartta: grup degil
    libro = groups["libro de familia"]
    assert {m["id"] for m in libro["mentions"]} == {"t:dogum", "t:aile", "t:s1"}
    assert libro["values"]["sure"] == ["10 gun", "30 gun"] and libro["farkli_deger"] == ["sure"]
    fee = groups["790-012"]
    assert fee["values"]["ucret"] == ["12 €", "16,08 €"] and fee["farkli_deger"] == ["ucret"]
    first = next(m for m in fee["mentions"] if m["id"] == "t:tie")
    assert (first["sheet"], first["row"]) == ("Yanıtlar", 10) and "16,08 €" in first["sentences"][0]
    assert d["groups"][0]["anchor"] in {"libro de familia", "790-012"}  # farkli degerli gruplar once
    assert code == 1 and s["groups"] == 2 and s["farkli_degerli"] == 2


@pytest.mark.parametrize("text, anchors", [
    ("EX-17 formu ve 790-012 harcı", {"ex-17", "790-012"}),
    ("Modelo 030 ile bildir; EX 15 değil", {"modelo 030", "ex-15"}),
    ("Ücret 790 euro", set()),  # kod degil
    ("EX15 formu ve 790 kodu 012 harç belgesi", {"ex-15", "790-012"}),  # Sevilla yazimi
    ("tasa 790 012, Modelo 790-107 ile", {"790-012", "790-107"}),
    ("Hizmet mi 12 Mart'ta? Başvuru ta 20 gün", set()),  # Turkce soru eki, kod degil
])
def test_code_anchors(text, anchors):
    assert set(kt.code_anchors(text)) == anchors


def test_values_are_normalised():
    assert kt.facts("Harç 16,08 € ve 16.08 euro; 10 gün, on gün; 4 Mart.") == {
        "ucret": {"16,08 €"}, "sure": {"10 gun"}, "tarih": {"4 mart"}}
    assert kt.facts("tasa 790 012,10 € değil") == {}  # basta sifir: kodun parcasi, ucret degil


def test_usage_error(tmp_path, capsys):
    assert kt.main([str(tmp_path / "yok.xlsx"), "--out", str(tmp_path / "o.json")]) == 2
    assert "error" in capsys.readouterr().out
