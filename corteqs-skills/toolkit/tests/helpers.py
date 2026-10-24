"""Test yardimcilari: sentetik mini sablon ve mini paket (gercek veriye dokunmaz)."""
from __future__ import annotations

import hashlib
from pathlib import Path

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image

# Sayfa -> teknik anahtarlar (gercek v7 sablonundan sadelestirilmis; sira onemlidir)
SHEET_KEYS: dict[str, list[str]] = {
    "Teslimat": ["schema_version", "delivery_id", "city_id", "editor", "reviewer", "scope", "updated_on", "editorial_status", "notes"],
    "Şehir": ["city_id", "name", "country", "country_code", "region", "kind", "latitude", "longitude", "aliases", "intro", "editorial_status"],
    "Konular": ["topic_id", "city_id", "slug", "title", "description", "icon", "order", "used_count"],
    "Yanıtlar": ["entry_id", "city_id", "slug", "title", "summary", "source_ids", "reviewed_on", "order", "featured", "editorial_status", "topic_id", "visual_json"],
    "Kaynaklar": ["source_id", "city_id", "title", "publisher", "url", "checked_on", "verified", "snapshot_file", "snapshot_sha256"],
    "Fotoğraflar": ["photo_id", "city_id", "role", "entry_id", "file", "sha256", "caption", "author", "license", "license_url", "source_url",
                    "capture_date", "source_checked_on", "modifications", "evidence_id", "completion_credit", "order", "editorial_status",
                    "unique_venue_count", "unique_hero_count"],
    "Kanıtlar": ["evidence_id", "city_id", "photo_id", "source_url", "author", "license", "license_url", "checked_on",
                 "metadata_review_status", "pixel_review_status", "rights_status"],
    "FotoğrafKaynakları": ["photo_id", "author_url", "attribution_url", "original_file_url", "downloaded_file_url", "downloaded_on",
                           "original_file", "original_sha256", "web_file", "web_sha256", "web_width", "web_height", "derivation",
                           "provenance_status", "provenance_note"],
}
# Veri dogrulama (aciilir liste) tanimlari: (sayfa, anahtar) -> secenekler
ENUMS: dict[tuple[str, str], str] = {
    ("Yanıtlar", "featured"): "Evet,Hayır",
    ("Yanıtlar", "editorial_status"): "Taslak,İnceleme,Düzeltme,Onaylandı",
    ("Fotoğraflar", "role"): "Mekân,Kapak",
    ("Kanıtlar", "metadata_review_status"): "Bekliyor,Doğrulandı",
    ("Kanıtlar", "pixel_review_status"): "Bekliyor,Doğrulandı",
    ("Kanıtlar", "rights_status"): "Bekliyor,Bu kare için uygun",
}
CAPACITY = 30  # her sayfada veri satiri kapasitesi (satir 8..37)


def make_template(path: Path) -> Path:
    """Ozet (formul), kimlik satirli Teslimat/Sehir, formul sutunlu Konular/Fotograflar iceren mini sablon."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ozet = wb.create_sheet("Özet")
    ozet["A1"] = "Özet"
    ozet["B8"] = "=COUNTA('Yanıtlar'!$A$8:$A$37)"
    for name, keys in SHEET_KEYS.items():
        ws = wb.create_sheet(name)
        for col, key in enumerate(keys, start=1):
            ws.cell(5, col, "Zorunlu" if col == 1 else "İsteğe bağlı")
            ws.cell(6, col, key.upper())
            ws.cell(7, col, key)
        ws.freeze_panes = "C8"
        if name == "Teslimat":
            for key, value in {"schema_version": "1.0", "delivery_id": "CITY-999-v001", "city_id": "testville", "editorial_status": "Taslak"}.items():
                ws.cell(8, keys.index(key) + 1, value)
        elif name == "Şehir":
            for key, value in {"city_id": "testville", "name": "Testville", "country": "Testland", "country_code": "TL",
                               "region": "EU", "kind": "city", "latitude": 51.4, "longitude": 5.5, "aliases": "Testville"}.items():
                ws.cell(8, keys.index(key) + 1, value)
        elif name == "Konular":
            col = keys.index("used_count") + 1
            letter = openpyxl.utils.get_column_letter(col)
            for r in range(8, 8 + CAPACITY):
                ws.cell(r, col, f'=IF(A{r}="",0,COUNTIF(\'Yanıtlar\'!$K$8:$K$37,A{r}))')
        elif name == "Fotoğraflar":
            c1, c2 = keys.index("unique_venue_count") + 1, keys.index("unique_hero_count") + 1
            for r in range(8, 8 + CAPACITY):
                ws.cell(r, c1, f'=IF(OR(A{r}="",C{r}<>"Mekân"),0,1)')
                ws.cell(r, c2, f'=IF(OR(A{r}="",C{r}<>"Kapak"),0,1)')
        for (sheet, key), options in ENUMS.items():
            if sheet == name:
                letter = openpyxl.utils.get_column_letter(keys.index(key) + 1)
                dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
                dv.add(f"{letter}8:{letter}{7 + CAPACITY}")
                ws.add_data_validation(dv)
    wb.save(path)
    return Path(path)


def make_image(path: Path, fmt: str, size: tuple[int, int] = (64, 48), color: tuple[int, int, int] = (200, 30, 30)) -> str:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", size, color).save(path, format=fmt)
    return hashlib.sha256(path.read_bytes()).hexdigest()


def make_package(root: Path, n: int = 2) -> dict:
    """Mini paket: kokte Excel, Fotograflar/ (WebP web kopyasi), Lisans kanitlari/orijinaller/ (JPEG orijinal), Teslimler/.

    Excel'de Fotograflar.file = original_file (v7 sozlesmesi), web kopyasi web_file alaninda.
    """
    root = Path(root)
    xlsx = make_template(root / "CITY-999_testville_v001.xlsx") if not root.exists() else make_template(root / "CITY-999_testville_v001.xlsx")
    (root / "Teslimler").mkdir(exist_ok=True)
    wb = openpyxl.load_workbook(xlsx)
    photos = []
    for i in range(1, n + 1):
        orig_rel = f"Lisans kanıtları/orijinaller/testville_yer_{i:02d}.jpg"
        web_rel = f"Fotoğraflar/testville_yer_{i:02d}.webp"
        orig_sha = make_image(root / orig_rel, "JPEG", (640, 480), (30 * i, 80, 160))
        web_sha = make_image(root / web_rel, "WEBP", (320, 240), (30 * i, 80, 160))
        photos.append({"photo_id": f"testville:foto:yer-{i:02d}", "orig_rel": orig_rel, "orig_sha": orig_sha, "web_rel": web_rel, "web_sha": web_sha,
                       "source_url": f"https://commons.wikimedia.org/wiki/File:Yer_{i}.jpg"})
        ws = wb["Fotoğraflar"]
        row = {"photo_id": photos[-1]["photo_id"], "city_id": "testville", "role": "Mekân", "entry_id": f"testville:yer-{i}", "file": orig_rel,
               "sha256": orig_sha, "caption": f"Yer {i}", "author": "Test Yazar", "license": "CC BY-SA 4.0",
               "license_url": "https://creativecommons.org/licenses/by-sa/4.0/", "source_url": photos[-1]["source_url"],
               "evidence_id": f"testville:kanit:yer-{i:02d}", "order": i, "editorial_status": "Taslak"}
        keys = SHEET_KEYS["Fotoğraflar"]
        for k, v in row.items():
            ws.cell(7 + i, keys.index(k) + 1, v)
        ws = wb["FotoğrafKaynakları"]
        keys = SHEET_KEYS["FotoğrafKaynakları"]
        for k, v in {"photo_id": photos[-1]["photo_id"], "original_file": orig_rel, "original_sha256": orig_sha, "web_file": web_rel,
                     "web_sha256": web_sha, "web_width": 320, "web_height": 240}.items():
            ws.cell(7 + i, keys.index(k) + 1, v)
    wb.save(xlsx)
    return {"root": root, "xlsx": xlsx, "photos": photos}
