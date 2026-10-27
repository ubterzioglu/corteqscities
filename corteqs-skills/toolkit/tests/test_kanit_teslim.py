"""kanit_teslim: uretimdeki inceleme kanitlari Teslimler'e (Sevilla dis denetimi: B3/B5/C2-C5/D4 kanitsiz 0 aldi)."""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import kanit_kapsami as kk  # noqa: E402
import kanit_teslim as kt  # noqa: E402

SHEETS = {"Mekânlar": ["entry_id"], "Mahalleler": ["entry_id"], "Rotalar": ["route_id"],
          "Fotoğraflar": ["photo_id", "entry_id", "role"]}
ROWS = {"Mekânlar": [{"entry_id": "t:m1"}, {"entry_id": "t:m2"}], "Mahalleler": [{"entry_id": "t:h1"}],
        "Rotalar": [{"route_id": "t:route:001"}, {"route_id": "t:route:002"}],
        "Fotoğraflar": [{"photo_id": f"t:photo:p{i}", "entry_id": "t:m1", "role": "Mekân"} for i in range(1, 5)]
        + [{"photo_id": "t:photo:kapak-01", "entry_id": "", "role": "Kapak"}]}
KEYS = {f"t:photo:p{i}": f"t:m1|File:P{i}.jpg" for i in range(1, 5)} | {"t:photo:kapak-01": "|File:K.jpg"}


def _xlsx(path: Path, rows: dict = ROWS) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, keys in SHEETS.items():
        ws = wb.create_sheet(sheet)
        for col, key in enumerate(keys, start=1):
            ws.cell(6, col, key.upper())
            ws.cell(7, col, key)
        for row, record in enumerate(rows[sheet], start=8):
            for key, value in record.items():
                ws.cell(row, keys.index(key) + 1, value)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _jsonl(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("".join(json.dumps(r, ensure_ascii=False) + "\n" for r in rows), encoding="utf-8")


def _work(tmp_path: Path) -> Path:
    work = tmp_path / "_is"
    _jsonl(work / "rota_plan.jsonl", [
        {"route_id": "t:route:001", "gun": "salı", "durak": "A", "varis_saati": "10:00", "acik_mi": "evet", "dayanak": "m1"},
        {"route_id": "t:route:002", "gun": "salı", "durak": "B", "varis_saati": "11:00", "acik_mi": "dışarıdan", "dayanak": "m2"},
        {"route_id": "t:route:999", "gun": "salı", "durak": "Z", "varis_saati": "12:00", "acik_mi": "evet", "dayanak": "silindi"}])
    _jsonl(work / kt.LEGS_NAME, [
        {"route_id": "t:route:001", "bacak": 1, "from": "t:m1", "to": "t:m2", "kus_ucusu_m": 480, "yuruyus_m": 563, "yuruyus_dk": 8},
        {"route_id": "t:route:002", "bacak": 1, "from": "t:m2", "to": "t:m1", "kus_ucusu_m": 480, "yuruyus_m": 570, "yuruyus_dk": 8}])
    _jsonl(work / "denetim" / "mekan_mekan-dogrulama.jsonl", [
        {"entry_id": "t:m1", "sonuc": "kaldı", "not": "saat portaldan"}, {"entry_id": "t:m2", "sonuc": "geçti"},
        {"entry_id": "t:h1", "sonuc": "geçti"}])
    _jsonl(work / "denetim" / "mekan-d5_mekan-dogrulama.jsonl", [{"entry_id": "t:m1", "sonuc": "geçti"}])  # son tur ezer
    _jsonl(work / "foto" / "secim_temiz.jsonl", [{"photo_id": p, "key": k} for p, k in KEYS.items()])
    ins = work / "foto" / "inceleme"
    _jsonl(ins / "review_tur1_1.jsonl", [{"key": KEYS["t:photo:p1"], "result": "keep", "reason": "cephe", "ek_hak": "yok"},
                                         {"key": KEYS["t:photo:p2"], "result": "keep"}])
    _jsonl(ins / "verify_tur1_1.jsonl", [{"key": KEYS["t:photo:p2"], "ok": False, "why": "maket"}])
    _jsonl(ins / "covers.jsonl", [{"key": KEYS["t:photo:kapak-01"], "role": "Kapak"}])
    _jsonl(work / "foto_d3" / "karar.jsonl", [
        {"photo_id": "eski", "entry_id": "t:m1", "karar": "degistir", "yeni_commons_file": "File:P3.jpg", "gorulen": "iç"},
        {"photo_id": "eski2", "entry_id": "t:m1", "karar": "cikar"}])
    _jsonl(ins / "denetim_1.jsonl", [{"photo_id": "t:photo:p4", "gorsel_sonuc": "keep", "ek_hak": "yok", "not": "açık alan"}])
    return work


def _run(capsys, tmp_path, quality=True, work=None):
    pkg = tmp_path / "PKG"
    xlsx = _xlsx(pkg / "CITY-1_t_v001.xlsx")
    work = work or _work(tmp_path)
    q = tmp_path / "gk" / kt.QUALITY_OUT
    q.parent.mkdir(parents=True, exist_ok=True)
    q.write_text(json.dumps({"photos": [{"photo_id": p, "flags": []} for p in KEYS]}), encoding="utf-8")
    argv = ["--xlsx", str(xlsx), "--work", str(work), "--pkg", str(pkg)] + (["--quality", str(q)] if quality else [])
    code = kt.main(argv)
    summary = json.loads(capsys.readouterr().out.strip().splitlines()[-1])
    return code, summary, pkg / "Teslimler", xlsx, pkg


def _read(path: Path) -> list[dict]:
    return [json.loads(l) for l in path.read_text(encoding="utf-8").splitlines() if l.strip()]


def test_writes_all_four_evidence_files_with_precedence(capsys, tmp_path):
    code, s, t, _, _ = _run(capsys, tmp_path)
    assert code == 0 and s["eksik"] == {} and (t / kt.QUALITY_OUT).is_file()
    route_out = _read(t / kt.ROUTE_OUT)
    assert {r["route_id"] for r in route_out} == {"t:route:001", "t:route:002"}  # Excel'de olmayan rota yok
    assert len(route_out) == 4 and sum(1 for r in route_out if "bacak" in r) == 2
    places = {r["entry_id"]: r for r in _read(t / kt.PLACE_OUT)}
    assert places["t:m1"]["sonuc"] == "geçti" and places["t:m1"]["tur"] == "mekan-d5" and s["mekan_gecti"] == 3
    photos = {r["photo_id"].split(":")[-1]: r for r in _read(t / kt.PHOTO_OUT)}
    assert (photos["p1"]["gorsel_sonuc"], photos["p1"]["dayanak"]) == ("keep", "inceleme")
    assert (photos["p2"]["gorsel_sonuc"], photos["p2"]["dayanak"]) == ("reject", "ikinci göz")
    assert (photos["p3"]["gorsel_sonuc"], photos["p3"]["dayanak"]) == ("keep", "ek tur kararı")
    assert (photos["p4"]["ek_hak"], photos["p4"]["dayanak"]) == ("yok", "son foto denetimi")
    assert photos["kapak-01"]["gorsel_sonuc"] == "keep" and photos["p1"]["ek_hak"] == "yok"
    assert photos["p3"]["ek_hak"] == kt.UNKNOWN_RIGHTS and s["foto_ek_hak_belirsiz"] == 3  # yazilmayan ek_hak uydurulmaz


def test_missing_inputs_are_reported_not_invented(capsys, tmp_path):
    work = _work(tmp_path)
    (work / "rota_plan.jsonl").write_text(json.dumps({"route_id": "t:route:001", "acik_mi": "evet"}) + "\n", encoding="utf-8")
    (work / "foto" / "inceleme" / "review_tur1_1.jsonl").unlink()
    _jsonl(work / kt.LEGS_NAME, [{"route_id": "t:route:001", "bacak": 1, "sorun": "koordinatsiz durak"}])
    code, s, t, _, _ = _run(capsys, tmp_path, quality=False, work=work)
    assert code == 1 and s["eksik"]["rota_plansiz"] == ["t:route:002"] and s["eksik"]["foto_incelemesiz"] == ["t:photo:p1"]
    assert s["eksik"]["rota_bacaksiz"] == ["t:route:001", "t:route:002"]
    assert s["eksik_sayim"]["gorsel_kalite"] == 1 and not (t / kt.QUALITY_OUT).exists()


def test_output_is_read_by_kanit_kapsami(capsys, tmp_path, monkeypatch):
    """Uretici ile tuketici ayni alan adlarini kullanir (Sevilla: gorsel_kalite 'photos' yazdi, kapsam 'rows' okudu)."""
    _, _, t, xlsx, pkg = _run(capsys, tmp_path)
    monkeypatch.setattr(kk, "TARGETS", {k: 1 for k in kk.TARGETS})
    by = {c["id"]: c for c in kk.evaluate(xlsx, t)}
    assert (by["B3"]["kanitli"], by["B3"]["payda"]) == (2, 2)
    assert (by["B5"]["kanitli"], by["B5"]["payda"]) == (3, 3)
    assert (by["C2"]["kanitli"], by["C2"]["payda"]) == (4, 5)  # p2 ikinci gözden kaldı
    assert (by["C4"]["kanitli"], by["D4"]["kanitli"]) == (5, 2)


def test_usage_error_on_missing_xlsx(capsys, tmp_path):
    assert kt.main(["--xlsx", str(tmp_path / "yok.xlsx"), "--work", str(tmp_path), "--pkg", str(tmp_path)]) == 2
    with pytest.raises(SystemExit):
        kt.main(["--xlsx", str(tmp_path / "yok.xlsx")])
