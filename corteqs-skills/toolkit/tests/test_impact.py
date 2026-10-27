"""impact.py testleri: bagimlilik grafigi, --ids etkisi, --diff, CLI sozlesmesi (K99/K118)."""
from __future__ import annotations

import copy
import json
from pathlib import Path
from typing import Any, Mapping

import openpyxl
import pytest

import helpers
import impact

C = "testville"
TOOLKIT = Path(__file__).resolve().parent.parent
REAL_XLSX = Path(r"C:\temp_private\corteqs\0000_newera\arsiv\eindhoven_qwen38_max\CITY-039_eindhoven_v001.xlsx")


def base_rows() -> dict[str, list[dict[str, Any]]]:
    js = lambda v: json.dumps(v, ensure_ascii=False)  # noqa: E731
    return {
        "Şehir": [{"hero_photo_id": "testville:foto:kapak-01"}],
        "Konular": [{"topic_id": "testville:konut", "city_id": C, "slug": "konut", "title": "Konut"}],
        "Yanıtlar": [
            {"entry_id": "testville:kira", "city_id": C, "slug": "kira", "title": "Kira nasıl ödenir?", "summary": "Kira sözleşmesi.",
             "source_ids": "testville:kaynak:a;testville:kaynak:b", "topic_id": "testville:konut"},
            {"entry_id": "testville:park-rehberi", "city_id": C, "slug": "park-rehberi", "title": "Hafta sonu nereye?",
             "summary": "Hafta sonu Stadspark'ta piknik yapılır.", "source_ids": "testville:kaynak:b", "topic_id": "testville:konut"},
        ],
        "Kaynaklar": [{"source_id": f"testville:kaynak:{s}", "city_id": C, "url": f"https://{s}.example/x"} for s in "abc"],
        "Mekânlar": [
            {"entry_id": "testville:stadspark", "city_id": C, "slug": "stadspark", "title": "Stadspark", "summary": "Büyük park.",
             "source_ids": "testville:kaynak:c", "location": "Parkweg 1"},
            {"entry_id": "testville:muze", "city_id": C, "slug": "muze", "title": "Testville Müzesi (Merkez)", "summary": "Şehir müzesi.",
             "source_ids": "testville:kaynak:a", "location": "Museumplein 2"},
        ],
        "Mahalleler": [{"entry_id": "testville:mah-merkez", "city_id": C, "slug": "mah-merkez", "title": "Merkez",
                        "summary": "Stadspark ve TESTVILLE MÜZESI yakınında.", "source_ids": "testville:kaynak:a"}],
        "Bağlantılar": [
            {"action_id": "testville:action:stadspark-1", "city_id": C, "entry_id": "testville:stadspark", "label": "Haritada aç",
             "url": "https://maps.example/?q=Stadspark"},
            {"action_id": "testville:action:kira-1", "city_id": C, "entry_id": "testville:kira", "label": "Başvur", "url": "https://a.example/kira"},
        ],
        "Rotalar": [{"route_id": "testville:route:001", "city_id": C, "title": "Park ve müze", "stops": "Stadspark\nTestville Müzesi",
                     "stop_entry_ids": "testville:stadspark;testville:muze", "entry_id": "testville:stadspark", "source_ids": "testville:kaynak:a"}],
        "Sözlük": [{"term_id": "testville:term:001", "city_id": C, "term": "Huur", "definition": "Kira.", "source_ids": "testville:kaynak:b"}],
        "Adımlar": [{"step_id": "testville:step:1", "city_id": C, "entry_id": "testville:kira", "title": "Kirala", "description": "Ev bul."}],
        "SayfaAlanları": [
            {"field_id": "testville:field:facts", "city_id": C, "key": "facts", "value_type": "JSON",
             "value": js([{"entry": "kira", "label": "Kira", "value": "1"}, {"entry": "stadspark", "label": "Park", "value": "2"}])},
            {"field_id": "testville:field:quick_links", "city_id": C, "key": "quick_links", "value_type": "JSON",
             "value": js([{"label": "Müze", "entry": "muze"}])},
            {"field_id": "testville:field:advice_entry", "city_id": C, "key": "advice_entry", "value_type": "Metin", "value": "kira"},
            {"field_id": "testville:field:timeline", "city_id": C, "key": "timeline", "value_type": "JSON",
             "value": js([{"offset": -30, "entry": "kira", "title": "Hazırlan"}])},
            {"field_id": "testville:field:city_overview", "city_id": C, "key": "city_overview", "value_type": "Metin",
             "value": "Yeşil Stadspark şehrin kalbidir."},
        ],
        "EkDetaylar": [{"detail_id": "testville:detail:muze-1", "city_id": C, "entry_id": "testville:muze", "key": "hours", "value": "10-17"}],
        "Fotoğraflar": [
            {"photo_id": "testville:foto:stadspark-01", "city_id": C, "role": "Mekân", "entry_id": "testville:stadspark",
             "evidence_id": "testville:kanit:stadspark-01", "caption": "Stadspark"},
            {"photo_id": "testville:foto:kapak-01", "city_id": C, "role": "Kapak", "evidence_id": "testville:kanit:kapak-01", "caption": "Kapak"},
        ],
        "Kanıtlar": [
            {"evidence_id": "testville:kanit:stadspark-01", "city_id": C, "photo_id": "testville:foto:stadspark-01"},
            {"evidence_id": "testville:kanit:kapak-01", "city_id": C, "photo_id": "testville:foto:kapak-01"},
        ],
        "FotoğrafKaynakları": [
            {"photo_id": "testville:foto:stadspark-01", "web_file": "Fotoğraflar/a.webp"},
            {"photo_id": "testville:foto:kapak-01", "web_file": "Fotoğraflar/b.webp"},
        ],
    }


def build_book(path: Path, rows: Mapping[str, list[dict[str, Any]]]) -> Path:
    """helpers.make_template uzerine eksik sayfa/anahtarlari ekler; kayitlar satir 8'den yazilir."""
    helpers.make_template(path)
    wb = openpyxl.load_workbook(path)
    for sheet, records in rows.items():
        ws = wb[sheet] if sheet in wb.sheetnames else wb.create_sheet(sheet)
        keys = [c.value for c in ws[7] if c.value]
        for rec in records:
            for key in rec:
                if key not in keys:
                    keys.append(key)
                    ws.cell(7, len(keys), key)
        for offset, rec in enumerate(records):
            for key, value in rec.items():
                ws.cell(8 + offset, keys.index(key) + 1, value)
    wb.save(path)
    return path


@pytest.fixture()
def book(tmp_path) -> Path:
    return build_book(tmp_path / "base.xlsx", base_rows())


@pytest.fixture()
def graph(book):
    return impact.build_graph(impact.load_sheets(book))


def rows_of(items) -> set[tuple[str, int]]:
    return {(i["sheet"], i["_row"]) for i in items}


def find(items, sheet: str, row: int) -> list[dict[str, Any]]:
    return [i for i in items if i["sheet"] == sheet and i["_row"] == row]


# --- birim: ayrisma -------------------------------------------------------------------------
def test_split_ids_handles_separators_and_blanks():
    assert impact.split_ids("a; b;;\nc ") == ["a", "b", "c"]
    assert impact.split_ids(None) == []
    assert impact.split_ids(42) == ["42"]


def test_json_entry_slugs_walks_nested_and_tolerates_bad_json():
    value = json.dumps([{"entry": "a"}, {"x": {"entry": "b"}}, {"entry": ""}, "c"])
    assert impact.json_entry_slugs(value) == ["a", "b"]
    assert impact.json_entry_slugs("{bozuk") == []
    assert impact.json_entry_slugs(None) == []
    assert impact.json_entry_slugs({"entry": "z"}) == ["z"]


# --- graf ve --ids ----------------------------------------------------------------------------
def test_graph_city_and_reverse_refs(graph):
    assert graph.city == C
    sources = {(key[0], ref.field) for key, ref in graph.reverse["testville:stadspark"]}
    assert {("Bağlantılar", "entry_id"), ("Rotalar", "entry_id"), ("Rotalar", "stop_entry_ids"),
            ("Fotoğraflar", "entry_id"), ("SayfaAlanları", "facts.entry")} <= sources


def test_resolve_slug_full_id_and_unknown(graph):
    assert impact.resolve(graph, "stadspark") == ["testville:stadspark"]
    assert impact.resolve(graph, "testville:kaynak:b") == ["testville:kaynak:b"]
    assert impact.resolve(graph, "konut") == ["testville:konut"]
    assert impact.resolve(graph, "yok") == []
    assert impact.resolve(graph, "testville:yok") == []
    assert impact.resolve(graph, "kaynak:b") == ["testville:kaynak:b"]  # sehir oneki eksik kisa kimlik tamamlanir
    assert impact.resolve(graph, "kaynak:yok") == []


def test_place_impact_structural_text_and_level_two(graph):
    items = impact.impact(graph, "testville:stadspark", depth=2)
    assert find(items, "Mekânlar", 8)[0]["relation"] == "self"
    route = {i["field"]: i["relation"] for i in find(items, "Rotalar", 8)}
    assert route == {"entry_id": "route_start", "stop_entry_ids": "route_stop"}
    assert find(items, "Bağlantılar", 8)[0]["relation"] == "action_of_entry"
    assert find(items, "Fotoğraflar", 8)[0]["relation"] == "photo_of_entry"
    assert find(items, "SayfaAlanları", 8)[0]["field"] == "facts.entry"
    fwd = find(items, "Kaynaklar", 10)[0]
    assert (fwd["direction"], fwd["relation"], fwd["field"]) == ("forward", "cites_source", "source_ids")
    mentions = {(i["sheet"], i["_row"], i["field"]) for i in items if i["direction"] == "text"}
    assert mentions == {("Yanıtlar", 9, "summary"), ("Mahalleler", 8, "summary"), ("SayfaAlanları", 12, "value")}
    level2 = {(i["sheet"], i["_row"], i["via"]) for i in items if i["level"] == 2}
    assert level2 == {("Kanıtlar", 8, "testville:foto:stadspark-01"), ("FotoğrafKaynakları", 8, "testville:foto:stadspark-01")}
    assert ("Yanıtlar", 8) not in rows_of(items)


def test_mention_uses_title_without_parenthetical_case_insensitive(graph):
    items = impact.impact(graph, "testville:muze", depth=1)
    assert ("Mahalleler", 8) in {(i["sheet"], i["_row"]) for i in items if i["direction"] == "text"}
    assert find(items, "Rotalar", 8)[0]["direction"] == "backward"  # yapisal bag metin bagini ezer
    assert find(items, "EkDetaylar", 8)[0]["relation"] == "detail_of_entry"
    assert find(items, "SayfaAlanları", 9)[0]["field"] == "quick_links.entry"


def test_mention_patterns_fold_turkish_i_without_partial_words():
    pats = impact.mention_patterns({"title": "Işık Parkı (Merkez)", "slug": "isik-parki"})
    texts = ("IŞIK PARKI'nda", "ışık parkı", "İŞIK  PARKI", "isik-parki", "Işıklar Parkı", "XIşık Parkı")
    assert [t for t in texts if any(p.search(t) for p in pats)] == list(texts[:4])


def test_empty_data_sheets_are_tolerated(tmp_path):
    g = impact.build_graph(impact.load_sheets(build_book(tmp_path / "e.xlsx", {**base_rows(), "Rotalar": [], "Sözlük": []})))
    sheets = {s for s, _ in rows_of(impact.impact(g, "testville:stadspark", depth=2))}
    assert "Mekânlar" in sheets and not sheets & {"Rotalar", "Sözlük"}


def test_source_impact_two_levels_and_depth_one(graph):
    items = impact.impact(graph, "testville:kaynak:b", depth=2)
    level1 = {(i["sheet"], i["_row"]) for i in items if i["level"] == 1}
    assert level1 == {("Yanıtlar", 8), ("Yanıtlar", 9), ("Sözlük", 8)}
    level2 = {(i["sheet"], i["_row"], i["field"]) for i in items if i["level"] == 2}
    assert {("Bağlantılar", 9, "entry_id"), ("Adımlar", 8, "entry_id"), ("SayfaAlanları", 8, "facts.entry"),
            ("SayfaAlanları", 10, "advice_entry"), ("SayfaAlanları", 11, "timeline.entry")} == level2
    assert all(i["via"] == "testville:kira" for i in items if i["level"] == 2)
    shallow = impact.impact(graph, "testville:kaynak:b", depth=1)
    assert max(i["level"] for i in shallow) == 1


def test_photo_cycle_is_bounded_and_hero_found(graph):
    items = impact.impact(graph, "testville:foto:kapak-01", depth=2)
    assert rows_of(items) == {("Fotoğraflar", 9), ("Şehir", 8), ("Kanıtlar", 9), ("FotoğrafKaynakları", 9)}
    assert find(items, "Şehir", 8)[0]["relation"] == "hero_photo"
    assert all(i["level"] <= 2 for i in items)


def test_topic_and_answer_forward_links(graph):
    topic = impact.impact(graph, "testville:konut", depth=1)
    assert {(i["sheet"], i["_row"], i["relation"]) for i in topic if i["level"] == 1} == {
        ("Yanıtlar", 8, "answer_topic"), ("Yanıtlar", 9, "answer_topic")}
    answer = impact.impact(graph, "testville:kira", depth=1)
    forward = {(i["sheet"], i["_row"]) for i in answer if i["direction"] == "forward"}
    assert forward == {("Kaynaklar", 8), ("Kaynaklar", 9), ("Konular", 8)}


def test_items_are_sorted_and_inputs_not_mutated(book):
    sheets = impact.load_sheets(book)
    snapshot = copy.deepcopy(sheets)
    g = impact.build_graph(sheets)
    first = impact.impact(g, "testville:stadspark", depth=2)
    assert first == impact.impact(g, "testville:stadspark", depth=2)
    impact.diff_sheets(sheets, sheets)
    assert sheets == snapshot
    levels = [i["level"] for i in first]
    assert levels == sorted(levels)


# --- diff -------------------------------------------------------------------------------------
def changed_rows() -> dict[str, list[dict[str, Any]]]:
    rows = base_rows()
    rows["Mekânlar"][0] = {**rows["Mekânlar"][0], "location": "Parkweg 99"}
    rows["Fotoğraflar"][0] = {**rows["Fotoğraflar"][0], "entry_id": "testville:muze"}
    rows["Kaynaklar"] = [r for r in rows["Kaynaklar"] if r["source_id"] != "testville:kaynak:c"]
    rows["Kaynaklar"].append({"source_id": "testville:kaynak:d", "city_id": C, "url": "https://d.example/x"})
    return rows


def test_diff_sheets_detects_changed_added_deleted(book, tmp_path):
    new = build_book(tmp_path / "new.xlsx", changed_rows())
    changes = impact.diff_sheets(impact.load_sheets(book), impact.load_sheets(new))
    got = {(c["sheet"], c["id"], c["change"], tuple(c["fields"])) for c in changes}
    assert got == {
        ("Mekânlar", "testville:stadspark", "changed", ("location",)),
        ("Fotoğraflar", "testville:foto:stadspark-01", "changed", ("entry_id",)),
        ("Kaynaklar", "testville:kaynak:c", "deleted", ()),
        ("Kaynaklar", "testville:kaynak:d", "added", ()),
    }
    loc = next(c for c in changes if c["sheet"] == "Mekânlar")
    assert loc["values"]["location"] == {"old": "Parkweg 1", "new": "Parkweg 99"}


def test_diff_duplicate_ids_and_row_moves_are_not_changes():
    old = {"Bağlantılar": [{"action_id": "x", "url": "u", "_row": 8}, {"action_id": "x", "url": "v", "_row": 9}]}
    new = {"Bağlantılar": [{"action_id": "x", "url": "u", "_row": 10}, {"action_id": "x", "url": "v", "_row": 11}]}
    assert impact.diff_sheets(old, new) == []
    keyless = {"Bağlantılar": [{"label": "a", "_row": 8}]}
    assert impact.diff_sheets(keyless, {"Bağlantılar": []})[0]["change"] == "deleted"


def test_analyze_diff_previous_refs_dangling_and_checks(book, tmp_path):
    new = build_book(tmp_path / "new.xlsx", changed_rows())
    report = impact.analyze(new, [], book, depth=2, catalog=impact.load_catalog(TOOLKIT / "checks.yaml"))
    photo = report["impacts"]["testville:foto:stadspark-01"]["items"]
    forward = {(i["sheet"], i["_row"], bool(i.get("previous"))) for i in photo if i["direction"] == "forward"}
    assert ("Mekânlar", 9, False) in forward and ("Mekânlar", 8, True) in forward
    deleted = report["impacts"]["testville:kaynak:c"]["items"]
    assert any(i.get("status") == "deleted" and i["sheet"] == "Kaynaklar" for i in deleted)
    dangling = [i for i in deleted if i.get("dangling")]
    assert rows_of(dangling) == {("Mekânlar", 8)}
    codes = [c["id"] for c in report["checks"]]
    assert {"K12", "K22", "K42", "K55", "K57", "K62", "K80", "K29", "K35", "K50", "K51"} <= set(codes)
    assert codes == sorted(codes, key=lambda c: int(c[1:]))
    assert report["counts"]["changes"] == {"changed": 2, "added": 1, "deleted": 1}


def test_diff_rename_finds_stale_mentions_of_old_title(book, tmp_path):
    rows = base_rows()
    rows["Mekânlar"][1] = {**rows["Mekânlar"][1], "title": "Kent Müzesi"}  # slug 'muze' kisa: yalniz eski baslik bulur
    new = build_book(tmp_path / "renamed.xlsx", rows)
    report = impact.analyze(new, [], book, depth=1, catalog={})
    items = report["impacts"]["testville:muze"]["items"]
    stale = {(i["sheet"], i["_row"], i["field"]) for i in items if i["direction"] == "text"}
    assert ("Mahalleler", 8, "summary") in stale


# --- kontrol onerisi --------------------------------------------------------------------------
def test_load_catalog_real_and_missing(tmp_path):
    cat = impact.load_catalog(TOOLKIT / "checks.yaml")
    assert cat["K46"]["title"] == "Rota veri ilişkileri" and cat["K46"]["exec"] == "S"
    assert impact.load_catalog(tmp_path / "yok.yaml") == {}
    bad = tmp_path / "bad.yaml"
    bad.write_text("checks: [", encoding="utf-8")
    assert impact.load_catalog(bad) == {}


def test_recommend_checks_route_and_unknown_catalog():
    items = [{"sheet": "Rotalar", "_row": 8}]
    rec = impact.recommend_checks(items, [], {"K46": {"title": "t", "exec": "S"}})
    assert [r["id"] for r in rec] == ["K46"]  # katalogda olmayan K47 dusurulur
    assert rec[0]["sheets"] == ["Rotalar"]
    raw = impact.recommend_checks(items, [], {})
    assert [r["id"] for r in raw] == ["K46", "K47"] and raw[0]["title"] is None


# --- CLI --------------------------------------------------------------------------------------
def run_cli(capsys, argv: list[str]) -> tuple[int, dict[str, Any]]:
    code = impact.main(argv)
    lines = capsys.readouterr().out.strip().splitlines()
    assert len(lines) == 1
    return code, json.loads(lines[0])


def test_cli_ids_writes_out_deterministically(book, tmp_path, capsys):
    out1, out2 = tmp_path / "a.json", tmp_path / "b.json"
    code, summary = run_cli(capsys, [str(book), "--ids", "stadspark, testville:kaynak:b", "--out", str(out1)])
    assert code == impact.common.EXIT_FINDINGS
    assert summary["seeds"] == 2 and summary["unknown_ids"] == []
    assert "K46" in summary["checks"] and summary["by_sheet"]["Rotalar"] == 1
    run_cli(capsys, [str(book), "--ids", "stadspark,testville:kaynak:b", "--out", str(out2)])
    assert out1.read_bytes() == out2.read_bytes()
    report = json.loads(out1.read_text(encoding="utf-8"))
    assert set(report["impacts"]) == {"testville:stadspark", "testville:kaynak:b"}
    assert report["impacts"]["testville:stadspark"]["tokens"] == ["stadspark"]


def test_no_text_skips_mentions_and_summary_counts_text_only(graph, book, capsys):
    items = impact.impact(graph, "testville:stadspark", depth=2, text=False)
    assert items and all(i["direction"] != "text" for i in items)
    _, full = run_cli(capsys, [str(book), "--ids", "stadspark"])
    _, lean = run_cli(capsys, [str(book), "--ids", "stadspark", "--no-text"])
    assert (full["text_only"], lean["text_only"], full["impacted"] - lean["impacted"]) == (3, 0, 3)


@pytest.mark.parametrize("target", ["xlsx", "diff", "dotted", "checks"])
def test_cli_out_never_overwrites_inputs(book, tmp_path, capsys, target):
    old, checks = tmp_path / "old.xlsx", tmp_path / "checks.yaml"
    old.write_bytes(book.read_bytes())
    checks.write_bytes((TOOLKIT / "checks.yaml").read_bytes())
    (tmp_path / "sub").mkdir()
    out = {"xlsx": book, "diff": old, "dotted": tmp_path / "sub" / ".." / book.name, "checks": checks}[target]
    before = [p.read_bytes() for p in (book, old, checks)]
    code, summary = run_cli(capsys, [str(book), "--diff", str(old), "--checks", str(checks), "--out", str(out)])
    assert code == impact.common.EXIT_USAGE and "--out" in summary["error"]
    assert [p.read_bytes() for p in (book, old, checks)] == before


def test_cli_partial_unknown_is_finding(book, capsys):
    code, summary = run_cli(capsys, [str(book), "--ids", "stadspark,yok", "--depth", "1"])
    assert code == 1 and summary["unknown_ids"] == ["yok"]


def test_cli_identical_diff_is_clean(book, tmp_path, capsys):
    copy_path = tmp_path / "copy.xlsx"
    copy_path.write_bytes(book.read_bytes())
    code, summary = run_cli(capsys, [str(copy_path), "--diff", str(book)])
    assert code == impact.common.EXIT_OK
    assert summary["changes"] == {"changed": 0, "added": 0, "deleted": 0} and summary["impacted"] == 0


def test_cli_cell_edit_on_copy_is_detected(book, tmp_path, capsys):
    edited = tmp_path / "edited.xlsx"
    edited.write_bytes(book.read_bytes())
    wb = openpyxl.load_workbook(edited)
    ws = wb["Fotoğraflar"]
    col = [c.value for c in ws[7]].index("entry_id") + 1
    ws.cell(8, col, "testville:muze")
    wb.save(edited)
    code, summary = run_cli(capsys, [str(edited), "--diff", str(book), "--ids", "kira"])
    assert code == 1 and summary["mode"] == "ids+diff"
    assert summary["changes"]["changed"] == 1


@pytest.mark.parametrize("argv_tail", [[], ["--ids", " , "], ["--ids", "yok,testville:hic"]])
def test_cli_usage_errors(book, capsys, argv_tail):
    code, summary = run_cli(capsys, [str(book), *argv_tail])
    assert code == impact.common.EXIT_USAGE and "error" in summary


def test_cli_unreadable_inputs(book, tmp_path, capsys):
    bad = tmp_path / "bad.xlsx"
    bad.write_text("excel degil", encoding="utf-8")
    assert run_cli(capsys, [str(bad), "--ids", "x"])[0] == 2
    assert run_cli(capsys, [str(book), "--diff", str(tmp_path / "yok.xlsx")])[0] == 2


def test_cli_duplicate_key_row_is_usage_error(tmp_path, capsys):
    path = build_book(tmp_path / "dup.xlsx", base_rows())
    wb = openpyxl.load_workbook(path)
    ws = wb["Rotalar"]
    ws.cell(7, ws.max_column + 1, "route_id")
    wb.save(path)
    assert run_cli(capsys, [str(path), "--ids", "stadspark"])[0] == 2


# --- gercek veri (salt okunur) ----------------------------------------------------------------
@pytest.mark.real
@pytest.mark.skipif(not REAL_XLSX.exists(), reason="gercek qwen paketi yok")
def test_real_wilhelminaplein_and_kvk():
    report = impact.analyze(REAL_XLSX, ["eindhoven:wilhelminaplein", "kaynak:kvk"], None, depth=2, catalog={})
    wil = rows_of(report["impacts"]["eindhoven:wilhelminaplein"]["items"])
    assert {("Mekânlar", 58), ("Mahalleler", 13), ("Rotalar", 12)} <= wil
    kvk = rows_of(report["impacts"]["eindhoven:kaynak:kvk"]["items"])
    assert {("Yanıtlar", 224), ("Yanıtlar", 234)} <= kvk
