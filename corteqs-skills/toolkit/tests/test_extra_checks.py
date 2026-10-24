"""extra_checks: resmi dogrulayicida olmayan proje kontrolleri ('ek kontrol')."""
from __future__ import annotations

import copy

import openpyxl
import pytest

import extra_checks as ec
import helpers


def _sources(*rows):
    return [{"source_id": sid, "url": url, **extra} for sid, url, extra in rows]


# --- placeholders -----------------------------------------------------------

@pytest.mark.parametrize("value", ["5621?", "[[Hoogstraat (Eindhoven)", "Bkz. ]]", "{{isim}}", "TODO: yaz", "Lorem ipsum dolor",
                                   "Adres: Markt 1?", "Hoogstraat?", "Boschdijk 525, 5621? Eindhoven",
                                   "Groene Loper 3? / Den Dolech, 5612 Eindhoven"])
def test_placeholders_flags_markers(value):
    sheets = {"Mekânlar": [{"_row": 9, "entry_id": "x:a", "summary": value}]}
    result = ec.placeholders(sheets)
    assert result["count"] == 1
    item = result["items"][0]
    assert item["sheet"] == "Mekânlar" and item["row"] == 9 and item["key"] == "summary" and item["id"] == "x:a"


@pytest.mark.parametrize("value", ["Eindhoven'a taşınmadan önce hangi sırayla plan yapmalıyım?", "Normal metin.",
                                   "https://example.org/a?b=1", "https://example.org/page2?id=7", "Todos los santos",
                                   "Saat 9'da mı açılıyor?", 42])
def test_placeholders_ignores_clean_values(value):
    sheets = {"Yanıtlar": [{"_row": 8, "entry_id": "x:b", "title": value}]}
    assert ec.placeholders(sheets)["count"] == 0


PILOT_STEPS = (  # Eindhoven pilotu, Yanıtlar 'belastingaangifte-yapma-yollari' steps (yanlis alarm: '...2025?')
    "Hazırlık yapın: DigiD'iniz (veya Avrupa girişiniz), BSN'iniz, IBAN'ınız, 2025 yıllık gelir özetiniz ve hesap özetleriniz elinizde olsun.\n"
    "Mijn Belastingdienst'e DigiD ile girip \"Inkomstenbelasting\" ve \"Belastingjaar 2025\" bölümünü açın.\n"
    "Vergi partneriniz varsa ortak beyan için \"Wilt u samen aangifte doen?\" sorusuna \"ja\" deyin.\n"
    "Süre yetmeyecekse \"Kan ik nog uitstel aanvragen voor de belastingaangifte 2025?\" başlığından erteleme seçeneğini inceleyin.")


@pytest.mark.parametrize("value", [  # A13: tirnak icindeki metin ve '?' ile biten cok kelimeli baslik yer tutucu degil
    PILOT_STEPS,
    "Belastingdienst'in “Wat verandert er in 2026?” sayfasına bakın.", "„Was ändert sich 2026?“ başlıklı sayfa.",
    "«Toeslagen 2026?» başlıklı duyuru.", "»Kindergeld 2026?« sayfası.", "Sayfa başlığı 'Huurtoeslag 2026?' olarak geçiyor.",
    "‘Zorgtoeslag 2026?’ sayfası.", "\"Waarom?\" adlı sergi.",
    "Kan ik nog uitstel aanvragen voor de belastingaangifte 2025?",
    "Süre yetmeyecekse Kan ik nog uitstel aanvragen voor de belastingaangifte 2025? başlığına bakın.",
])
def test_placeholders_ignore_quoted_text_and_question_titles(value):
    sheets = {"Yanıtlar": [{"_row": 177, "entry_id": "x:belastingaangifte", "steps": value}]}
    assert ec.placeholders(sheets)["count"] == 0


@pytest.mark.parametrize("value", [  # V4 inceleme: baslik bicimden taninir (soru sozcugu / yardimci fiil + ozne), uzunluktan degil
    "Zorgtoeslag: wat verandert in 2026?", "wat verandert in 2026?", "Hoeveel kost een paspoort in 2026?",
    "Moet ik aangifte doen over 2025?", "Is er een wachtlijst in 2026?", "Was ändert sich 2026?", "Gibt es Kindergeld 2026?",
    "How much is the fee in 2026?", "Bu konuda Mag ik mijn huis verhuren in 2026? sayfası yol gösterir.",
    "’Kan ik uitstel aanvragen voor 2025?’ başlığına bakın.",  # yanlis eslesen tirnak cifti: sozcukteki tirnak yok sayilir
])
def test_placeholders_ignore_short_question_titles(value):
    assert ec.is_placeholder(value) is False


@pytest.mark.parametrize("value", [  # V4 inceleme: Turkce cumlenin sonundaki belirsiz sayi uzunluguna bakilmadan yer tutucudur
    "Kayıt ücreti şu an 85?", "Kayıt ücreti bu yıl yaklaşık 85?", "Müzeye giriş ücreti yetişkinler için yaklaşık 15?",
    "Müzeye giriş ücreti hafta sonları yetişkinler ve öğrenciler için yaklaşık 15?", "Johannes van der Waalsweg 4?",
    "Kan testi ücreti yaklaşık 85?", "Hat 401?",  # Turkce 'kan'/'hat': ardindan ozne yok, yardimci fiil sayilmaz
    "WHO verilerine göre oran 2025?",  # kisaltma soru sozcugu sayilmaz
    "Aangifte doen over 2025?",  # soru sozcugu yok: resmi baslik tirnak icine alinmali
])
def test_placeholders_flag_uncertain_number_regardless_of_clause_length(value):
    assert ec.is_placeholder(value) is True


@pytest.mark.parametrize("value", [  # istisna gercek yer tutucuyu gizlemez
    "Eindhoven'a yakın: Markt 1? ve Belastingdienst'in ofisi",  # Turkce kesme isaretleri tirnak cifti sayilmaz
    "\"Kan ik uitstel aanvragen 2025?\" sayfası; adres Markt 1?",  # tirnak disindaki yer tutucu
    "Adres: Willem de Zwijgerlaan 3?", "Jan van Lieshoutstraat 3?",  # kisa adres parcasi soru basligi degil
    "Başlık: \"{{isim}}\"", "'TODO' yazılı alan",  # acik isaretler tirnak icinde de yer tutucudur
])
def test_placeholders_still_flag_real_markers_outside_quotes_and_short_clauses(value):
    sheets = {"Yanıtlar": [{"_row": 8, "entry_id": "x:c", "steps": value}]}
    assert ec.placeholders(sheets)["count"] == 1


def test_placeholders_only_scans_content_sheets_and_lines():
    sheets = {"Kaynaklar": [{"_row": 8, "source_id": "s", "title": "TODO"}],
              "Sözlük": [{"_row": 8, "term_id": "t", "definition": "Birinci satır.\nİkinci 12?"}]}
    result = ec.placeholders(sheets)
    assert result["count"] == 1 and result["items"][0]["sheet"] == "Sözlük"


# --- wikipedia_labeled_official ----------------------------------------------

def test_wikipedia_labeled_official_links_and_sources():
    sheets = {
        "Bağlantılar": [
            {"_row": 8, "action_id": "a1", "label": "Resmî işlem ve bilgi", "url": "https://nl.wikipedia.org/wiki/Eindhoven"},
            {"_row": 9, "action_id": "a2", "label": "Resmî işlem ve bilgi", "url": "https://www.eindhoven.nl/x"},
            {"_row": 10, "action_id": "a3", "label": "Ansiklopedi", "url": "https://en.wikipedia.org/wiki/X"},
        ],
        "Kaynaklar": _sources(("s1", "https://nl.wikipedia.org/wiki/Y", {"_row": 8, "publisher": "Resmi kaynak"}),
                              ("s2", "https://nl.wikipedia.org/wiki/Z", {"_row": 9, "publisher": "Wikipedia"}),
                              ("s3", "https://gemeente.nl/", {"_row": 10, "publisher": "Official site"})),
    }
    result = ec.wikipedia_labeled_official(sheets)
    assert result["count"] == 2
    assert result["by_sheet"] == {"Bağlantılar": 1, "Kaynaklar": 1}
    assert {i["id"] for i in result["items"]} == {"a1", "s1"}


def test_wikipedia_labeled_official_negative():
    sheets = {"Bağlantılar": [{"_row": 8, "action_id": "a", "label": "Resmî", "url": "notaurl"}]}
    assert ec.wikipedia_labeled_official(sheets)["count"] == 0
    assert ec.wikipedia_labeled_official({})["count"] == 0


# --- verified_without_snapshot ---------------------------------------------

def test_verified_without_snapshot():
    sheets = {"Kaynaklar": [
        {"_row": 8, "source_id": "s1", "verified": "Evet"},
        {"_row": 9, "source_id": "s2", "verified": "Evet", "snapshot_file": "Kaynak/s2.pdf"},
        {"_row": 10, "source_id": "s3", "verified": "Evet", "snapshot_file": "Kaynak/s3.pdf", "snapshot_sha256": "a" * 64},
        {"_row": 11, "source_id": "s4", "verified": "Hayır"},
        {"_row": 12, "source_id": "s5", "verified": True, "snapshot_sha256": "b" * 64},
    ]}
    result = ec.verified_without_snapshot(sheets)
    assert result["count"] == 3
    assert [i["id"] for i in result["items"]] == ["s1", "s2", "s5"]
    assert result["verified_total"] == 4


# --- shared_coordinates -------------------------------------------------------

def test_shared_coordinates_groups_distinct_entries():
    rows = [
        {"_row": 8, "entry_id": "x:klokgebouw", "latitude": 51.448, "longitude": 5.457},
        {"_row": 9, "entry_id": "x:centraal", "latitude": 51.448, "longitude": 5.457},
        {"_row": 10, "entry_id": "x:catharinaplein", "latitude": 51.448, "longitude": 5.457},
        {"_row": 11, "entry_id": "x:other", "latitude": 51.4, "longitude": 5.4},
        {"_row": 12, "entry_id": "x:nocoord"},
    ]
    result = ec.shared_coordinates({"Mekânlar": rows})
    assert result["count"] == 1
    assert result["entries"] == 3
    assert result["items"][0]["entry_ids"] == ["x:catharinaplein", "x:centraal", "x:klokgebouw"]


def test_shared_coordinates_same_entry_twice_is_not_shared():
    rows = [{"_row": 8, "entry_id": "x:a", "latitude": 1.0, "longitude": 2.0},
            {"_row": 9, "entry_id": "x:a", "latitude": 1.0, "longitude": 2.0}]
    assert ec.shared_coordinates({"Mekânlar": rows})["count"] == 0


# --- homepage_only_sources ----------------------------------------------------

def test_homepage_only_sources():
    sheets = {
        "Kaynaklar": _sources(("home", "https://www.eindhoven.nl/", {}), ("bare", "https://www.tue.nl", {}),
                              ("deep", "https://www.eindhoven.nl/stad/verhuizen", {}), ("query", "https://site.nl/?p=1", {})),
        "Yanıtlar": [
            {"_row": 8, "entry_id": "y:1", "source_ids": "home;bare"},
            {"_row": 9, "entry_id": "y:2", "source_ids": "home;deep"},
            {"_row": 10, "entry_id": "y:3", "source_ids": "deep"},
            {"_row": 11, "entry_id": "y:4"},
            {"_row": 12, "entry_id": "y:5", "source_ids": "missing"},
            {"_row": 13, "entry_id": "y:6", "source_ids": "query"},
        ],
    }
    result = ec.homepage_only_sources(sheets)
    assert result["count"] == 1
    assert result["items"][0]["id"] == "y:1"
    assert result["items"][0]["urls"] == ["https://www.eindhoven.nl/", "https://www.tue.nl"]


@pytest.mark.parametrize("url,expected", [
    ("https://www.tue.nl", True), ("https://www.tue.nl/", True), ("https://www.tue.nl/en", True), ("https://x.nl/nl/", True),
    ("https://x.nl/en-gb", True), ("https://x.nl/home.html", True), ("https://x.nl/index.htm", True),
    ("https://x.nl/inloggen", False), ("https://x.nl/en/study", False), ("https://x.nl/?p=1", False), ("not a url", False),
    ("http://[::1", False),
])
def test_is_homepage(url, expected):
    assert ec.is_homepage(url) is expected


# --- history_share --------------------------------------------------------------

def test_history_share_flags_above_threshold():
    rows = [{"entry_id": f"x:{i}", "group": "Tarih ve mimari" if i < 4 else "Müzeler"} for i in range(10)]
    result = ec.history_share({"Mekânlar": rows})
    assert result["history"] == 4 and result["total"] == 10
    assert result["ratio"] == pytest.approx(0.4)
    assert result["flagged"] is True and result["count"] == 1
    assert "öneri" in result["threshold_note"]


def test_history_share_below_threshold_and_empty():
    rows = [{"entry_id": f"x:{i}", "group": "Tarih ve mimari" if i < 3 else "Parklar"} for i in range(10)]
    assert ec.history_share({"Mekânlar": rows})["count"] == 0
    empty = ec.history_share({})
    assert empty["count"] == 0 and empty["ratio"] == 0.0


# --- run_all / labeling / purity -----------------------------------------------

def test_run_all_labels_every_check_and_does_not_mutate_input():
    sheets = {"Mekânlar": [{"_row": 8, "entry_id": "x:a", "summary": "5621?", "latitude": 1.0, "longitude": 2.0,
                            "group": "Tarih ve mimari"}],
              "Kaynaklar": [{"_row": 8, "source_id": "s", "url": "https://a.nl/", "verified": "Evet"}],
              "Yanıtlar": [{"_row": 8, "entry_id": "y", "source_ids": "s"}]}
    before = copy.deepcopy(sheets)
    results = ec.run_all(sheets)
    assert sheets == before
    assert set(results) == set(ec.CHECKS)
    assert all(r["label"] == ec.LABEL for r in results.values())
    assert ec.counts(results) == {"placeholders": 1, "wikipedia_labeled_official": 0, "verified_without_snapshot": 1,
                                  "shared_coordinates": 0, "homepage_only_sources": 1, "history_share": 1}


def test_load_sheets_reads_row7_keys_and_skips_missing_sheets(tmp_path):
    path = helpers.make_template(tmp_path / "t.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb["Kaynaklar"]
    keys = helpers.SHEET_KEYS["Kaynaklar"]
    ws.cell(8, keys.index("source_id") + 1, "s1")
    ws.cell(8, keys.index("verified") + 1, "Evet")
    wb.save(path)
    sheets = ec.load_sheets(path)
    assert sheets["Kaynaklar"] == [{"source_id": "s1", "verified": "Evet", "_row": 8}]
    assert sheets["Mekânlar"] == []  # sablonda yok -> bos
    assert ec.run_all(sheets)["verified_without_snapshot"]["count"] == 1
