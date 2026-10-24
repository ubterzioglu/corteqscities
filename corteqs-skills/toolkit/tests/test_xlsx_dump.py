"""xlsx_dump.py testleri: K10 ham/tekil sayimlari, JSONL dokumu, salt-okunur davranis."""
from __future__ import annotations

import datetime as dt
import json
import os
from pathlib import Path

import openpyxl
import pytest

import common
import xlsx_dump
from helpers import SHEET_KEYS


def _put(xlsx: Path, sheet: str, row: int, values: dict) -> None:
    wb = openpyxl.load_workbook(xlsx)
    ws = wb[sheet]
    keys = SHEET_KEYS[sheet]
    for key, value in values.items():
        ws.cell(row, keys.index(key) + 1, value)
    wb.save(xlsx)


def _run(capsys, argv: list[str]) -> tuple[int, dict]:
    code = xlsx_dump.main(argv)
    out = capsys.readouterr().out.strip().splitlines()
    assert len(out) == 1, "stdout tek satir olmali"
    return code, json.loads(out[0])


def _tree_state(root: Path) -> list[tuple[str, int, int]]:
    return sorted((str(p.relative_to(root)), p.stat().st_size, p.stat().st_mtime_ns) for p in root.rglob("*") if p.is_file())


def test_counts_empty_template_ignores_formula_only_rows(template, capsys):
    code, summary = _run(capsys, [str(template)])
    assert code == 0
    assert "Özet" not in summary["sheets"]
    assert set(summary["sheets"]) == set(SHEET_KEYS)
    konular = summary["sheets"]["Konular"]
    # satir 8..37 yalniz formul (used_count) tasir -> kayit sayilmaz
    assert konular["rows"] == 0 and konular["unique"] == 0 and konular["dup_ids"] == 0
    assert konular["formula_cols"] == ["used_count"]
    assert summary["sheets"]["Teslimat"]["rows"] == 1
    assert summary["total_rows"] == 2  # Teslimat + Sehir


def test_duplicate_and_blank_ids(template, capsys):
    _put(template, "Yanıtlar", 8, {"entry_id": "a", "title": "A1"})
    _put(template, "Yanıtlar", 9, {"entry_id": "a", "title": "A2"})
    _put(template, "Yanıtlar", 10, {"entry_id": "b", "title": "B"})
    _put(template, "Yanıtlar", 11, {"title": "kimliksiz"})
    code, summary = _run(capsys, [str(template)])
    assert code == 0
    yanitlar = summary["sheets"]["Yanıtlar"]
    assert (yanitlar["rows"], yanitlar["unique"], yanitlar["dup_ids"], yanitlar["blank_key"]) == (4, 2, 1, 1)
    assert summary["total_rows"] == 2 + 4


def test_package_photo_sheets_and_formula_cols(package, capsys):
    code, summary = _run(capsys, [str(package["xlsx"])])
    assert code == 0
    fotos = summary["sheets"]["Fotoğraflar"]
    assert fotos["rows"] == 2 and fotos["unique"] == 2 and fotos["dup_ids"] == 0
    assert fotos["formula_cols"] == ["unique_hero_count", "unique_venue_count"]
    assert summary["sheets"]["FotoğrafKaynakları"]["rows"] == 2


def test_jsonl_roundtrip_matches_iter_records(package, tmp_path, capsys):
    _put(package["xlsx"], "Yanıtlar", 8, {"entry_id": "a", "title": "Başlık ğüşiöç", "reviewed_on": dt.datetime(2026, 5, 4), "order": 3.0})
    out_dir = tmp_path / "dump"
    code, summary = _run(capsys, [str(package["xlsx"]), "--out-dir", str(out_dir)])
    assert code == 0
    assert json.loads((out_dir / "summary.json").read_text(encoding="utf-8")) == summary
    wb = openpyxl.load_workbook(package["xlsx"])
    for sheet in summary["sheets"]:
        expected = list(common.iter_records(wb[sheet]))
        text = (out_dir / f"{sheet}.jsonl").read_text(encoding="utf-8")
        got = [json.loads(line) for line in text.splitlines()]
        assert got == expected
    raw = (out_dir / "Yanıtlar.jsonl").read_text(encoding="utf-8")
    assert "Başlık ğüşiöç" in raw  # ensure_ascii=False
    row = json.loads(raw.splitlines()[0])
    assert row["_row"] == 8 and row["reviewed_on"] == "2026-05-04" and row["order"] == 3
    foto_rows = [json.loads(x) for x in (out_dir / "Fotoğraflar.jsonl").read_text(encoding="utf-8").splitlines()]
    assert [r["_row"] for r in foto_rows] == [8, 9]
    assert "unique_venue_count" not in foto_rows[0]  # formul sutunu yok


def test_summary_flag_skips_jsonl(template, tmp_path, capsys):
    out_dir = tmp_path / "dump"
    code, _ = _run(capsys, [str(template), "--out-dir", str(out_dir), "--summary"])
    assert code == 0
    assert (out_dir / "summary.json").is_file()
    assert not list(out_dir.glob("*.jsonl"))


def test_sheets_filter_and_unknown_sheet(template, capsys):
    code, summary = _run(capsys, [str(template), "--sheets", "Konular,Yanıtlar"])
    assert code == 0 and set(summary["sheets"]) == {"Konular", "Yanıtlar"}
    code, summary = _run(capsys, [str(template), "--sheets", "YokBoyleSayfa"])
    assert code == 2 and "error" in summary


def test_unreadable_workbook_exit_2(tmp_path, capsys):
    bad = tmp_path / "bozuk.xlsx"
    bad.write_bytes(b"bu bir xlsx degil")
    code, summary = _run(capsys, [str(bad)])
    assert code == 2 and "error" in summary
    code, summary = _run(capsys, [str(tmp_path / "yok.xlsx")])
    assert code == 2 and "error" in summary


def test_duplicate_technical_key_is_structure_error(template, capsys):
    wb = openpyxl.load_workbook(template)
    wb["Konular"].cell(7, 2, "topic_id")  # satir 7'de yinelenen anahtar
    wb.save(template)
    code, summary = _run(capsys, [str(template)])
    assert code == 2 and "topic_id" in summary["error"]


def test_missing_sheet_is_reported_not_fatal(tmp_path, capsys):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Konular"
    for col, key in enumerate(SHEET_KEYS["Konular"], start=1):
        ws.cell(7, col, key)
    ws.cell(8, 1, "k1")
    path = tmp_path / "kucuk.xlsx"
    wb.save(path)
    code, summary = _run(capsys, [str(path)])
    assert code == 0
    assert summary["sheets"]["Konular"]["rows"] == 1
    assert "Yanıtlar" in summary["missing_sheets"]


def test_deterministic_and_read_only(package, tmp_path, capsys):
    before = _tree_state(package["root"])
    out_dir = tmp_path / "dump"
    _, first = _run(capsys, [str(package["xlsx"]), "--out-dir", str(out_dir)])
    snap = {p.name: p.read_bytes() for p in out_dir.iterdir()}
    _, second = _run(capsys, [str(package["xlsx"]), "--out-dir", str(out_dir)])
    assert first == second
    assert snap == {p.name: p.read_bytes() for p in out_dir.iterdir()}
    assert _tree_state(package["root"]) == before


def test_runs_as_script(template, tmp_path):
    import subprocess
    import sys

    script = Path(xlsx_dump.__file__)
    env = {**os.environ, "PYTHONIOENCODING": "utf-8"}
    proc = subprocess.run([sys.executable, str(script), str(template)], capture_output=True, text=True, encoding="utf-8", env=env)
    assert proc.returncode == 0
    assert json.loads(proc.stdout)["total_rows"] == 2


@pytest.mark.parametrize("argv", [[], ["--bogus"]])
def test_usage_errors_exit_2(argv, capsys):
    with pytest.raises(SystemExit) as exc:
        xlsx_dump.main(argv)
    assert exc.value.code == 2
