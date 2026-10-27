"""preview_checks.py testleri (city-preview araci).

drive-compare (K01/K03/K109/K110): sentetik Drive listesi <-> mini paket envanteri (inventory.py). Iki kayit bicimi:
{path, size, md5Checksum} ve Drive baglayicisinin gercek ham kaydi (2026-10 Eindhoven denemesi: title, fileSize metin,
parentId, viewUrl; md5Checksum, headRevisionId ve path YOK). Eksik, fazla, boyut (yarim yukleme), boyutsuz kayit, sayfa
sinirinda yinelenen kayit, MD5, SHA-256 yalniz indirilen bayttan, ad+boyutun hash sayilmamasi, kanit ZIP'i (ekip akisi
G.2), klasor yapisi, kok adi, kok kimligi bagimsiz kaynakla, onceki kosunun kimlikleri, arac sonucu klasoru, kullanim
hatalari. warning-scan (K103/K104): Berlin'in ve staging sayfasinin gercek cumleleri, bant alanlarinda ve kapak
metinlerinde her emir, SayfaAlanlari; tarihli somut kapanis ve normal islem adimi yanlis pozitif degil, tek basina yil ya
da 'bilet' muaf degil. report: ilk satir, kanit, gerekce, karar ifadesi, kisisel veri, kapsam, K93/K112 zorunlu islevleri,
--drive yokken Drive maddeleri, Drive kanitiyla tutarlilik. Ag yok; paket salt okunur (agac once/sonra ayni).
"""
from __future__ import annotations

import base64
import hashlib
import json
import shutil
import unicodedata
from pathlib import Path

import openpyxl
import pytest

import inventory
import preview_checks as pc

FOLDER = "application/vnd.google-apps.folder"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FOLDERS = ("Fotoğraflar", "Lisans kanıtları", "Lisans kanıtları/orijinaller", "Teslimler")
EXCEL = "CITY-999_testville_v001.xlsx"
JPG1, JPG2 = "Lisans kanıtları/orijinaller/testville_yer_01.jpg", "Lisans kanıtları/orijinaller/testville_yer_02.jpg"
WEB1, WEB2 = "Fotoğraflar/testville_yer_01.webp", "Fotoğraflar/testville_yer_02.webp"
SCOPE = ("K01", "K03", "K92", "K93", "K94", "K103", "K109", "K110", "K111", "K112")
DRIVE_IDS = ("K01", "K03", "K109", "K110")
ROOT_OK = ("--root-id", "kok-1", "--root-source", "prompt")
ROOT_META = {  # get_file_metadata sonucu (Eindhoven kokunun alanlari, degerler sentetik)
    "canAddChildren": True, "createdTime": "2026-10-12T14:31:48.977Z", "id": "kok-1", "inheritedPermissionsDisabled": False,
    "mimeType": FOLDER, "modifiedTime": "2026-10-13T15:03:09.832Z", "owner": "sahip@example.org", "parentId": "uretilecek",
    "readOnly": False, "resourceUri": "googledrive://file/kok-1", "starred": False,
    "title": "CITY-999 · Testville · Testland · testville", "viewUrl": "https://drive.google.com/drive/folders/kok-1"}


def _run(capsys, *argv) -> tuple[int, dict]:
    code = pc.main([str(a) for a in argv])
    lines = capsys.readouterr().out.strip().splitlines()
    assert len(lines) == 1, f"stdout tek satir olmali: {lines}"
    return code, json.loads(lines[0])


def _write(path: Path, payload) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return path


def _read(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _tree(root: Path) -> list[tuple[str, int, int]]:
    return sorted((str(p.relative_to(root)), p.stat().st_size, p.stat().st_mtime_ns) for p in root.rglob("*") if p.is_file())


def _md5(path: Path) -> str:
    return hashlib.md5(path.read_bytes()).hexdigest()


# --- drive-compare ---

@pytest.fixture()
def drive(package, tmp_path, capsys):
    return _drive_setup(package, tmp_path, capsys)


def _drive_setup(package, tmp_path, capsys) -> dict:
    """Mini paket + inventory.py envanteri + paketle birebir ayni {path, md5Checksum} listesi (OUT paket disinda)."""
    inv = tmp_path / "is" / "dosya-envanteri.json"
    assert inventory.main([str(package["root"]), "--xlsx", str(package["xlsx"]), "--out", str(inv)]) == 0
    capsys.readouterr()
    entries = [{"path": p, "id": f"klasor-{i}", "mimeType": FOLDER} for i, p in enumerate(FOLDERS)]
    for i, item in enumerate(_read(inv)["files"]):
        entries.append({"path": item["path"], "id": f"dosya-{i}", "size": str(item["size"]),
                        "md5Checksum": _md5(package["root"] / item["path"]), "mimeType": "application/octet-stream",
                        "modifiedTime": "2026-10-23T10:00:00.000Z"})
    root = {"id": "kok-1", "name": "CITY-999 · Testville", "url": "https://drive.google.com/drive/folders/kok-1"}
    return {"pkg": package, "inv": inv, "listing": {"root": root, "files": entries}, "out": tmp_path / "OUT"}


def _compare(capsys, drive, listing=None, *extra) -> tuple[int, dict, dict]:
    remote = _write(drive["out"] / "drive_listesi.json", drive["listing"] if listing is None else listing)
    out = drive["out"] / "drive_karsilastirma.json"
    out.unlink(missing_ok=True)
    code, summary = _run(capsys, "drive-compare", "--local", drive["inv"], "--remote", remote,
                         "--pkg", drive["pkg"]["root"], "--out", out, *extra)
    return code, summary, (_read(out) if out.exists() else {})


def _files(drive, change=None, drop=(), add=()) -> dict:
    """Listenin yeni kopyasi: drop'taki yollar cikar, change {yol: {alan: deger}} uygulanir, add eklenir."""
    change = change or {}
    files = [{**e, **change.get(e["path"], {})} for e in drive["listing"]["files"] if e["path"] not in drop]
    return {**drive["listing"], "files": files + list(add)}


def _connector(drive, change=None, drop=(), add=()) -> dict:
    """Baglayicinin gercek bicimi: root = get_file_metadata, files = search_files kayitlari (yol ve md5 yok, yol
    parentId zincirinden). change {yol: {alan: deger}}; deger None ise alan cikar (ornek: fileSize'siz kayit)."""
    change, ids, records = change or {}, {"": ROOT_META["id"]}, []
    for i, folder in enumerate(FOLDERS):
        parent, _, name = folder.rpartition("/")
        ids[folder] = f"k{i}"
        records.append({"id": f"k{i}", "title": name, "mimeType": FOLDER, "parentId": ids[parent], "_yol": folder})
    for i, item in enumerate(_read(drive["inv"])["files"]):
        parent, _, name = item["path"].rpartition("/")
        records.append({"id": f"d{i}", "title": name, "fileSize": str(item["size"]), "parentId": ids[parent],
                        "mimeType": "application/octet-stream", "createdTime": "2026-10-12T16:32:43.541Z",
                        "modifiedTime": "2026-10-21T06:34:35.604Z", "viewUrl": f"https://drive.google.com/file/d/d{i}/view",
                        "_yol": item["path"]})
    files = [{k: v for k, v in {**r, **change.get(r["_yol"], {})}.items() if v is not None and k != "_yol"}
             for r in records if r["_yol"] not in drop]
    return {"root": ROOT_META, "files": files + list(add)}


def test_drive_exact_match_is_clean(drive, capsys):
    code, s, d = _compare(capsys, drive)
    assert code == 0, s
    assert (s["local_files"], s["remote_files"], s["eslesen"], s["md5_eslesti"], s["icerik_dogrulanan"]) == (5, 5, 5, 5, 5)
    assert s["findings"] == 0 and s["yalniz_ad_boyut"] == 0 and d["yapi"] == [] and d["checked_at"]
    assert d["id_map"]["kok"]["id"] == "kok-1" and d["id_map"]["excel"]["path"] == EXCEL
    assert set(d["id_map"]["klasorler"]) == {"Fotoğraflar", "Lisans kanıtları", "Teslimler"}
    assert {r["result"] for r in d["files"]} == {"md5_eslesti"} and {r["method"] for r in d["files"]} == {"md5"}
    assert d["excel_yerel"]["sha256"] == hashlib.sha256(drive["pkg"]["xlsx"].read_bytes()).hexdigest()
    assert d["sha256_beklenen"] == [EXCEL]  # ana Excel icin indirme yapilmadi: SHA-256 eksigi ayri listede


def test_drive_reads_raw_connector_records_and_builds_paths(drive, capsys):
    code, s, d = _compare(capsys, drive, _connector(drive))
    by_path = {r["path"]: r for r in d["files"]}
    assert code == 0 and (s["eslesen"], s["eksik"], s["fazla"], s["findings"]) == (5, 0, 0, 0), s
    assert set(by_path) == {EXCEL, JPG1, JPG2, WEB1, WEB2} and by_path[JPG1]["remote_size"] == by_path[JPG1]["local_size"]
    assert (s["icerik_dogrulanan"], s["yalniz_ad_boyut"], s["md5_eslesti"]) == (0, 5, 0)  # baglayici md5 vermiyor
    assert d["id_map"]["kok"]["name"] == ROOT_META["title"] and d["id_map"]["excel"]["url"].startswith("https://")
    assert d["id_map"]["klasorler"] == {"Fotoğraflar": ["k0"], "Lisans kanıtları": ["k1"], "Teslimler": ["k3"]}
    assert d["id_map"]["excel"]["createdTime"] == "2026-10-12T16:32:43.541Z" and "owner" not in json.dumps(d)


def test_drive_partial_upload_is_a_size_finding_with_connector_names(drive, capsys):
    code, s, d = _compare(capsys, drive, _connector(drive, change={WEB1: {"fileSize": "10"}}))  # yarim kalmis yukleme
    assert code == 1 and (s["boyut_farki"], s["findings"]) == (1, 1) and d["farkli"] == [WEB1], s


def test_drive_sizeless_record_is_name_only_and_counted(drive, capsys):
    code, s, d = _compare(capsys, drive, _connector(drive, change={JPG1: {"fileSize": None}}))
    row = next(r for r in d["files"] if r["path"] == JPG1)
    assert code == 0 and (row["method"], row["result"], s["yalniz_ad"]) == ("ad", "yalniz_ad", 1), s
    assert JPG1 in d["dogrulanmayan"] and "ID/ad/boyut" in d["method"]["ad"]


def test_drive_page_boundary_repeat_is_one_file(drive, capsys):
    listing = _connector(drive)
    repeat = next(r for r in listing["files"] if r["title"] == "testville_yer_02.webp")
    code, s, d = _compare(capsys, drive, {**listing, "files": listing["files"] + [dict(repeat)]})
    assert code == 0 and (s["yinelenen_dosya"], s["yinelenen_kayit"], s["remote_files"]) == (0, 1, 5), s


def test_drive_record_outside_the_listed_root_is_a_usage_error(drive, capsys):
    stray = {"id": "x1", "title": "yabanci.jpg", "fileSize": "3", "parentId": "liste-disi", "mimeType": "image/jpeg"}
    code, s, d = _compare(capsys, drive, _connector(drive, add=[stray]))
    assert code == 2 and "koke baglanmayan" in s["error"] and d == {}, s


def test_drive_listing_may_be_a_folder_of_raw_tool_results(drive, capsys):
    files = _connector(drive)["files"]
    pages = drive["out"] / "_drive"
    _write(pages / "00_kok.json", ROOT_META)
    _write(pages / "01_kok_sayfa1.json", {"files": files[:4], "nextPageToken": "~!!~sayfa2"})  # arac sonucu aynen
    wrapped = [{"type": "text", "text": json.dumps({"files": files[4:]}, ensure_ascii=False)}]
    _write(pages / "02_kayitli_sonuc.txt", wrapped)  # harness'in dosyaya kaydettigi buyuk sonuc
    (pages / "03_son_sayfa.json").write_text("{}", encoding="utf-8")  # bos yanit: sayfalama bitti
    out = drive["out"] / "k.json"
    code, s = _run(capsys, "drive-compare", "--local", drive["inv"], "--remote", pages, "--pkg", drive["pkg"]["root"],
                   "--out", out)
    assert code == 0 and (s["eslesen"], s["yalniz_ad_boyut"], s["remote_folders"]) == (5, 5, 4), s
    assert _read(out)["id_map"]["kok"]["id"] == "kok-1"
    _write(pages / "04_baska_kok.json", {**ROOT_META, "id": "kok-2"})
    code, s = _run(capsys, "drive-compare", "--local", drive["inv"], "--remote", pages, "--out", out)
    assert code == 2 and "kok kaydi" in s["error"]


def test_drive_missing_extra_size_and_md5_are_findings(drive, capsys):
    extra = {"path": "Teslimler/eski_surum.txt", "id": "fazla-1", "size": "3", "mimeType": "text/plain"}
    listing = _files(drive, drop=(WEB2,), add=[extra],
                     change={WEB1: {"size": "1"}, JPG2: {"md5Checksum": "0" * 32}})
    code, s, d = _compare(capsys, drive, listing)
    assert code == 1
    assert (s["eksik"], s["fazla"], s["boyut_farki"], s["md5_farki"], s["md5_eslesti"]) == (1, 1, 1, 1, 2)
    assert s["findings"] == 4 and s["eslesen"] == 4
    assert d["eksik"] == [WEB2] and d["fazla"] == ["Teslimler/eski_surum.txt"] and sorted(d["farkli"]) == sorted([WEB1, JPG2])
    by_path = {r["path"]: r for r in d["files"]}
    assert by_path[JPG2]["result"] == "md5_farki" and by_path[JPG2]["remote_md5"] == "0" * 32
    assert by_path[WEB1]["result"] == "boyut_farki" and by_path[WEB1]["remote_size"] == 1


def test_drive_name_and_size_only_is_never_called_hash_verified(drive, capsys):
    listing = {**drive["listing"], "files": [{k: v for k, v in e.items() if k != "md5Checksum"}
                                             for e in drive["listing"]["files"]]}
    code, s, d = _compare(capsys, drive, listing)
    assert code == 0 and s["findings"] == 0  # eksik/fazla yok ...
    assert (s["icerik_dogrulanan"], s["yalniz_ad_boyut"]) == (0, 5)  # ... ama icerik dogrulanmadi
    assert sorted(d["dogrulanmayan"]) == sorted([EXCEL, JPG1, JPG2, WEB1, WEB2])
    assert {r["method"] for r in d["files"]} == {"ad+boyut"} and "DEĞİL" in d["method"]["ad+boyut"]


def test_drive_sha256_comes_only_from_downloaded_bytes(drive, capsys, tmp_path):
    root, dl = drive["pkg"]["root"], tmp_path / "indirme"
    dl.mkdir()
    shutil.copyfile(root / EXCEL, dl / "excel.xlsx")  # ham indirme (kullanici tarayicidan indirdi)
    payload = base64.b64encode((root / JPG1).read_bytes()).decode()
    tool_result = [{"type": "text", "text": json.dumps({"name": "a.jpg", "content": payload})}]  # harness'in kaydettigi sonuc
    (dl / "sonuc.txt").write_text(json.dumps(tool_result), encoding="utf-8")
    (dl / "kesik.txt").write_text(base64.b64encode((root / WEB1).read_bytes()[:-10]).decode(), encoding="utf-8")
    (dl / "bozuk.webp").write_bytes((root / WEB2).read_bytes()[:-1] + b"X")  # ayni boyut, farkli bayt
    mapping = [{"path": EXCEL, "file": str(dl / "excel.xlsx"), "format": "raw"},
               {"path": JPG1, "file": str(dl / "sonuc.txt"), "format": "base64"},
               {"path": WEB1, "file": str(dl / "kesik.txt"), "format": "base64"},
               {"path": WEB2, "file": str(dl / "bozuk.webp"), "format": "raw"}]
    downloads = _write(drive["out"] / "indirilenler.json", mapping)
    code, s, d = _compare(capsys, drive, None, "--downloads", downloads)
    by_path = {r["path"]: r["result"] for r in d["files"]}
    assert code == 1 and (by_path[EXCEL], by_path[JPG1]) == ("sha256_dogrulandi", "sha256_dogrulandi")
    assert by_path[WEB1] == "md5_eslesti"  # kesik indirme sahte 'farkli' vermez: MD5'e duser, sorun ayri yazilir
    assert by_path[WEB2] == "sha256_farki"
    assert (s["sha256_dogrulandi"], s["sha256_farki"], s["indirme_sorunu"], s["icerik_dogrulanan"]) == (2, 1, 1, 4)
    assert d["sha256_beklenen"] == [] and [p["path"] for p in d["indirme_sorunlari"]] == [WEB1]


def test_drive_evidence_zip_is_neither_missing_nor_verified(drive, capsys):
    bundle = {"path": "Lisans kanıtları/kanitlar.zip", "id": "z1", "size": "999", "mimeType": "application/zip"}
    listing = _files(drive, drop=(JPG1, JPG2, "Lisans kanıtları/orijinaller"), add=[bundle])  # ekip akisi G.2
    code, s, d = _compare(capsys, drive, listing)
    by_path = {r["path"]: r["result"] for r in d["files"]}
    assert code == 0 and [by_path[p] for p in (JPG1, JPG2, bundle["path"])] == ["zipte_olabilir"] * 2 + ["kanit_zip"]
    assert (s["eksik"], s["fazla"], s["zipte_olabilir"], s["kanit_zip"]) == (0, 0, 2, 1) and JPG1 in d["dogrulanmayan"]
    code, s, d = _compare(capsys, drive, _files(drive, drop=(WEB1,), add=[bundle]))  # ZIP Fotoğraflar'i kapsamaz
    assert code == 1 and d["eksik"] == [WEB1] and s["kanit_zip"] == 1


def test_drive_structure_rules(drive, capsys):
    extra = [{"path": "Codex teslim v2", "id": "k-x", "mimeType": FOLDER},
             {"path": "Fotoğraflar", "id": "k-ikinci", "mimeType": FOLDER},
             {"path": "Lisans kanitlari", "id": "k-benzer", "mimeType": FOLDER},
             {"path": "CITY-999_testville_v002.xlsx", "id": "x-2", "size": "10", "mimeType": XLSX_MIME}]
    listing = _files(drive, drop=("Teslimler",), add=extra)
    listing = {**listing, "root": {**listing["root"], "id": "baska-kok"}}
    code, s, d = _compare(capsys, drive, listing, *ROOT_OK)
    kinds = {p["kind"] for p in d["yapi"]}
    assert code == 1 and s["yapi"] == len(d["yapi"])
    assert {"fazla_ust_klasor", "yinelenen_klasor", "benzer_adli_klasor", "birden_fazla_excel", "eksik_klasor",
            "kok_kimligi_farkli"} <= kinds, kinds
    duplicate = next(p for p in d["yapi"] if p["kind"] == "yinelenen_klasor")
    assert duplicate["yol"] == "Fotoğraflar" and sorted(duplicate["id"]) == ["k-ikinci", "klasor-0"]
    assert next(p for p in d["yapi"] if p["kind"] == "eksik_klasor")["ad"] == "Teslimler"


def test_drive_root_name_must_carry_city_code_and_slug(drive, capsys):
    for name in ("CITY-998 · Başkaville", "CITY-9990 · Testville"):  # baska kod; CITY-9990 CITY-999 degildir
        listing = {**drive["listing"], "root": {**drive["listing"]["root"], "name": name}}
        code, s, d = _compare(capsys, drive, listing)
        problem = next((p for p in d["yapi"] if p["kind"] == "kok_adi_uyumsuz"), None)
        assert code == 1 and problem and problem["beklenen"] == ["CITY-999", "testville"], (name, d["yapi"])


def test_drive_root_identity_needs_an_independent_source(drive, capsys):
    code, s, d = _compare(capsys, drive)
    assert code == 0 and d["kok_kimligi"]["durum"] == "beklenen_yok"  # bulgu degil; K109 GEÇTİ olamaz (report)
    code, s, d = _compare(capsys, drive, None, "--root-id", "kok-1")
    assert code == 2 and "--root-source" in s["error"] and d == {}
    code, s, d = _compare(capsys, drive, None, *ROOT_OK)
    assert code == 0 and d["kok_kimligi"] == {"durum": "eslesti", "beklenen": "kok-1", "kaynak": "prompt", "bulunan": "kok-1"}
    code, s, d = _compare(capsys, drive, None, "--root-id", "baska-kok", "--root-source", "yonetim")
    assert code == 1 and d["kok_kimligi"]["durum"] == "farkli" and "kok_kimligi_farkli" in {p["kind"] for p in d["yapi"]}


def test_drive_ids_are_compared_with_the_previous_run(drive, capsys, tmp_path):
    _compare(capsys, drive)
    previous = tmp_path / "onceki_ONIZLEME" / "drive_karsilastirma.json"
    previous.parent.mkdir()
    shutil.copyfile(drive["out"] / "drive_karsilastirma.json", previous)
    code, s, d = _compare(capsys, drive, None, "--previous", previous)
    assert code == 0 and d["onceki_kimlik"]["degisen"] == 0, d["onceki_kimlik"]
    old = next(e["id"] for e in drive["listing"]["files"] if e["path"] == EXCEL)
    code, s, d = _compare(capsys, drive, _files(drive, change={EXCEL: {"id": "yeni-excel"}}), "--previous", previous)
    changed = [p for p in d["yapi"] if p["kind"] == "kimlik_degisti"]  # yeni surum yerine yeni dosya yuklenmis (G.3)
    assert code == 1 and changed == [{"kind": "kimlik_degisti", "oge": "excel", "onceki": old, "simdi": "yeni-excel"}]


def test_drive_duplicate_upload_and_google_sheet(drive, capsys):
    excel = next(e for e in drive["listing"]["files"] if e["path"] == EXCEL)
    sheet = {"path": "CITY-999_testville_v001", "id": "g-1", "mimeType": "application/vnd.google-apps.spreadsheet"}
    code, s, d = _compare(capsys, drive, _files(drive, add=[{**excel, "id": "x-ikinci"}, sheet]))
    by_path = {r["path"]: r for r in d["files"]}
    assert code == 1 and (s["yinelenen_dosya"], s["google_bicimi"], s["fazla"]) == (1, 1, 0)
    assert by_path[EXCEL]["result"] == "yinelenen_dosya" and sorted(by_path[EXCEL]["remote_id"]) == [excel["id"], "x-ikinci"]


def test_drive_paths_are_unicode_and_slash_normalized(drive, capsys):
    change = {path: {"path": unicodedata.normalize("NFD", path)} for path in (WEB1, WEB2, JPG1)}
    change[JPG2] = {"path": JPG2.replace("/", "\\")}
    code, s, _ = _compare(capsys, drive, _files(drive, change=change))
    assert code == 0 and (s["eslesen"], s["eksik"], s["fazla"], s["md5_eslesti"]) == (5, 0, 0, 5)


def test_drive_reads_nfd_named_local_files_from_disk(package, tmp_path, capsys):
    nfd = unicodedata.normalize("NFD", "Teslimler/Özet notu.txt")  # macOS'tan kopyalanmis ad: diskte NFD
    (package["root"] / nfd).write_text("not", encoding="utf-8")
    drive = _drive_setup(package, tmp_path, capsys)
    listing = _files(drive, change={nfd: {"path": unicodedata.normalize("NFC", nfd)}})  # Drive NFC dondurur
    code, s, d = _compare(capsys, drive, listing)
    assert code == 0 and (s["eslesen"], s["md5_eslesti"], s["eksik"]) == (6, 6, 0), s


def test_drive_usage_errors_exit_2(drive, capsys):
    out, pkg = drive["out"] / "k.json", drive["pkg"]["root"]
    base = ["drive-compare", "--local", drive["inv"], "--pkg", pkg]
    good = _write(drive["out"] / "r.json", drive["listing"])
    pathless = _write(drive["out"] / "np.json", {**drive["listing"], "files": [{"id": "x", "size": "1"}]})
    bad_download = _write(drive["out"] / "dl.json", [{"path": EXCEL, "file": str(drive["inv"]), "format": "zip"}])
    cases = {"uzak_liste_yok": [*base, "--remote", drive["out"] / "yok.json", "--out", out],
             "yolsuz_kayit": [*base, "--remote", pathless, "--out", out],
             "paket_icine_cikti": [*base, "--remote", good, "--out", pkg / "k.json"],
             "gecersiz_indirme": [*base, "--remote", good, "--downloads", bad_download, "--out", out],
             "onceki_gecersiz": [*base, "--remote", good, "--previous", drive["inv"], "--out", out]}
    for name, argv in cases.items():
        code, summary = _run(capsys, *argv)
        assert code == 2 and summary.get("error"), (name, summary)
    assert not (pkg / "k.json").exists() and not out.exists()
    (pkg / WEB1).write_bytes(b"degisti")  # envanter alindiktan sonra degisen paket
    code, summary = _run(capsys, *base, "--remote", good, "--out", out)
    assert code == 2 and "envanter" in summary["error"]


# --- warning-scan ---

SCAN_SHEETS = {
    "Yanıtlar": ["entry_id", "title", "subtitle", "summary", "steps", "note"],
    "Mekânlar": ["entry_id", "title", "summary", "visit_status", "note", "overlay", "status_note"],
    "Fotoğraflar": ["photo_id", "entry_id", "caption"],
    "EkDetaylar": ["detail_id", "entry_id", "key", "label", "value_type", "value"],
    "SayfaAlanları": ["field_id", "key", "label", "value_type", "value"],
}
POSITIVE = {  # (kayit, alan) -> (kontrol, en az bu kaliplar)
    "Yanıtlar": [
        {"entry_id": "t:y1", "summary": "Başvuru ücreti her yıl değişir. Bilgileri kurumdan doğrulayın."},
        {"entry_id": "t:y2", "note": "Tekrar kontrol edin."},
        {"entry_id": "t:y3", "summary": "Gitmeden açılış saatlerine bak."},
        {"entry_id": "t:y4", "note": "Ücretleri resmî siteden doğrula."},
        {"entry_id": "t:y5", "note": "Lütfen kontrol edin."},
        {"entry_id": "t:y6", "steps": "1. Formu doldur. 2. TODO: kaynak eklenecek."},
        {"entry_id": "t:y7", "note": "Saatleri ziyaret öncesinde kontrol et."},  # Berlin orneginde 84 kartta tekrar
        {"entry_id": "t:y8", "note": "Güncel saatler için resmî siteyi kontrol edin."},
        {"entry_id": "t:y9", "note": "Bilet fiyatları için resmî siteye göz atın."},
        {"entry_id": "t:y10", "summary": "Ziyaret etmeden önce güncel durumu teyit edin."},
        {"entry_id": "t:y11", "note": "2026 sezonu biletleri satışta; gitmeden sitesinden kontrol edin."},  # yil+bilet muaf degil
        {"entry_id": "t:y12", "summary": "Bağlantılı kurumda güncel koşulları ve kendi durumuna uygun işlem yolunu "
                                         "kontrol et."},  # Berlin sayfasinda 168 kez
        {"entry_id": "t:y13", "steps": "Ulaşım, asansör ve koltuk koşullarını kurumla doğrula."},  # Berlin
    ],
    "Mekânlar": [
        {"entry_id": "t:m1", "overlay": "Kapanmış olabilir, ziyaretten önce teyit edin"},
        {"entry_id": "t:m2", "status_note": "Bilgi güncel olmayabilir."},
        {"entry_id": "t:m3", "visit_status": "Ziyaret durumunu kontrol et"},  # Berlin Mekânlar.visit_status
        {"entry_id": "t:m4", "visit_status": "Geçici sergi yerini kontrol et"},  # Berlin Mekânlar.visit_status
        {"entry_id": "t:m5", "note": "Güncel kurum duyurusunu kontrol et."},  # staging Berlin sayfasi
        {"entry_id": "t:m6", "status_note": "Tadilat nedeniyle 1 Kasım 2026'dan 15 Mart 2027'ye kadar kapalı; yeniden "
                                            "açılış tarihini müzenin sitesinden kontrol edin."},  # bantta muaflik yok
    ],
    "Fotoğraflar": [{"photo_id": "t:foto:1", "entry_id": "t:m1", "caption": "Arşiv fotoğrafı olabilir; tekrar kontrol et"}],
    "EkDetaylar": [{"detail_id": "t:d1", "entry_id": "t:m1", "key": "venue_editorial",
                    "value": "Saatleri müzenin sitesinden kontrol edin."}],
    "SayfaAlanları": [{"field_id": "t:f1", "key": "hero_note", "value_type": "text",
                       "value": "Ziyaret öncesi güncel saatleri kontrol edin."},  # kapak fotografinin ustundeki yazi
                      {"field_id": "t:f2", "key": "explore_intro", "value_type": "text",
                       "value": "Müzeler için gitmeden saatleri kontrol edin."}],
}
EXPECTED_HITS = {
    ("t:y1", "summary"): ("K104", {"kurumdan_dogrula"}), ("t:y2", "note"): ("K104", {"tekrar_kontrol"}),
    ("t:y3", "summary"): ("K104", {"gitmeden_kontrol"}), ("t:y4", "note"): ("K104", {"kaynaktan_dogrula"}),
    ("t:y5", "note"): ("K104", {"yalin_uyari"}), ("t:y6", "steps"): ("K104", {"editor_notu"}),
    ("t:y7", "note"): ("K104", {"gitmeden_kontrol"}), ("t:y8", "note"): ("K104", {"kaynaktan_dogrula", "guncel_kontrol"}),
    ("t:y9", "note"): ("K104", {"kaynaktan_dogrula"}), ("t:y10", "summary"): ("K104", {"gitmeden_kontrol"}),
    ("t:y11", "note"): ("K104", {"gitmeden_kontrol"}), ("t:y12", "summary"): ("K104", {"guncel_kontrol"}),
    ("t:y13", "steps"): ("K104", {"kurumdan_dogrula"}),
    ("t:m1", "overlay"): ("K103", {"belirsiz_durum", "gitmeden_kontrol", "bant_emri"}),
    ("t:m2", "status_note"): ("K103", {"belirsiz_durum"}), ("t:m3", "visit_status"): ("K103", {"bant_emri"}),
    ("t:m4", "visit_status"): ("K103", {"bant_emri"}), ("t:m5", "note"): ("K104", {"guncel_kontrol"}),
    ("t:m6", "status_note"): ("K103", {"bant_emri"}), ("t:foto:1", "caption"): ("K103", {"arsiv_olabilir", "tekrar_kontrol"}),
    ("t:d1", "venue_editorial"): ("K103", {"siteden_bilgi_dogrula", "bant_emri"}),
    ("t:f1", "hero_note"): ("K103", {"bant_emri", "gitmeden_kontrol"}),
    ("t:f2", "explore_intro"): ("K104", {"gitmeden_kontrol"}),
}
NEGATIVE = {
    "Yanıtlar": [
        {"entry_id": "t:n1", "summary": "Lodewijkstraat yenileme nedeniyle kapalı (kapanış tarihi 30-11-2025); Acht'a git."},
        {"entry_id": "t:n2", "summary": "Sözleşmede kira bedelini, depozitoyu ve giriş tarihini kontrol edin."},
        {"entry_id": "t:n3", "note": "Kimliğini DigiD ile doğrula. AB hukuku doğrulaması 6 ay sürer."},
        {"entry_id": "t:n4", "summary": "Mağazalar 6:00-22:00 açık olabilir; başvuru kurum tarafından kontrol edilir."},
        {"entry_id": "t:n5", "note": "Kesin yılı mektubundaki tarihten teyit et."},
        {"entry_id": "t:n6", "steps": "Gitmeden önce online bilet al. Ziyaret öncesi rezervasyon yaptır."},
        {"entry_id": "t:n7", "subtitle": "Tazminat hakkı ve son maaşta kontrol edilecekler"},
        {"entry_id": "t:n10", "note": "Bisikletin belediye tarafından yanlış park nedeniyle kaldırılmış olabilir: bu "
                                      "bisikletler fotoğraf ve açıklamayla verlorenofgevonden.nl'e konuyor."},  # Eindhoven
        {"entry_id": "t:n11", "summary": "Sözleşmeyi imzalamadan önce kira bedelini kontrol edin."},
        {"entry_id": "t:n12", "note": "Bu bilgi resmî kaynaktan doğrulanmıştır."},
        {"entry_id": "t:n13", "steps": "Kalan tahmini tüketimi fiyat farkıyla çarparak hesabı kontrol edin."},  # Eindhoven
        {"entry_id": "t:n14", "note": "Bildirim gelince wooniezie.nl'e giriş yapıp kaydının güncel ve yenilenmiş "
                                      "olduğunu kontrol et."},  # Eindhoven: kendi kaydini denetlemek is adimidir
        {"entry_id": "t:n15", "steps": "9292 ya da NS yolculuk planlayıcısında güncel toplu taşıma bilgisine bak."},
        {"entry_id": "t:n16", "steps": "Satın aldığın kurumun hesabından şartları kontrol et; iptal onayını sakla."},
    ],
    "Mekânlar": [
        {"entry_id": "t:n9", "note": "Tadilat nedeniyle 1 Kasım 2026'dan 15 Mart 2027'ye kadar kapalı; biletler üç "
                                     "hafta önceden satışa çıkar."},  # kapanis notta: bant degil
    ],
    "Fotoğraflar": [
        {"photo_id": "t:foto:n", "caption": "1931'de açılan istasyonun arşiv kaydı yok; cephe 2019'da yenilendi."},
        {"photo_id": "t:foto:n8", "caption": "Tadilat nedeniyle 1 Kasım 2026'dan 15 Mart 2027'ye kadar kapalı; yeniden "
                                             "açılış tarihini müzenin sitesinden kontrol edin."},  # tarihli somut: muaf
    ],
    "SayfaAlanları": [{"field_id": "t:f9", "key": "photo_credit", "value_type": "text",
                       "value": "Kapak: Ana Yazar / CC BY-SA 4.0 / Wikimedia Commons"},
                      {"field_id": "t:f8", "key": "_db_original_guide_content", "value_type": "text",
                       "value": "Arşiv: tekrar kontrol edin."}],  # Berlin ham arsiv alani: kapsam disi
}


def _scan_xlsx(path: Path, rows: dict) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, keys in SCAN_SHEETS.items():
        ws = wb.create_sheet(sheet)
        for col, key in enumerate(keys, start=1):
            ws.cell(6, col, key.upper())
            ws.cell(7, col, key)
        for row, record in enumerate(rows.get(sheet, []), start=8):
            for key, value in record.items():
                ws.cell(row, keys.index(key) + 1, value)
    wb.save(path)
    return path


def _scan(capsys, tmp_path, rows) -> tuple[int, dict, dict]:
    out = tmp_path / "OUT" / "uyari-tarama.json"
    code, summary = _run(capsys, "warning-scan", "--xlsx", _scan_xlsx(tmp_path / "t.xlsx", rows), "--out", out)
    return code, summary, (_read(out) if out.exists() else {})


def test_warning_scan_flags_real_band_texts_and_reader_audit_work(tmp_path, capsys):
    code, s, d = _scan(capsys, tmp_path, POSITIVE)
    found = {(h["id"], h["field"]): (h["check"], set(h["patterns"])) for h in d["hits"]}
    assert code == 1 and set(found) == set(EXPECTED_HITS), set(found) ^ set(EXPECTED_HITS)
    wrong = {key: (found[key], want) for key, want in EXPECTED_HITS.items()
             if found[key][0] != want[0] or not want[1] <= found[key][1]}
    assert wrong == {}, wrong
    assert s["by_check"] == {"K103": 8, "K104": 15, "K105": 0} and s["hits"] == 23 and "t:y2" in d["samples"]["tekrar_kontrol"]
    card = [h for h in d["hits"] if h["field"] in pc.CARD_BAND_FIELDS]
    assert card and all("kart_bandi" in h["patterns"] for h in card)
    assert next(h for h in d["hits"] if h["id"] == "t:foto:1")["entry_id"] == "t:m1"  # kare + bagli kart
    assert {b["field"] for b in d["band_values"]} == {"overlay", "status_note", "venue_editorial", "visit_status", "hero_note"}
    assert s["exempt"] == 0  # bant alaninda tarihli kapanis cumlesi de muaf degil (t:m6)


def test_warning_scan_keeps_concrete_closures_bookings_and_normal_steps(tmp_path, capsys):
    code, s, d = _scan(capsys, tmp_path, NEGATIVE)
    assert code == 0 and s["hits"] == 0 and d["hits"] == [], d["hits"]
    assert s["exempt"] == 1 and d["exempt"][0]["id"] == "t:foto:n8"  # tarihli kapanis cumlesindeki 'sitesinden kontrol'
    assert d["exempt"][0]["patterns"] == ["siteden_bilgi_dogrula"]
    assert {(b["id"], b["field"]) for b in d["band_values"]} == {("t:f9", "photo_credit")}


def test_warning_scan_flags_any_text_in_card_band_fields(tmp_path, capsys):
    """Sevilla dis denetimi S06: dolu visit_status sablonda fotograf ustune bant olur; tarihli, somut ya da 'Açık' fark etmez."""
    rows = {"Mekânlar": [
        {"entry_id": "t:b1", "visit_status": "Yapım çalışmaları nedeniyle kapalı"},  # Sevilla Mekânlar!41
        {"entry_id": "t:b2", "status_note": "Tadilat nedeniyle 1 Kasım 2026'dan 15 Mart 2027'ye kadar kapalı."},
        {"entry_id": "t:b3", "visit_status": "Açık"},
        {"entry_id": "t:b4", "visit_status": "   "},  # bosluk = bos
    ], "EkDetaylar": [{"detail_id": "t:b5", "entry_id": "t:b1", "key": "venue_editorial", "value": "Pazartesi kapalı."}],
        "SayfaAlanları": [{"field_id": "t:b6", "key": "hero_kicker", "value_type": "text", "value": "İspanya · Sevilla"}]}
    code, s, d = _scan(capsys, tmp_path, rows)
    found = {(h["id"], h["field"]): (h["check"], h["patterns"]) for h in d["hits"]}
    assert code == 1 and found == {("t:b1", "visit_status"): ("K103", ["kart_bandi"]),
                                   ("t:b2", "status_note"): ("K103", ["kart_bandi"]),
                                   ("t:b3", "visit_status"): ("K103", ["kart_bandi"]),
                                   ("t:b5", "venue_editorial"): ("K103", ["kart_bandi"])}, found
    assert s["by_pattern"] == {"kart_bandi": 4}


@pytest.mark.parametrize("sentence, exempt", [
    ("Müze 1 Kasım'da tadilat için kapanıyor; açılış tarihini sitesinden kontrol edin.", True),
    ("Müze 2026'da tadilat için kapanıyor; açılış tarihini sitesinden kontrol edin.", False),  # tek basina yil
    ("1 Kasım konser biletleri satışta; saatleri sitesinden kontrol edin.", False),  # 'bilet' somut durum degil
    ("28 Ocak 2022 haberine göre kaydı her yıl yenilemen gerekir; güncel kuralı sitesinden kontrol et.", False),  # yenileme
])
def test_warning_scan_exemption_needs_day_month_and_a_closure_word(sentence, exempt):
    hits, freed = pc.sentence_patterns(sentence)
    assert (bool(freed), bool(hits)) == (exempt, not exempt), (hits, freed)


@pytest.mark.parametrize("sentence, pattern", [
    ("4-7 Aralık kongresi nedeniyle ziyaret yalnız Giralda ile sınırlı.", "yilsiz_istisna"),  # Sevilla S21
    ("4-8 Aralık 2024 kongresi nedeniyle katedral ziyareti sınırlıydı.", "gecmis_tarihli_istisna"),
    ("12 Kasım 2026 itibarıyla yenileme nedeniyle kapalı (müzenin sitesi).", None),  # bu yil
    ("Her yıl 28 Şubat'ta Endülüs Günü nedeniyle resmî daireler kapalıdır.", None),  # yinelenen
    ("Başvurular 7 Nisan-6 Mayıs arasında yapılır.", None),  # istisna sozcugu yok
    ("Pazartesi günleri kapalı.", None),  # tarih yok
    ("Okul yılı 2024'te değişen takvimle başladı.", None),  # gun-ay yok
    ("Saray 1 Ocak, 6 Ocak, Kutsal Cuma ve 25 Aralık'ta kapalı.", None),  # Sevilla: tatil listesi, neden yok
    ("Fernando 23 Kasım'ın San Clemente günü olması nedeniyle 1248'de manastırı kurdurdu.", None),  # tarih bilgisi
])
def test_dated_exception_pattern(sentence, pattern):
    assert pc.dated_exception(sentence, 2026) == pattern


def test_warning_scan_reports_dated_exceptions_as_k105(tmp_path, capsys):
    rows = {"Yanıtlar": [{"entry_id": "t:k1", "note": "Ziyaret saatleri değişebilir. 4-7 Aralık kongresi nedeniyle "
                                                      "ziyaret yalnız Giralda ile sınırlı."}]}
    code, s, d = _scan(capsys, tmp_path, rows)
    assert code == 1 and s["by_check"]["K105"] == 1
    hit = d["hits"][0]
    assert (hit["id"], hit["check"], hit["patterns"]) == ("t:k1", "K105", ["yilsiz_istisna"])
    assert hit["sentences"] == ["4-7 Aralık kongresi nedeniyle ziyaret yalnız Giralda ile sınırlı."]


def test_warning_scan_usage_errors(tmp_path, capsys):
    fake = tmp_path / "sahte.xlsx"
    fake.write_text("excel degil", encoding="utf-8")
    for xlsx in (tmp_path / "yok.xlsx", fake):
        code, s = _run(capsys, "warning-scan", "--xlsx", xlsx, "--out", tmp_path / "u.json")
        assert code == 2 and s.get("error") and not (tmp_path / "u.json").exists()


def test_tools_never_write_into_the_package(drive, capsys, tmp_path):
    before = _tree(drive["pkg"]["root"])
    _compare(capsys, drive)
    _run(capsys, "warning-scan", "--xlsx", drive["pkg"]["xlsx"], "--out", tmp_path / "u.json")
    code, s = _run(capsys, "warning-scan", "--xlsx", drive["pkg"]["xlsx"], "--out", drive["pkg"]["root"] / "u.json")
    assert code == 2 and _tree(drive["pkg"]["root"]) == before


# --- report ---

def _results(out: Path) -> dict:
    """Gecerli taban: sayfa/yonetim maddeleri GEÇTİ (K93/K112 her zorunlu islevin GEÇTİ senaryosuyla), Drive maddeleri
    KONTROL EDİLEMEDİ (Drive araci yok, --drive verilmez)."""
    (out / "ekran").mkdir(parents=True, exist_ok=True)
    for name in ("masaustu.png", "mobil.png", "senaryo.png"):
        (out / "ekran" / name).write_bytes(b"\x89PNG")
    checks = [{"id": cid, "sonuc": "KONTROL EDİLEMEDİ", "kanit": [], "gerekce": "Drive aracı yok"} if cid in DRIVE_IDS
              else {"id": cid, "sonuc": "GEÇTİ", "kanit": ["ekran/masaustu.png"], "gozlem": "sorun görülmedi"}
              for cid in SCOPE]
    scenarios = [{"kontrol": cid, "islev": flow, "senaryo": flow.replace("_", " "), "girdi": "sentetik: gelir 3000",
                  "beklenen": "kalan 1800", "gozlenen": "kalan 1800", "sonuc": "GEÇTİ", "ekran": "ekran/senaryo.png"}
                 for cid, flows in pc.FLOWS.items() for flow in flows]
    return {"sehir": "Testville", "city_id": "testville", "ortam": "staging", "zaman": "2026-10-23T15:00:00+02:00",
            "adresler": {"yonetim": "https://staging.example/yonetim/uretim/"},
            "kapsam": {"yonetim": True, "masaustu": True, "mobil": True}, "kontroller": checks, "senaryolar": scenarios}


def _set(data: dict, check_id: str, **fields) -> dict:
    return {**data, "kontroller": [{**c, **fields} if c["id"] == check_id else c for c in data["kontroller"]]}


def _drive_passed(data: dict) -> dict:
    for cid in DRIVE_IDS:
        data = _set(data, cid, sonuc="GEÇTİ", kanit=["drive_karsilastirma.json"], gerekce="", gozlem="Drive kanıtı")
    return data


def _clean_drive(capsys, drive, tmp_path) -> Path:
    """Temiz Drive kaniti: MD5'li liste, ana Excel ham indirmeyle SHA-256, kok kimligi promptla eslesti."""
    (tmp_path / "indirme_excel").mkdir()
    shutil.copyfile(drive["pkg"]["root"] / EXCEL, tmp_path / "indirme_excel" / "excel.xlsx")
    downloads = _write(drive["out"] / "indirilenler.json",
                       [{"path": EXCEL, "file": str(tmp_path / "indirme_excel" / "excel.xlsx"), "format": "raw"}])
    code, s, _ = _compare(capsys, drive, None, "--downloads", downloads, *ROOT_OK)
    assert code == 0, s
    return drive["out"] / "drive_karsilastirma.json"


def _report(capsys, tmp_path: Path, data: dict, *extra) -> tuple[int, dict, Path]:
    out = tmp_path / "OUT"
    results = _write(out / "sonuclar.json", data)
    (out / "ONIZLEME_OZETI.md").unlink(missing_ok=True)
    code, summary = _run(capsys, "report", "--results", results, "--out-dir", out, *extra)
    return code, summary, out / "ONIZLEME_OZETI.md"


def test_report_first_line_counts_and_drive_coverage(drive, capsys, tmp_path):
    evidence = _clean_drive(capsys, drive, tmp_path)
    data = _set(_drive_passed(_results(tmp_path / "OUT")), "K94", sonuc="KONTROL EDİLEMEDİ", kanit=[],
                gerekce="uygulama sayfası yok")
    data = _set(_set(data, "K103", sonuc="KALDI", gozlem="kapakta uyarı bandı"),
                "K112", sonuc="KONTROL EDİLEMEDİ", kanit=[], gerekce="onay yok")
    code, s, md = _report(capsys, tmp_path, data, "--drive", evidence)
    text = md.read_text(encoding="utf-8")
    assert text.splitlines()[0] == ("Önizleme kontrolü: Testville | Ortam: staging | Kapsam: Drive 5/5, yönetim ekranı "
                                    "okundu, şehir sayfası (masaüstü var/mobil var) | Sonuç: 7/10, KALDI 1, "
                                    "KONTROL EDİLEMEDİ 2")
    assert code == 1 and (s["gecti"], s["kaldi"], s["kontrol_edilemedi"], s["toplam"]) == (7, 1, 2, 10)
    assert "| K103 | Fotoğrafların üzerinde uyarı bandı olmaması | KALDI |" in text
    assert "kapakta uyarı bandı" in text and "uygulama sayfası yok" in text and EXCEL in text
    assert "Kök kimliği: eslesti (beklenen kok-1, kaynak prompt, bulunan kok-1)" in text and "önceki koşu: yok" in text
    assert pc.DECISION.search(pc.fold(text)) is None  # aracin kendi metninde karar ifadesi yok


def test_report_clean_run_quotes_and_not_applicable(drive, capsys, tmp_path):
    evidence = _clean_drive(capsys, drive, tmp_path)
    data = _set(_drive_passed(_results(tmp_path / "OUT")), "K111",
                gozlem="Panel aşaması «Yayına alınacak»; alanda “Onaylandı” yazıyor")
    data = {**_set(data, "K112", sonuc="UYGULANMAZ", kanit=[], gerekce="önizlemede kişisel araçlar yok"),
            "ortam": "staging | UAT"}
    code, s, md = _report(capsys, tmp_path, data, "--drive", evidence)
    first = md.read_text(encoding="utf-8").splitlines()[0]
    assert code == 0 and first.endswith("| Sonuç: 9/9, KALDI 0, KONTROL EDİLEMEDİ 0, UYGULANMAZ 1"), first
    assert "| Ortam: staging / UAT |" in first and "Kapsam: Drive 5/5," in first and s["uygulanmaz"] == 1


def test_report_without_drive_never_passes_drive_checks(tmp_path, capsys):
    code, s, md = _report(capsys, tmp_path, _results(tmp_path / "OUT"))
    first = md.read_text(encoding="utf-8").splitlines()[0]
    assert code == 1 and "Kapsam: Drive ölçülmedi," in first and first.endswith("Sonuç: 6/10, KALDI 0, KONTROL EDİLEMEDİ 4")
    (tmp_path / "OUT" / "drive_karsilastirma.json").write_text("{}", encoding="utf-8")  # var olan ama olculmemis kanit
    code, s, md = _report(capsys, tmp_path, _drive_passed(_results(tmp_path / "OUT")))
    assert code == 2 and sorted({p.split()[0] for p in s["problems"]}) == sorted(DRIVE_IDS) and not md.exists(), s


def _without(data: dict, key: str, cid: str) -> dict:
    return {**data, key: [item for item in data[key] if item.get("id", item.get("kontrol")) != cid]}


def _only(data: dict, keep) -> dict:
    return {**data, "senaryolar": [s for s in data["senaryolar"] if keep(s)]}


BAD_RESULTS = {
    "gecersiz_sonuc": lambda d: _set(d, "K92", sonuc="PASS"),
    "kanitsiz_gecti": lambda d: _set(d, "K92", kanit=[]),
    "kanitsiz_kaldi": lambda d: _set(d, "K92", sonuc="KALDI", kanit=[]),
    "olmayan_kanit": lambda d: _set(d, "K92", kanit=["ekran/yok.png"]),
    "gerekcesiz": lambda d: _set(d, "K94", sonuc="KONTROL EDİLEMEDİ", kanit=[], gerekce=" "),
    "eksik_kontrol": lambda d: _without(d, "kontroller", "K109"),
    "yinelenen": lambda d: {**d, "kontroller": d["kontroller"] + d["kontroller"][:1]},
    "bilinmeyen": lambda d: {**d, "kontroller": d["kontroller"] + [{"id": "K50", "sonuc": "UYGULANMAZ", "gerekce": "x"}]},
    "sehirsiz": lambda d: {**d, "sehir": ""},
    "karar_hazir": lambda d: _set(d, "K92", gozlem="Sayfa yayına hazır"),
    "karar_uat": lambda d: _set(d, "K111", gozlem="Rehber UAT'de yayımlandı"),
    "karar_yayinda": lambda d: _set(d, "K111", gozlem="Rehber UAT'de yayında"),
    "karar_alindi": lambda d: _set(d, "K111", gozlem="Sayfa yayına alındı"),
    "karar_canli": lambda d: _set(d, "K111", gozlem="Şehir canlıya alındı"),
    "karar_uygun": lambda d: _set(d, "K111", gozlem="Rehber yayın için uygun"),
    "karar_kabul": lambda d: _set(d, "K111", gozlem="Paket kabul edildi"),
    "karar_onay": lambda d: {**d, "notlar": ["Fotoğraflar onaylandı"]},
    "eposta_gozlem": lambda d: _set(d, "K111", gozlem="giriş: kisi@example.com"),
    "token_gozlem": lambda d: _set(d, "K111", gozlem="adres çubuğunda token=abc123"),
    "eposta_not": lambda d: {**d, "notlar": ["sorumlu: kisi@example.com"]},
    "uygulanmaz_k94": lambda d: _set(d, "K94", sonuc="UYGULANMAZ", kanit=[], gerekce="sayfada atıf alanı yok"),
    "uygulanmaz_k110": lambda d: _set(d, "K110", sonuc="UYGULANMAZ", gerekce="Drive yok"),
    "yonetim_okunmadi": lambda d: {**d, "kapsam": {**d["kapsam"], "yonetim": False}},
    "mobil_yok": lambda d: {**d, "kapsam": {**d["kapsam"], "mobil": False}},
    "sayfa_yok": lambda d: _set(_set({**d, "kapsam": {"yonetim": True}}, "K92", sonuc="KONTROL EDİLEMEDİ", kanit=[],
                                     gerekce="sayfa yok"), "K93", sonuc="KONTROL EDİLEMEDİ", gerekce="sayfa yok"),
    "senaryosuz": lambda d: _only(d, lambda s: s["kontrol"] != "K112"),
    "k112_ics_yok": lambda d: _only(d, lambda s: s["islev"] != "takvim_ics"),
    "k93_adimlar_yok": lambda d: _only(d, lambda s: s["islev"] != "adimlar"),
    "k112_bakilamayan": lambda d: {**d, "senaryolar": d["senaryolar"] + [
        {**d["senaryolar"][-1], "sonuc": "KONTROL EDİLEMEDİ", "ekran": "", "gozlenen": "ICS indirme onayı yok"}]},
    "kaldi_senaryo": lambda d: {**d, "senaryolar": d["senaryolar"] + [{**d["senaryolar"][0], "sonuc": "KALDI"}]},
    "gecersiz_islev": lambda d: {**d, "senaryolar": d["senaryolar"] + [{**d["senaryolar"][0], "islev": "uydurma"}]},
    "kisisel_veri": lambda d: {**d, "senaryolar": [{**s, "girdi": "ad: ali@example.com"} for s in d["senaryolar"]]},
    "ekransiz_senaryo": lambda d: {**d, "senaryolar": [{**s, "ekran": ""} for s in d["senaryolar"]]},
}


def test_report_base_results_are_valid(tmp_path, capsys):
    code, s, md = _report(capsys, tmp_path, _results(tmp_path / "OUT"))
    assert code == 1 and md.exists() and s["gecti"] == 6, s  # BAD_RESULTS yalniz kendi kusuruyla reddedilir


@pytest.mark.parametrize("case", sorted(BAD_RESULTS))
def test_report_rejects_bad_results_without_writing(case, tmp_path, capsys):
    code, s, md = _report(capsys, tmp_path, BAD_RESULTS[case](_results(tmp_path / "OUT")))
    assert code == 2 and s.get("error") and s.get("problems") and not md.exists(), s


def test_report_drive_claims_must_match_the_evidence(drive, capsys, tmp_path):
    evidence = _clean_drive(capsys, drive, tmp_path)
    downloads = drive["out"] / "indirilenler.json"
    bundle = {"path": "Lisans kanıtları/kanitlar.zip", "id": "z1", "size": "99", "mimeType": "application/zip"}
    renamed = {**drive["listing"], "root": {**drive["listing"]["root"], "name": "CITY-998 · Başka"}}
    cases = {  # Drive durumu -> GEÇTİ'si reddedilen maddeler
        "baglayici_hash_yok": ((_connector(drive), *ROOT_OK), {"K01", "K03"}),
        "drive_excel_eski": ((_connector(drive, change={EXCEL: {"fileSize": "605912"}}), *ROOT_OK),
                             {"K01", "K03", "K110"}),  # Eindhoven: Drive'daki Excel yerelden eski
        "eksik_dosya": ((_files(drive, drop=(WEB2,)), *ROOT_OK), {"K110"}),
        "kok_kimligi_tek_kaynak": ((None,), {"K109"}),
        "boyutsuz_kayit": ((_connector(drive, change={JPG1: {"fileSize": None}}), "--downloads", downloads, *ROOT_OK),
                           {"K110"}),
        "zipte_dogrulanmadi": ((_files(drive, drop=(JPG1, JPG2, "Lisans kanıtları/orijinaller"), add=[bundle]),
                                *ROOT_OK), {"K110"}),
        "kok_adi_baska_sehir": ((renamed, *ROOT_OK), {"K01", "K109"}),
    }
    data = _drive_passed(_results(tmp_path / "OUT"))
    for name, ((listing, *extra), rejected) in cases.items():
        _compare(capsys, drive, listing, *extra)
        code, s, _ = _report(capsys, tmp_path, data, "--drive", evidence)
        assert code == 2 and {p.split()[0] for p in s["problems"]} == rejected, (name, s)


def test_report_scope_follows_checks_yaml(tmp_path, capsys):
    assert list(pc.load_scope()) == list(SCOPE) and pc.read_checks()[2] == list(DRIVE_IDS)
    code, s, _ = _report(capsys, tmp_path, _results(tmp_path / "OUT"), "--checks", tmp_path / "yok.yaml")
    assert code == 2 and "checks" in s["error"]
