"""text_quality: K27/K28/K104/K21/K33/K25 metin kalitesi kontrolleri (sentetik veri)."""
from __future__ import annotations

import copy
import hashlib
import itertools
import json
import random
import subprocess
import sys
import time
from pathlib import Path

import openpyxl
import pytest

import text_quality as tq

LONG = ("Belediye kaydı için randevu alırken pasaportunu, kira sözleşmeni ve doğum belgeni yanında götür; "
        "kayıt tamamlandıktan sonra BSN numaran posta ile gelir ve bankada, işte, sigortada bu numarayı kullanırsın. "
        "Randevu sistemi yoğun dönemlerde birkaç hafta sonrasını gösterebilir.")
HEDGE = "Saat, ücret ve kapalı günleri ziyaret öncesi kurumdan doğrula; bilet kurumun kendi kanalındadır."


def _ans(i: int, **fields) -> dict:
    base = {"_row": 8 + i, "entry_id": f"x:a{i}", "title": f"Benzersiz başlık {hashlib.sha1(str(i).encode()).hexdigest()[:16]}",
            "summary": f"Özgün özet {i} " + "kelime" * (i % 5)}
    return {**base, **fields}


def _place(i: int, **fields) -> dict:
    return {"_row": 8 + i, "entry_id": f"x:p{i}", "title": f"Mekan {i}", **fields}


def _link(entry: str, label: str = "Başvur", url: str = "https://example.org/basvur") -> dict:
    return {"_row": 8, "action_id": f"a:{entry}:{label}", "entry_id": entry, "label": label, "url": url}


def _xlsx(path: Path, sheets: dict[str, list[dict]]) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        keys = list(dict.fromkeys(k for r in rows for k in r if not k.startswith("_"))) or ["entry_id"]
        for col, key in enumerate(keys, start=1):
            ws.cell(7, col, key)
        for i, rec in enumerate(rows):
            for key, value in rec.items():
                if not key.startswith("_"):
                    ws.cell(8 + i, keys.index(key) + 1, value)
    wb.save(path)
    return path


# --- temel yardimcilar ---------------------------------------------------------

def test_normalize_lowercases_strips_accents_and_punctuation():
    assert tq.normalize("  İstanbul'da ŞEHİR, ılık! Çğü ö ") == "istanbul da sehir ilik cgu o"
    assert tq.normalize("a\n\n-b—c") == "a b c"
    assert tq.normalize(None) == ""


@pytest.mark.parametrize("value", ["-", " Yok ", "N/A", "bilinmiyor", "TODO", "—", "?", "...", "n.a."])
def test_placeholder_values(value):
    assert tq.is_placeholder_value(value)


@pytest.mark.parametrize("value", ["Ücretsiz", "yoklama", "Yok sayılmaz bir not", 0, None, "1-2 ay"])
def test_non_placeholder_values(value):
    assert not tq.is_placeholder_value(value)


def test_near_pairs_matches_brute_force():
    rng = random.Random(3)
    words = ["kayit", "randevu", "belge", "kira", "sigorta", "banka", "okul", "is", "bsn", "gemeente"]
    base = [" ".join(rng.choice(words) for _ in range(40)) for _ in range(12)]
    texts = base + [t[:-3] + "xyz" for t in base[:6]] + [t.replace("kira", "kara", 1) for t in base[6:]]
    sets = [tq.shingles(t) for t in texts]
    for threshold in (0.9, 0.7, 0.5):
        brute = sorted((i, j) for i, j in itertools.combinations(range(len(sets)), 2)
                       if tq.jaccard(sets[i], sets[j]) >= threshold)
        assert tq.near_pairs(sets, threshold) == brute


def test_shingles_and_jaccard_edges():
    assert tq.shingles("abc") == frozenset({"abc"})
    assert tq.shingles("") == frozenset()
    assert tq.jaccard(frozenset(), frozenset()) == 0.0


# --- K27 birebir tekrar ----------------------------------------------------------

def test_exact_long_duplicates_clustered_short_and_placeholders_apart():
    answers = [_ans(0, summary=LONG, cost="Ücretsiz"), _ans(1, summary=LONG.upper(), cost="Ücretsiz"),
               _ans(2, cost="Ücretsiz"), _ans(3, cost="-"), _ans(4, cost="-")]
    report = tq.analyze({"Yanıtlar": answers})
    exact = [c for c in report["exact_duplicates"] if c["field"] == "summary"]
    assert len(exact) == 1 and exact[0]["size"] == 2
    assert [m["row"] for m in exact[0]["members"]] == [8, 9]
    assert not [c for c in report["exact_duplicates"] if c["field"] == "cost"]
    short = [s for s in report["short_repeats"] if s["field"] == "cost"]
    assert short == [{"sheet": "Yanıtlar", "field": "cost", "value": "Ücretsiz", "size": 3}]
    assert report["placeholders"]["by_field"] == {"Yanıtlar.cost": 2}
    assert report["summary"]["placeholders"] == 2


def test_combined_card_copy_detected():
    card = {"summary": LONG, "note": LONG[::-1]}
    report = tq.analyze({"Yanıtlar": [_ans(0, **card), _ans(1, **card)]})
    assert any(c["field"] == tq.COMBINED for c in report["exact_duplicates"])


def test_photo_fields_and_caption_is_title():
    places = [_place(0, title="Van Abbemuseum"), _place(1, title="De Blob")]
    mod = "Corteqs: Commons 1600px türev kopyası indirildi; WebP'ye dönüştürüldü."
    photos = [{"_row": 8 + i, "photo_id": f"f{i}", "entry_id": f"x:p{i % 2}", "caption": c, "modifications": mod}
              for i, c in enumerate(["Van Abbemuseum", "De Blob", "Müzenin Dommel kıyısındaki yeni kanadı"])]
    report = tq.analyze({"Mekânlar": places, "Fotoğraflar": photos})
    mods = [c for c in report["exact_duplicates"] if c["field"] == "modifications"]
    assert mods and mods[0]["size"] == 3
    assert report["caption_is_title"]["count"] == 2 and report["caption_is_title"]["total"] == 3


# --- K28 yakin / kalip tekrar -----------------------------------------------------

def test_near_duplicates_respect_min_cluster_and_threshold():
    variants = [LONG, LONG.replace("hafta", "gün"), LONG.replace("posta", "e-posta")]
    answers = [_ans(i, summary=v) for i, v in enumerate(variants)]
    report = tq.analyze({"Yanıtlar": answers}, min_cluster=3, near=0.9)
    near = [c for c in report["near_duplicates"] if c["field"] == "summary"]
    assert len(near) == 1 and near[0]["size"] == 3 and near[0]["distinct"] == 3
    assert tq.analyze({"Yanıtlar": answers}, min_cluster=4)["near_duplicates"] == []
    assert tq.analyze({"Yanıtlar": answers}, near=0.999)["near_duplicates"] == []


def test_template_numbered_titles_and_masked_subtitles():
    answers = [_ans(i, title=f"Sık sorulan soru #{i + 1}") for i in range(3)]
    tail = " Ev bakarken sokağı farklı saatlerde gör; bina bazında gürültü, park ve ulaşımı yerinde kontrol et."
    hoods = [{"_row": 8 + i, "entry_id": f"x:m{i}", "title": f"Semt {i}", "subtitle": sub, "note": f"Kimine uygun: {sub}.{tail}"}
             for i, sub in enumerate(["Aileler ve çocuklu haneler", "Öğrenciler, genç profesyoneller", "Sakinlik arayan emekliler"])]
    report = tq.analyze({"Yanıtlar": answers, "Mahalleler": hoods})
    fields = {(c["sheet"], c["field"]): c["size"] for c in report["template_duplicates"]}
    assert fields[("Yanıtlar", "title")] == 3
    assert fields[("Mahalleler", "note")] == 3
    assert report["filler"]["count"] == 3
    sentences = {s["sentence"]: s["records"] for s in report["repeated_sentences"]}
    assert sentences["ev bakarken sokagi farkli saatlerde gor"] == 3


def test_template_ignores_numeric_only_values():
    answers = [_ans(i, cost=f"{10 * (i + 1)} euro aylık abonelik ücreti") for i in range(3)]
    report = tq.analyze({"Yanıtlar": [*answers]})
    assert [c for c in report["template_duplicates"] if c["field"] == "cost"]
    words_only = [_ans(i, timing=f"{i + 2} hafta") for i in range(3)]
    assert tq.analyze({"Yanıtlar": words_only})["template_duplicates"] == []


def test_shared_steps_three_lines():
    common = "Pasaportunu hazırla\nRandevu al\nFormu doldur"
    answers = [_ans(0, steps=common + "\nÖde"), _ans(1, steps="Başla\n" + common), _ans(2, steps="Pasaportunu hazırla\nRandevu al")]
    report = tq.analyze({"Yanıtlar": answers})
    assert len(report["shared_steps"]) == 1
    assert [m["id"] for m in report["shared_steps"][0]["members"]] == ["x:a0", "x:a1"]
    assert report["shared_steps"][0]["shared"] == 3


# --- K104 / K25 ------------------------------------------------------------------

def test_hedge_patterns_detect_generic_verification():
    assert tq.hedge_hits(HEDGE) == ["ziyaret_oncesi_dogrula", "kurumdan_dogrula"]
    assert tq.hedge_hits("Güncel bilgiyi kurumdan alın.") == ["guncel_bilgi_kurum"]
    assert tq.hedge_hits("Resmî siteden doğrulayın.") == ["resmi_kaynaktan_dogrula"]
    assert tq.hedge_hits("Gitmeden kontrol edin.") == ["gitmeden_kontrol"]
    assert tq.hedge_hits("Randevu onayını kontrol et ve bilet al.") == []
    assert tq.hedge_hits("Şifreni kimseyle paylaşma; hiçbir kurum şifreni sormaz.") == []  # gercek veri yanlis pozitifi
    assert tq.hedge_hits("Bekleme süresini kuruma sorun.") == ["kurumdan_dogrula"]


def test_hedge_share_per_card():
    places = [_place(i, note=HEDGE) for i in range(3)] + [_place(3, note="Pazar günü kapalıdır.")]
    report = tq.analyze({"Mekânlar": places})
    hedges = report["hedges"]
    assert hedges["cards_total"] == 4 and hedges["cards_with_hedge"] == 3
    assert hedges["share"] == 0.75 and hedges["max_per_card"] == 2
    assert hedges["by_pattern"] == {"kurumdan_dogrula": 3, "ziyaret_oncesi_dogrula": 3}
    assert report["summary"]["hedge_cards"] == 3
    two = [_place(i, note=f"{HEDGE} Giriş kapısı {side}.") for i, side in enumerate(["kuzeyde", "güneyde"])]
    below = tq.analyze({"Mekânlar": two})  # 2 kart < min_cluster; Bağlantılar sayfası yok -> K21 atlanir
    assert below["summary"]["hedge_cards"] == 2 and below["summary"]["findings"] == 0
    assert below["action_links"]["skipped"] is True


def test_evasive_first_sentence_for_question_titles():
    answers = [_ans(0, title="Kira ne kadar?", summary="Güncel bilgi için resmî siteye bakın. Sonra plan yap."),
               _ans(1, title="Kira ne kadar?", summary="Ortalama oda kirası 700-900 euro arasıdır. Resmî siteye bakın."),
               _ans(2, title="Kira tablosu", summary="Resmî siteye bakın."),
               _ans(3, title="Ücret?", summary="Duruma göre değişir.")]
    report = tq.analyze({"Yanıtlar": answers})
    assert [i["id"] for i in report["evasive_answers"]["items"]] == ["x:a0", "x:a3"]


# --- K21 eylem baglantisi -----------------------------------------------------------

def test_action_links_missing_and_map_only():
    answers = [_ans(0), _ans(1)]
    places = [_place(0), _place(1)]
    links = [_link("x:a0"), _link("x:p0", "Konum (Google Haritalar)", "https://www.google.com/maps/search/?api=1&query=x"),
             _link("x:p1", "Bilgi", "https://maps.app.goo.gl/abc"), _link("x:p1", "Bölgeyi haritada gör", "https://osm.org/x")]
    report = tq.analyze({"Yanıtlar": answers, "Mekânlar": places, "Bağlantılar": links})
    links_report = report["action_links"]
    assert links_report["checked"] == 4 and links_report["missing"] == 1 and links_report["map_only"] == 2
    assert {(i["id"], i["reason"]) for i in links_report["items"]} == {("x:a1", "baglanti_yok"), ("x:p0", "yalniz_harita"),
                                                                     ("x:p1", "yalniz_harita")}
    assert report["summary"]["without_action"] == 3


# --- K33 visual_json ------------------------------------------------------------------

def test_visual_json_errors_and_empty():
    good = json.dumps({"kind": "zones", "items": [{"key": "A", "title": "Merkez"}]})
    process = json.dumps({"kind": "steps", "items": [{"title": "Bir"}, {"title": "İki"}]})
    truncated = json.dumps({"kind": "zones", "items": [{"title": "Belediyeye kayıt ve…"}]})
    answers = [_ans(0, visual_json=good), _ans(1, visual_json="{bozuk"), _ans(2, visual_json="[1, 2]"),
               _ans(3, visual_json=process, steps="Bir\nİki\nÜç"), _ans(4), _ans(5, visual_json=json.dumps({"items": []})),
               _ans(6, visual_json=json.dumps({"kind": "zones", "items": "x"})), _ans(7, visual_json=truncated)]
    vj = tq.analyze({"Yanıtlar": answers})["visual_json"]
    assert vj["total"] == 8 and vj["empty"] == 1 and vj["valid"] == 1
    assert vj["by_error"] == {"parse_error": 1, "not_object": 1, "step_count_mismatch": 1, "missing_kind": 1,
                              "items_not_list": 1, "truncated_label": 1}


# --- rapor butunlugu ---------------------------------------------------------------------

def _rich_sheets() -> dict:
    answers = [_ans(i, summary=LONG, cost="-", steps="Pasaportunu hazırla\nRandevu al\nFormu doldur") for i in range(3)]
    places = [_place(i, note=HEDGE) for i in range(4)]
    return {"Yanıtlar": answers, "Mekânlar": places, "Bağlantılar": [_link("x:a0")]}


def test_analyze_is_pure_and_deterministic():
    sheets = _rich_sheets()
    snapshot = copy.deepcopy(sheets)
    first = tq.analyze(sheets)
    assert sheets == snapshot
    assert json.dumps(first, sort_keys=True) == json.dumps(tq.analyze(copy.deepcopy(sheets)), sort_keys=True)
    top = first["top_clusters"]
    assert 1 <= len(top) <= 5 and top[0]["size"] >= top[-1]["size"]
    assert all(len(c["sample"]) <= tq.SAMPLE_CHARS + 1 for c in top)
    assert first["summary"]["findings"] > 0


def test_scales_to_300_records_under_10_seconds(monkeypatch):
    """310 kayit (~47 bin kelime) 10 sn icinde; yakin tekrar aramasi kaba kuvvete (n^2 jaccard) donmez. Olcum
    (skill-son-hal-5, ayni veri): kapsamsiz 1,6-1,9 sn (hitap olcumu 0,36-0,45 sn), kapsamla tek basina 4,0-4,6 sn
    (hitap 1,2-1,4 sn), kapsamli tam kosuda 5,26 sn: eski 5 sn siniri dogru kodda ara sira kirmiziydi. Kaba kuvvet
    near_pairs kapsamsiz yalniz 2,9-3,1 sn surer, yani duvar saati bu boyutta karmasiklik bekcisi degildir; o is
    kapsam ve yukten bagimsiz jaccard cagri sayisindadir (57 bin; kaba kuvvet 320 bin). 10 sn: en yavas olcumun ~2
    kati, yalniz buyuk yavaslamayi yakalar. Hitap olcumu hizlandirilmadi: uretim oncesi sezgisel araca dokunmak risk,
    kazanc kucuk."""
    rng = random.Random(7)
    vocab = [f"kelime{n}" for n in range(400)]
    text = lambda n: " ".join(rng.choice(vocab) for _ in range(n))  # noqa: E731
    answers = [_ans(i, summary=text(60), steps="\n".join(text(8) for _ in range(5)), note=text(30), documents=text(10),
                    timing=text(6), cost=text(5)) for i in range(300)]
    answers += [_ans(300 + i, summary=answers[i]["summary"] + " ek") for i in range(10)]
    start = time.perf_counter()
    report = tq.analyze({"Yanıtlar": answers})
    assert time.perf_counter() - start < 10
    assert any(c["field"] == "summary" for c in report["near_duplicates"]) is False  # cift = 2 < min_cluster
    calls, jaccard = [], tq.jaccard

    def counted(a: frozenset[str], b: frozenset[str]) -> float:
        calls.append(1)
        return jaccard(a, b)

    monkeypatch.setattr(tq, "jaccard", counted)
    assert tq.analyze({"Yanıtlar": answers}, min_cluster=2)["summary"]["near_clusters"] == 10
    distinct = [{tq.normalize(a.get(field)) for a in answers} for field in tq.FIELDS["Yanıtlar"]]
    sizes = [sum(len(norm) > tq.SHORT_LIMIT for norm in norms) for norms in distinct]
    brute = sum(n * (n - 1) // 2 for n in sizes)  # alan basina butun ciftler (kaba kuvvet)
    assert 0 < len(calls) < brute / 3, (len(calls), brute)


# --- CLI ------------------------------------------------------------------------------------

def test_cli_reports_findings_and_writes_out(tmp_path, capsys):
    sheets = _rich_sheets()
    xlsx = _xlsx(tmp_path / "paket.xlsx", sheets)
    out = tmp_path / "rapor" / "text-quality.json"
    code = tq.main([str(xlsx), "--out", str(out)])
    lines = capsys.readouterr().out.strip().splitlines()
    assert code == 1 and len(lines) == 1
    summary = json.loads(lines[0])
    assert summary["tool"] == "text_quality" and summary["placeholders"] == 3 and summary["hedge_cards"] == 4
    report = json.loads(out.read_text(encoding="utf-8"))
    assert report["summary"]["exact_clusters"] >= 1 and report["params"] == {"min_cluster": 3, "near": 0.9}
    assert report["exact_duplicates"][0]["members"][0]["row"] == 8


def test_cli_clean_workbook_exit_zero(tmp_path, capsys):
    xlsx = _xlsx(tmp_path / "temiz.xlsx", {"Yanıtlar": [_ans(0, steps="Tek adım")], "Bağlantılar": [_link("x:a0")]})
    assert tq.main([str(xlsx), "--out", str(tmp_path / "o.json")]) == 0
    assert json.loads(capsys.readouterr().out)["findings"] == 0


def test_cli_usage_errors(tmp_path, capsys):
    assert tq.main([str(tmp_path / "yok.xlsx"), "--out", str(tmp_path / "o.json")]) == 2
    bad = tmp_path / "bozuk.xlsx"
    bad.write_text("excel degil", encoding="utf-8")
    assert tq.main([str(bad), "--out", str(tmp_path / "o.json")]) == 2
    xlsx = _xlsx(tmp_path / "t.xlsx", {"Yanıtlar": [_ans(0)]})
    assert tq.main([str(xlsx), "--out", str(tmp_path / "o.json"), "--near", "1.5"]) == 2
    assert tq.main([str(xlsx), "--out", str(tmp_path / "o.json"), "--min-cluster", "1"]) == 2
    empty = _xlsx(tmp_path / "bos.xlsx", {"Konular": [{"topic_id": "t"}]})
    assert tq.main([str(empty), "--out", str(tmp_path / "o.json")]) == 2
    with pytest.raises(SystemExit) as exc:
        tq.main([str(xlsx)])
    assert exc.value.code == 2
    for line in capsys.readouterr().out.strip().splitlines():
        assert json.loads(line)["error"]


def test_cli_refuses_out_equal_to_input(tmp_path, capsys):
    xlsx = _xlsx(tmp_path / "paket.xlsx", _rich_sheets())
    before = xlsx.read_bytes()
    variants = [xlsx, tmp_path / "alt" / ".." / "paket.xlsx", tmp_path / "PAKET.XLSX"]
    for out in variants if sys.platform == "win32" else variants[:2]:
        assert tq.main([str(xlsx), "--out", str(out)]) == 2
        assert xlsx.read_bytes() == before
    for line in capsys.readouterr().out.strip().splitlines():
        assert "ayni dosya" in json.loads(line)["error"]


def test_cli_unwritable_out_is_usage_error(tmp_path, capsys):
    xlsx = _xlsx(tmp_path / "paket.xlsx", _rich_sheets())
    folder = tmp_path / "klasor"
    folder.mkdir()
    assert tq.main([str(xlsx), "--out", str(folder)]) == 2
    assert tq.main([str(xlsx), "--out", str(tmp_path / "o\0.json")]) == 2  # gecersiz yol: ValueError da kullanim hatasi
    lines = capsys.readouterr().out.strip().splitlines()
    assert len(lines) == 2 and all("yazilamadi" in json.loads(line)["error"] for line in lines)
    assert tq.same_file(tmp_path / "a\0b", xlsx) is False


@pytest.mark.parametrize("count", [51, 60])
def test_shared_steps_mass_copied_block_not_skipped(count):
    block = "Pasaportunu hazırla\nRandevu al\nFormu doldur"
    answers = [_ans(i, steps=f"{block}\nÖzgün adım {i}") for i in range(count)]
    found = tq.analyze({"Yanıtlar": answers})["shared_steps"]
    assert [(c["size"], c["shared"]) for c in found] == [(count, 3)]


def test_shared_steps_mass_lines_combine_with_rare_lines():
    common = "Pasaportunu hazırla\nRandevu al"
    answers = [_ans(i, steps=f"{common}\nÖzgün adım {i}") for i in range(55)]
    answers += [_ans(55 + i, steps=f"{common}\nKira sözleşmesini getir") for i in range(2)]
    found = tq.analyze({"Yanıtlar": answers})["shared_steps"]
    assert [([m["id"] for m in c["members"]], c["shared"]) for c in found] == [(["x:a55", "x:a56"], 3)]


@pytest.mark.parametrize("text", ["Ziyaretten önce bakkaldan su al.", "Gitmeden önce bakımlı bir bisiklet kirala.",
                                  "Ziyaret öncesi kontrollü otopark var.", "Belediye kurumu binayı kontrol eder.",
                                  "Gitmeden önce bakanlık binasının önünden geç.", "Resmî sitede öğrenci indirimi kontrol edilir."])
def test_hedge_patterns_ignore_word_prefixes_and_statements(text):
    assert tq.hedge_hits(text) == []


@pytest.mark.parametrize(("text", "name"), [("Gitmeden saatleri kontrol etmek iyi olur.", "gitmeden_kontrol"),
                                            ("Ziyaretten önce saatlere bakın.", "ziyaret_oncesi_dogrula"),
                                            ("Ziyaret öncesi biletleri teyit edin.", "ziyaret_oncesi_dogrula"),
                                            ("Resmî sitesinden güncel ücreti öğrenin.", "resmi_kaynaktan_dogrula")])
def test_hedge_patterns_keep_advisory_forms(text, name):
    assert tq.hedge_hits(text) == [name]


def test_template_masks_other_city_names_with_suffix():
    tail = "yaşam maliyeti Hollanda ortalamasına yakındır; kira ve ulaşım bütçeni buna göre planla."
    answers = [_ans(i, summary=f"{city}'{suffix} {tail}")
               for i, (city, suffix) in enumerate([("Eindhoven", "da"), ("Amsterdam", "da"), ("Rotterdam", "da")])]
    report = tq.analyze({"Şehir": [{"_row": 8, "name": "Eindhoven"}], "Yanıtlar": answers})
    found = [c for c in report["template_duplicates"] if c["field"] == "summary"]
    assert [(c["size"], c["distinct"]) for c in found] == [(3, 3)]


def test_action_links_card_without_entry_id_not_credited():
    answers = [{"_row": 8, "title": "Kimliksiz kart"}, _ans(1)]
    links = [{"_row": 8, "label": "Başvur", "url": "https://e.org"}, _link("x:a1")]
    report = tq.analyze({"Yanıtlar": answers, "Bağlantılar": links})["action_links"]
    assert [(i["row"], i["reason"]) for i in report["items"]] == [(8, "kimlik_yok")]
    assert report["no_id"] == 1 and report["missing"] == 0 and report["checked"] == 2


def test_script_runs_standalone(tmp_path):
    xlsx = _xlsx(tmp_path / "s.xlsx", _rich_sheets())
    script = Path(tq.__file__)
    proc = subprocess.run([sys.executable, "-B", str(script), str(xlsx), "--out", str(tmp_path / "s.json")],
                          capture_output=True, text=True, encoding="utf-8", timeout=60)
    assert proc.returncode == 1, proc.stderr
    assert json.loads(proc.stdout)["tool"] == "text_quality"


# --- B16 okura hitap (Baris karari 23.10.2026, duzeltilmis: "sen", Berlin gibi; city-pack kapilari) -----------

def _address_sheets() -> dict:
    answers = [_ans(0, summary="Başvuru belediyeye yapılır."), _ans(1, summary="Başvurunuzu belediyeye yapabilirsiniz."),
               _ans(2, summary="Kayıttan sonra randevunuzu alın; kartı da yanında götür.")]
    terms = [{"_row": 8, "term_id": "x:t0", "term": "zorgtoeslag", "definition": "Sağlık sigortası katkısı."}]
    return {"Yanıtlar": answers, "Sözlük": terms, "Bağlantılar": [_link(f"x:a{i}") for i in range(3)]}


def test_address_is_measured_but_not_a_finding_by_default():
    report = tq.analyze(_address_sheets())
    assert report["summary"]["address"] == {"standard": "off", "records": 4, "hitapsiz": 2, "siz": 1, "sen": 0,
                                            "karisik": 1, "findings": 0,
                                            "labels": {"records": 3, "hitapsiz": 0, "siz": 0, "sen": 3, "karisik": 0},
                                            "skipped": []}
    assert report["summary"]["findings"] == 0 and report["address"]["findings"] == []
    assert [i["id"] for i in report["address"]["items"]] == ["x:a1", "x:a2"]


def _duplicate_key(xlsx: Path, sheet: str, key: str) -> Path:
    wb = openpyxl.load_workbook(xlsx)
    ws = wb[sheet]
    ws.cell(7, ws.max_column + 1, key)  # satir 7'de yinelenen teknik anahtar: common.sheet_keys ValueError
    wb.save(xlsx)
    return xlsx


def test_unreadable_address_only_sheet_is_skipped_not_fatal(tmp_path, capsys):
    xlsx = _duplicate_key(_xlsx(tmp_path / "dup.xlsx", _address_sheets()), "Sözlük", "term_id")
    out = tmp_path / "tq.json"
    assert tq.main([str(xlsx), "--out", str(out)]) == 0  # diger kontroller ve hitap olcumu kosmaya devam eder
    summary = json.loads(capsys.readouterr().out)
    assert summary["address"]["skipped"] == ["Sözlük"] and summary["address"]["records"] == 3
    assert summary["without_action"] == 0 and summary["findings"] == 0
    report = json.loads(out.read_text(encoding="utf-8"))
    assert [s["sheet"] for s in report["address"]["skipped"]] == ["Sözlük"]
    assert "yinelenen teknik anahtar" in report["address"]["skipped"][0]["error"]
    assert "Sözlük" not in report["address"]["by_sheet"]
    base = _duplicate_key(_xlsx(tmp_path / "dup2.xlsx", _address_sheets()), "Yanıtlar", "entry_id")
    assert tq.main([str(base), "--out", str(out)]) == 2  # metin kontrollerinin sayfasi okunamazsa kullanim hatasi
    assert "yinelenen" in json.loads(capsys.readouterr().out)["error"]


@pytest.mark.parametrize(("standard", "expected"), [("hitapsiz", [("x:a1", "standart_disi"), ("x:a2", "karisik")]),
                                                    ("siz", [("x:a2", "karisik")]),
                                                    ("sen", [("x:a1", "standart_disi"), ("x:a2", "karisik")])])
def test_address_standard_turns_records_into_findings(standard, expected):
    report = tq.analyze(_address_sheets(), address_standard=standard)
    found = report["address"]["findings"]
    assert [(f["id"], f["reason"]) for f in found] == expected
    assert report["summary"]["address"]["findings"] == len(expected) == report["summary"]["findings"]
    mixed = found[-1]
    assert (mixed["siz"], mixed["sen"], mixed["fields"]) == (["randevunuzu", "alın"], ["götür"], {"summary": "karisik"})


def test_cli_address_standard_hitapsiz_exit_one_and_report(tmp_path, capsys):
    sheets = _address_sheets()
    sheets["Sözlük"][0]["definition"] = "Sağlık sigortası için başvurursun."
    xlsx = _xlsx(tmp_path / "hitap.xlsx", sheets)
    out = tmp_path / "tq.json"
    assert tq.main([str(xlsx), "--out", str(out)]) == 0  # varsayilan off: yalniz olcum
    off = json.loads(capsys.readouterr().out)
    assert off["address"]["standard"] == "off" and off["address"]["sen"] == 1 and off["findings"] == 0
    assert tq.main([str(xlsx), "--out", str(out), "--address-standard", "hitapsiz"]) == 1
    lines = capsys.readouterr().out.strip().splitlines()
    assert len(lines) == 1 and json.loads(lines[0])["address"]["findings"] == 3
    report = json.loads(out.read_text(encoding="utf-8"))
    assert report["address"]["standard"] == "hitapsiz" and report["params"] == {"min_cluster": 3, "near": 0.9}
    assert [(f["sheet"], f["id"]) for f in report["address"]["findings"]] == [
        ("Yanıtlar", "x:a1"), ("Yanıtlar", "x:a2"), ("Sözlük", "x:t0")]
    assert report["address"]["by_sheet"]["Sözlük"] == {"records": 1, "hitapsiz": 0, "siz": 0, "sen": 1, "karisik": 0}


def test_cli_address_standard_sen_flags_siz_and_mixed_but_not_impersonal(tmp_path, capsys):
    """city-pack kapilarinin standardi 'sen': okura 'siz' diyen ve 'siz' ile 'sen'i karistiran kayit bulgu; 'sen'
    kaydi ve kisisiz kayitlar (kisisiz kart, Sozluk tanimi, Konular aciklamasi, foto alt yazisi) bulgu degil; buton
    metni ayri sayilir, kayit bicimine girmez."""
    answers = [_ans(0, summary="Kayıt için belediyeye başvur; randevunu çevrimiçi al."),
               _ans(1, summary="Başvurunuzu belediyeye yapabilirsiniz."),
               _ans(2, summary="Kayıttan sonra randevunuzu alın; kartı da yanında götür."),
               _ans(3, summary="Başvuru belediyeye yapılır.")]
    sheets = {"Yanıtlar": answers, "Bağlantılar": [_link(f"x:a{i}") for i in range(4)],
              "Konular": [{"_row": 8, "topic_id": "x:resmi", "title": "Resmî işler", "description": "Kayıt ve vergi."}],
              "Sözlük": [{"_row": 8, "term_id": "x:term:001", "term": "zorgtoeslag",
                          "definition": "Sağlık sigortası katkısı."}],
              "Fotoğraflar": [{"_row": 8, "photo_id": "x:photo:a-01", "entry_id": "x:a0",
                               "caption": "Belediye binasının girişi."}]}
    out = tmp_path / "tq.json"
    assert tq.main([str(_xlsx(tmp_path / "sen.xlsx", sheets)), "--out", str(out), "--address-standard", "sen"]) == 1
    summary = json.loads(capsys.readouterr().out)
    assert (summary["findings"], summary["address"]["standard"], summary["address"]["findings"]) == (2, "sen", 2)
    address = json.loads(out.read_text(encoding="utf-8"))["address"]
    assert [(f["sheet"], f["id"], f["reason"]) for f in address["findings"]] == [
        ("Yanıtlar", "x:a1", "standart_disi"), ("Yanıtlar", "x:a2", "karisik")]
    by_sheet = {sheet: (counts["records"], counts["hitapsiz"]) for sheet, counts in address["by_sheet"].items()}
    assert by_sheet == {"Yanıtlar": (4, 1), "Konular": (1, 1), "Sözlük": (1, 1), "Fotoğraflar": (1, 1)}
    assert (address["sen"], address["labels"]["sen"], address["skipped"]) == (1, 4, [])


def test_cli_rejects_unknown_address_standard(tmp_path):
    xlsx = _xlsx(tmp_path / "t.xlsx", {"Yanıtlar": [_ans(0)]})
    with pytest.raises(SystemExit) as exc:
        tq.main([str(xlsx), "--out", str(tmp_path / "o.json"), "--address-standard", "biz"])
    assert exc.value.code == 2
    with pytest.raises(ValueError):
        tq.analyze({"Yanıtlar": [_ans(0)]}, address_standard="biz")
