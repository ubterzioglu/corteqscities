"""claim_extract: kart metinlerinden degisken iddialarin (para, yuzde, sure, yas, gun, saat, telefon, tarih) cikarilmasi."""
from __future__ import annotations

import copy
import json
import re

import openpyxl
import pytest

import claim_extract as ce
import helpers


def _kinds(sentence: str) -> list[tuple[str, list[dict]]]:
    return [(c["kind"], c["values"]) for c in ce.extract_values(sentence)]


def _vals(sentence: str, kind: str) -> list:
    return [v["value"] for c in ce.extract_values(sentence) if c["kind"] == kind for v in c["values"]]


# --- sayi bicimleri (Turkce / Hollandaca / Ingilizce) ---------------------------------------

@pytest.mark.parametrize("text,expected", [
    ("85,15", 85.15), ("85.15", 85.15), ("1.234,56", 1234.56), ("1,234.56", 1234.56), ("2.500", 2500),
    ("10,4", 10.4), ("4.4", 4.4), ("80", 80), ("1.500.000", 1500000), ("0,5", 0.5), ("12,50", 12.5),
])
def test_parse_number_handles_tr_nl_and_en_formats(text, expected):
    assert ce.parse_number(text) == expected


def test_parse_number_returns_int_for_whole_values():
    assert isinstance(ce.parse_number("2.500"), int)
    assert isinstance(ce.parse_number("80"), int)


def test_number_regex_is_the_inverse_of_parse_number():
    rx = re.compile(ce.number_regex(85.15))
    assert rx.search("€ 85,15") and rx.search("85.15") and not rx.search("185,15") and not rx.search("85,159")
    rx = re.compile(ce.number_regex(2500))
    assert all(rx.search(t) for t in ("2.500", "2,500", "2500", "2 500", "2.500,00")) and not rx.search("12.500")
    rx = re.compile(ce.number_regex(80))
    assert rx.search("80,-") and rx.search("80,00") and not rx.search("80,50") and not rx.search("800")
    rx = re.compile(ce.number_regex(4.4))
    assert rx.search("4,40%") and rx.search("4.4") and not rx.search("4,45")
    assert all(re.compile(ce.number_regex(ce.parse_number(t))).search(t) for t in ("1.234,56", "10,4", "12,50", "1.500.000"))


# --- para ------------------------------------------------------------------------------------

@pytest.mark.parametrize("sentence,expected", [
    ("KVK kayıt ücreti € 85,15 tutar.", [85.15]), ("Kayıt €85,15.", [85.15]), ("Ücret EUR 80 civarında.", [80]),
    ("Yaklaşık 50 euro öde.", [50]), ("Ceza üst sınırı €50+KDV'dir.", [50]), ("Muhasebe aylık €75-150 arası.", [75, 150]),
    ("Depozito € 1.234,56 olabilir.", [1234.56]), ("Bilet 3,40 € tutar.", [3.4]), ("Kira €2.500 civarı.", [2500]),
    ("Kayıt €85,15; binlerce kişi.", [85.15]),
    ("Aracının emisyon sınıfını (Euro 4/5/6, EV) ruhsattan/RDW'den öğren.", []),  # emisyon sinifi, para degil
    ("Euro 6 dizel araçlar girebilir.", []), ("EURO 5 norm geçerli.", []), ("Euro-6 sınıfı şart.", []),
])
def test_money_variants(sentence, expected):
    claims = [c for c in ce.extract_values(sentence) if c["kind"] == "money"]
    assert [v["value"] for c in claims for v in c["values"]] == expected
    assert all(v["currency"] == "EUR" and "scale" not in v for c in claims for v in c["values"])


def test_money_plus_vat_is_not_an_age():
    assert [k for k, _ in _kinds("Ceza üst sınırı €50+KDV'dir.")] == ["money"]


@pytest.mark.parametrize("sentence,values,scale", [
    ("Bütçe €1,5 milyon.", [1500000], 1000000), ("Fonds van 2,5 miljoen euro.", [2500000], 1000000),
    ("Peşinat €5 bin civarı.", [5000], 1000), ("Proje €1-2 milyon arası.", [1000000, 2000000], 1000000),
    ("Total EUR 3 billion.", [3000000000], 1000000000),
])
def test_money_magnitude_words_scale_the_value(sentence, values, scale):
    found = [v for c in ce.extract_values(sentence) if c["kind"] == "money" for v in c["values"]]
    assert found == [{"kind": "money", "value": v, "currency": "EUR", "scale": scale} for v in values]


# --- yuzde -------------------------------------------------------------------------------------

@pytest.mark.parametrize("sentence,expected", [
    ("Kira artışı en fazla %4,4 olabilir.", [4.4]), ("Max 4.4% increase.", [4.4]), ("Yatırımda %10,4 vergi.", [10.4]),
    ("Alıcı masrafları ~%4-6 arası.", [4, 6]), ("İlk 2 ay maaşın %75'i ödenir.", [75]), ("Yüzde 21 KDV uygulanır.", [21]),
    ("Kendi oturumunda %2, yatırımda %10,4.", [2, 10.4]),
])
def test_percent_variants(sentence, expected):
    assert _vals(sentence, "percent") == expected


@pytest.mark.parametrize("sentence", ["Tarief 35,75% 2 schijven.", "Tarief 35,75 % 2 schijven."])
def test_suffix_percent_sign_is_not_reused_as_prefix(sentence):
    assert _vals(sentence, "percent") == [35.75]


# --- sure --------------------------------------------------------------------------------------

@pytest.mark.parametrize("sentence,value,unit", [
    ("En az 26 hafta çalışmış olmalısın.", 26, "week"), ("Je moet 36 weken gewerkt hebben.", 36, "week"),
    ("Başvuru 14 gün içinde yapılır.", 14, "day"), ("Binnen 90 dagen aanvragen.", 90, "day"),
    ("Son 5 yılda çalışmış olmak gerekir.", 5, "year"), ("İlk 2 ay ödeme yapılır.", 2, "month"),
    ("Ödeme 3 aylık olarak yapılır.", 3, "month"), ("Destek 24 saat açıktır.", 24, "hour"), ("Karar 8 weeks içinde.", 8, "week"),
    # Turkce ek beyaz listesi: gecerli ekler sure kalir
    ("Taşınmadan 10 günü önce bildir.", 10, "day"), ("90 günden kısa kalış.", 90, "day"), ("25 güne kadar izin.", 25, "day"),
    ("5 iş günü sürer.", 5, "day"), ("14 günlük süre.", 14, "day"), ("2 günlüğüne kapalı.", 2, "day"),
    ("3 günde biter.", 3, "day"), ("10 günler sürer.", 10, "day"), ("3 yıllık oturum.", 3, "year"),
    ("5 yılı doldur.", 5, "year"), ("1 yıla kadar.", 1, "year"), ("2 yıldır burada.", 2, "year"),
    ("2 senelik sözleşme.", 2, "year"), ("1 sene sürer.", 1, "year"), ("3 senedir bekliyor.", 3, "year"),
])
def test_duration_variants(sentence, value, unit):
    durations = [v for c in ce.extract_values(sentence) if c["kind"] == "duration" for v in c["values"]]
    assert durations == [{"kind": "duration", "value": value, "unit": unit}]


def test_duration_range_gives_two_values():
    assert _vals("Opzegtermijn genelde 1-2 ay sürer.", "duration") == [1, 2]


@pytest.mark.parametrize("sentence", [  # birimle ayni onekli kelimeler ve takvim yillari sure degildir
    "2025 yılında kural değişti.", "9292 güncel saatleri gösterir.", "2029 güncellemeleriyle bölge genişler.",
    "3 Güncel bilgi.", "Güney Eindhoven 3 güneyde.", "2 güneşli gün.", "2026 aylık katkı payı.", "5 yıldızlı otel.",
    "1 yılbaşı tatili.", "2 senet imzala.", "2026 haftasında açılır.", "2025 dagen geleden.",
])
def test_unit_prefix_words_and_calendar_years_are_not_durations(sentence):
    assert _vals(sentence, "duration") == []


def test_same_value_twice_in_sentence_is_one_claim_per_occurrence_deduped_later():
    claims = [c for c in ce.extract_values("1 hafta önce ila 1 hafta sonra başvur.") if c["kind"] == "duration"]
    assert len(claims) == 2  # cumle duzeyinde iki eslesme; kart duzeyinde tekillestirilir


# --- yas ---------------------------------------------------------------------------------------

@pytest.mark.parametrize("sentence,expected", [
    ("16–18 yaş arası gençler.", [16, 18]),
    ("18+ kiracılara verilir.", [18]),
    ("Çocuk 5 yaşında okula başlar.", [5]),
    ("16 yaşına kadar tam zorunlu.", [16]),
])
def test_age_variants(sentence, expected):
    assert _vals(sentence, "age") == expected


@pytest.mark.parametrize("sentence,expected", [
    ("4+ ziyarette Museumkaart kârlıdır.", []), ("3+ kişi için grup bileti.", []), ("6+ çocuklara ücretsiz yüzme.", [6])])
def test_plus_count_is_not_an_age(sentence, expected):
    assert _vals(sentence, "age") == expected


# --- gun / saat -----------------------------------------------------------------------------------

@pytest.mark.parametrize("sentence,expected", [
    ("Koopavond genelde perşembe.", ["thursday"]), ("Pazar her cuma kurulur.", ["friday"]),
    ("Cumartesi Feelgood Market var.", ["saturday"]),
    ("Alışverişi perşembe akşamına veya cumartesi sabahına planla.", ["thursday", "saturday"]),
    ("Pazar günü merkez açık.", ["sunday"]), ("Salıları kapalı.", ["tuesday"]), ("Markt op zaterdag.", ["saturday"]),
])
def test_weekday_variants(sentence, expected):
    assert _vals(sentence, "weekday") == expected


def test_monthly_ordinal_weekday():
    claims = [c for c in ce.extract_values("Pazar ayın üçüncü pazarı kurulur.") if c["kind"] == "weekday"]
    assert claims[0]["values"] == [{"kind": "weekday", "value": "sunday", "ordinal": 3}]
    assert len(claims) == 1


@pytest.mark.parametrize("sentence", ["Pazar yerinde taze sebze var.", "Pazarlama bütçesi.", "Salıncak parkı.", "Cumhuriyet bayramı."])
def test_non_weekday_words_are_ignored(sentence):
    assert _vals(sentence, "weekday") == []


def test_clock_times():
    assert _vals("Saatler 9/10.00-18.00 arası.", "time") == ["10:00", "18:00"]
    assert _vals("Kapanış 21:30.", "time") == ["21:30"]


def test_money_decimal_is_not_a_time():
    assert [k for k, _ in _kinds("Bilet €10.50.")] == ["money"]


# --- telefon / tarih ---------------------------------------------------------------------------

@pytest.mark.parametrize("sentence,expected", [
    ("Randevu 14 040 (yurt dışından +31 40 2386000) ile alınır.", ["14040", "0402386000"]),
    ("Hat: 040-2386000.", ["0402386000"]),
    ("Bilgi hattı 0800-1234.", ["08001234"]),
    ("Ara: 0031 (0)20 123 4567.", ["0201234567"]),
])
def test_phone_numbers(sentence, expected):
    assert _vals(sentence, "phone") == expected


def test_phone_digits_do_not_leak_into_other_kinds():
    kinds = {k for k, _ in _kinds("Hat: 040-2386000.")}
    assert kinds == {"phone"}


def test_dates_with_and_without_year():
    claims = [c for c in ce.extract_values("Son gün 31 Aralık; zam genelde 1 Temmuz'da.") if c["kind"] == "date"]
    assert [c["values"][0] for c in claims] == [{"kind": "date", "day": 31, "month": 12}, {"kind": "date", "day": 1, "month": 7}]
    claims = [c for c in ce.extract_values("1 Ocak 2026 itibarıyla.") if c["kind"] == "date"]
    assert claims[0]["values"] == [{"kind": "date", "day": 1, "month": 1, "year": 2026}]


def _dates(sentence: str) -> list[dict]:
    return [v for c in ce.extract_values(sentence) if c["kind"] == "date" for v in c["values"]]


@pytest.mark.parametrize("sentence", [  # A5: ay adi yer adinin parcasi ('18 Septemberplein' meydan adidir, tarih degil)
    "Eindhoven'da 18 Septemberplein (yaklaşık 1400 yer) var.", "18 Septemberplein: Pzt-Cmt 07:00-01:00",
    "18 Septemberplein'e bisikletle gel.", "5 Meiplein'de buluşun.", "Adres: 4 Meistraat 12.", "1 Oktoberlaan üzerinde.",
    "17 Septemberweg tarafında.", "18 September-plein girişi.", "2 Decembersingel 4.", "18 Septemberpleingarage kapalı.",
])
def test_month_inside_place_name_is_not_a_date(sentence):
    assert _dates(sentence) == []


@pytest.mark.parametrize("sentence,expected", [
    ("Stalling 18 september'de açılır.", [{"kind": "date", "day": 18, "month": 9}]),
    ("18 Septemberplein 18 Eylül'de kapalı.", [{"kind": "date", "day": 18, "month": 9}]),
    ("Açılış 18 september 2026.", [{"kind": "date", "day": 18, "month": 9, "year": 2026}]),
    ("Meivakantie 5 mei'de başlar.", [{"kind": "date", "day": 5, "month": 5}]),
    ("Son gün 31 Aralık'ta.", [{"kind": "date", "day": 31, "month": 12}]),
    ("Op 1 meiviering in het park.", [{"kind": "date", "day": 1, "month": 5}]),  # 'viering' 'ring' ile biter, yer adi degil
    ("Na de 5 meivergadering.", [{"kind": "date", "day": 5, "month": 5}]),
])
def test_real_dates_next_to_place_names_are_kept(sentence, expected):
    assert _dates(sentence) == expected


def test_place_name_sentence_yields_no_claims():
    assert ce.extract_values("18 Septemberplein (yaklaşık 1400 yer), Heuvel (900 yer).") == []


@pytest.mark.parametrize("sentence", [  # A5: Almanca bilesik ve ayri yazilan, buyuk harfle baslayan yer adi sozcugu
    "1 Oktoberstraße 5 numara.", "4 Septemberplatz'ta buluşun.", "2 Novemberallee boyunca.", "17 Juni Straße tarafında.",
    "5 Mei Laan üzerinde.", "19 Mayıs Caddesi'nde buluşun.", "Adres: 1 Nisan Sokak No: 5.", "12 Eylül Sk. 4 numara.",
    "29 Ekim Meydanı kalabalık.", "23 Nisan Parkı'na gidin.",
])
def test_month_in_german_or_separate_place_name_is_not_a_date(sentence):
    assert _dates(sentence) == []


@pytest.mark.parametrize("sentence,expected", [  # V4 inceleme: yer adi sozcugu ay adindan HEMEN sonra gelmeli
    ("Anma 4 Mayıs Dam Meydanı'nda.", [{"kind": "date", "day": 4, "month": 5}]),  # tarih + baska bir yerin adi
    ("Koningsdag 27 Nisan Markt Meydanı'nda.", [{"kind": "date", "day": 27, "month": 4}]),
    ("Maç 18 Ağustos Philips Stadyumu'nda.", [{"kind": "date", "day": 18, "month": 8}]),
    ("Maç 12 Haziran Westfalen Stadyumu'nda.", [{"kind": "date", "day": 12, "month": 6}]),
    ("Tören 4 Mayısta Dam Meydanı'nda.", [{"kind": "date", "day": 4, "month": 5}]),  # ek almis ay yer adinin parcasi olamaz
    ("1 Mayısta Sokak şenliği var.", [{"kind": "date", "day": 1, "month": 5}]),  # ekli ayda sonraki sozcuge hic bakilmaz
    # Karar: araya ad giren Turk sokak/meydan adi tarih sayilir (NL/DE paketlerinde nadir; gercek tarihi silmemek oncelikli)
    ("30 Ağustos Zafer Bulvarı üzerinde.", [{"kind": "date", "day": 30, "month": 8}]),
    ("15 Temmuz Demokrasi Meydanı'nda tören var.", [{"kind": "date", "day": 15, "month": 7}]),
])
def test_date_before_another_place_name_is_kept(sentence, expected):
    assert _dates(sentence) == expected


@pytest.mark.parametrize("sentence,expected", [  # yer adi baglami yoksa tarih korunur (pilot: '20 Kasım Cuma', '1 Ocak Cuma')
    ("Ödeme 20 Kasım Cuma günü saat 24.00'ten önce yatar.", [{"kind": "date", "day": 20, "month": 11}]),
    ("2027'de yılbaşı 1 Ocak Cuma.", [{"kind": "date", "day": 1, "month": 1}]),
    ("20 Kasım Cuma Plein'de pazar kurulur.", [{"kind": "date", "day": 20, "month": 11}]),  # gun adi ozel ad sayilmaz
    ("Konser 18 Eylül Markt'ta.", [{"kind": "date", "day": 18, "month": 9}]),  # 'Markt' Eindhoven meydani: tarih + yer
    ("Tören 4 Mayıs Plein'de.", [{"kind": "date", "day": 4, "month": 5}]),  # Den Haag 'Plein' meydani
    ("18 Eylül Ring'de koşu var.", [{"kind": "date", "day": 18, "month": 9}]),  # Eindhoven 'Ring' cevre yolu
    # yilli tarih yer adi olmaz: yil kosulu olmasa 'Sokak' (ayri yer sozcugu) tarihi silerdi
    ("1 Mayıs 2026 Sokak şenliği yapılır.", [{"kind": "date", "day": 1, "month": 5, "year": 2026}]),
    ("1 Mayıs Emek ve Dayanışma Günü kutlanır.", [{"kind": "date", "day": 1, "month": 5}]),
    ("18 Eylül meydanında tören var.", [{"kind": "date", "day": 18, "month": 9}]),  # kucuk harf: genel ad, ozel ad degil
])
def test_real_dates_without_place_context_are_kept(sentence, expected):
    assert _dates(sentence) == expected


def test_no_claims_in_plain_text():
    assert ce.extract_values("Belge dosyanı hazırla ve randevu al.") == []


def test_claim_text_is_the_matched_span():
    claim = ce.extract_values("KVK kayıt ücreti € 85,15 tutar.")[0]
    assert claim["text"] == "€ 85,15"


# --- cumle bolme --------------------------------------------------------------------------------

def test_split_sentences_keeps_decimals_and_splits_lines():
    parts = ce.split_sentences("Ücret €85,15. Süre 2.5 hafta.\nİkinci satır; üçüncü parça")
    assert parts == ["Ücret €85,15.", "Süre 2.5 hafta.", "İkinci satır;", "üçüncü parça"]
    assert ce.split_sentences("") == []
    assert ce.split_sentences(None) == []


# --- kayit duzeyi ------------------------------------------------------------------------------

def _sheets():
    return {
        "Yanıtlar": [
            {"_row": 8, "entry_id": "t:kvk", "topic_id": "t:is", "summary": "KVK kaydı €80 tutar.",
             "cost": "€ 85,15; muhasebe €75-150", "source_ids": "t:kaynak:kvk; t:kaynak:yok", "reviewed_on": "2026-10-20"},
            {"_row": 9, "entry_id": "t:ww", "topic_id": "t:sosyal",
             "summary": "1 hafta önce ila 1 hafta sonra başvur.", "timing": "Perşembe 18.00'e kadar", "source_ids": "t:kaynak:uwv"},
            {"_row": 10, "entry_id": "t:genel", "topic_id": "t:is", "summary": "Sayı yok.", "source_ids": "t:kaynak:uwv"},
            {"_row": 11, "entry_id": "t:nosrc", "topic_id": "t:is", "note": "Ücret 50 euro."},
        ],
        "Sözlük": [{"_row": 8, "term_id": "t:term:1", "definition": "Huurtoeslag 18+ kiracılara.", "source_ids": "t:kaynak:uwv"}],
        "Kaynaklar": [{"_row": 8, "source_id": "t:kaynak:kvk", "url": "https://www.kvk.nl/"},
                      {"_row": 9, "source_id": "t:kaynak:uwv", "url": "https://www.uwv.nl/nl/ww"}],
    }


def test_extract_claims_builds_ids_rows_and_sources():
    claims = ce.extract_claims(_sheets())
    kvk = [c for c in claims if c["entry_id"] == "t:kvk"]
    assert [c["claim_id"] for c in kvk] == ["t:kvk:#1", "t:kvk:#2", "t:kvk:#3"]
    assert kvk[0]["_row"] == 8 and kvk[0]["sheet"] == "Yanıtlar" and kvk[0]["field"] == "summary"
    assert kvk[0]["sentence"] == "KVK kaydı €80 tutar."
    assert kvk[0]["values"] == [{"kind": "money", "value": 80, "currency": "EUR"}]
    assert kvk[0]["sources"] == [{"source_id": "t:kaynak:kvk", "url": "https://www.kvk.nl/"},
                                 {"source_id": "t:kaynak:yok", "url": None}]
    assert kvk[0]["missing_source_ids"] == ["t:kaynak:yok"]
    assert kvk[2]["values"] == [{"kind": "money", "value": 75, "currency": "EUR"}, {"kind": "money", "value": 150, "currency": "EUR"}]


def test_extract_claims_dedupes_identical_claims_within_sentence():
    ww = [c for c in ce.extract_claims(_sheets()) if c["entry_id"] == "t:ww"]
    assert [(c["field"], c["kind"]) for c in ww] == [("summary", "duration"), ("timing", "weekday"), ("timing", "time")]


def test_extract_claims_includes_glossary_and_no_source_cards():
    claims = ce.extract_claims(_sheets())
    term = [c for c in claims if c["entry_id"] == "t:term:1"]
    assert term and term[0]["sheet"] == "Sözlük" and term[0]["field"] == "definition"
    nosrc = [c for c in claims if c["entry_id"] == "t:nosrc"]
    assert nosrc[0]["sources"] == [] and nosrc[0]["source_ids"] == []


def test_extract_claims_topic_filter_accepts_suffix_and_skips_glossary():
    claims = ce.extract_claims(_sheets(), topics={"sosyal"})
    assert {c["entry_id"] for c in claims} == {"t:ww"}
    claims = ce.extract_claims(_sheets(), topics={"t:is"})
    assert {c["entry_id"] for c in claims} == {"t:kvk", "t:nosrc"}


def test_extract_claims_does_not_mutate_input():
    sheets = _sheets()
    before = copy.deepcopy(sheets)
    ce.extract_claims(sheets)
    assert sheets == before


def test_read_claims_round_trip(tmp_path):
    claims = ce.extract_claims(_sheets())
    path = tmp_path / "claims.jsonl"
    path.write_text("".join(json.dumps(c, ensure_ascii=False) + "\n" for c in claims) + "\n", encoding="utf-8")
    assert ce.read_claims(path) == claims


MONEY_80 = [{"kind": "money", "value": 80}]


@pytest.mark.parametrize("content", [b"{not json}\n", b'{"claim_id": "x"}\n', b"[1, 2]\n", b'\xff\xfe{"claim_id": "\xfc"}\n']
                         + [(json.dumps(claim) + "\n").encode() for claim in (  # bozuk ogeler + UTF-8 olmayan dosya: ClaimError
        {"claim_id": "x", "values": MONEY_80, "sources": ["https://kvk.nl"]}, {"claim_id": "x", "values": ["80"], "sources": []},
        {"claim_id": "x", "values": [{"value": 80}], "sources": []},
        {"claim_id": "x", "values": MONEY_80, "sources": [{"source_id": "s", "url": 5}]})])
def test_read_claims_rejects_bad_lines_items_and_encoding(tmp_path, content):
    path = tmp_path / "bad.jsonl"
    path.write_bytes(content)
    with pytest.raises(ce.ClaimError):
        ce.read_claims(path)


def test_read_claims_missing_file(tmp_path):
    with pytest.raises(ce.ClaimError):
        ce.read_claims(tmp_path / "yok.jsonl")


def test_summarize_counts():
    claims = ce.extract_claims(_sheets())
    summary = ce.summarize(claims)
    assert summary["claims"] == len(claims)
    assert summary["by_kind"]["money"] == 4
    assert summary["entries_with_claims"] == 4
    assert summary["claims_without_source"] == 1
    assert summary["claims_with_missing_source_ids"] == 3


# --- Excel + CLI --------------------------------------------------------------------------------

def _workbook(tmp_path, *, with_missing: bool = True):
    path = helpers.make_template(tmp_path / "CITY-999_testville_v001.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb["Yanıtlar"]
    base = len(helpers.SHEET_KEYS["Yanıtlar"])
    for i, key in enumerate(("steps", "documents", "timing", "cost", "note"), start=1):
        ws.cell(7, base + i, key)
    gl = wb.create_sheet("Sözlük")
    for col, key in enumerate(("term_id", "city_id", "term", "definition", "source_ids"), start=1):
        gl.cell(7, col, key)
    data = _sheets()
    if not with_missing:
        data["Yanıtlar"] = [dict(r, source_ids="t:kaynak:kvk") for r in data["Yanıtlar"]]
    for name in ("Yanıtlar", "Sözlük", "Kaynaklar"):
        sheet = wb[name]
        keys = {c.value: c.column for c in sheet[7] if c.value}
        for rec in data[name]:
            for key, value in rec.items():
                if key != "_row":
                    sheet.cell(rec["_row"], keys[key], value)
    wb.save(path)
    return path


def test_main_writes_jsonl_and_one_line_summary(tmp_path, capsys):
    xlsx = _workbook(tmp_path)
    out = tmp_path / "out" / "claims.jsonl"
    code = ce.main([str(xlsx), "--out", str(out)])
    lines = capsys.readouterr().out.strip().splitlines()
    assert len(lines) == 1
    summary = json.loads(lines[0])
    assert code == 1  # kaynaksiz kart + Kaynaklar'da olmayan source_id = bulgu
    records = [json.loads(line) for line in out.read_text(encoding="utf-8").splitlines()]
    assert summary["claims"] == len(records) > 0
    assert records[0]["claim_id"] == "t:kvk:#1" and records[0]["_row"] == 8
    assert any(r["sheet"] == "Sözlük" for r in records)


def test_main_exit_zero_when_all_sources_resolve(tmp_path, capsys):
    xlsx = _workbook(tmp_path, with_missing=False)
    out = tmp_path / "claims.jsonl"
    assert ce.main([str(xlsx), "--out", str(out)]) == 0


def test_main_is_deterministic(tmp_path, capsys):
    xlsx = _workbook(tmp_path)
    a, b = tmp_path / "a.jsonl", tmp_path / "b.jsonl"
    ce.main([str(xlsx), "--out", str(a)])
    ce.main([str(xlsx), "--out", str(b)])
    assert a.read_bytes() == b.read_bytes()


def test_main_topics_option(tmp_path, capsys):
    xlsx = _workbook(tmp_path)
    out = tmp_path / "claims.jsonl"
    ce.main([str(xlsx), "--out", str(out), "--topics", "sosyal"])
    ids = {json.loads(line)["entry_id"] for line in out.read_text(encoding="utf-8").splitlines()}
    assert ids == {"t:ww"}


def test_main_usage_error_for_unreadable_xlsx(tmp_path, capsys):
    bad = tmp_path / "bad.xlsx"
    bad.write_text("not a workbook", encoding="utf-8")
    code = ce.main([str(bad), "--out", str(tmp_path / "c.jsonl")])
    assert code == 2
    assert "error" in json.loads(capsys.readouterr().out)


def test_main_usage_error_for_missing_answers_sheet(tmp_path, capsys):
    path = tmp_path / "x.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Kaynaklar"
    wb.save(path)
    assert ce.main([str(path), "--out", str(tmp_path / "c.jsonl")]) == 2
