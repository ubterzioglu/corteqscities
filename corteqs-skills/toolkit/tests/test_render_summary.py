import json
import unicodedata

import openpyxl
import pytest

import common
import helpers
import render_summary
import summary_metrics

TARGET = {"answers": 240, "places": 120, "neighborhoods": 16, "topics": 12, "photos": 250, "hero_photos": 8, "routes": 12, "glossary": 60}
MINIMUM = {"answers": 212, "places": 92, "neighborhoods": 12, "topics": 12, "photos": 187, "hero_photos": 6, "routes": 9, "glossary": 48}


def inventory(**overrides):
    summary = {"files": 10, "total_mb": 1.5, "images": 4, "missing": 0, "sha_mismatch": 0, "file_ne_original": 0,
               "bad_ext": 0, "ext_mismatch": 0, "orphans": 0, "dup_sha": 0, "errors": 0}
    summary.update(overrides)
    findings = {k: [f"ornek-{k}-{i}" for i in range(v)] for k, v in summary.items() if k in
                ("missing", "sha_mismatch", "file_ne_original", "bad_ext", "ext_mismatch", "orphans", "dup_sha", "errors")}
    return {"summary": summary, "totals": {"files": summary["files"], "total_bytes": 1_500_000, "total_mb": summary["total_mb"]},
            "folders": {".": {"files": 1, "bytes": 10}, "Fotoğraflar": {"files": 2, "bytes": 20}}, "findings": findings}


def validation(counts=None, classes=(), media=(), gaps=None, complete=True, extras=None):
    counts = counts or dict(TARGET)
    return {"summary": {"structure": {"errors": sum(c["count"] for c in classes), "warnings": 5},
                        "media": {"errors": sum(c["count"] for c in media)},
                        "minimum_gaps": gaps or {}, "error_classes": list(classes), "media_classes": list(media),
                        "tolerated_duplicate_list_cells": 0, "extras": extras or {}, "technically_complete": complete},
            "structure": {"counts": counts, "minimum": MINIMUM, "target": TARGET, "errors": [], "warnings": []},
            "media": {"different_files": 2, "different_source_pages": 2, "errors": []}}


def write(tmp_path, name, payload):
    path = tmp_path / name
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return path


def run(tmp_path, package, inv=None, val=None, *extra):
    out = tmp_path / "out"
    args = ["--xlsx", str(package["xlsx"]), "--out-dir", str(out), "--date", "2026-10-22", *extra]
    if inv is not None:
        args += ["--inventory", str(write(tmp_path, "inv.json", inv))]
    if val is not None:
        args += ["--validation", str(write(tmp_path, "val.json", val))]
    return render_summary.main(args), out


def test_clean_run_writes_three_files_without_v7_suffix_and_required_sentence(tmp_path, package, capsys):
    package_with_gaps = package  # mini paket: 0 yanit; dogrulayici sayilari hedefte oldugu icin eksik sayilmaz
    code, out = run(tmp_path, package_with_gaps, inventory(), validation())
    assert sorted(p.name for p in out.iterdir()) == ["ACIK_KONULAR.txt", "ILERLEME_KAYDI.txt", "TESLIM_OZETI.txt"]
    open_issues = (out / "ACIK_KONULAR.txt").read_text(encoding="utf-8")
    assert "Araştırmada saptanan açık konu yok; editör/hak kabulü ayrıca gereklidir" in open_issues
    summary = (out / "TESLIM_OZETI.txt").read_text(encoding="utf-8")
    assert "city_id: testville" in summary and "CITY-999" in summary and "Testville" in summary
    assert "YEREL DENEME" in summary and "Drive'a yüklenmedi" in summary
    assert "Editör tarafından onaylandı: HAYIR" in summary
    assert code == 0
    assert json.loads(capsys.readouterr().out.strip().splitlines()[-1])["verdict"].startswith("TEKNİK KONTROL GEÇTİ")


def test_issues_are_listed_with_counts_and_exit_1(tmp_path, package):
    classes = [{"class": "Kanıtlar: geçersiz seçim", "count": 272, "examples": ["Kanıtlar!8"]}]
    media = [{"class": "ana fotoğraf kaydı ile kaynak dosyası uyuşmuyor", "count": 272, "examples": ["testville:foto:yer-01"]}]
    code, out = run(tmp_path, package, inventory(bad_ext=6, file_ne_original=272), validation(classes=classes, media=media, complete=False,
                                                                                                extras={"placeholders": 3, "wikipedia_labeled_official": 0}))
    text = (out / "ACIK_KONULAR.txt").read_text(encoding="utf-8")
    assert code == 1
    assert "Kanıtlar: geçersiz seçim — 272" in text and "kaynak dosyası uyuşmuyor — 272" in text
    assert "bad_ext" in text and "6" in text and "file_ne_original" in text
    assert "ek kontrol" in text and "placeholders" in text and "wikipedia_labeled_official" not in text  # sifir olan listelenmez
    assert "Araştırmada saptanan açık konu yok" not in text
    assert "DÜZELTME GEREKLİ" in (out / "TESLIM_OZETI.txt").read_text(encoding="utf-8")


def test_berlin_gaps_below_minimum_are_blocking_and_below_target_listed(tmp_path, package):
    counts = dict(TARGET, answers=200, places=100)   # 200 < min 212 ; 100 < target 120 ama >= min 92
    code, out = run(tmp_path, package, inventory(), validation(counts=counts, gaps={"answers": {"actual": 200, "minimum": 212}}, complete=False))
    text = (out / "ACIK_KONULAR.txt").read_text(encoding="utf-8")
    assert code == 1
    assert text.count("Pratik yanıt 200/212") == 1          # ayni eksik tek satir (dogrulayici minimum_gaps ile cift sayilmaz)
    assert "Gerçek mekân 100/120" in text and "Hedef eksik" in text
    table = (out / "TESLIM_OZETI.txt").read_text(encoding="utf-8")
    assert "alt sınırın altında" in table and "hedefin altında" in table


def test_missing_inventory_and_validation_are_reported_not_invented(tmp_path, package, capsys):
    code, out = run(tmp_path, package)
    summary = (out / "TESLIM_OZETI.txt").read_text(encoding="utf-8")
    assert "ölçülmedi" in summary
    assert json.loads(capsys.readouterr().out.strip().splitlines()[-1])["verdict"].startswith("KANIT EKSİK")
    assert code == 1


def test_notes_are_appended_to_open_issues(tmp_path, package):
    notes = tmp_path / "notes.md"
    notes.write_text("- Cafe Wilhelmina adresi resmi sayfadan doğrulanamadı\n\n- İkinci not\n", encoding="utf-8")
    code, out = run(tmp_path, package, inventory(), validation(), "--notes", str(notes))
    text = (out / "ACIK_KONULAR.txt").read_text(encoding="utf-8")
    assert "Cafe Wilhelmina adresi" in text and "İkinci not" in text
    assert "Araştırmada saptanan açık konu yok" not in text and code == 1


def _log(tmp_path):
    return (tmp_path / "out" / "ILERLEME_KAYDI.txt").read_text(encoding="utf-8")


def _last_json(capsys):
    return json.loads(capsys.readouterr().out.strip().splitlines()[-1])


def test_progress_log_same_day_same_excel_is_updated_not_repeated(tmp_path, package, capsys):
    """A14: ayni gun + ayni Excel (SHA) icin girdi yeniden eklenmez, son hali yazilir (pilotta ayni girdi 4 kez)."""
    manifest = tmp_path / "manifest.json"
    manifest.write_text(json.dumps({"phase": "assemble", "next_step": "ilk adım", "last_ids": ["testville:yer-2"]}), encoding="utf-8")
    run(tmp_path, package, inventory(), validation(), "--manifest", str(manifest))
    first = _last_json(capsys)
    manifest.write_text(json.dumps({"phase": "deliver", "next_step": "fotoğraf incelemesi", "last_ids": ["testville:yer-2"]}), encoding="utf-8")
    run(tmp_path, package, inventory(), validation(), "--manifest", str(manifest))
    second = _last_json(capsys)
    log = _log(tmp_path)
    assert log.count("=== 2026-10-22") == 1 and "Faz: deliver" in log and "fotoğraf incelemesi" in log and "ilk adım" not in log
    assert f"Excel SHA-256: {common.sha256_file(package['xlsx'])}" in log and "testville:yer-2" in log
    assert first["progress"] == "eklendi" and second["progress"] == "güncellendi"


def test_progress_log_new_day_adds_entry_and_keeps_old(tmp_path, package):
    run(tmp_path, package, inventory(), validation())
    render_summary.main(["--xlsx", str(package["xlsx"]), "--out-dir", str(tmp_path / "out"), "--date", "2026-10-23"])
    log = _log(tmp_path)
    assert log.count("=== 2026-10-22") == 1 and log.count("=== 2026-10-23") == 1
    assert log.index("=== 2026-10-22") < log.index("=== 2026-10-23")


def _add_sources(xlsx, ids):
    wb = openpyxl.load_workbook(xlsx)
    ws, keys = wb["Kaynaklar"], helpers.SHEET_KEYS["Kaynaklar"]
    for r, sid in enumerate(ids, start=8):
        ws.cell(r, keys.index("source_id") + 1, sid)
        ws.cell(r, keys.index("verified") + 1, "Evet")
    wb.save(xlsx)


def test_progress_log_same_day_changed_excel_adds_entry(tmp_path, package):
    run(tmp_path, package, inventory(), validation())
    old_sha = common.sha256_file(package["xlsx"])
    _add_sources(package["xlsx"], ["s1", "s2"])  # durum degisti: yeni SHA, yeni sayilar
    run(tmp_path, package, inventory(), validation())
    log = _log(tmp_path)
    assert log.count("=== 2026-10-22") == 2
    assert f"Excel SHA-256: {old_sha}" in log and f"Excel SHA-256: {common.sha256_file(package['xlsx'])}" in log


def test_progress_log_rebuilt_excel_with_same_state_is_not_repeated(tmp_path, package):
    """Excel yeniden derlenince SHA degisir (openpyxl zaman damgasi); sayilar ayni ise ayni gunun son girdisi guncellenir."""
    run(tmp_path, package, inventory(), validation())
    old_sha = common.sha256_file(package["xlsx"])
    wb = openpyxl.load_workbook(package["xlsx"])
    wb.properties.keywords = "yeniden derlendi"  # yalniz dosya baytlari degisir, icerik ayni
    wb.save(package["xlsx"])
    assert common.sha256_file(package["xlsx"]) != old_sha
    run(tmp_path, package, inventory(), validation())
    log = _log(tmp_path)
    assert log.count("=== 2026-10-22") == 1 and old_sha not in log and common.sha256_file(package["xlsx"]) in log


def test_progress_log_keeps_foreign_text_and_legacy_entries(tmp_path, package):
    out = tmp_path / "out"
    out.mkdir()
    legacy = "Elle yazılmış başlık\n\n=== 2026-10-21 · eski.xlsx (v001) ===\nFaz: photos\nİndirilmiş fotoğraflar: 0\n\n"
    (out / "ILERLEME_KAYDI.txt").write_text("﻿" + legacy, encoding="utf-8")  # BOM'lu elle duzenlenmis dosya
    run(tmp_path, package, inventory(), validation())
    run(tmp_path, package, inventory(), validation())
    log = _log(tmp_path)
    assert log.startswith(legacy) and log.count("=== 2026-10-22") == 1 and "﻿" not in log


def _line(tmp_path, prefix):
    return next(ln for ln in _log(tmp_path).splitlines() if ln.startswith(prefix))


PACKAGE_SOURCES = "· pakette 3 (ölçüm: Kaynaklar sekmesindeki satır)"


@pytest.mark.parametrize("claimed,researched,mismatch", [  # V4 inceleme: manifest = kaynak_dizini satiri (pilot 408), pakette 397
    (352, "352 (manifest: kaynak dizini satırı)", []), (3, "3 (manifest: kaynak dizini satırı)", []),
    (2, "en az 3 (manifest 2 diyor; pakettekinden az olamaz)", ["sources_researched"]),  # pilotun eski 352 < 397 hatasi
    ("çok", "en az 3 (manifest çok diyor; sayı değil)", ["sources_researched"]),
])
def test_progress_sources_researched_vs_package_count(tmp_path, package, capsys, claimed, researched, mismatch):
    """Arastirilan kaynak sayisi manifest'ten (bu arac kaynak dizinini olcemez), pakettekiler olcumden; yalniz manifest
    pakettekinden azsa (ya da sayi degilse) celiski. Foto: klasordeki dosya, manifest ayniysa not yok."""
    _add_sources(package["xlsx"], ["s1", "s2", "s3"])
    inv = inventory()
    inv["folders"]["Fotoğraflar"]["files"] = 3
    manifest = tmp_path / "manifest.json"
    manifest.write_text(json.dumps({"phase": "deliver", "sources_researched": claimed, "photos_downloaded": 3}), encoding="utf-8")
    run(tmp_path, package, inv, validation(), "--manifest", str(manifest))
    assert _line(tmp_path, "Araştırılan resmî kaynaklar:") == f"Araştırılan resmî kaynaklar: {researched} {PACKAGE_SOURCES}"
    photos = _line(tmp_path, "İndirilmiş fotoğraflar:")
    assert photos.startswith("İndirilmiş fotoğraflar: 3 (ölçüm: Fotoğraflar klasör") and "manifest" not in photos
    assert _last_json(capsys)["manifest_mismatch"] == mismatch


def test_progress_counts_without_inventory_or_manifest_fields(tmp_path, package, capsys):
    manifest = tmp_path / "manifest.json"
    manifest.write_text(json.dumps({"phase": "photos", "photos_downloaded": 0}), encoding="utf-8")  # pilot: eski 'foto 0'
    run(tmp_path, package, None, None, "--manifest", str(manifest))
    photos = _line(tmp_path, "İndirilmiş fotoğraflar:")
    assert photos.startswith("İndirilmiş fotoğraflar: 2 (ölçüm: Excel Fotoğraflar") and "manifest 0" in photos
    assert _line(tmp_path, "Araştırılan resmî kaynaklar:") == ("Araştırılan resmî kaynaklar: belirtilmedi (manifest'te yok) "
                                                               "· pakette 0 (ölçüm: Kaynaklar sekmesindeki satır)")
    assert _last_json(capsys)["manifest_mismatch"] == ["photos_downloaded"]


def test_progress_log_same_day_legacy_entry_of_same_excel_is_replaced(tmp_path, package, capsys):
    """V4 inceleme (a): eski bicimli (SHA'siz) girdi ayni gun + ayni baslikla (Excel adi + surum) ayni Excel'in girdisidir."""
    out = tmp_path / "out"
    out.mkdir()
    name = package["xlsx"].name
    head = f"=== 2026-10-22 · {name} (v001) ==="
    earlier = f"=== 2026-10-21 · {name} (v001) ===\nFaz: photos\n\n"
    other = "=== 2026-10-22 · CITY-999_testville_v002.xlsx (v002) ===\nFaz: deliver\n\n"
    legacy = f"{head}\nFaz: deliver\nAraştırılan resmî kaynaklar: 408\n\n"
    (out / "ILERLEME_KAYDI.txt").write_text(earlier + other + legacy + legacy, encoding="utf-8")  # pilot: ayni girdi tekrarli
    run(tmp_path, package, inventory(), validation())
    log = _log(tmp_path)
    assert log.startswith(earlier + other) and log.count(head) == 1 and "kaynaklar: 408" not in log
    assert _last_json(capsys)["progress"] == "güncellendi"


def test_progress_log_keeps_hand_note_written_after_an_entry(tmp_path, package):
    """V4 inceleme (b): girdiden bos satirla ayrilan elle yazilmis not girdinin parcasi degildir; ayni gun yeniden kosu silmez."""
    run(tmp_path, package, inventory(), validation())
    path = tmp_path / "out" / "ILERLEME_KAYDI.txt"
    path.write_text(path.read_text(encoding="utf-8") + "NOT (elle): Barış ile konuşuldu.\n", encoding="utf-8")
    run(tmp_path, package, inventory(), validation())
    log = _log(tmp_path)
    assert log.count("NOT (elle): Barış ile konuşuldu.") == 1 and log.count("=== 2026-10-22") == 1


def test_progress_entry_stays_one_block_when_manifest_text_has_blank_lines(tmp_path, package):
    """Girdi ilk bos satirda biter: manifest metnindeki satir sonlari bosluga cevrilir, yoksa artik metin her kosuda birikir."""
    manifest = tmp_path / "manifest.json"
    manifest.write_text(json.dumps({"phase": "deliver", "next_step": "gönder\n\nsonra denetim", "sources_researched": "çok\n\nkaynakX",
                                    "photos_downloaded": "iki\n\nfotoX"}), encoding="utf-8")
    run(tmp_path, package, inventory(), validation(), "--manifest", str(manifest))
    run(tmp_path, package, inventory(), validation(), "--manifest", str(manifest))
    log = _log(tmp_path)
    assert [log.count(t) for t in ("sonra denetim", "kaynakX", "fotoX")] == [1, 1, 1]
    assert "Sonraki somut adım: gönder sonra denetim" in log and "manifest çok kaynakX diyor" in log and "manifest iki fotoX" in log


def test_progress_photo_folder_name_is_unicode_normalized(tmp_path, package):
    inv = inventory()
    inv["folders"] = {unicodedata.normalize("NFD", "Fotoğraflar"): {"files": 5, "bytes": 50}}  # macOS'ta NFD klasor adi
    run(tmp_path, package, inv, validation())
    assert "İndirilmiş fotoğraflar: 5 (ölçüm: Fotoğraflar klasör" in _log(tmp_path)


def test_unique_photos_use_sha_or_source_union(package):
    wb = openpyxl.load_workbook(package["xlsx"])
    ws = wb["Fotoğraflar"]
    keys = helpers.SHEET_KEYS["Fotoğraflar"]
    ws.cell(9, keys.index("sha256") + 1, ws.cell(8, keys.index("sha256") + 1).value)  # 2. foto ayni SHA -> tek kare
    wb.save(package["xlsx"])
    metrics = summary_metrics.collect(package["xlsx"], None, None)
    assert metrics["photos"]["rows"] == 2 and metrics["photos"]["venue_unique"] == 1 and metrics["photos"]["cover_unique"] == 0


def test_human_pending_counted_from_kanitlar(package):
    wb = openpyxl.load_workbook(package["xlsx"])
    ws = wb["Kanıtlar"]
    keys = helpers.SHEET_KEYS["Kanıtlar"]
    for r, (a, b, c) in enumerate([("Bekliyor", "Doğrulandı", "Bekliyor"), ("Tamam", "Bekliyor", "Bu kare için uygun")], start=8):
        ws.cell(r, keys.index("evidence_id") + 1, f"e{r}")
        for key, value in zip(("metadata_review_status", "pixel_review_status", "rights_status"), (a, b, c)):
            ws.cell(r, keys.index(key) + 1, value)
    wb.save(package["xlsx"])
    human = summary_metrics.collect(package["xlsx"], None, None)["human"]
    assert human["rows"] == 2 and human["pending_fields"] == 3 and human["invalid_values"] == 1


def test_sources_verified_without_snapshot(package):
    wb = openpyxl.load_workbook(package["xlsx"])
    ws = wb["Kaynaklar"]
    keys = helpers.SHEET_KEYS["Kaynaklar"]
    rows = [("s1", "Evet", None, None), ("s2", "Evet", "Teslimler/a.html", "ab" * 32), ("s3", "Hayır", None, None)]
    for r, (sid, ver, snap, sha) in enumerate(rows, start=8):
        for key, value in (("source_id", sid), ("verified", ver), ("snapshot_file", snap), ("snapshot_sha256", sha)):
            ws.cell(r, keys.index(key) + 1, value)
    wb.save(package["xlsx"])
    sources = summary_metrics.collect(package["xlsx"], None, None)["sources"]
    assert sources == {"total": 3, "verified_yes": 2, "with_snapshot": 1, "verified_without_snapshot": 1}


def test_places_without_photo_listed(tmp_path, package):
    wb = openpyxl.load_workbook(package["xlsx"])
    ws = wb.create_sheet("Mekânlar")
    ws.cell(7, 1, "entry_id")
    for r, eid in enumerate(["testville:yer-1", "testville:yer-9"], start=8):
        ws.cell(r, 1, eid)
    wb.save(package["xlsx"])
    metrics = summary_metrics.collect(package["xlsx"], None, None)
    assert metrics["places"] == {"total": 2, "with_photo": 1, "without_photo": ["testville:yer-9"]}


def test_usage_errors_exit_2(tmp_path, capsys):
    assert render_summary.main(["--xlsx", str(tmp_path / "yok.xlsx"), "--out-dir", str(tmp_path / "o")]) == 2
    assert "error" in json.loads(capsys.readouterr().out.strip().splitlines()[-1])


@pytest.mark.parametrize("date", ["22.10.2026", "bugün", "2026-13-01"])
def test_bad_date_exit_2_and_nothing_written(tmp_path, package, capsys, date):
    out = tmp_path / "o"
    assert render_summary.main(["--xlsx", str(package["xlsx"]), "--out-dir", str(out), "--date", date]) == 2
    assert "error" in _last_json(capsys) and not out.exists()


def test_unreadable_progress_log_exit_2_and_nothing_overwritten(tmp_path, package, capsys):
    out = tmp_path / "out"
    out.mkdir()
    (out / "ILERLEME_KAYDI.txt").write_bytes(b"=== 2026-10-21 \xff\xfe bozuk")  # UTF-8 degil
    assert render_summary.main(["--xlsx", str(package["xlsx"]), "--out-dir", str(out), "--date", "2026-10-22"]) == 2
    assert "ilerleme" in _last_json(capsys)["error"] and sorted(p.name for p in out.iterdir()) == ["ILERLEME_KAYDI.txt"]


def test_bad_json_input_exit_2(tmp_path, package, capsys):
    bad = tmp_path / "bad.json"
    bad.write_text("{", encoding="utf-8")
    code = render_summary.main(["--xlsx", str(package["xlsx"]), "--out-dir", str(tmp_path / "o"), "--inventory", str(bad)])
    assert code == 2


def test_city_code_and_version_parsing(package):
    metrics = summary_metrics.collect(package["xlsx"], None, None)
    assert metrics["city_code"] == "CITY-999" and metrics["version"] == "v001" and metrics["city_id"] == "testville"
