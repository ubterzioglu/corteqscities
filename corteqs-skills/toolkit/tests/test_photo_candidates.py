"""photo_candidates: Commons imageinfo normalizasyonu, kanit dosyasi yukleyici, Excel'den plan."""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

import photo_candidates as pc


def _info(**over) -> dict:
    ext = {
        "ObjectName": {"value": "Van Abbemuseum front"},
        "ImageDescription": {"value": "<p>The <b>museum</b> &amp; river</p>"},
        "Categories": {"value": "Van Abbemuseum|2019 in Eindhoven"},
        "DateTimeOriginal": {"value": "2019-05-04 10:00:00"},
        "DateTime": {"value": "2020-01-01 00:00:00"},
        "GPSLatitude": {"value": "51.4347"},
        "GPSLongitude": {"value": "5.4819"},
        "Artist": {"value": '<a href="x">Jane Doe</a>'},
        "LicenseShortName": {"value": "CC BY-SA 4.0"},
        "LicenseUrl": {"value": "https://creativecommons.org/licenses/by-sa/4.0"},
        "UsageTerms": {"value": "Creative Commons Attribution-Share Alike 4.0"},
        "Restrictions": {"value": ""},
        "AttributionRequired": {"value": "true"},
        "Copyrighted": {"value": "True"},
        "Credit": {"value": "<span>Own work</span>"},
    }
    info = {"width": 4000, "height": 3000, "size": 123, "mime": "image/jpeg", "url": "https://upload.example/a.jpg",
            "descriptionurl": "https://commons.wikimedia.org/wiki/File:A.jpg", "thumburl": "https://thumb.example/a.jpg",
            "sha1": "abc", "extmetadata": ext}
    info.update(over)
    return info


def test_strip_html_decodes_entities_and_collapses_space():
    assert pc.strip_html("<p>The <b>museum</b> &amp;\n river</p>") == "The museum & river"
    assert pc.strip_html(None) == ""


@pytest.mark.parametrize("raw,year", [
    ("2019-05-04 10:00:00", 2019), ("Tue Mar 20 14:18:00 CET 2018", 2018), ("29 September 2015, 13:54", 2015),
    ("1746", 1746), ("1966-05", 1966), ("16th of July", None), ("", None), (None, None),
])
def test_parse_year(raw, year):
    assert pc.parse_year(raw) == year


def test_normalize_imageinfo_flat_fields():
    cand = pc.normalize_imageinfo(_info(), candidate_id="c1")
    assert cand["candidate_id"] == "c1"
    assert cand["title"] == "Van Abbemuseum front"
    assert cand["description_text"] == "The museum & river"
    assert cand["categories"] == ["Van Abbemuseum", "2019 in Eindhoven"]
    assert cand["year"] == 2019 and cand["date_original"] == "2019-05-04 10:00:00"
    assert cand["lat"] == pytest.approx(51.4347) and cand["lon"] == pytest.approx(5.4819)
    assert (cand["width"], cand["height"], cand["mime"]) == (4000, 3000, "image/jpeg")
    assert cand["license"] == "CC BY-SA 4.0" and cand["author_text"] == "Jane Doe"
    assert cand["source_url"].endswith("File:A.jpg") and cand["original_url"].endswith("a.jpg")
    assert cand["thumb_url"] and cand["sha1"] == "abc" and cand["restrictions"] == ""


def test_normalize_imageinfo_missing_fields_and_fallbacks():
    info = _info(mime=None, url="https://upload.example/b.PNG")
    info["extmetadata"] = {"DateTime": {"value": "2001-02-03"}, "GPSLatitude": {"value": "x"}}
    cand = pc.normalize_imageinfo(info)
    assert cand["candidate_id"] == "https://commons.wikimedia.org/wiki/File:A.jpg"
    assert cand["title"] == "A" and cand["categories"] == []
    assert cand["year"] == 2001 and cand["lat"] is None and cand["lon"] is None
    assert cand["mime"] == "image/png"


def test_normalize_does_not_mutate_input():
    info = _info()
    before = json.dumps(info, sort_keys=True)
    pc.normalize_imageinfo(info)
    assert json.dumps(info, sort_keys=True) == before


def _write_evidence(folder: Path, photo_id: str, info: dict) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{photo_id.split(':')[-1]}.source-evidence.json"
    path.write_text(json.dumps({"photo_id": photo_id, "commons_file": "File:A.jpg", "source_page": info["descriptionurl"],
                                "original_file_url": info["url"], "downloaded_file_url": "https://thumb/x.jpg",
                                "api_imageinfo": info}), encoding="utf-8")
    return path


def _sheet(wb, name: str, keys: list[str], rows: list[dict]) -> None:
    ws = wb.create_sheet(name)
    for col, key in enumerate(keys, start=1):
        ws.cell(7, col, key)
    for r, row in enumerate(rows, start=8):
        for k, v in row.items():
            ws.cell(r, keys.index(k) + 1, v)


def _make_xlsx(root: Path) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Şehir", ["city_id", "name", "aliases", "latitude", "longitude"],
           [{"city_id": "tv", "name": "Testville", "aliases": "Testville;Testburg", "latitude": 51.43, "longitude": 5.5}])
    _sheet(wb, "Mekânlar", ["entry_id", "title", "group", "latitude", "longitude"],
           [{"entry_id": "tv:museum", "title": "Van Abbemuseum", "group": "Müzeler", "latitude": 51.4347, "longitude": 5.4819}])
    _sheet(wb, "Mahalleler", ["entry_id", "title"], [{"entry_id": "tv:mah-centrum", "title": "Centrum (Merkez)"}])
    _sheet(wb, "Fotoğraflar", ["photo_id", "role", "entry_id", "file"], [
        {"photo_id": "tv:foto:kapak-merkez-01", "role": "Kapak", "file": "F/a.webp"},
        {"photo_id": "tv:foto:museum-01", "role": "Mekân", "entry_id": "tv:museum", "file": "F/b.webp"},
        {"photo_id": "tv:foto:centrum-01", "role": "Mekân", "entry_id": "tv:mah-centrum", "file": "F/c.webp"},
    ])
    path = root / "CITY-999_testville_v001.xlsx"
    wb.save(path)
    return path


@pytest.fixture()
def evidence_pkg(tmp_path):
    root = tmp_path / "pkg"
    folder = root / "Lisans kanıtları"
    _write_evidence(folder, "tv:foto:museum-01", _info())
    _write_evidence(folder, "tv:foto:kapak-merkez-01", _info(sha1="def"))
    (folder / "broken.source-evidence.json").write_text("{not json", encoding="utf-8")
    _make_xlsx(root)
    return root


def test_load_evidence_dir_attaches_package_rows(evidence_pkg):
    cands, errors = pc.load_evidence_dir(evidence_pkg)
    assert len(cands) == 2 and len(errors) == 1 and "broken" in errors[0]
    by_id = {c["candidate_id"]: c for c in cands}
    museum = by_id["tv:foto:museum-01"]
    assert museum["entry_id"] == "tv:museum" and museum["photo_index"] == 2 and museum["role"] == "venue"
    cover = by_id["tv:foto:kapak-merkez-01"]
    assert cover["entry_id"] == "tv:kapak-merkez" and cover["role"] == "cover" and cover["photo_index"] == 1
    assert [c["photo_index"] for c in cands] == [1, 2]


def test_build_plan_from_xlsx(evidence_pkg):
    plan = pc.build_plan(next(evidence_pkg.glob("*.xlsx")))
    by_id = {row["entry_id"]: row for row in plan}
    museum = by_id["tv:museum"]
    assert museum["role"] == "venue" and museum["municipality"] == "Testville" and museum["group"] == "Müzeler"
    assert museum["lat"] == pytest.approx(51.4347) and "Testburg" in museum["municipality_aliases"]
    assert by_id["tv:mah-centrum"]["role"] == "district" and "Centrum" in by_id["tv:mah-centrum"]["aliases"]
    cover = by_id["tv:kapak-merkez"]
    assert cover["role"] == "cover" and cover["name"] == "Testville"
    assert museum["city_lat"] == pytest.approx(51.43)


def test_cli_from_evidence_and_plan(evidence_pkg, tmp_path, capsys):
    out = tmp_path / "c.jsonl"
    assert pc.main(["from-evidence", str(evidence_pkg), "--out", str(out)]) == 1  # bir bozuk dosya -> bulgu
    summary = json.loads(capsys.readouterr().out.strip())
    assert summary["candidates"] == 2 and summary["errors"] == 1
    lines = out.read_text(encoding="utf-8").splitlines()
    assert len(lines) == 2 and json.loads(lines[0])["candidate_id"]
    plan_out = tmp_path / "plan.jsonl"
    assert pc.main(["plan-from-xlsx", str(evidence_pkg), "--out", str(plan_out)]) == 0
    assert json.loads(capsys.readouterr().out.strip())["entries"] == 3


def test_from_search_keeps_entry_and_query_kind(tmp_path, capsys):
    rows = [{"entry_id": "tv:museum", "query_kind": "geosearch", "title": "File:A.jpg", "imageinfo": _info()},
            {"entry_id": "tv:museum", "query_kind": "category", "title": "File:B.jpg", "imageinfo": _info(sha1="b")},
            {"entry_id": "tv:museum", "title": "File:C.jpg"}]
    src = tmp_path / "search.jsonl"
    src.write_text("\n".join(json.dumps(r) for r in rows) + "\n", encoding="utf-8")
    cands, errors = pc.load_search_rows(src)
    assert [c["query_kind"] for c in cands] == ["geosearch", "category"] and len(errors) == 1
    assert cands[0]["entry_id"] == "tv:museum" and cands[0]["candidate_id"] == "tv:museum|File:A.jpg"
    out = tmp_path / "c.jsonl"
    assert pc.main(["from-search", str(src), "--out", str(out)]) == 1
    assert json.loads(capsys.readouterr().out.strip())["candidates"] == 2
    assert pc.main(["from-search", str(tmp_path / "nope.jsonl"), "--out", str(out)]) == 2


def test_cli_usage_errors(tmp_path, capsys):
    assert pc.main(["from-evidence", str(tmp_path / "missing"), "--out", str(tmp_path / "x.jsonl")]) == 2
    assert pc.main(["plan-from-xlsx", str(tmp_path), "--out", str(tmp_path / "p.jsonl")]) == 2
    with pytest.raises(SystemExit) as exc:
        pc.main([])
    assert exc.value.code == 2
