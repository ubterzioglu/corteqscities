"""rota_mesafe: ardisik durak bacaklari (kus ucusu + OSM yaya), onbellek, istek araligi, sorunlu bacaklar."""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import rota_mesafe as rm  # noqa: E402

PLACES = [("t:a", "A", 37.3861, -5.9926), ("t:b", "B", 37.3826, -5.9893), ("t:c", "C", 37.3800, -5.9850),
          ("t:x", "X", None, None)]


def _xlsx(path: Path, routes: list[tuple[str, str]]) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, keys, rows in (("Rotalar", ["route_id", "stop_entry_ids"], routes),
                              ("Mekânlar", ["entry_id", "title", "latitude", "longitude"], PLACES)):
        ws = wb.create_sheet(sheet)
        for col, key in enumerate(keys, start=1):
            ws.cell(6, col, key.upper())
            ws.cell(7, col, key)
        for r, values in enumerate(rows, start=8):
            for c, value in enumerate(values, start=1):
                ws.cell(r, c, value)
    wb.save(path)
    return path


class FakeOsrm:
    def __init__(self, fail_for: str | None = None):
        self.urls: list[str] = []
        self.fail_for = fail_for

    def __call__(self, url: str) -> dict:
        self.urls.append(url)
        if self.fail_for and self.fail_for in url:
            return {"code": "NoRoute", "routes": []}
        return {"code": "Ok", "routes": [{"distance": 562.8, "duration": 461.6}]}


def _run(capsys, tmp_path, routes, transport, sleeps=None):
    xlsx = _xlsx(tmp_path / "t.xlsx", routes)
    out = tmp_path / "rota_mesafe.jsonl"
    code = rm.main(["--xlsx", str(xlsx), "--out", str(out), "--cache-dir", str(tmp_path / "cache")], transport=transport,
                   sleep=(sleeps.append if sleeps is not None else lambda s: None), today="2026-10-26")
    summary = json.loads(capsys.readouterr().out.strip().splitlines()[-1])
    rows = [json.loads(l) for l in out.read_text(encoding="utf-8").splitlines()]
    return code, summary, rows


def test_haversine_known_distance():
    assert rm.haversine_m(37.3861, -5.9926, 37.3826, -5.9893) == pytest.approx(486, abs=5)


def test_legs_measured_and_cached(capsys, tmp_path):
    fake = FakeOsrm()
    code, s, rows = _run(capsys, tmp_path, [("t:route:001", "t:a\nt:b\nt:c")], fake)
    assert code == 0 and s["legs"] == 2 and s["measured"] == 2 and s["requests"] == 2
    assert rows[0] == {"route_id": "t:route:001", "bacak": 1, "from": "t:a", "to": "t:b", "from_title": "A",
                       "to_title": "B", "kaynak": rm.SOURCE, "sorgu_tarihi": "2026-10-26", "kus_ucusu_m": rows[0]["kus_ucusu_m"],
                       "yuruyus_m": 563, "yuruyus_dk": 8}
    assert "/foot/-5.992600,37.386100;-5.989300,37.382600" in fake.urls[0]  # OSRM: boylam,enlem sirasi
    code, s, _ = _run(capsys, tmp_path, [("t:route:001", "t:a\nt:b\nt:c")], fake)
    assert s["requests"] == 0 and s["cache_hits"] == 2 and len(fake.urls) == 2


def test_uncached_requests_are_spaced(capsys, tmp_path):
    sleeps: list[float] = []
    _run(capsys, tmp_path, [("t:route:001", "t:a, t:b, t:c")], FakeOsrm(), sleeps)
    assert len(sleeps) == 1 and 0 < sleeps[0] <= rm.MIN_INTERVAL_S


def test_missing_coordinates_and_failed_route_are_problems(capsys, tmp_path):
    fake = FakeOsrm(fail_for="-5.985000")
    code, s, rows = _run(capsys, tmp_path, [("t:route:001", "t:a\nt:x"), ("t:route:002", "t:b\nt:c")], fake)
    assert code == 1 and s["problems"] == 2
    assert rows[0]["sorun"] == "koordinatsiz durak" and "yuruyus_m" not in rows[0]
    assert rows[1]["sorun"] == "rota yok: NoRoute" and rows[1]["kus_ucusu_m"] > 0
    assert not list((tmp_path / "cache").glob("*.json"))  # basarisiz yanit saklanmaz


def test_usage_errors(capsys, tmp_path):
    assert rm.main(["--xlsx", str(tmp_path / "yok.xlsx"), "--out", str(tmp_path / "o.jsonl"),
                    "--cache-dir", str(tmp_path / "c")]) == 2
    wb = openpyxl.Workbook()
    wb.save(tmp_path / "bos.xlsx")
    assert rm.main(["--xlsx", str(tmp_path / "bos.xlsx"), "--out", str(tmp_path / "o.jsonl"),
                    "--cache-dir", str(tmp_path / "c")]) == 2
