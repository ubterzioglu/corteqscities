"""kanit_kapsami: 25 olcutun teslimdeki kanit paydasi ve dis denetim modeliyle alt sinir (Sevilla %44,7 dersi)."""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import kanit_kapsami as kk  # noqa: E402

SHEETS = {"Yanıtlar": ["entry_id"], "Mekânlar": ["entry_id"], "Mahalleler": ["entry_id"], "Rotalar": ["route_id"],
          "Sözlük": ["term_id"], "Fotoğraflar": ["photo_id", "entry_id", "role"], "SayfaAlanları": ["field_id"]}
ROWS = {"Yanıtlar": [{"entry_id": "t:y1"}, {"entry_id": "t:y2"}], "Mekânlar": [{"entry_id": "t:m1"}],
        "Mahalleler": [{"entry_id": "t:h1"}], "Rotalar": [{"route_id": "t:route:001"}], "Sözlük": [{"term_id": "t:s1"}],
        "Fotoğraflar": [{"photo_id": "t:p1", "entry_id": "t:m1", "role": "Mekân"}, {"photo_id": "t:p2", "entry_id": "t:h1", "role": "Kapak"}],
        "SayfaAlanları": [{"field_id": "t:f1"}]}


@pytest.fixture(autouse=True)
def small_targets(monkeypatch):
    monkeypatch.setattr(kk, "TARGETS", {k: 1 for k in kk.TARGETS})


def _xlsx(path: Path) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, keys in SHEETS.items():
        ws = wb.create_sheet(sheet)
        for col, key in enumerate(keys, start=1):
            ws.cell(6, col, key.upper())
            ws.cell(7, col, key)
        for row, record in enumerate(ROWS[sheet], start=8):
            for key, value in record.items():
                ws.cell(row, keys.index(key) + 1, value)
    wb.save(path)
    return path


def _write(path: Path, payload) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.suffix == ".jsonl":
        path.write_text("".join(json.dumps(r, ensure_ascii=False) + "\n" for r in payload), encoding="utf-8")
    else:
        path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def _check(structure_errors: list | None = None, media_errors: list | None = None, complete: bool = True) -> dict:
    """validate_all --mode emit ciktisinin (paket-kontrol-v7.json) ilgili alanlari."""
    return {"technically_complete": complete, "structure": {"valid": True, "errors": structure_errors or []},
            "media": {"structural_passed": True, "errors": media_errors or []}, "minimum_gaps": {}}


def _full_evidence(t: Path) -> None:
    _write(t / "iddia-defteri-ozet-v7.json", {
        "son_durum": {"t:y1": "temiz", "t:y2": "düzeltildi", "t:s1": "temiz", "t:m1": "temiz", "t:h1": "temiz"},
        "iddiasiz_kayit": {}, "denetlenmemis_kayit": {}, "kopyasiz_iddia": [], "duzelt_kayitsiz": [], "dogrulama_sorunlu": []})
    _write(t / "rota-plani-v7.jsonl", [{"route_id": "t:route:001", "durak": "A", "acik_mi": "açık"}])
    _write(t / "mekan-dogrulama-v7.jsonl", [{"entry_id": "t:m1", "sonuc": "geçti"}, {"entry_id": "t:h1", "sonuc": "geçti"}])
    _write(t / "foto-inceleme-v7.jsonl", [{"photo_id": "t:p1", "gorsel_sonuc": "keep", "ek_hak": "yok"},
                                          {"photo_id": "t:p2", "gorsel_sonuc": "keep", "ek_hak": "yok"}])
    # ureticilerin gercek bicimi (Sevilla 26.10.2026: sahte bicimle C4/C5 ve E1-E3 hep 0 kaliyordu, provenance listesi cokertiyordu)
    _write(t / "gorsel-kalite-v7.json", {"photos": [{"photo_id": "t:p1", "flags": []}, {"photo_id": "t:p2", "flags": []}]})
    _write(t / "paket-kontrol-v7.json", _check())
    _write(t / "photo-provenance-v7.json", [{"photo_id": "t:p1"}, {"photo_id": "t:p2"}])
    for name in ("TESLIM_OZETI.txt", "ACIK_KONULAR.txt"):
        (t / name).write_text("özet", encoding="utf-8")


def _run(tmp_path: Path, capsys, prepare=None) -> tuple[int, dict, dict]:
    pkg = tmp_path / "PKG"
    xlsx = _xlsx(pkg / "CITY-1_t_v001.xlsx") if (pkg.mkdir() or True) else None
    _full_evidence(pkg / "Teslimler")
    if prepare:
        prepare(pkg / "Teslimler")
    out = tmp_path / "OUT" / "kanit-kapsami.json"
    code = kk.main(["--xlsx", str(xlsx), "--pkg", str(pkg), "--out", str(out)])
    summary = json.loads(capsys.readouterr().out.strip().splitlines()[-1])
    return code, summary, json.loads(out.read_text(encoding="utf-8"))


def test_full_local_evidence_closes_everything_but_preview_stage(tmp_path, capsys):
    code, s, d = _run(tmp_path, capsys)
    status = {c["id"]: c["durum"] for c in d["olcutler"]}
    assert status["D5"] == status["E4"] == "onizleme"
    assert all(v == "tam" for k, v in status.items() if k not in {"D5", "E4"}), status
    assert d["alt_sinir"] == pytest.approx(95.0) and d["onizleme_ile_azami"] == pytest.approx(100.0)
    assert code == 0 and s["eksik"] == []


def test_missing_ledger_and_unaudited_cards_lower_the_bound(tmp_path, capsys):
    def prepare(t):
        _write(t / "iddia-defteri-ozet-v7.json", {"son_durum": {"t:y1": "temiz"}, "iddiasiz_kayit": {"Sözlük": ["t:s1"]},
                                                  "denetlenmemis_kayit": {"Yanıtlar": ["t:y2"]}, "kopyasiz_iddia": [],
                                                  "duzelt_kayitsiz": [], "dogrulama_sorunlu": []})
        (t / "rota-plani-v7.jsonl").unlink()
    code, s, d = _run(tmp_path, capsys, prepare)
    by = {c["id"]: c for c in d["olcutler"]}
    assert (by["A1"]["kanitli"], by["A1"]["payda"], by["A1"]["durum"]) == (1, 2, "kismi")
    assert by["B3"]["durum"] == "yok" and by["B4"]["durum"] == "yok"
    assert d["alt_sinir"] < 95 and code == 1 and set(s["eksik"]) >= {"A1", "B3", "B4"}


def test_flagged_photos_and_validator_errors_count_against(tmp_path, capsys):
    def prepare(t):
        _write(t / "gorsel-kalite-v7.json", {"photos": [{"photo_id": "t:p1", "flags": ["bulanik"]}, {"photo_id": "t:p2", "flags": []}]})
        _write(t / "paket-kontrol-v7.json", _check(structure_errors=[{"code": "x"}], media_errors=[{"code": "y"}]))
    code, s, d = _run(tmp_path, capsys, prepare)
    by = {c["id"]: c for c in d["olcutler"]}
    assert (by["C4"]["kanitli"], by["C4"]["payda"]) == (1, 2) and by["E1"]["durum"] == "yok" and code == 1


@pytest.mark.parametrize("check", [_check(complete=False), {"errors": 0}, ["liste"]])
def test_incomplete_or_unknown_validator_output_is_not_evidence(tmp_path, capsys, check):
    code, _, d = _run(tmp_path, capsys, lambda t: _write(t / "paket-kontrol-v7.json", check))
    by = {c["id"]: c for c in d["olcutler"]}
    assert by["E1"]["durum"] == by["E2"]["durum"] == by["E3"]["durum"] == "yok" and code == 1


def test_provenance_dict_form_still_read(tmp_path, capsys):
    code, _, d = _run(tmp_path, capsys,
                      lambda t: _write(t / "photo-provenance-v7.json", {"photos": [{"photo_id": "t:p1"}, {"photo_id": "t:p2"}]}))
    assert {c["id"]: c["durum"] for c in d["olcutler"]}["D1"] == "tam" and code == 0


def test_usage_error(tmp_path, capsys):
    assert kk.main(["--xlsx", str(tmp_path / "yok.xlsx"), "--pkg", str(tmp_path), "--out", str(tmp_path / "o.json")]) == 2
