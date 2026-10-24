"""inventory.py testleri: dosya envanteri, Excel capraz kontrolu, kurcalama senaryolari, salt-okunur davranis."""
from __future__ import annotations

import json
import os
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest
from PIL import Image

import inventory
from helpers import SHEET_KEYS, make_image


def _set_cell(xlsx: Path, sheet: str, row: int, key: str, value) -> None:
    wb = openpyxl.load_workbook(xlsx)
    wb[sheet].cell(row, SHEET_KEYS[sheet].index(key) + 1).value = value  # cell(..., None) degeri silmez
    wb.save(xlsx)


def _tree_state(root: Path) -> list[tuple[str, int, int]]:
    return sorted((str(p.relative_to(root)), p.stat().st_size, p.stat().st_mtime_ns) for p in root.rglob("*") if p.is_file())


def _run(capsys, package, tmp_path, *extra: str, with_xlsx: bool = True) -> tuple[int, dict, dict]:
    out = tmp_path / "out" / "envanter.json"
    argv = [str(package["root"]), "--out", str(out), *extra]
    if with_xlsx:
        argv += ["--xlsx", str(package["xlsx"])]
    code = inventory.main(argv)
    lines = capsys.readouterr().out.strip().splitlines()
    assert len(lines) == 1, "stdout tek satir olmali"
    details = json.loads(out.read_text(encoding="utf-8")) if out.exists() else {}
    return code, json.loads(lines[0]), details


def test_clean_package_exit_0(package, tmp_path, capsys):
    code, s, details = _run(capsys, package, tmp_path)
    assert code == 0
    # xlsx + 2 orijinal + 2 web (Teslimler bos)
    assert s["files"] == 5 and s["images"] == 4
    assert (s["missing"], s["sha_mismatch"], s["file_ne_original"], s["bad_ext"], s["ext_mismatch"], s["errors"]) == (0, 0, 0, 0, 0, 0)
    assert s["orphans"] == 0 and s["dup_sha"] == 0
    assert s["total_mb"] == round(details["totals"]["total_bytes"] / 1_000_000, 2)
    by_path = {f["path"]: f for f in details["files"]}
    orig = by_path[package["photos"][0]["orig_rel"]]
    assert (orig["format"], orig["width"], orig["height"]) == ("JPEG", 640, 480)
    assert orig["sha256"] == package["photos"][0]["orig_sha"]
    web = by_path[package["photos"][0]["web_rel"]]
    assert (web["format"], web["width"], web["height"]) == ("WEBP", 320, 240)
    assert details["folders"]["Fotoğraflar"]["files"] == 2
    assert details["folders"]["Lisans kanıtları"]["files"] == 2
    assert details["folders"]["."]["files"] == 1


def test_without_xlsx_only_file_level_checks(package, tmp_path, capsys):
    code, s, details = _run(capsys, package, tmp_path, with_xlsx=False)
    assert code == 0 and s["files"] == 5 and s["orphans"] == 0
    assert details["xlsx_checked"] is False


def test_missing_file(package, tmp_path, capsys):
    (package["root"] / package["photos"][0]["orig_rel"]).unlink()
    code, s, details = _run(capsys, package, tmp_path)
    assert code == 1 and s["missing"] == 1  # ayni yol iki yerde referans -> tekil sayim
    assert {m["path"] for m in details["findings"]["missing"]} == {package["photos"][0]["orig_rel"]}
    assert len(details["findings"]["missing"]) == 2


def test_modified_bytes_is_sha_mismatch(package, tmp_path, capsys):
    target = package["root"] / package["photos"][1]["orig_rel"]
    target.write_bytes(target.read_bytes() + b"\x00\x00")
    code, s, _ = _run(capsys, package, tmp_path)
    assert code == 1 and s["sha_mismatch"] == 1 and s["missing"] == 0


def test_bad_extension(package, tmp_path, capsys):
    web = package["root"] / package["photos"][0]["web_rel"]
    web.rename(web.with_name("testville_yer_01.jpg&utm=1"))
    code, s, details = _run(capsys, package, tmp_path)
    assert code == 1 and s["bad_ext"] == 1
    assert s["ext_mismatch"] == 0  # bad_ext ayri bulgu, cift sayilmaz
    assert s["missing"] == 1  # web_file artik yok
    assert details["findings"]["bad_ext"][0]["path"].endswith("testville_yer_01.jpg&utm=1")


def test_png_saved_as_jpg_is_ext_mismatch(package, tmp_path, capsys):
    target = package["root"] / package["photos"][1]["orig_rel"]
    make_image(target, "PNG", (64, 48))
    code, s, details = _run(capsys, package, tmp_path)
    assert code == 1 and s["ext_mismatch"] == 1
    assert details["findings"]["ext_mismatch"][0]["format"] == "PNG"


def test_file_ne_original(package, tmp_path, capsys):
    photo = package["photos"][0]
    _set_cell(package["xlsx"], "Fotoğraflar", 8, "file", photo["web_rel"])
    _set_cell(package["xlsx"], "Fotoğraflar", 8, "sha256", photo["web_sha"])
    code, s, details = _run(capsys, package, tmp_path)
    assert code == 1 and s["file_ne_original"] == 1
    assert s["sha_mismatch"] == 0 and s["missing"] == 0
    assert details["findings"]["file_ne_original"][0]["photo_id"] == photo["photo_id"]


def test_unsafe_path_is_error(package, tmp_path, capsys):
    _set_cell(package["xlsx"], "Fotoğraflar", 8, "file", "../kacak.jpg")
    code, s, details = _run(capsys, package, tmp_path)
    assert code == 1 and s["errors"] >= 1
    assert any(e["kind"] == "unsafe_path" for e in details["findings"]["errors"])


def test_orphans_and_exclusions(package, tmp_path, capsys):
    root = package["root"]
    make_image(root / "Fotoğraflar" / "yetim.webp", "WEBP", (16, 16), (1, 2, 3))
    (root / "Teslimler" / "paket.zip").write_bytes(b"zip")
    (root / "durum.json").write_text("{}", encoding="utf-8")
    code, s, details = _run(capsys, package, tmp_path)
    assert code == 0  # yetim dosya cikis kodunu etkilemez
    assert s["orphans"] == 1
    assert details["findings"]["orphans"] == ["Fotoğraflar/yetim.webp"]
    assert details["findings"]["orphans_excluded"] == {"teslimler": 1, "xlsx": 1, "json": 1}


def test_snapshot_file_reference_is_not_orphan(package, tmp_path, capsys):
    snap = package["root"] / "Kaynak anlık görüntüleri" / "k1.html"
    snap.parent.mkdir()
    snap.write_text("<html></html>", encoding="utf-8")
    _set_cell(package["xlsx"], "Kaynaklar", 8, "source_id", "k1")
    _set_cell(package["xlsx"], "Kaynaklar", 8, "snapshot_file", "Kaynak anlık görüntüleri/k1.html")
    _set_cell(package["xlsx"], "Kaynaklar", 9, "source_id", "k2")
    _set_cell(package["xlsx"], "Kaynaklar", 9, "snapshot_file", "../disari.html")  # sessizce yok sayilir, orphan listesini bozmaz
    code, s, _ = _run(capsys, package, tmp_path)
    assert code == 0 and s["orphans"] == 0


def test_duplicate_sha_groups(package, tmp_path, capsys):
    root = package["root"]
    src = root / package["photos"][0]["web_rel"]
    (root / "Fotoğraflar" / "kopya.webp").write_bytes(src.read_bytes())
    code, s, details = _run(capsys, package, tmp_path)
    assert code == 0 and s["dup_sha"] == 1
    assert len(details["findings"]["dup_sha"][0]["files"]) == 2


def test_corrupt_image_is_error(package, tmp_path, capsys):
    target = package["root"] / package["photos"][0]["web_rel"]
    target.write_bytes(b"bozuk bayt dizisi")
    code, s, details = _run(capsys, package, tmp_path)
    assert code == 1 and s["errors"] >= 1
    assert any(e["kind"] == "image_unreadable" for e in details["findings"]["errors"])


def test_non_image_with_unknown_ext_is_not_error(package, tmp_path, capsys):
    (package["root"] / "Fotoğraflar" / "notlar.txt").write_text("merhaba", encoding="utf-8")
    code, s, _ = _run(capsys, package, tmp_path)
    assert code == 0 and s["errors"] == 0 and s["images"] == 4


def test_jpeg_with_jpeg_ext_and_uppercase(package, tmp_path, capsys):
    path = package["root"] / "Fotoğraflar" / "BUYUK.JPEG"
    Image.new("RGB", (8, 8)).save(path, format="JPEG")
    code, s, _ = _run(capsys, package, tmp_path)
    assert s["ext_mismatch"] == 0 and s["bad_ext"] == 0 and s["images"] == 5


def test_empty_file_reference_is_error(package, tmp_path, capsys):
    _set_cell(package["xlsx"], "Fotoğraflar", 8, "file", None)
    code, s, details = _run(capsys, package, tmp_path)
    assert code == 1
    assert any(e["kind"] == "empty_reference" for e in details["findings"]["errors"])


def test_package_untouched_and_default_out_outside_package(package, tmp_path, capsys, monkeypatch):
    before = _tree_state(package["root"])
    workdir = tmp_path / "calisma"
    workdir.mkdir()
    monkeypatch.chdir(workdir)
    code = inventory.main([str(package["root"]), "--xlsx", str(package["xlsx"])])
    capsys.readouterr()
    assert code == 0
    assert (workdir / "dosya-envanteri.json").is_file()
    assert _tree_state(package["root"]) == before


def test_default_out_inside_package_refused(package, capsys, monkeypatch):
    monkeypatch.chdir(package["root"])
    code = inventory.main([str(package["root"])])
    out = json.loads(capsys.readouterr().out.strip())
    assert code == 2 and "error" in out
    assert not (package["root"] / "dosya-envanteri.json").exists()


def test_explicit_out_inside_package_allowed_and_not_counted(package, capsys):
    out = package["root"] / "dosya-envanteri.json"
    for _ in range(2):  # ikinci kosuda onceki cikti envantere girmemeli
        code = inventory.main([str(package["root"]), "--xlsx", str(package["xlsx"]), "--out", str(out)])
        s = json.loads(capsys.readouterr().out.strip())
        assert code == 0 and s["files"] == 5


def test_deterministic_output(package, tmp_path, capsys):
    _, first, d1 = _run(capsys, package, tmp_path)
    _, second, d2 = _run(capsys, package, tmp_path)
    assert first == second and d1 == d2


def test_bad_root_and_bad_xlsx_exit_2(package, tmp_path, capsys):
    code = inventory.main([str(tmp_path / "yok"), "--out", str(tmp_path / "o.json")])
    assert code == 2 and "error" in json.loads(capsys.readouterr().out.strip())
    bad = tmp_path / "bozuk.xlsx"
    bad.write_bytes(b"degil")
    code = inventory.main([str(package["root"]), "--xlsx", str(bad), "--out", str(tmp_path / "o.json")])
    assert code == 2 and "error" in json.loads(capsys.readouterr().out.strip())


def test_xlsx_without_photo_sheets_exit_2(package, tmp_path, capsys):
    wb = openpyxl.load_workbook(package["xlsx"])
    del wb["Fotoğraflar"]
    wb.save(package["xlsx"])
    code, s, _ = _run(capsys, package, tmp_path)
    assert code == 2 and "error" in s


def test_runs_as_script(package, tmp_path):
    out = tmp_path / "o.json"
    env = {**os.environ, "PYTHONIOENCODING": "utf-8"}
    proc = subprocess.run([sys.executable, inventory.__file__, str(package["root"]), "--xlsx", str(package["xlsx"]), "--out", str(out)],
                          capture_output=True, text=True, encoding="utf-8", env=env)
    assert proc.returncode == 0
    assert json.loads(proc.stdout)["files"] == 5


@pytest.mark.parametrize("ext,expected", [(".jpg", False), (".jpg&utm=1", True), (".org&utm_campaign=imageinfo&utm_content=original", True),
                                          ("", False), (".webp", False), (".jp g", True)])
def test_is_bad_ext(ext, expected):
    assert inventory.is_bad_ext(ext) is expected
