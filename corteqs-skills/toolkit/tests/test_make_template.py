"""make_template.py: Berlin orneginden bos sehir sablonu (zip/XML duzeyinde, kayipsiz).

Birim testler sentetik mini calisma kitabiyla (openpyxl + elle eklenen x14 extLst) calisir; entegrasyon testleri gercek
Berlin ornegini (salt okunur) Atinc'in Eindhoven sablonuyla (altin olcut) karsilastirir. Ciktilar yalniz tmp_path'e yazilir.
Sentetik ornek ve zip/XML inceleme yardimcilari make_template_helpers.py'de.
"""
from __future__ import annotations

import json
import re
import subprocess
import sys
import zipfile
from copy import copy as shallow
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
import pytest

import build_xlsx
import common
import make_template as mt
import xlsx_fidelity
from make_template_helpers import (DATA_FILL, DOCX_PARTS, MAIN, REL, SPACER_FILL, cells_of, content_types, leaked_tokens,
                                   make_base, make_tools, opc_problems, patch_part, read_part, sheet_parts, sheet_root)

# openpyxl x14 uzantilarini okurken atar ve uyarir; aracin zip/XML duzeyinde calismasinin nedeni budur (beklenen uyari)
pytestmark = pytest.mark.filterwarnings("ignore:.*extension is not supported and will be removed:UserWarning")

ROOT = Path(r"C:\temp_private\corteqs\0000_newera")
REAL_BASE = ROOT / "ekler" / "CITY-000_berlin_ornek_v001.xlsx"
REAL_ORACLE = ROOT / "ekler_eindhoven" / "CITY-039_eindhoven_v001.xlsx"
REAL_TOOLS = ROOT / "araclar" / "arac"
VALIDATE_ALL = Path(__file__).resolve().parent.parent / "scripts" / "validate_all.py"

NOTE_LOCAL = "Berlin örneğinden türetildi (make_template); içerik üretilmedi. YEREL DENEME; resmî katalog numarası değil."


def argv(base: Path, out: Path, tools: Path, number: str = "901", *extra: str) -> list[str]:
    return ["--base", str(base), "--out", str(out), "--city-id", "tilburg", "--city-name", "Tilburg", "--country", "Hollanda",
            "--country-code", "NL", "--region", "EU", "--city-number", number, "--updated-on", "2026-10-20",
            "--tools", str(tools), *extra]


def run(capsys, args: list[str]) -> tuple[int, dict]:
    code = mt.main(args)
    lines = capsys.readouterr().out.strip().splitlines()
    assert len(lines) == 1, lines
    return code, json.loads(lines[0])


@pytest.fixture()
def tools(tmp_path):
    return make_tools(tmp_path / "arac")


@pytest.fixture()
def base(tmp_path):
    return make_base(tmp_path)


@pytest.fixture()
def derived(tmp_path, base, tools, capsys):
    out = tmp_path / "out" / "CITY-901_tilburg_v001.xlsx"
    report = tmp_path / "out" / "rapor.json"
    code, summary = run(capsys, argv(base, out, tools, "901", "--report", str(report)))
    assert code == common.EXIT_OK, summary
    return {"out": out, "summary": summary, "report": json.loads(report.read_text(encoding="utf-8"))}


# --- birim testler (sentetik) ---------------------------------------------------------------------------

def test_keeps_only_contract_sheets_and_drops_sample_media(derived, base):
    out, summary = derived["out"], derived["summary"]
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Özet", "Teslimat", "Şehir", "Konular", "Yanıtlar", mt.PROVENANCE_SHEET]
    assert summary["removed_sheets"] == ["DBKayıtları"] and summary["sheets"] == 6
    assert summary["removed_media"] == 1 and summary["self_check"] == "ok"
    with zipfile.ZipFile(out) as archive:
        names = archive.namelist()
        types_xml = archive.read("[Content_Types].xml").decode("utf-8")
    assert not [n for n in names if n.startswith(("xl/media/", "xl/drawings/"))]
    assert "png" not in types_xml.casefold() and "drawing" not in types_xml
    # silinen sekmenin eski kaydi (sheet6/table5), o adi yeniden alan FotografKaynaklari'ninkiyle yinelenmez; yinelenen
    # kayit Excel'de "acilmiyor", .NET'te "same key" hatasidir (kabul SORUN-1)
    overrides = [key for tag, key, _ in content_types(out) if tag == "Override"]
    assert opc_problems(out) == [] and overrides.count("/xl/worksheets/sheet6.xml") == overrides.count("/xl/tables/table5.xml") == 1
    app = read_part(out, "docProps/app.xml")
    assert "DBKayıtları" not in app and "Microsoft Excel" in app  # silinen sekme adi listesi kalmaz, diger ozellikler kalir
    assert sorted(n for n in names if n.startswith("xl/worksheets/sheet")) == [f"xl/worksheets/sheet{i}.xml" for i in range(1, 7)]
    assert sorted(n for n in names if n.startswith("xl/tables/")) == [f"xl/tables/table{i}.xml" for i in range(1, 6)]
    assert "BerlinAlan" not in wb.defined_names and "KonuListesi" in wb.defined_names
    view = ET.fromstring(read_part(out, "xl/workbook.xml").encode("utf-8")).find(f".//{{{MAIN}}}workbookView")
    assert view is not None and view.get("activeTab", "0") == "0"  # aktif sekme silinen DBKayitlari'ni gostermez


def test_data_rows_emptied_styles_kept_heights_reset(derived):
    out = derived["out"]
    wb = openpyxl.load_workbook(out)
    for name in ("Konular", "Yanıtlar", mt.PROVENANCE_SHEET):
        ws = wb[name]
        assert list(common.iter_records(ws)) == [], name  # formul sutunlari kayit sayilmaz
        assert ws["A8"].fill.fgColor.rgb == DATA_FILL.fgColor.rgb  # bos hucrenin bicimi korunur
        assert ws.row_dimensions[8].height == 54  # icerige gore buyumus satir sablon yuksekligine doner
    assert wb["Konular"]["D9"].value.startswith("=IF(A9")
    konular = cells_of(sheet_root(out, "Konular"))
    assert konular["D8"].find(f"{{{MAIN}}}v").text == "0"  # bos sablonda formul onbellegi 0
    assert wb["Yanıtlar"].column_dimensions["A"].width == 30 and wb["Yanıtlar"].freeze_panes == "C8"


def test_tables_fit_capacity_with_dv_cf_and_x14_extensions(derived):
    out, summary = derived["out"], derived["summary"]
    wb = openpyxl.load_workbook(out)
    ws = wb["Yanıtlar"]
    assert {t.name: t.ref for t in ws.tables.values()} == {"CT_5": "A7:E11"}
    assert summary["tables_resized"] == {"Yanıtlar": "A7:E14->A7:E11"}
    assert ws.max_row == 11
    assert [str(dv.sqref) for dv in ws.data_validations.dataValidation] == ["D8:D11"]
    assert sorted(str(cf.sqref) for cf in ws.conditional_formatting) == ["A8:A11", "E8:E11"]
    assert {t.ref for t in wb[mt.PROVENANCE_SHEET].tables.values()} == {"A7:B12"}  # kapasitesi olmayan tablo aynen
    xml = read_part(out, sheet_parts(out)["Yanıtlar"])
    assert "<xm:sqref>E8:E11</xm:sqref>" in xml and "<xm:sqref>C8:C11</xm:sqref>" in xml  # x14 DV + CF korunur, kirpilir
    assert "x14:dataBar" in xml and "Konular!$A$8:$A$10" in xml
    assert 'xmlns:x14ac="http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac"' in xml  # mc:Ignorable oneki kaybolmaz
    table_xml = read_part(out, "xl/tables/table4.xml")  # sekme sirasiyla: Teslimat 1, Sehir 2, Konular 3, Yanitlar 4
    assert 'ref="A7:E11"' in table_xml and '<autoFilter ref="A7:E11"' in table_xml.replace("x:autoFilter", "autoFilter")


def test_cross_sheet_formula_ranges_follow_target_table(derived):
    out = derived["out"]
    wb = openpyxl.load_workbook(out)
    assert all("'Yanıtlar'!$E$8:$E$11" in wb["Konular"][f"D{r}"].value for r in range(8, 11))
    assert wb["Özet"]["B6"].value == "=COUNTA('Yanıtlar'!$A$8:$A$11)"  # kucultulen tabloya izleyen aralik
    assert wb["Özet"]["B7"].value == "=SUM('Konular'!$D$8:$D$10)"  # zaten dogru aralik degismez
    assert derived["summary"]["formulas_refit"] == 4


def test_ozet_header_notes_and_cached_counts(derived):
    out, summary = derived["out"], derived["summary"]
    wb = openpyxl.load_workbook(out)
    ws = wb["Özet"]
    assert ws["A3"].value == "Tilburg · Hollanda · CITY-901"
    assert ws["A9"].value == "Dosyayı doldurma sırası"
    assert summary["ozet_removed"] == ["A12", "A15"]
    assert [str(r) for r in ws.merged_cells.ranges] == ["A9:F9"]
    assert ws.max_row == 15  # son notun tasan satirlari (16-17) atilir
    for row in range(12, 16):
        assert all(ws.cell(row, c).value is None for c in range(1, 7)), row
        assert ws.cell(row, 1).fill.fgColor.rgb == SPACER_FILL.fgColor.rgb and ws.row_dimensions[row].height == 30
    cached = {ref: c.find(f"{{{MAIN}}}v").text for ref, c in cells_of(sheet_root(out, "Özet")).items() if c.find(f"{{{MAIN}}}f") is not None}
    assert cached == {"B6": "0", "E6": "3", "F6": "4", "B7": "0", "E7": "2", "F7": "2"}


def style_of(cell) -> tuple:
    """openpyxl StyleProxy'leri birbirine esit cikmaz; karsilastirma icin asil nesnelerin kopyasi."""
    return (shallow(cell.font), shallow(cell.fill), shallow(cell.border), shallow(cell.alignment), cell.number_format,
            shallow(cell.protection))


def row8(path: Path, sheet: str) -> dict:
    ws = openpyxl.load_workbook(path)[sheet]
    keys = common.sheet_keys(ws)
    return {k: common.normalize_value(ws.cell(8, c).value) for k, c in keys.items()}


def test_teslimat_and_sehir_identity_rows(derived):
    out, summary = derived["out"], derived["summary"]
    assert row8(out, "Teslimat") == {"schema_version": "1.0", "delivery_id": "city-901-v001", "city_id": "tilburg",
                                     "editor": "Claude Code city-pack", "reviewer": None, "scope": "Tam teslim",
                                     "updated_on": "2026-10-20", "editorial_status": "Taslak", "notes": NOTE_LOCAL}
    assert row8(out, "Şehir") == {"city_id": "tilburg", "name": "Tilburg", "country": "Hollanda", "country_code": "NL",
                                  "region": "EU", "kind": "city", "latitude": None, "longitude": None, "source_position": 901,
                                  "priority": None, "aliases": None, "editorial_status": None}
    assert summary["written"] == {"Teslimat": ["schema_version", "delivery_id", "city_id", "editor", "scope", "updated_on",
                                               "editorial_status", "notes"],
                                  "Şehir": ["city_id", "name", "country", "country_code", "region", "kind", "source_position"]}
    items = " | ".join(summary["open_items"])
    assert "priority" in items and "yerel deneme" in items.casefold()
    assert derived["report"]["city_required_blank"] == ["latitude", "longitude", "priority", "editorial_status"]


def test_catalog_number_below_900_has_no_local_trial_note(tmp_path, base, tools, capsys):
    out = tmp_path / "CITY-039_tilburg_v001.xlsx"
    code, summary = run(capsys, argv(base, out, tools, "39"))
    assert code == common.EXIT_OK, summary
    values = row8(out, "Teslimat")
    assert values["delivery_id"] == "city-039-v001"
    assert values["notes"] == "Berlin örneğinden türetildi (make_template); içerik üretilmedi."
    assert row8(out, "Şehir")["source_position"] == 39
    assert openpyxl.load_workbook(out)["Özet"]["A3"].value == "Tilburg · Hollanda · CITY-039"
    assert "yerel deneme" not in " ".join(summary["open_items"]).casefold()


def test_no_sample_tokens_anywhere_and_shared_strings_compacted(derived):
    out = derived["out"]
    notes_cell = next(w["cell"] for w in derived["report"]["written_cells"] if w["sheet"] == "Teslimat" and w["key"] == "notes")
    assert leaked_tokens(out, {("Teslimat", notes_cell)}) == []
    sst = read_part(out, "xl/sharedStrings.xml")
    assert "Berlin sorusu" not in sst and "berlin:" not in sst


def test_base_read_only_and_existing_out_needs_overwrite(tmp_path, base, tools, capsys):
    before = common.sha256_file(base)
    out = tmp_path / "out.xlsx"
    out.write_bytes(b"eski")
    code, summary = run(capsys, argv(base, out, tools))
    assert code == common.EXIT_USAGE and "--overwrite" in summary["error"] and out.read_bytes() == b"eski"
    code, summary = run(capsys, argv(base, out, tools, "901", "--overwrite"))
    assert code == common.EXIT_OK and zipfile.is_zipfile(out)
    code, summary = run(capsys, argv(base, base, tools, "901", "--overwrite"))
    assert code == common.EXIT_USAGE and "--base" in summary["error"]
    for inside in (base.parent / "yeni.xlsx", base.parent / "alt" / "yeni.xlsx"):  # ornek klasoru ve alti salt okunur
        code, summary = run(capsys, argv(base, inside, tools))
        assert code == common.EXIT_USAGE and "klasör" in summary["error"] and not inside.exists()
    assert common.sha256_file(base) == before


@pytest.mark.parametrize("flag,case", [*[("--report", c) for c in ("base", "out", "sample_dir", "not_json", "is_dir", "parent_is_file")],
                                       ("--out", "is_dir"), ("--out", "parent_is_file"), ("--base", "base")])
def test_bad_output_paths_are_usage_errors(tmp_path, base, tools, capsys, flag, case):
    """--out/--report ornegi, sablonu, ornek klasorunu, bir klasoru ezmez; yol hatasinda rapor da yazilmaz (cikis 2, JSON)."""
    before = common.sha256_file(base)
    (tmp_path / "dosya").write_text("x", encoding="utf-8")
    out, report = tmp_path / "out" / "CITY-901_tilburg_v001.xlsx", tmp_path / "rapor.json"
    suffix = ".json" if flag == "--report" else ".xlsx"
    target = {"base": base, "out": out, "sample_dir": base.parent / "alt" / f"x{suffix}", "not_json": tmp_path / "rapor.txt",
              "is_dir": tmp_path / f"klasor{suffix}", "parent_is_file": tmp_path / "dosya" / f"x{suffix}"}[case]
    if case == "is_dir":
        target.mkdir()
    out, report = (target, report) if flag == "--out" else (out, target) if flag == "--report" else (base, base)
    code, summary = run(capsys, argv(base, out, tools, "901", "--overwrite", "--report", str(report)))
    assert code == common.EXIT_USAGE and summary["out"] is None and flag in summary["error"], summary
    assert common.sha256_file(base) == before and (tmp_path / "dosya").read_text(encoding="utf-8") == "x"
    assert not (tmp_path / "out").exists() and not (tmp_path / "rapor.json").exists() and not list(tmp_path.glob("**/*.tmp-*"))
    if case == "is_dir":
        assert not any(target.iterdir())
    elif case != "base":
        assert not target.exists()


def test_output_is_deterministic(tmp_path, base, tools, capsys):
    first, second = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    assert run(capsys, argv(base, first, tools))[0] == common.EXIT_OK
    assert run(capsys, argv(base, second, tools))[0] == common.EXIT_OK
    assert common.sha256_file(first) == common.sha256_file(second)


@pytest.mark.parametrize("change,needle", [
    (("--city-id", "Tilburg"), "--city-id"),
    (("--country-code", "nl"), "--country-code"),
    (("--region", "EUR"), "--region"),
    (("--city-number", "0"), "--city-number"),
    (("--city-number", "1000"), "--city-number"),
    (("--city-number", "13"), "örnek"),
    (("--updated-on", "2999-01-01"), "--updated-on"),
    (("--updated-on", "20.10.2026"), "--updated-on"),
    (("--city-name", "=HYPERLINK(1)"), "--city-name"),
])
def test_bad_arguments_are_usage_errors(tmp_path, base, tools, capsys, change, needle):
    report = tmp_path / "rapor.json"  # yol hatasi degil: rapor hatayla yazilir (yol hatasinda yazilmaz, ustteki test)
    args = argv(base, tmp_path / "out.xlsx", tools, "901", "--report", str(report))
    args[args.index(change[0]) + 1] = change[1]
    code, summary = run(capsys, args)
    assert code == common.EXIT_USAGE and needle in summary["error"], summary
    assert not (tmp_path / "out.xlsx").exists() and json.loads(report.read_text(encoding="utf-8"))["exit_code"] == code


def test_missing_tools_or_contract_sheet_is_usage_error(tmp_path, base, tools, capsys, monkeypatch):
    monkeypatch.delenv("CORTEQS_TOOLS", raising=False)
    args = argv(base, tmp_path / "out.xlsx", tools)
    del args[args.index("--tools"):args.index("--tools") + 2]
    code, summary = run(capsys, args)
    assert code == common.EXIT_USAGE and "--tools" in summary["error"]
    monkeypatch.setenv("CORTEQS_TOOLS", str(tools))
    assert run(capsys, args)[0] == common.EXIT_OK
    other = tmp_path / "eksik"
    other.mkdir()
    no_prov = make_base(other, "no_provenance")
    code, summary = run(capsys, argv(no_prov, tmp_path / "out2.xlsx", tools))
    assert code == common.EXIT_USAGE and mt.PROVENANCE_SHEET in summary["error"]


def test_unsupported_sheet_content_is_usage_error(tmp_path, tools, capsys):
    base = make_base(tmp_path, "hyperlink")
    code, summary = run(capsys, argv(base, tmp_path / "out.xlsx", tools))
    assert code == common.EXIT_USAGE and "hyperlinks" in summary["error"]
    assert not (tmp_path / "out.xlsx").exists()


@pytest.mark.parametrize("damage", ["docx", "truncated_workbook", "latin1_workbook", "missing_sheet_part"])
def test_unreadable_base_is_usage_error(tmp_path, tools, capsys, damage):
    """Okunamayan/beklenmedik ornek: hata dokumu yok, cikis 2 ve tek satir JSON (cikis 1 'oz denetim bulgusu' demektir)."""
    base = make_base(tmp_path)
    if damage == "docx":
        with zipfile.ZipFile(base, "w") as archive:
            for name, text in DOCX_PARTS.items():
                archive.writestr(name, text)
    elif damage == "truncated_workbook":
        patch_part(base, "xl/workbook.xml", lambda xml: xml[: len(xml) // 2])
    elif damage == "latin1_workbook":
        patch_part(base, "xl/workbook.xml", lambda data: data.replace("Özet".encode(), "Özet".encode("latin-1")), binary=True)
    else:
        patch_part(base, "xl/_rels/workbook.xml.rels", lambda xml: xml.replace("worksheets/sheet3.xml", "worksheets/yok.xml"))
    out = tmp_path / "out.xlsx"
    code, summary = run(capsys, argv(base, out, tools))
    assert code == common.EXIT_USAGE and summary["error"] and not out.exists(), summary


def test_self_check_finding_blocks_output(tmp_path, tools, capsys):
    base = make_base(tmp_path, "header_token")  # satir 6 basligi korunur; icinde ornek belirteci kalir
    out = tmp_path / "out.xlsx"
    code, summary = run(capsys, argv(base, out, tools))
    assert code == common.EXIT_FINDINGS and not out.exists() and summary["out"] is None
    assert any("Konular" in f and "berlin" in f.casefold() for f in summary["self_check"]), summary
    assert not list(tmp_path.glob("*.tmp*"))


def test_self_check_blocks_broken_package(tmp_path, base, tools, capsys, monkeypatch):
    """Oz denetim paket duzeyine bakar (Excel ve .NET buna takilir): yinelenen kayit, parcasi olmayan kayit, yinelenen
    uzanti, turu olmayan parca, hedefi olmayan iliski. Biri bile varsa sablon yazilmaz (cikis 1)."""
    to_bytes = mt.Package.to_bytes

    def broken(pkg):
        types = ('<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/xml" /><Override '
                 'PartName="/xl/worksheets/sheet99.xml" ContentType="application/xml" /><Default Extension="XML" '
                 'ContentType="application/xml" /></Types>')
        pkg.data["[Content_Types].xml"] = pkg.data["[Content_Types].xml"].replace(b"</Types>", types.encode())
        rel = f'<Relationship Id="rIdX" Type="{REL}/worksheet" Target="worksheets/sheet77.xml" /></Relationships>'
        pkg.data["xl/_rels/workbook.xml.rels"] = pkg.data["xl/_rels/workbook.xml.rels"].replace(b"</Relationships>", rel.encode())
        pkg.infos["xl/media/foto.bin"], pkg.data["xl/media/foto.bin"] = zipfile.ZipInfo("xl/media/foto.bin"), b"\0"
        return to_bytes(pkg)

    monkeypatch.setattr(mt.Package, "to_bytes", broken)
    out = tmp_path / "out.xlsx"
    code, summary = run(capsys, argv(base, out, tools))
    assert code == common.EXIT_FINDINGS and not out.exists(), summary
    found = [f for f in summary["self_check"] if f.startswith("paket:")]
    text = " | ".join(found)
    assert len(found) == 5 and all(n in text for n in ("[Content_Types].xml", "sheet1.xml", "sheet99.xml", "foto.bin", "sheet77.xml"))


def test_sample_tokens_match_turkish_capitals(tmp_path, tools, capsys):
    assert mt.find_tokens("BERLİN NOTU", ["berlin"]) == ["berlin"]
    base = make_base(tmp_path, "upper_case")  # Ozet notu "BERLİN TAM ..." (Turkce buyuk İ); buyuk harfli parca kaydi
    out = tmp_path / "out.xlsx"
    code, summary = run(capsys, argv(base, out, tools))
    assert code == common.EXIT_OK and summary["ozet_removed"] == ["A12", "A15"], summary
    assert leaked_tokens(out, {("Teslimat", "I8")}) == []
    assert ("Override", "/xl/worksheets/sheet2.xml") in [(tag, key) for tag, key, _ in content_types(out)]


def test_template_works_with_build_xlsx(tmp_path, derived, capsys):
    data = tmp_path / "data"
    data.mkdir()
    record = {"entry_id": "tilburg:soru-1", "city_id": "tilburg", "title": "Soru", "featured": "Evet", "topic_id": "tilburg:planlama"}
    (data / "Yanıtlar.jsonl").write_text(json.dumps(record, ensure_ascii=False) + "\n", encoding="utf-8")
    (data / "Şehir.jsonl").write_text(json.dumps({"latitude": 51.56, "longitude": 5.09}) + "\n", encoding="utf-8")
    out = tmp_path / "ara.xlsx"
    code = build_xlsx.main(["--template", str(derived["out"]), "--data-dir", str(data), "--out", str(out)])
    summary = json.loads(capsys.readouterr().out.strip().splitlines()[-1])
    assert code == common.EXIT_OK and summary["fidelity"] == "ok", summary
    assert summary["rows_written"] == {"Yanıtlar": 1, "Şehir": 1}


def test_help_describes_interface(capsys):
    with pytest.raises(SystemExit) as exc:
        mt.main(["--help"])
    assert exc.value.code == 0
    text = capsys.readouterr().out
    for flag in ("--base", "--out", "--city-id", "--city-number", "--updated-on", "--tools", "--report", "--overwrite"):
        assert flag in text


def test_code_is_split_into_three_files_under_the_line_limit():
    """V2 bolmesi: zip/XML + yeniden numaralama (make_template_zip), sayfa/Ozet/formul (make_template_sheets), komut
    satiri/sozlesme/oz denetim (make_template). Eski adlar make_template'ten ayni nesneyle erisilir (monkeypatch ayni
    sinifa gider); betikler, bu test dosyasi ve yardimcilari kodlama sinirinin (800 satir) altinda."""
    import make_template_sheets as sheets
    import make_template_zip as zipxml

    assert mt.Package is zipxml.Package and mt.TemplateError is zipxml.TemplateError and mt.find_tokens is zipxml.find_tokens
    assert mt.Sheet is sheets.Sheet and mt.Book is sheets.Book and mt.Stats is sheets.Stats
    assert zipxml.renumber.__module__ == "make_template_zip" and sheets.write_cached.__module__ == "make_template_sheets"
    paths = [Path(module.__file__) for module in (mt, zipxml, sheets)]
    for path in paths + [Path(__file__), Path(__file__).with_name("make_template_helpers.py")]:
        assert len(path.read_text(encoding="utf-8").splitlines()) < 800, path.name


# --- entegrasyon: gercek Berlin ornegi -> Atinc'in Eindhoven sablonu ------------------------------------

EIND_ARGS = ["--city-id", "eindhoven", "--city-name", "Eindhoven", "--country", "Hollanda", "--country-code", "NL",
             "--region", "EU", "--city-number", "39", "--updated-on", "2026-10-12"]
ATINC_ONLY = {"Teslimat": {"delivery_id", "editor", "notes"},
              "Şehir": {"latitude", "longitude", "coordinate_source_url", "coordinate_checked_on", "source_url",
                        "source_checked_on", "priority", "aliases", "editorial_status"}}


@pytest.fixture(scope="module")
def real(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("real")
    out = tmp / "CITY-039_eindhoven_v001.xlsx"
    report = tmp / "rapor.json"
    before = common.sha256_file(REAL_BASE)
    code = mt.main(["--base", str(REAL_BASE), "--out", str(out), "--tools", str(REAL_TOOLS), "--report", str(report), *EIND_ARGS])
    payload = json.loads(report.read_text(encoding="utf-8"))
    assert code == common.EXIT_OK, payload["summary"]
    assert common.sha256_file(REAL_BASE) == before
    return {"out": out, "report": payload, "tmp": tmp,
            "ours": openpyxl.load_workbook(out), "atinc": openpyxl.load_workbook(REAL_ORACLE)}


def test_real_summary_matches_measured_differences(real):
    summary = real["report"]["summary"]
    assert summary["removed_sheets"] == ["DBKayıtları", "FotoAlbümü"] and summary["sheets"] == 17
    assert summary["removed_media"] == 187 and summary["self_check"] == "ok"
    assert summary["tables_resized"] == {"EkDetaylar": "A7:G381->A7:G207"}
    assert summary["ozet_removed"] == ["A35", "A39"]
    assert summary["formulas_refit"] == 30  # Konular!H8:H37 -> 'Yanıtlar'!$N$8:$N$307
    with zipfile.ZipFile(real["out"]) as ours, zipfile.ZipFile(REAL_ORACLE) as atinc:
        assert sorted(ours.namelist()) == sorted(atinc.namelist())  # 56 parca, ayni adlar (medya/cizim yok)


def test_real_content_types_equal_atinc(real):
    """Excel ve .NET'in acabilmesi icin: her parca tek kayit; kayit listesi (sirasi dahil) Atinc'inkiyle birebir ayni."""
    assert opc_problems(real["out"]) == []
    assert content_types(real["out"]) == content_types(REAL_ORACLE)
    assert opc_problems(REAL_ORACLE) == []  # olcut dosya da ayni kurala uyar (kuralin kendisi dogru)


def test_real_structure_equals_atinc_template(real):
    ours, atinc = real["ours"], real["atinc"]
    assert ours.sheetnames == atinc.sheetnames
    for name in atinc.sheetnames:
        a, b = atinc[name], ours[name]
        width = max(a.max_column, b.max_column)
        for row in range(1, common.KEY_ROW + 1):
            assert [b.cell(row, c).value for c in range(1, width + 1)] == [a.cell(row, c).value for c in range(1, width + 1)], (name, row)
        assert {t.name: t.ref for t in b.tables.values()} == {t.name: t.ref for t in a.tables.values()}, name
        assert sorted((str(d.sqref), d.formula1) for d in b.data_validations.dataValidation) == \
            sorted((str(d.sqref), d.formula1) for d in a.data_validations.dataValidation), name
        assert sorted(str(cf.sqref) for cf in b.conditional_formatting) == sorted(str(cf.sqref) for cf in a.conditional_formatting), name
        assert sorted(map(str, b.merged_cells.ranges)) == sorted(map(str, a.merged_cells.ranges)), name
        assert b.freeze_panes == a.freeze_panes, name
    ozet_a = {c.coordinate: c.value for row in atinc["Özet"].iter_rows() for c in row if c.value is not None}
    ozet_b = {c.coordinate: c.value for row in ours["Özet"].iter_rows() for c in row if c.value is not None}
    assert ozet_b == ozet_a
    for name in ("Fotoğraflar", "Konular"):
        assert xlsx_fidelity.formula_cells(ours[name]) == xlsx_fidelity.formula_cells(atinc[name]), name
    assert len(xlsx_fidelity.formula_cells(ours["Fotoğraflar"])) == 1200


def test_real_toolkit_fidelity_and_styles_match_atinc(real):
    ours, atinc = real["ours"], real["atinc"]
    snap = lambda wb, path: {"sheets": wb.sheetnames, "defined names": sorted(wb.defined_names), "zip parts": sorted(zipfile.ZipFile(path).namelist()),
                             "per_sheet": {ws.title: xlsx_fidelity.sheet_snapshot(ws) for ws in wb.worksheets}}
    diffs, _ = xlsx_fidelity.compare_snapshots(snap(atinc, REAL_ORACLE), snap(ours, real["out"]))
    assert diffs == []  # satir yukseklikleri, sutun genislikleri, tum formuller, DV/CF, tablolar dahil
    seen: dict[tuple, bool] = {}
    for ws_a in atinc.worksheets:
        ws_b = ours[ws_a.title]
        for row in ws_a.iter_rows():
            for cell_a in row:
                cell_b = ws_b[cell_a.coordinate]
                key = (type(cell_a), type(cell_b), tuple(cell_a._style or ()), tuple(cell_b._style or ()))  # MergedCell: _style None
                if key not in seen:
                    seen[key] = style_of(cell_a) == style_of(cell_b)
                assert seen[key], (ws_a.title, cell_a.coordinate)


def test_real_identity_rows(real):
    ours, atinc = real["ours"], real["atinc"]
    for sheet, own in ATINC_ONLY.items():
        a = {k: common.normalize_value(atinc[sheet].cell(8, c).value) for k, c in common.sheet_keys(atinc[sheet]).items()}
        b = {k: common.normalize_value(ours[sheet].cell(8, c).value) for k, c in common.sheet_keys(ours[sheet]).items()}
        assert {k: v for k, v in b.items() if k not in own} == {k: v for k, v in a.items() if k not in own}, sheet
    teslimat = {k: ours["Teslimat"].cell(8, c).value for k, c in common.sheet_keys(ours["Teslimat"]).items()}
    assert teslimat["delivery_id"] == "city-039-v001" and teslimat["editor"] == "Claude Code city-pack"
    assert teslimat["notes"] == "Berlin örneğinden türetildi (make_template); içerik üretilmedi."
    sehir = {k: ours["Şehir"].cell(8, c).value for k, c in common.sheet_keys(ours["Şehir"]).items()}
    assert sehir["source_position"] == 39 and sehir["latitude"] is None and sehir["priority"] is None
    assert sehir["aliases"] is None and sehir["editorial_status"] is None


def test_real_template_feeds_build_xlsx(real, capsys):
    data = real["tmp"] / "data"
    data.mkdir(exist_ok=True)
    (data / "Şehir.jsonl").write_text(json.dumps({"latitude": 51.4416, "longitude": 5.4697}) + "\n", encoding="utf-8")
    topic = {"topic_id": "eindhoven:planlama", "city_id": "eindhoven", "slug": "planlama", "title": "Planlama",
             "description": "Taşınmadan önce", "icon": "book", "order": 1}
    (data / "Konular.jsonl").write_text(json.dumps(topic, ensure_ascii=False) + "\n", encoding="utf-8")
    code = build_xlsx.main(["--template", str(real["out"]), "--data-dir", str(data), "--out", str(real["tmp"] / "ara.xlsx"),
                            "--tools", str(REAL_TOOLS)])
    summary = json.loads(capsys.readouterr().out.strip().splitlines()[-1])
    assert code == common.EXIT_OK and summary["fidelity"] == "ok", summary  # sonraki adim (KUR) sablonu sorunsuz kullanir
    assert summary["rows_written"] == {"Şehir": 1, "Konular": 1}


def test_real_no_berlin_tokens_and_official_validator_reads_it(real):
    notes_cell = next(w["cell"] for w in real["report"]["written_cells"] if w["sheet"] == "Teslimat" and w["key"] == "notes")
    assert leaked_tokens(real["out"], {("Teslimat", notes_cell)}) == []
    pkg = real["tmp"] / "paket"
    pkg.mkdir(exist_ok=True)
    results = {}
    for label, path in (("ours", real["out"]), ("atinc", REAL_ORACLE)):
        proc = subprocess.run([sys.executable, str(VALIDATE_ALL), str(path), str(pkg), "--tools", str(REAL_TOOLS), "--out",
                               str(real["tmp"] / f"val_{label}.json")], capture_output=True, text=True, encoding="utf-8", timeout=120)
        assert proc.returncode == common.EXIT_FINDINGS, proc.stdout + proc.stderr  # bos sablon: bulgu var, yapi hatasi yok
        results[label] = set(json.loads((real["tmp"] / f"val_{label}.json").read_text(encoding="utf-8"))["structure"]["errors"])
    assert list(pkg.iterdir()) == []  # audit modu salt okunur
    only_atinc = results["atinc"] - results["ours"]
    only_ours = results["ours"] - results["atinc"]
    assert len(only_atinc) == 1 and "Teslimat kimliği" in next(iter(only_atinc))  # pilot B10: buyuk harf delivery_id
    labels = {f["key"]: f["label"] for f in json.loads((REAL_TOOLS / "base-schema.json").read_text(encoding="utf-8"))["tables"]["city"]["fields"]}
    expected = {labels[k] for k in ("latitude", "longitude", "coordinate_source_url", "coordinate_checked_on", "source_url",
                                    "source_checked_on", "priority", "editorial_status")}
    assert {re.sub(r"^Şehir!\?: (.*) zorunlu\.$", r"\1", e) for e in only_ours} == expected, only_ours
