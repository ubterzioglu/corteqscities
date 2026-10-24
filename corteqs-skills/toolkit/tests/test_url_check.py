"""url_check: URL cikarma, bicim (K49), HTTP siniflama (K50-K52), satirlara geri dagitim (K56). Ag yok."""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

import helpers
import url_check as uc
from test_web_fetch import FakeTransport, ok, redirect, status
from web_fetch import NetworkError, RawResponse

LINK_KEYS = ["action_id", "city_id", "entry_id", "label", "url", "checked_on", "order"]
FIELD_KEYS = ["field_id", "city_id", "key", "label", "value_type", "value", "reviewed_on"]
SOURCE_KEYS = helpers.SHEET_KEYS["Kaynaklar"]
DEEP_HTML = b"<html><head><title>Inschrijven | Gemeente</title></head><body><main><h1>Inschrijven</h1><p>Uitleg</p></main></body></html>"


def _add_sheet(wb, name: str, keys: list[str]):
    ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
    for col, key in enumerate(keys, start=1):
        ws.cell(7, col, key)
    return ws


def _fill(ws, keys: list[str], rows: list[dict]) -> None:
    for offset, row in enumerate(rows):
        for key, value in row.items():
            ws.cell(8 + offset, keys.index(key) + 1, value)


def make_workbook(path: Path, sources: list[str], links: list[tuple[str, str]], fields: list[tuple[str, str]] = (),
                  city_url: str | None = None) -> Path:
    """Kaynaklar (url listesi), Baglantilar ((etiket, url)), SayfaAlanlari ((anahtar, deger)) dolu mini Excel."""
    helpers.make_template(path)
    wb = openpyxl.load_workbook(path)
    _fill(wb["Kaynaklar"], SOURCE_KEYS, [{"source_id": f"t:kaynak:{i}", "city_id": "testville", "title": f"K{i}",
                                          "url": u} for i, u in enumerate(sources, start=1)])
    links_ws = _add_sheet(wb, "Bağlantılar", LINK_KEYS)
    _fill(links_ws, LINK_KEYS, [{"action_id": f"t:action:{i}", "city_id": "testville", "entry_id": f"t:entry-{i}",
                                 "label": label, "url": u, "order": i} for i, (label, u) in enumerate(links, start=1)])
    fields_ws = _add_sheet(wb, "SayfaAlanları", FIELD_KEYS)
    _fill(fields_ws, FIELD_KEYS, [{"field_id": f"t:field:{k}", "city_id": "testville", "key": k, "value": v} for k, v in fields])
    if city_url:
        city = wb["Şehir"]
        col = city.max_column + 1
        city.cell(7, col, "source_url")
        city.cell(8, col, city_url)
    wb.save(path)
    return path


def run(tmp_path, xlsx: Path, routes: dict, *extra: str, capsys=None) -> tuple[int, dict, dict, FakeTransport]:
    fake = FakeTransport(routes)
    out = tmp_path / "url-checks.json"
    code = uc.main([str(xlsx), "--out", str(out), "--cache-dir", str(tmp_path / "cache"), *extra],
                   transport=fake, sleep=lambda s: None, clock=lambda: "2026-10-22")
    summary = json.loads(capsys.readouterr().out) if capsys else {}
    return code, json.loads(out.read_text(encoding="utf-8")), summary, fake


def by_url(report: dict) -> dict:
    return {rec["ham_URL"]: rec for rec in report["urls"]}


# ---------------------------------------------------------------- URL cikarma

def test_collect_relations_default_sheets(tmp_path):
    xlsx = make_workbook(tmp_path / "w.xlsx", ["https://a.example.org/x", "https://b.example.org/y"],
                         [("Başvur", "https://a.example.org/x"), ("Bilgi", "https://c.example.org/z")])
    relations = uc.collect_relations(xlsx, uc.DEFAULT_SHEETS)
    assert len(relations) == 4 and len({r["url"] for r in relations}) == 3
    first = relations[0]
    assert (first["sheet"], first["row"], first["id"], first["field"]) == ("Kaynaklar", 8, "t:kaynak:1", "url")
    link = next(r for r in relations if r["sheet"] == "Bağlantılar")
    assert (link["id"], link["label"], link["entry_id"], link["row"]) == ("t:action:1", "Başvur", "t:entry-1", 8)


def test_collect_relations_optional_sheets(tmp_path):
    quick = json.dumps([{"label": "Gemeente", "url": "https://g.example.org/"}, {"label": "IND", "url": "https://i.example.org/en"}])
    xlsx = make_workbook(tmp_path / "w.xlsx", [], [], fields=[("quick_links", quick), ("emergency_url", "https://e.example.org/112")],
                         city_url="https://s.example.org/eindhoven")
    relations = uc.collect_relations(xlsx, ("Şehir", "SayfaAlanları"))
    found = {(r["sheet"], r["url"]) for r in relations}
    assert found == {("Şehir", "https://s.example.org/eindhoven"), ("SayfaAlanları", "https://g.example.org/"),
                     ("SayfaAlanları", "https://i.example.org/en"), ("SayfaAlanları", "https://e.example.org/112")}
    city = next(r for r in relations if r["sheet"] == "Şehir")
    assert city["id"] == "testville" and city["field"] == "source_url"


def test_collect_relations_unknown_sheet_raises(tmp_path):
    xlsx = make_workbook(tmp_path / "w.xlsx", [], [])
    with pytest.raises(ValueError):
        uc.collect_relations(xlsx, ("Yok",))


# ---------------------------------------------------------------- K49 bicim

@pytest.mark.parametrize("url, expected", [
    ("https://www.ind.nl/en/residence-permits", set()),
    ("http://www.ind.nl/en/x", {"not_https"}),
    ("https://user:pw@x.example.org/a", {"credentials"}),
    ("https://x.example.org/page?sessionid=abc", {"session_param"}),
    ("https://x.example.org/page?access_token=abc", {"session_param"}),
    ("https://x.example.org/a;jsessionid=XYZ", {"session_param"}),
    ("https://www.google.com/search?q=eindhoven", {"search_results"}),
    ("https://x.example.org/zoeken?q=afval", {"search_results"}),
    ("https://www.google.com/maps/search/?api=1&query=Evoluon", {"map_search"}),
    ("https://maps.google.com/?q=Evoluon", {"map_search"}),
    ("https://www.openstreetmap.org/search?query=Evoluon", {"map_search"}),
    ("https://example.com/", {"placeholder"}),
    ("https://localhost/a", {"placeholder"}),
    (" https://x.example.org/a", {"whitespace"}),
    ("https://x.example.org/straße", {"non_ascii"}),
    ("https://x.example.org/a?utm_source=nl", {"tracking_param"}),
    ("ftp://x.example.org/", {"bad_scheme"}),
    ("geen url", {"bad_scheme"}),
    ("https:///pad", {"no_host"}),
])
def test_format_issues(url, expected):
    assert set(uc.format_issues(url)) == expected


@pytest.mark.parametrize("url, expected", [
    ("https://x.example.org", True), ("https://x.example.org/", True), ("https://x.example.org/en", True),
    ("https://x.example.org/nl/", True), ("https://x.example.org/home", True), ("https://x.example.org/home.html", True),
    ("https://x.example.org/index.php", True), ("https://x.example.org/en/visa", False),
    ("https://x.example.org/?page=2", False), ("https://x.example.org/contact", False),
    ("https://x.example.org/nl-nl", True), ("https://x.example.org/en_GB/", True), ("https://x.example.org/de", True),
    ("https://www.example.nl/ov", False), ("https://gemeente.nl/tv", False), ("https://x.example.org/ov-nl", False),
])
def test_is_homepage(url, expected):
    assert uc.is_homepage(url) is expected


@pytest.mark.parametrize("title, text, hit", [
    ("Pagina niet gevonden | Gemeente", "", True),
    ("Gemeente", "Helaas, deze pagina bestaat niet meer.", True),
    ("Error 404", "", True),
    ("Page not found", "", True),
    ("Afval", "Bel 040-238 4040 voor info. Pagina 404 van het archief? nee", False),
    ("Inschrijven", "Uitleg over inschrijven.", False),
])
def test_soft_404_signal(title, text, hit):
    assert (uc.soft_404_signal(title, text) is not None) is hit


# ---------------------------------------------------------------- siniflama + geri dagitim

CASES = {
    "https://ok.example.org/a": ok(DEEP_HTML),
    "https://moved.example.org/a": redirect("https://moved.example.org/b"),
    "https://moved.example.org/b": ok(DEEP_HTML),
    "https://gone.example.org/404": status(404),
    "https://gone.example.org/410": status(410),
    "https://shield.example.org/a": status(403),
    "https://busy.example.org/a": status(429, Retry_After="600"),
    "https://down.example.org/a": status(500),
    "https://down.example.org/b": status(502),
    "https://tls.example.org/a": NetworkError("tls", "certificate verify failed"),
    "https://soft.example.org/a": ok(b"<html><head><title>Pagina niet gevonden</title></head><body>x</body></html>"),
    "https://home.example.org/": ok(DEEP_HTML),
    "https://nl.wikipedia.org/wiki/Eindhoven": ok(DEEP_HTML),
    "https://auth.example.org/a": status(401),
}


def _case_workbook(tmp_path) -> Path:
    urls = [u for u in CASES if u != "https://moved.example.org/b"]
    sources = urls[:7]
    links = [("Bilgi", u) for u in urls[7:]] + [("Resmî bilgi", "https://nl.wikipedia.org/wiki/Eindhoven"),
                                                 ("Bilgi", "https://gone.example.org/404")]
    return make_workbook(tmp_path / "w.xlsx", sources, links)


def test_classes_verdicts_and_flags(tmp_path, capsys):
    code, report, summary, fake = run(tmp_path, _case_workbook(tmp_path), CASES, capsys=capsys)
    recs = by_url(report)
    classes = {u: r["class"] for u, r in recs.items()}
    assert classes == {
        "https://ok.example.org/a": "ok_200", "https://moved.example.org/a": "redirect_ok",
        "https://gone.example.org/404": "not_found_404", "https://gone.example.org/410": "gone_410",
        "https://shield.example.org/a": "forbidden_403", "https://busy.example.org/a": "rate_limited_429",
        "https://down.example.org/a": "server_5xx", "https://down.example.org/b": "server_5xx",
        "https://tls.example.org/a": "network_error", "https://soft.example.org/a": "soft_404",
        "https://home.example.org/": "ok_200", "https://nl.wikipedia.org/wiki/Eindhoven": "ok_200",
        "https://auth.example.org/a": "other_4xx"}
    verdict = {u: r["erişim_sonucu"] for u, r in recs.items()}
    assert verdict["https://gone.example.org/404"] == verdict["https://soft.example.org/a"] == "bozuk"
    assert verdict["https://shield.example.org/a"] == verdict["https://tls.example.org/a"] == "erişim teyidi açık"
    assert verdict["https://busy.example.org/a"] == verdict["https://down.example.org/a"] == "erişim teyidi açık"
    assert verdict["https://ok.example.org/a"] == verdict["https://moved.example.org/a"] == "çalışıyor"
    assert "homepage_only" in recs["https://home.example.org/"]["flags"]
    assert {"wikipedia", "wikipedia_labeled_official"} <= set(recs["https://nl.wikipedia.org/wiki/Eindhoven"]["flags"])
    assert recs["https://tls.example.org/a"]["error_kind"] == "tls"
    assert summary["verdicts"] == {"çalışıyor": 4, "bozuk": 3, "erişim teyidi açık": 6, "test edilmedi": 0, "geçersiz": 0}
    assert summary["classes"]["server_5xx"] == 2 and summary["flags"]["homepage_only"] == 1
    assert code == uc.EXIT_FINDINGS


def test_url_record_matches_section16_fields(tmp_path, capsys):
    _, report, _, _ = run(tmp_path, _case_workbook(tmp_path), CASES, capsys=capsys)
    rec = by_url(report)["https://moved.example.org/a"]
    assert set(uc.URL_RECORD_KEYS) <= set(rec)
    assert list(uc.URL_RECORD_KEYS) == ["ham_URL", "istek_URL", "bağlı_IDler", "Excel_satırları", "checked_on", "HTTP_durumu",
                                        "son_URL", "sayfa_başlığı", "erişim_sonucu", "anlamsal_hedef_sonucu", "tekrar_kaydı", "açıklama"]
    assert rec["son_URL"] == "https://moved.example.org/b" and rec["HTTP_durumu"] == 200
    assert rec["sayfa_başlığı"] == "Inschrijven | Gemeente" and rec["checked_on"] == "2026-10-22"
    assert [h["status"] for h in rec["redirect_chain"]] == [301, 200]
    busy = by_url(report)["https://busy.example.org/a"]
    assert busy["tekrar_kaydı"] == [{"status": 429, "retry_after_seconds": 600.0, "gave_up": True}]


def test_results_are_distributed_to_every_linked_row(tmp_path, capsys):
    _, report, summary, _ = run(tmp_path, _case_workbook(tmp_path), CASES, capsys=capsys)
    rec = by_url(report)["https://gone.example.org/404"]
    assert rec["Excel_satırları"] == ["Bağlantılar!15", "Kaynaklar!10"]
    assert rec["bağlı_IDler"] == ["t:action:8", "t:kaynak:3"]
    rows = [r for r in report["relations"] if r["url"] == "https://gone.example.org/404"]
    assert len(rows) == 2 and all(r["class"] == "not_found_404" and r["erişim_sonucu"] == "bozuk" for r in rows)
    assert summary["unique_urls"] == 13 and summary["relations"] == 15
    assert report["summary"]["relation_classes"]["not_found_404"] == 2


def test_errors_grouped_by_domain_and_class(tmp_path, capsys):
    _, report, _, _ = run(tmp_path, _case_workbook(tmp_path), CASES, capsys=capsys)
    groups = report["summary"]["error_groups"]
    assert groups["down.example.org"] == {"server_5xx": 2}
    assert groups["gone.example.org"] == {"gone_410": 1, "not_found_404": 1}
    assert "ok.example.org" not in groups


def test_clean_workbook_exits_zero_with_single_json_line(tmp_path, capsys):
    xlsx = make_workbook(tmp_path / "w.xlsx", ["https://ok.example.org/a"], [("Başvur", "https://ok.example.org/b")])
    code, _, _, _ = run(tmp_path, xlsx, {"https://ok.example.org/a": ok(DEEP_HTML), "https://ok.example.org/b": ok(DEEP_HTML)})
    lines = capsys.readouterr().out.strip().splitlines()
    assert code == uc.EXIT_OK and len(lines) == 1
    summary = json.loads(lines[0])
    assert summary["status"] == "ok" and summary["unique_urls"] == 2 and summary["tested"] == 2


def test_format_issue_alone_is_a_finding(tmp_path, capsys):
    xlsx = make_workbook(tmp_path / "w.xlsx", ["http://ok.example.org/a"], [])
    code, report, _, _ = run(tmp_path, xlsx, {"http://ok.example.org/a": ok(DEEP_HTML)}, capsys=capsys)
    assert code == uc.EXIT_FINDINGS and report["urls"][0]["format_issues"] == ["not_https"]


def test_invalid_url_is_not_fetched(tmp_path, capsys):
    xlsx = make_workbook(tmp_path / "w.xlsx", ["ftp://files.example.org/a"], [])
    code, report, summary, fake = run(tmp_path, xlsx, {}, capsys=capsys)
    assert fake.calls == [] and report["urls"][0]["class"] == "invalid_url"
    assert report["urls"][0]["erişim_sonucu"] == "geçersiz" and code == uc.EXIT_FINDINGS


def test_max_urls_reports_untested_openly(tmp_path, capsys):
    xlsx = make_workbook(tmp_path / "w.xlsx", ["https://a.example.org/1", "https://b.example.org/2", "https://c.example.org/3"], [])
    code, report, summary, fake = run(tmp_path, xlsx, {u: ok(DEEP_HTML) for u in ("https://a.example.org/1",)},
                                      "--max-urls", "1", capsys=capsys)
    assert summary["tested"] == 1 and summary["verdicts"]["test edilmedi"] == 2 and len(fake.calls) == 1
    assert report["summary"]["coverage"] == "1/3" and code == uc.EXIT_FINDINGS


def test_second_run_is_identical_and_served_from_cache(tmp_path, capsys):
    xlsx = _case_workbook(tmp_path)
    _, first, s1, _ = run(tmp_path, xlsx, CASES, capsys=capsys)
    out_bytes = (tmp_path / "url-checks.json").read_bytes()
    _, second, s2, fake2 = run(tmp_path, xlsx, {}, capsys=capsys)
    assert (tmp_path / "url-checks.json").read_bytes() == out_bytes
    assert fake2.calls == [] and s2["requests_made"] == 0 and s2["cache_hits"] >= 13 and s1["requests_made"] > 0


def test_labeled_official_flag_stays_on_its_own_labeled_relation(tmp_path, capsys):
    wiki = "https://nl.wikipedia.org/wiki/Eindhoven"
    xlsx = make_workbook(tmp_path / "w.xlsx", [wiki], [("Resmî bilgi", wiki), ("Bilgi", wiki)])
    _, report, _, _ = run(tmp_path, xlsx, {wiki: ok(DEEP_HTML)}, capsys=capsys)
    assert report["urls"][0]["flags"] == ["wikipedia", "wikipedia_labeled_official"]
    flags = {(r["sheet"], r.get("label")): r["flags"] for r in report["relations"]}
    assert flags[("Bağlantılar", "Resmî bilgi")] == ["wikipedia", "wikipedia_labeled_official"]
    assert flags[("Bağlantılar", "Bilgi")] == ["wikipedia"] and flags[("Kaynaklar", None)] == ["wikipedia"]
    assert report["summary"]["relation_flags"] == {"Bağlantılar": {"wikipedia": 2, "wikipedia_labeled_official": 1},
                                                   "Kaynaklar": {"wikipedia": 1}}


def test_retry_transient_refetches_cached_server_errors(tmp_path, capsys):
    url = "https://down.example.org/a"
    xlsx = make_workbook(tmp_path / "w.xlsx", [url], [])
    assert run(tmp_path, xlsx, {url: status(500)}, capsys=capsys)[1]["urls"][0]["class"] == "server_5xx"
    assert run(tmp_path, xlsx, {url: ok(DEEP_HTML)}, capsys=capsys)[1]["urls"][0]["class"] == "server_5xx"
    code, report, _, _ = run(tmp_path, xlsx, {url: ok(DEEP_HTML)}, "--retry-transient", capsys=capsys)
    assert report["urls"][0]["class"] == "ok_200" and code == uc.EXIT_OK


def test_head_only_mode_uses_head_and_skips_body_checks(tmp_path, capsys):
    xlsx = make_workbook(tmp_path / "w.xlsx", ["https://soft.example.org/a"], [])
    _, report, _, fake = run(tmp_path, xlsx, {("HEAD", "https://soft.example.org/a"): ok(b""),
                                              ("GET", "https://soft.example.org/a"): CASES["https://soft.example.org/a"]},
                             "--head-only", capsys=capsys)
    assert [c[0] for c in fake.calls] == ["HEAD"] and report["urls"][0]["class"] == "ok_200"


def test_redirect_to_homepage_is_flagged(tmp_path, capsys):
    xlsx = make_workbook(tmp_path / "w.xlsx", ["https://r.example.org/deep/page"], [])
    _, report, _, _ = run(tmp_path, xlsx, {"https://r.example.org/deep/page": redirect("https://r.example.org/"),
                                           "https://r.example.org/": ok(DEEP_HTML)}, capsys=capsys)
    assert "redirected_to_homepage" in report["urls"][0]["flags"] and report["urls"][0]["class"] == "redirect_ok"


def test_redirect_to_ip_block_page_is_not_working(tmp_path, capsys):
    url = "https://dbre.example.org/home.html"
    block = "http://20.23.2.223:6080/php/urlblock.php?args=x&url=https://dbre.example.org"
    xlsx = make_workbook(tmp_path / "w.xlsx", [url], [])
    _, report, _, _ = run(tmp_path, xlsx, {url: redirect(block, 302), block: ok(b"<title>Web Page Blocked</title>")}, capsys=capsys)
    rec = report["urls"][0]
    assert rec["class"] == "intercepted" and rec["erişim_sonucu"] == "erişim teyidi açık"
    assert "IP" in rec["açıklama"] and report["summary"]["error_groups"]["dbre.example.org"] == {"intercepted": 1}


def test_map_search_is_reported_but_not_a_finding(tmp_path, capsys):
    url = "https://www.google.com/maps/search/?api=1&query=Evoluon"
    xlsx = make_workbook(tmp_path / "w.xlsx", [], [("Haritada aç", url)])
    code, report, _, _ = run(tmp_path, xlsx, {url: ok(DEEP_HTML)}, capsys=capsys)
    assert report["urls"][0]["format_issues"] == ["map_search"] and report["summary"]["format_issues"] == {"map_search": 1}
    assert code == uc.EXIT_OK


def test_redirect_loop_and_tls_notes_explain_the_open_result(tmp_path, capsys):
    loop, tls = "https://loop.example.org/", "https://tls.example.org/a"
    xlsx = make_workbook(tmp_path / "w.xlsx", [loop, tls], [])
    routes = {loop: redirect("https://loop.example.org/csq/", 307), "https://loop.example.org/csq/": redirect(loop, 308),
              tls: NetworkError("tls", "URLError: certificate verify failed: unable to get local issuer certificate")}
    _, report, _, _ = run(tmp_path, xlsx, routes, capsys=capsys)
    recs = by_url(report)
    assert recs[loop]["erişim_sonucu"] == "erişim teyidi açık" and "çerez" in recs[loop]["açıklama"]
    assert "ara sertifika" in recs[tls]["açıklama"]


@pytest.mark.parametrize("argv", [["yok.xlsx"], ["{xlsx}", "--sheets", "Yok"], ["{xlsx}", "--max-urls", "0"]])
def test_usage_errors_exit_2(tmp_path, capsys, argv):
    xlsx = make_workbook(tmp_path / "w.xlsx", [], [])
    args = [a.replace("{xlsx}", str(xlsx)) for a in argv]
    assert uc.main(args + ["--out", str(tmp_path / "o.json")], transport=FakeTransport({})) == uc.EXIT_USAGE


def test_input_workbook_is_not_modified(tmp_path, capsys):
    xlsx = _case_workbook(tmp_path)
    before = xlsx.read_bytes()
    run(tmp_path, xlsx, CASES, capsys=capsys)
    assert xlsx.read_bytes() == before


def test_main_raw_response_headers_case_insensitive(tmp_path, capsys):
    xlsx = make_workbook(tmp_path / "w.xlsx", ["https://ok.example.org/a"], [])
    route = {"https://ok.example.org/a": RawResponse(200, {"content-type": "text/html"}, DEEP_HTML)}
    _, report, _, _ = run(tmp_path, xlsx, route, capsys=capsys)
    assert report["urls"][0]["sayfa_başlığı"] == "Inschrijven | Gemeente"
