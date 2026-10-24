"""Commons imageinfo kaydini duz bir foto adayi sozlugune cevirir; kanit dosyalarini ve Excel planini okur.

Ag erisimi yok. Kullanim:
  photo_candidates.py from-evidence <paket_kok> --out candidates.jsonl
  photo_candidates.py from-search <commons_search.jsonl> --out candidates.jsonl
  photo_candidates.py plan-from-xlsx <paket_kok> --out venues.jsonl
Stdout: tek satir JSON ozet. Cikis: 0 temiz, 1 okunamayan kanit dosyasi var, 2 kullanim hatasi.
"""
from __future__ import annotations

import argparse
import html
import json
import re
import sys
from pathlib import Path
from typing import Any, Iterable

import openpyxl

from common import EXIT_FINDINGS, EXIT_OK, EXIT_USAGE, emit, iter_records, setup_utf8

EVIDENCE_DIR = "Lisans kanıtları"
EVIDENCE_GLOB = "*.source-evidence.json"
ROLE_FROM_EXCEL = {"Mekân": "venue", "Kapak": "cover"}
MIME_BY_EXT = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".webp": "image/webp",
               ".gif": "image/gif", ".tif": "image/tiff", ".tiff": "image/tiff", ".svg": "image/svg+xml"}
_TAG = re.compile(r"<[^>]+>")
_SPACE = re.compile(r"\s+")
_YEAR = re.compile(r"(?<!\d)(1[0-9]{3}|20[0-9]{2})(?!\d)")


def strip_html(value: Any) -> str:
    """HTML etiketlerini atar, entity'leri cozer, bosluklari teke indirir."""
    if value is None:
        return ""
    return _SPACE.sub(" ", html.unescape(_TAG.sub(" ", str(value)))).strip()


def parse_year(value: Any) -> int | None:
    """Serbest bicimli Commons tarihinden ilk 4 haneli yili cikarir ('Tue Mar 20 ... 2018', '1966-05', '1746')."""
    if value is None:
        return None
    match = _YEAR.search(str(value))
    return int(match.group(1)) if match else None


def _float(value: Any) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _int(value: Any) -> int | None:
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _guess_mime(url: str) -> str | None:
    suffix = Path(url.split("?", 1)[0]).suffix.lower()
    return MIME_BY_EXT.get(suffix)


def _title_from_url(url: str) -> str:
    name = url.rsplit("File:", 1)[-1] if "File:" in url else url.rsplit("/", 1)[-1]
    return Path(name).stem.replace("_", " ")


def normalize_imageinfo(info: dict[str, Any], candidate_id: str | None = None) -> dict[str, Any]:
    """Commons API imageinfo (extmetadata dahil) -> duz aday sozlugu. Girdi degistirilmez."""
    ext = info.get("extmetadata") or {}

    def meta(key: str) -> str:
        return strip_html((ext.get(key) or {}).get("value"))

    source_url = info.get("descriptionurl") or ""
    original_url = info.get("url") or ""
    date_original = meta("DateTimeOriginal") or None
    categories = [part.strip() for part in meta("Categories").split("|") if part.strip()]
    return {
        "candidate_id": candidate_id or source_url or original_url,
        "title": meta("ObjectName") or _title_from_url(source_url or original_url),
        "description_text": meta("ImageDescription"),
        "categories": categories,
        "date_original": date_original,
        "year": parse_year(date_original) or parse_year(meta("DateTime")),
        "lat": _float(meta("GPSLatitude")),
        "lon": _float(meta("GPSLongitude")),
        "width": _int(info.get("width")),
        "height": _int(info.get("height")),
        "mime": info.get("mime") or _guess_mime(original_url),
        "license": meta("LicenseShortName"),
        "license_url": meta("LicenseUrl"),
        "usage_terms": meta("UsageTerms"),
        "author_text": meta("Artist"),
        "credit_text": meta("Credit"),
        "restrictions": meta("Restrictions"),
        "attribution_required": meta("AttributionRequired"),
        "copyrighted": meta("Copyrighted"),
        "source_url": source_url,
        "original_url": original_url,
        "thumb_url": info.get("thumburl"),
        "sha1": info.get("sha1"),
        "provider": info.get("provider") or "Wikimedia Commons",
    }


def load_evidence_file(path: Path) -> dict[str, Any]:
    """Tek bir <id>.source-evidence.json dosyasini aday sozlugune cevirir."""
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    if not isinstance(data, dict) or not isinstance(data.get("api_imageinfo"), dict):
        raise ValueError("api_imageinfo yok")
    candidate = normalize_imageinfo(data["api_imageinfo"], candidate_id=data.get("photo_id"))
    return {**candidate, "photo_id": data.get("photo_id"), "commons_file": data.get("commons_file"),
            "downloaded_file_url": data.get("downloaded_file_url"), "evidence_file": Path(path).name}


def find_xlsx(root: Path) -> Path | None:
    found = sorted(p for p in Path(root).glob("*.xlsx") if not p.name.startswith("~$"))
    return found[0] if found else None


def cover_entry_id(photo_id: str) -> str:
    """'eindhoven:foto:kapak-merkez-01' -> 'eindhoven:kapak-merkez' (kapak kartinin kimligi)."""
    head, _, slug = photo_id.rpartition(":")
    city = head.split(":", 1)[0] if head else ""
    slug = re.sub(r"-\d+$", "", slug)
    return f"{city}:{slug}" if city else slug


def _photo_rows(xlsx: Path) -> dict[str, dict[str, Any]]:
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    try:
        rows = list(iter_records(wb["Fotoğraflar"])) if "Fotoğraflar" in wb.sheetnames else []
    finally:
        wb.close()
    out: dict[str, dict[str, Any]] = {}
    for index, row in enumerate(rows, start=1):
        photo_id = str(row.get("photo_id", ""))
        role = ROLE_FROM_EXCEL.get(str(row.get("role", "")), "venue")
        entry_id = row.get("entry_id") or (cover_entry_id(photo_id) if role == "cover" else None)
        out[photo_id] = {"photo_index": index, "role": role, "entry_id": entry_id, "package_file": row.get("file")}
    return out


def load_evidence_dir(root: Path) -> tuple[list[dict[str, Any]], list[str]]:
    """Paketin kanit klasorundeki tum adaylar (+ varsa Excel'den entry_id/rol/sira). Doner: (adaylar, hatalar)."""
    root = Path(root)
    xlsx = find_xlsx(root)
    package = _photo_rows(xlsx) if xlsx else {}
    candidates: list[dict[str, Any]] = []
    errors: list[str] = []
    for path in sorted((root / EVIDENCE_DIR).glob(EVIDENCE_GLOB)):
        try:
            candidate = load_evidence_file(path)
        except (ValueError, OSError) as exc:
            errors.append(f"{path.name}: {exc}")
            continue
        candidates.append({**candidate, **package.get(str(candidate.get("photo_id")), {})})
    candidates.sort(key=lambda c: (c.get("photo_index") or 10**9, c["candidate_id"]))
    return candidates, errors


def load_search_rows(path: Path) -> tuple[list[dict[str, Any]], list[str]]:
    """commons_search.py ciktisi ({entry_id, query_kind, title, imageinfo}) -> adaylar (entry_id ve query_kind korunur)."""
    candidates: list[dict[str, Any]] = []
    errors: list[str] = []
    for number, line in enumerate(Path(path).read_text(encoding="utf-8").splitlines(), start=1):
        if not line.strip():
            continue
        try:
            row = json.loads(line)
        except json.JSONDecodeError as exc:
            errors.append(f"satir {number}: gecersiz JSON ({exc})")
            continue
        if not isinstance(row, dict) or not isinstance(row.get("imageinfo"), dict) or not row.get("entry_id"):
            errors.append(f"satir {number}: entry_id/imageinfo eksik")
            continue
        cid = f"{row['entry_id']}|{row.get('title') or row['imageinfo'].get('descriptionurl', number)}"
        candidate = normalize_imageinfo(row["imageinfo"], candidate_id=cid)
        candidates.append({**candidate, "entry_id": row["entry_id"], "query_kind": row.get("query_kind"),
                           "commons_file": row.get("title")})
    return candidates, errors


def _aliases(text: Any) -> list[str]:
    return [part.strip() for part in re.split(r"[;\n|]", str(text or "")) if part.strip()]


def _name_aliases(title: str) -> list[str]:
    """'Binnenstad (Merkez)' -> ['Binnenstad'] (parantez ici genelde Turkce aciklamadir)."""
    bare = re.sub(r"\s*\([^)]*\)", "", title).strip()
    return [bare] if bare and bare != title else []


def _plan_rows(sheet_rows: Iterable[dict[str, Any]], role: str, city: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for row in sheet_rows:
        title = str(row.get("title", "")).strip()
        rows.append({"entry_id": row["entry_id"], "role": role, "name": title, "aliases": _name_aliases(title),
                     "group": row.get("group"), "lat": _float(row.get("latitude")), "lon": _float(row.get("longitude")),
                     **city})
    return rows


def build_plan(xlsx: Path) -> list[dict[str, Any]]:
    """Excel'den prefilter plani: Mekânlar -> venue, Mahalleler -> district, Fotoğraflar.Kapak -> cover."""
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    try:
        sheet = lambda name: list(iter_records(wb[name])) if name in wb.sheetnames else []  # noqa: E731
        city_row = (sheet("Şehir") or [{}])[0]
        venues, districts = sheet("Mekânlar"), sheet("Mahalleler")
    finally:
        wb.close()
    name = str(city_row.get("name", "")).strip()
    city = {"municipality": name, "municipality_aliases": sorted({name, *_aliases(city_row.get("aliases"))} - {""}),
            "city_lat": _float(city_row.get("latitude")), "city_lon": _float(city_row.get("longitude"))}
    plan = _plan_rows(venues, "venue", city) + _plan_rows(districts, "district", city)
    covers = sorted({row["entry_id"] for row in _photo_rows(xlsx).values() if row["role"] == "cover" and row["entry_id"]})
    plan += [{"entry_id": cid, "role": "cover", "name": name, "aliases": [], "group": None, "lat": None, "lon": None, **city}
             for cid in covers]
    return plan


def write_jsonl(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="\n") as stream:
        for row in rows:
            stream.write(json.dumps(row, ensure_ascii=False) + "\n")


def _cmd_from_evidence(args: argparse.Namespace) -> int:
    root = Path(args.package_root)
    if not (root / EVIDENCE_DIR).is_dir():
        emit({"ok": False, "error": f"kanit klasoru yok: {root / EVIDENCE_DIR}"})
        return EXIT_USAGE
    candidates, errors = load_evidence_dir(root)
    write_jsonl(Path(args.out), candidates)
    emit({"ok": not errors, "candidates": len(candidates), "errors": len(errors), "error_files": errors[:5], "out": str(args.out)})
    return EXIT_FINDINGS if errors else EXIT_OK


def _cmd_from_search(args: argparse.Namespace) -> int:
    source = Path(args.package_root)
    if not source.is_file():
        emit({"ok": False, "error": f"arama dosyasi yok: {source}"})
        return EXIT_USAGE
    candidates, errors = load_search_rows(source)
    write_jsonl(Path(args.out), candidates)
    kinds: dict[str, int] = {}
    for c in candidates:
        kinds[str(c.get("query_kind"))] = kinds.get(str(c.get("query_kind")), 0) + 1
    emit({"ok": not errors, "candidates": len(candidates), "errors": len(errors), "query_kinds": kinds, "out": str(args.out)})
    return EXIT_FINDINGS if errors else EXIT_OK


def _cmd_plan(args: argparse.Namespace) -> int:
    xlsx = find_xlsx(Path(args.package_root))
    if xlsx is None:
        emit({"ok": False, "error": f"Excel yok: {args.package_root}"})
        return EXIT_USAGE
    plan = build_plan(xlsx)
    write_jsonl(Path(args.out), plan)
    roles = {role: sum(1 for row in plan if row["role"] == role) for role in ("venue", "district", "cover")}
    emit({"ok": True, "entries": len(plan), **roles, "out": str(args.out)})
    return EXIT_OK


def main(argv: list[str] | None = None) -> int:
    setup_utf8()
    parser = argparse.ArgumentParser(description="Commons foto adaylarini normalize eder / Excel'den plan cikarir.")
    sub = parser.add_subparsers(dest="command", required=True)
    for name, func in (("from-evidence", _cmd_from_evidence), ("from-search", _cmd_from_search), ("plan-from-xlsx", _cmd_plan)):
        cmd = sub.add_parser(name)
        cmd.add_argument("package_root")
        cmd.add_argument("--out", required=True)
        cmd.set_defaults(func=func)
    args = parser.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
