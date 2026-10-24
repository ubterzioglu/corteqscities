"""Metin kalitesi: K27 birebir, K28 yakin/kalip tekrar, K104 genel denetim uyarisi, K21 eylem baglantisi, K33 visual_json,
K25 kacamak ilk cumle, B16 okura hitap (address_forms; olcum her zaman, bulgu yalniz --address-standard ile; buton
metni Baglantilar.label ayri sayilir, bulgu degil; yalniz hitap icin okunan sayfa okunamazsa address.skipped'e yazilir).
Kullanim: text_quality.py <xlsx> --out text-quality.json [--min-cluster 3] [--near 0.9]
          [--address-standard {hitapsiz,siz,sen,off}]   (varsayilan off = yalniz olcum)
Excel salt okunur, ag yok. Cikis: 0 temiz, 1 bulgu, 2 kullanim/yapi hatasi."""
import argparse
import itertools
import json
import math
import os
import re
import unicodedata
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Callable, Iterable, Mapping, NamedTuple
from urllib.parse import urlsplit

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from address_forms import ADDRESS_FIELDS, STANDARDS, address_report, address_summary
from common import EXIT_FINDINGS, EXIT_OK, EXIT_USAGE, emit, iter_records, setup_utf8, write_json

TOOL = "text_quality"
Sheets = Mapping[str, list[dict[str, Any]]]
FIELDS: dict[str, tuple[str, ...]] = {
    "Yanıtlar": ("title", "summary", "steps", "documents", "timing", "cost", "note"), "Mekânlar": ("subtitle", "summary", "note"),
    "Mahalleler": ("subtitle", "summary", "note"), "Fotoğraflar": ("caption", "modifications")}
CARD_SHEETS = ("Yanıtlar", "Mekânlar", "Mahalleler")
BASE_SHEETS = tuple(dict.fromkeys(("Şehir", *FIELDS, "Bağlantılar")))  # okunamazsa kosu durur (cikis 2)
ADDRESS_ONLY = tuple(s for s in ADDRESS_FIELDS if s not in BASE_SHEETS)  # yalniz hitap icin; okunamazsa atlanir
ID_KEYS, TITLE_FIELDS = ("entry_id", "photo_id", "action_id"), ("title", "subtitle", "caption")
COMBINED = "*birlesik*"  # kart govdesi (baslik disi alanlar) birlesik karsilastirma
MASK_KEYS, MASK, MIN_MASK, MIN_TEMPLATE_WORDS = ("title", "subtitle", "group", "district"), "¤", 4, 3
SHORT_LIMIT = 25  # <= 25 karakterlik tekrarlar (belge adi, 'Ücretsiz') kopya kume sayilmaz, ayri sayilir
SAMPLE_CHARS, TOP_N, SHARED_STEP_MIN, MAX_LINE_FANOUT, EPS = 120, 5, 3, 50, 1e-9
PLACEHOLDER_WORDS = frozenset({"", "yok","n a", "na", "bilinmiyor", "todo", "tbd", "belirtilmemis", "tanimsiz"})
PROCESS_KINDS = frozenset({"steps", "process", "timeline", "flow", "sequence", "adimlar"})
MAP_HOSTS = frozenset({"maps.google.com", "maps.app.goo.gl", "osm.org", "openstreetmap.org", "maps.apple.com"})
MAP_LABEL = re.compile(r"\b(harita\w*|konum|map|maps)\b")
DIGITS, TEMPLATE_WORD = re.compile(r"\d+"), re.compile(r"[^\W\d_]{3,}")
PROPER_SUFFIXED = re.compile(r"(?<!\w)([A-ZÂÇĞÎİÖŞÛÜ]\w+)['’]\w")  # Turkcede ozel ad + kesme + ek: Amsterdam'da
SENTENCE_SPLIT, FIRST_SENTENCE = re.compile(r"[.!?;\n]+"), re.compile(r"(?<=[.!?])\s+|\n")
NOT_STATEMENT = r"(?!er\b|en\b|ir\b|iyor|ildi|ilir\b)"  # 'kontrol eder/eden/ediyor/edilir' tavsiye degil, durum bildirir
CHECK = rf"(?:(?:kontrol|teyit) e[td]{NOT_STATEMENT}\w*|dogrula(?!r\b|nir\b|yor|ndi)\w*)"  # kelime oneki degil fiil kalibi
VERIFY = rf"(?:{CHECK}|bak(?:in|iniz|abilir\w*|mali\w*|mayi|mak)?\b)"  # 'bakkal', 'bakim', 'bakanlik' disarida
HEDGES = {
    "gitmeden_kontrol": re.compile(rf"\b(gitmeden|gelmeden|yola cikmadan)\b.{{0,40}}\b{VERIFY}"),
    "ziyaret_oncesi_dogrula": re.compile(rf"\bziyaret(ten)? (once|oncesi)\b.{{0,40}}\b{VERIFY}"),
    "resmi_kaynaktan_dogrula": re.compile(rf"\bresmi\b.{{0,30}}\b(site|kaynak|sayfa|kanal)\w*\b.{{0,30}}\b({CHECK}|ogren(?!ci|im)\w*)"),
    "guncel_bilgi_kurum": re.compile(r"\bguncel bilgi\w*\b.{0,40}\b(kurum|resmi|site)"),
    "kurumdan_dogrula": re.compile(rf"\bkurum\w*\b.{{0,20}}\b({CHECK}|sor(un|abilirsin)?\b)")}
EVASIVE = (re.compile(r"\bresmi\b.{0,20}\b(site|sayfa|kaynak)\w*\s+(bak|goz at|incele)"),
           re.compile(r"\b(duruma|kisiye|statune|kosullara) gore degis"))
FILLER = {"sss_numarali": re.compile(r"\bsik sorulan soru\b"), "lorem": re.compile(r"\blorem ipsum\b"),
          "detayli_islem_adimlari": re.compile(r"\bdetayli islem adimlari\b")}
Value = NamedTuple("Value", [("sheet", str), ("field", str), ("rec", Mapping[str, Any]), ("raw", str), ("norm", str)])

def normalize(text: Any) -> str:
    """Kucuk harf, aksansiz (Turkce dahil), noktalamasiz, tek bosluklu metin."""
    lowered = unicodedata.normalize("NFKD", ("" if text is None else str(text)).lower())
    plain = "".join(ch for ch in lowered if not unicodedata.combining(ch)).replace("ı", "i")
    return " ".join(re.sub(r"[^\w\s]|_", " ", plain).split())

def is_placeholder_value(value: Any) -> bool:
    """'-', 'yok', 'N/A', 'bilinmiyor', 'TODO', '...' gibi icerik yerine konmus degerler."""
    return isinstance(value, str) and bool(value.strip()) and normalize(value) in PLACEHOLDER_WORDS

def shingles(text: str, n: int = 5) -> frozenset[str]:  # n'den kisa metin tek parca
    return frozenset(text[i:i + n] for i in range(max(len(text) - n + 1, 1))) if text else frozenset()

def jaccard(a: frozenset[str], b: frozenset[str]) -> float:
    return len(a & b) / len(a | b) if a or b else 0.0

def near_pairs(sets: list[frozenset[str]], threshold: float) -> list[tuple[int, int]]:
    """Jaccard >= esik ciftleri; n^2 yerine onek-filtreli ters dizin (kayipsiz, deterministik)."""
    freq = Counter(tok for s in sets for tok in s)
    index: dict[str, list[int]] = defaultdict(list)
    pairs: set[tuple[int, int]] = set()
    for i in sorted((k for k in range(len(sets)) if sets[k]), key=lambda k: (len(sets[k]), k)):
        size = len(sets[i])
        prefix = sorted(sets[i], key=lambda tok: (freq[tok], tok))[: size - math.ceil(threshold * size - EPS) + 1]
        for j in sorted({j for tok in prefix for j in index[tok]}):
            if len(sets[j]) >= threshold * size - EPS and jaccard(sets[i], sets[j]) >= threshold:
                pairs.add((min(i, j), max(i, j)))
        for tok in prefix:
            index[tok].append(i)
    return sorted(pairs)

def hedge_hits(text: Any) -> list[str]:
    return [name for norm in [normalize(text)] for name, pattern in HEDGES.items() if pattern.search(norm)]

def _components(n: int, pairs: Iterable[tuple[int, int]]) -> list[list[int]]:
    parent = list(range(n))

    def find(i: int) -> int:
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i
    for a, b in pairs:
        ra, rb = find(a), find(b)
        parent[max(ra, rb)] = min(ra, rb)
    return [g for g in _groups(range(n), find).values() if len(g) > 1]

def _groups(items: Iterable[Any], key: Callable[[Any], Any]) -> dict[Any, list[Any]]:
    out: dict[Any, list[Any]] = defaultdict(list)
    for item in items:
        out[key(item)].append(item)
    return out

def _clip(text: Any) -> str:
    flat = " ".join(str(text).split())
    return flat if len(flat) <= SAMPLE_CHARS else flat[:SAMPLE_CHARS] + "…"

def _member(rec: Mapping[str, Any]) -> dict[str, Any]:
    return {"row": rec.get("_row"), "id": next((rec[k] for k in ID_KEYS if rec.get(k)), None)}

def _ref(sheet: str, rec: Mapping[str, Any], **extra: Any) -> dict[str, Any]:
    return {"sheet": sheet, **_member(rec), **extra}

def _cluster(members: list[Value], **extra: Any) -> dict[str, Any]:
    ordered = sorted(members, key=lambda v: v.rec.get("_row") or 0)
    return {"sheet": ordered[0].sheet, "field": ordered[0].field, "size": len(ordered), **extra,
            "sample": _clip(ordered[0].raw), "members": [_member(v.rec) for v in ordered]}

def _sorted(clusters: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return sorted(clusters, key=lambda c: (-c["size"], c["sheet"], c.get("field", ""), c["members"][0]["row"] or 0))

def _by_field(values: list[Value], combined: bool = False) -> dict[tuple[str, str], list[Value]]:
    return _groups((v for v in values if combined or v.field != COMBINED), lambda v: (v.sheet, v.field))

def collect(sheets: Sheets) -> tuple[list[Value], list[dict[str, Any]]]:
    """Metin alanlarini (ve kart govdesi birlesimini) toplar; yer tutucular ayri listelenir."""
    values: list[Value] = []
    holders: list[dict[str, Any]] = []
    for sheet, fields in FIELDS.items():
        for rec in sheets.get(sheet, []):
            body: list[Value] = []
            for field, raw in ((f, rec[f]) for f in fields if rec.get(f) is not None):
                if is_placeholder_value(raw):
                    holders.append(_ref(sheet, rec, field=field, value=str(raw)))
                    continue
                values.append(Value(sheet, field, rec, str(raw), normalize(raw)))
                body += [values[-1]] if field not in TITLE_FIELDS else []
            if len(body) >= 2:
                values.append(Value(sheet, COMBINED, rec, " | ".join(v.raw for v in body), " ".join(v.norm for v in body)))
    return values, holders

def city_names(sheets: Sheets) -> tuple[str, ...]:
    parts = (normalize(part) for rec in sheets.get("Şehir", [])
             for part in re.split(r"[;,\n]", f"{rec.get('name') or ''};{rec.get('aliases') or ''}"))
    return tuple(sorted({p for p in parts if len(p) >= MIN_MASK}))

def exact_duplicates(values: list[Value], min_cluster: int) -> tuple[list[dict], list[dict]]:
    """K27: normalize sonrasi ozdes degerler; uzunlar >= 2 kayitta kume, kisalar (<= 25) ayri sayilir."""
    clusters, short = [], []
    for (sheet, field), items in _by_field(values, combined=True).items():
        for norm, members in _groups(items, lambda v: v.norm).items():
            if len(members) >= 2 and len(norm) > SHORT_LIMIT:
                clusters.append(_cluster(members, length=len(norm)))
            elif len(members) >= max(2, min_cluster) and len(norm) <= SHORT_LIMIT:
                short.append({"sheet": sheet, "field": field, "value": members[0].raw, "size": len(members)})
    return _sorted(clusters), sorted(short, key=lambda s: (-s["size"], s["sheet"], s["field"], s["value"]))

def near_duplicates(values: list[Value], threshold: float, min_cluster: int) -> list[dict]:
    """K28: karakter 5-gram Jaccard >= esik olan farkli degerlerin kumeleri."""
    clusters = []
    for items in _by_field(values).values():
        groups = _groups((v for v in items if len(v.norm) > SHORT_LIMIT), lambda v: v.norm)
        norms, sets = list(groups), [shingles(n) for n in groups]
        for comp in _components(len(norms), near_pairs(sets, threshold)):
            members = [v for i in comp for v in groups[norms[i]]]
            if len(members) >= min_cluster:
                low = min(jaccard(sets[a], sets[b]) for a, b in itertools.combinations(comp[:20], 2))
                clusters.append(_cluster(members, distinct=len(comp), min_jaccard=round(low, 3),
                                         samples=[_clip(groups[norms[i]][0].raw) for i in comp[:3]]))
    return _sorted(clusters)

def template_key(value: Value, cities: Iterable[str]) -> str:
    """Rakamlar, sehir adlari, kesmeli ozel adlar (Amsterdam'da) ve kaydin ad alanlari (baslik, grup, semt) maskelenir."""
    proper = {normalize(name) for name in PROPER_SUFFIXED.findall(value.raw)}
    masks = {normalize(value.rec.get(k)) for k in MASK_KEYS if k != value.field} | set(cities) | proper
    text = value.norm
    for mask in sorted((m for m in masks if len(m) >= MIN_MASK), key=lambda m: (-len(m), m)):
        text = re.sub(rf"(?<!\w){re.escape(mask)}(?!\w)", MASK, text)
    return DIGITS.sub("0", text)

def template_duplicates(values: list[Value], cities: tuple[str, ...], min_cluster: int) -> list[dict]:
    """K28: numarali doldurma / ad degistirilmis genel metin (maske sonrasi ozdes, ham hali farkli)."""
    clusters = []
    for items in _by_field(values).values():
        for key, members in _groups(items, lambda v: template_key(v, cities)).items():
            distinct = len({v.norm for v in members})
            if distinct >= 2 and len(members) >= min_cluster and len(TEMPLATE_WORD.findall(key)) >= MIN_TEMPLATE_WORDS:
                clusters.append(_cluster(members, distinct=distinct, template=_clip(key)))
    return _sorted(clusters)

def repeated_sentences(values: list[Value], min_cluster: int) -> list[dict]:
    """K28 kalip: ayni sayfada >= min_cluster kayitta gecen uzun cumleler (bilgi amacli)."""
    parts = ((v, part.strip(), normalize(part)) for v in values if v.field != COMBINED for part in SENTENCE_SPLIT.split(v.raw))
    groups = _groups((p for p in parts if len(p[2]) > SHORT_LIMIT), lambda p: (p[0].sheet, p[2]))
    items = [{"sheet": sheet, "sentence": norm, "records": len({v.rec.get("_row") for v, _, _ in hits}),
              "fields": sorted({v.field for v, _, _ in hits}), "sample": _clip(hits[0][1])} for (sheet, norm), hits in groups.items()]
    return sorted((i for i in items if i["records"] >= min_cluster), key=lambda i: (-i["records"], i["sheet"], i["sentence"]))

def _step_edges(index: Mapping[str, list[int]], n: int) -> dict[tuple[int, int], int]:
    """Cift -> ortak adim satiri; yaygin (> MAX_LINE_FANOUT) satirlar cift yerine imza grubuyla sayilir (kayipsiz, n^2 yok)."""
    common = [(line, frozenset(ix)) for line, ix in index.items() if len(ix) > MAX_LINE_FANOUT]
    mass = [frozenset(line for line, owners in common if i in owners) for i in range(n)]
    rare = Counter(p for ix in index.values() if len(ix) <= MAX_LINE_FANOUT for p in itertools.combinations(ix, 2))
    edges = {p: c + len(mass[p[0]] & mass[p[1]]) for p, c in rare.items()}
    sigs = _groups((i for i in range(n) if len(mass[i]) >= SHARED_STEP_MIN), lambda i: mass[i])
    for x, y in itertools.combinations_with_replacement(sorted(sigs, key=lambda s: sigs[s][0]), 2):
        if len(x & y) >= SHARED_STEP_MIN:  # imza gruplari ardisik kenarla zincirlenir: bilesen ayni, cift sayisi dogrusal
            chain = sorted({*sigs[x], *sigs[y]})
            edges.update({p: max(edges.get(p, 0), len(x & y)) for p in zip(chain, chain[1:])})
    return {p: c for p, c in edges.items() if c >= SHARED_STEP_MIN}

def shared_steps(sheets: Sheets) -> list[dict]:
    """K28 'ayni uc adim': en az 3 ozdes adim satirini paylasan kayitlar (toplu kopyalanmis blok dahil)."""
    out = []
    for sheet in (s for s, fields in FIELDS.items() if "steps" in fields):
        recs = [r for r in sheets.get(sheet, []) if isinstance(r.get("steps"), str)]
        index: dict[str, list[int]] = defaultdict(list)
        for i, rec in enumerate(recs):
            for line in (ln for ln in dict.fromkeys(normalize(p) for p in rec["steps"].split("\n")) if ln):
                index[line].append(i)
        edges = _step_edges(index, len(recs))
        for comp in _components(len(recs), sorted(edges)):
            shared = max(c for (a, _), c in edges.items() if a in set(comp))
            out.append({"sheet": sheet, "size": len(comp), "shared": shared, "members": [_member(recs[i]) for i in comp]})
    return _sorted(out)

def filler(values: list[Value]) -> dict[str, Any]:
    items = [_ref(v.sheet, v.rec, field=v.field, pattern=name) for v in values if v.field != COMBINED
             for name, pattern in FILLER.items() if pattern.search(v.norm)]
    return {"count": len(items), "items": items}

def placeholder_report(holders: list[dict[str, Any]]) -> dict[str, Any]:
    by_field = Counter(f"{h['sheet']}.{h['field']}" for h in holders)
    return {"count": len(holders), "by_field": dict(sorted(by_field.items())), "items": holders}

def _cards(sheets: Sheets) -> Iterable[tuple[str, Mapping[str, Any]]]:
    return ((sheet, rec) for sheet in CARD_SHEETS for rec in sheets.get(sheet, []))

def hedges(sheets: Sheets) -> dict[str, Any]:
    """K104: okuyucuya denetim isi birakan kaliplarin kart basina sayisi ve toplam payi."""
    cards = list(_cards(sheets))
    found = [(sheet, rec, [h for f in FIELDS[sheet] if isinstance(rec.get(f), str) for h in hedge_hits(rec[f])])
             for sheet, rec in cards]
    items = [_ref(sheet, rec, count=len(hits), patterns=sorted(set(hits))) for sheet, rec, hits in found if hits]
    by_pattern, total = Counter(p for i in items for p in i["patterns"]), len(cards)
    return {"cards_total": total, "cards_with_hedge": len(items), "share": round(len(items) / total, 4) if total else 0.0,
            "max_per_card": max((i["count"] for i in items), default=0), "by_pattern": dict(sorted(by_pattern.items())), "items": items}

def evasive_answers(sheets: Sheets) -> dict[str, Any]:
    """K25 yardimcisi: baslik soruysa summary'nin ilk cumlesi 'resmi siteye bakin' turu kacamak mi."""
    items = []
    for sheet, rec in _cards(sheets):
        title, summary = str(rec.get("title") or "").strip(), rec.get("summary")
        first = FIRST_SENTENCE.split(summary.strip(), maxsplit=1)[0] if title.endswith("?") and isinstance(summary, str) else ""
        if first and (hedge_hits(first) or any(p.search(normalize(first)) for p in EVASIVE)):
            items.append(_ref(sheet, rec, title=title, first_sentence=_clip(first)))
    return {"count": len(items), "items": items}

def is_map_link(label: Any, url: Any) -> bool:
    if MAP_LABEL.search(normalize(label)):
        return True
    try:
        parts = urlsplit(str(url or "").strip())
    except ValueError:
        return False
    host, path = (parts.hostname or "").lower().removeprefix("www."), parts.path.lower()
    google = (host.startswith("google.") or host == "goo.gl") and path.startswith("/maps")
    return host in MAP_HOSTS or host.endswith(".openstreetmap.org") or google

def action_links(sheets: Sheets) -> dict[str, Any]:
    """K21: Yanit/Mekan/Mahalle kaydinin harita disi en az bir eylem baglantisi var mi."""
    if "Bağlantılar" not in sheets:
        return {"skipped": True, "checked": 0, "missing": 0, "map_only": 0, "no_id": 0, "items": []}
    kinds = _groups((link for link in sheets["Bağlantılar"] if link.get("entry_id")), lambda link: link["entry_id"])
    items, cards = [], list(_cards(sheets))
    for sheet, rec in cards:
        flags = [is_map_link(link.get("label"), link.get("url")) for link in kinds.get(rec.get("entry_id"), [])]
        if not rec.get("entry_id") or not flags or all(flags):  # kimliksiz kart kimliksiz baglantiyla eslesmez
            reason = "kimlik_yok" if not rec.get("entry_id") else "yalniz_harita" if flags else "baglanti_yok"
            items.append(_ref(sheet, rec, reason=reason))
    count = Counter(i["reason"] for i in items)
    return {"skipped": False, "checked": len(cards), "missing": count["baglanti_yok"], "map_only": count["yalniz_harita"],
            "no_id": count["kimlik_yok"], "items": items}

def visual_error(raw: str, steps: Any) -> str | None:
    try:
        data = json.loads(raw)
    except ValueError:
        return "parse_error"
    if not isinstance(data, dict):
        return "not_object"
    if not isinstance(data.get("kind"), str) or not data["kind"].strip():
        return "missing_kind"
    if not isinstance(items := data.get("items"), list):
        return "items_not_list"
    lines = [line for line in str(steps or "").split("\n") if line.strip()]
    if data["kind"].strip().lower() in PROCESS_KINDS and lines and len(items) != len(lines):
        return "step_count_mismatch"
    labels = [v for item in items if isinstance(item, dict) for v in item.values() if isinstance(v, str)]
    return "truncated_label" if any(v.rstrip().endswith(("…", "...")) for v in labels) else None

def visual_json(sheets: Sheets) -> dict[str, Any]:
    """K33: visual_json parse edilebiliyor mu, adim listesiyle celisiyor mu; bos olanlar ayrica sayilir."""
    records = sheets.get("Yanıtlar", [])
    filled = [r for r in records if str(r.get("visual_json") or "").strip()]
    checked = [(r, visual_error(str(r["visual_json"]), r.get("steps"))) for r in filled]
    errors = [_ref("Yanıtlar", r, error=e) for r, e in checked if e]
    return {"total": len(records), "empty": len(records) - len(filled), "valid": len(filled) - len(errors),
            "by_error": dict(sorted(Counter(e["error"] for e in errors).items())), "errors": errors}

def caption_is_title(sheets: Sheets) -> dict[str, Any]:
    titles = {rec.get("entry_id"): normalize(rec.get("title")) for _, rec in _cards(sheets) if rec.get("entry_id")}
    photos = [r for r in sheets.get("Fotoğraflar", []) if isinstance(r.get("caption"), str) and normalize(r["caption"])]
    items = [_ref("Fotoğraflar", r, entry_id=r.get("entry_id"), caption=r["caption"])
             for r in photos if normalize(r["caption"]) == titles.get(r.get("entry_id"))]
    return {"count": len(items), "total": len(photos), "items": items}

def _summary(report: Mapping[str, Any], min_cluster: int) -> dict[str, Any]:
    hedge_cards = report["hedges"]["cards_with_hedge"]
    findings = {"exact_clusters": len(report["exact_duplicates"]), "near_clusters": len(report["near_duplicates"]),
                "template_clusters": len(report["template_duplicates"]), "shared_steps": len(report["shared_steps"]),
                "placeholders": report["placeholders"]["count"], "filler": report["filler"]["count"],
                "evasive": report["evasive_answers"]["count"], "without_action": len(report["action_links"]["items"]),
                "visual_json_errors": len(report["visual_json"]["errors"]), "caption_is_title": report["caption_is_title"]["count"]}
    address = address_summary(report["address"])  # B16: standart disi + karisik kayit (off: 0)
    hedge_finding = hedge_cards if hedge_cards >= min_cluster else 0  # K104: tekrar eden kalip
    total = sum(findings.values()) + hedge_finding + address["findings"]
    return {**findings, "hedge_cards": hedge_cards, "hedge_share": report["hedges"]["share"],
            "exact_records": sum(c["size"] for c in report["exact_duplicates"]),
            "short_repeat_groups": len(report["short_repeats"]), "repeated_sentences": len(report["repeated_sentences"]),
            "cards_total": report["hedges"]["cards_total"], "visual_json_empty": report["visual_json"]["empty"],
            "address": address, "findings": total}

def _top(report: Mapping[str, Any]) -> list[dict[str, Any]]:
    pool = [{"kind": kind, **{k: c[k] for k in ("sheet", "field", "size", "distinct", "sample") if k in c}}
            for kind in ("exact", "near", "template") for c in report[f"{kind}_duplicates"] if c["field"] != COMBINED]
    return sorted(pool, key=lambda c: (-c["size"], c["kind"], c["sheet"], c["field"]))[:TOP_N]

def analyze(sheets: Sheets, min_cluster: int = 3, near: float = 0.9, address_standard: str = "off",
            address_skipped: Mapping[str, str] | None = None) -> dict[str, Any]:
    """Saf analiz: girdiyi degistirmez, ayni girdi -> ayni rapor. Bilinmeyen hitap standardi ValueError.
    address_skipped: okunamayan hitap sayfasi -> hata (load_sheets); address.skipped'de listelenir."""
    values, holders = collect(sheets)
    exact, short = exact_duplicates(values, min_cluster)
    report = {"exact_duplicates": exact, "near_duplicates": near_duplicates(values, near, min_cluster),
              "template_duplicates": template_duplicates(values, city_names(sheets), min_cluster),
              "short_repeats": short, "repeated_sentences": repeated_sentences(values, min_cluster),
              "shared_steps": shared_steps(sheets), "placeholders": placeholder_report(holders), "filler": filler(values),
              "hedges": hedges(sheets), "evasive_answers": evasive_answers(sheets), "action_links": action_links(sheets),
              "visual_json": visual_json(sheets), "caption_is_title": caption_is_title(sheets),
              "address": address_report(sheets, address_standard, address_skipped)}
    return {"summary": _summary(report, min_cluster), "top_clusters": _top(report), **report}

def load_sheets(workbook: Path) -> tuple[dict[str, list[dict[str, Any]]], dict[str, str]]:
    """Excel'i salt okur; yalniz var olan sayfalar doner. Metin sayfasi yoksa ValueError. Yalniz hitap icin okunan
    sayfa (Rotalar, Adimlar, Konular, Sozluk, SayfaAlanlari) okunamazsa kosu durmaz: (sayfalar, {sayfa: hata})."""
    wb = openpyxl.load_workbook(workbook, data_only=False)
    try:
        sheets = {name: list(iter_records(wb[name])) for name in BASE_SHEETS if name in wb.sheetnames}
        skipped: dict[str, str] = {}
        for name in (n for n in ADDRESS_ONLY if n in wb.sheetnames):
            try:
                sheets[name] = list(iter_records(wb[name]))
            except ValueError as exc:  # satir 7'de yinelenen teknik anahtar: K27/K28/K104/K21/K33 sonuclari kaybolmasin
                skipped[name] = str(exc)
    finally:
        wb.close()
    if not any(name in sheets for name in FIELDS):
        raise ValueError(f"metin sayfasi yok (beklenen: {', '.join(FIELDS)})")
    return sheets, skipped

def same_file(a: Path, b: Path) -> bool:  # ayni yol, harf ya da '..' farki, sert baglanti: --out girdiyi ezmesin
    try:
        return a.resolve() == b.resolve() or (a.exists() and b.exists() and os.path.samefile(a, b))
    except (OSError, ValueError):  # cozulemeyen yol ayni sayilmaz; okunamayan girdi/cikti zaten kullanim hatasi
        return False

def main(argv: list[str] | None = None) -> int:
    setup_utf8()
    parser = argparse.ArgumentParser(description="Metin kalitesi: K27/K28/K104/K21/K33/K25 + B16 hitap (ag yok, Excel salt okunur).")
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("--out", type=Path, required=True, help="ayrintili JSON rapor")
    parser.add_argument("--min-cluster", type=int, default=3, help="yakin/kalip kume, kisa tekrar, tekrar eden cumle icin en az kayit")
    parser.add_argument("--near", type=float, default=0.9, help="karakter 5-gram Jaccard esigi (0-1]")
    parser.add_argument("--address-standard", choices=STANDARDS, default="off",
                        help="B16 okura hitap: hitapsiz/siz/sen standart disi ve karisik kaydi bulgu sayar; off yalniz olcer")
    args = parser.parse_args(argv)
    checks = ((not 0 < args.near <= 1 or args.min_cluster < 2, "--near (0,1] ve --min-cluster >= 2 olmali"),
              (same_file(args.xlsx, args.out), "--out girdi Excel ile ayni dosya; girdi asla uzerine yazilmaz"))
    if error := next((message for failed, message in checks if failed), None):
        emit({"tool": TOOL, "error": error})
        return EXIT_USAGE
    try:
        sheets, skipped = load_sheets(args.xlsx)
    except (OSError, ValueError, KeyError, zipfile.BadZipFile, InvalidFileException) as exc:
        emit({"tool": TOOL, "error": f"{type(exc).__name__}: {exc}"})
        return EXIT_USAGE
    report = {"tool": TOOL, "workbook": args.xlsx.name, "params": {"min_cluster": args.min_cluster, "near": args.near},
              **analyze(sheets, args.min_cluster, args.near, args.address_standard, skipped)}
    try:
        write_json(args.out, report)
    except (OSError, ValueError) as exc:  # klasor, kilitli dosya, izin, gecersiz yol: kullanim hatasi, 'bulgu' (1) degil
        emit({"tool": TOOL, "error": f"--out yazilamadi: {type(exc).__name__}: {exc}"})
        return EXIT_USAGE
    emit({"tool": TOOL, "workbook": args.xlsx.name, **report["summary"], "out": str(args.out)})
    return EXIT_FINDINGS if report["summary"]["findings"] else EXIT_OK

if __name__ == "__main__":
    raise SystemExit(main())
