"""xlsx_dump.py <xlsx> [--out-dir D] [--summary] [--sheets A,B]

Excel'i salt-okunur dokumler (formuller degerlendirilmez; formul hucreleri kayda girmez).
K10 sayimlari: ham satir (anahtarli hucresi dolu), tekil kimlik, tekrar eden kimlik, anahtarsiz satir, formul sutunlari.

stdout: tek satir JSON {sheets:{ad:{rows,unique,dup_ids,blank_key,formula_cols}}, total_rows, missing_sheets}.
--out-dir: her sayfa icin <Sayfa>.jsonl (kayit basina bir JSON, `_row` = Excel satiri) + summary.json.
--summary: --out-dir ile birlikte yalniz summary.json yazilir (JSONL yazilmaz).
Cikis: 0 temiz / 2 kullanim veya yapi hatasi (okunamayan xlsx, bilinmeyen sayfa, yinelenen teknik anahtar).
dup_ids = ayni kimligi tasiyan FAZLA satir sayisi (satir - tekil - anahtarsiz).
"""
from __future__ import annotations

import argparse
import json
import sys
import zipfile
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

import common


class DumpError(Exception):
    """Kullanim veya yapi hatasi (cikis kodu 2)."""


def open_workbook(path: Path):
    """Formulleri degerlendirmeden, degisiklik yapmadan acar."""
    try:
        return openpyxl.load_workbook(path, data_only=False)
    except (OSError, zipfile.BadZipFile, InvalidFileException, KeyError, ValueError) as exc:
        raise DumpError(f"xlsx okunamadi: {path} ({type(exc).__name__}: {exc})") from exc


def sheet_stats(records: list[dict[str, Any]], id_key: str | None, formula_cols: set[str]) -> dict[str, Any]:
    ids = [r[id_key] for r in records if id_key and id_key in r]
    blank = len(records) - len(ids)
    unique = len(set(map(str, ids)))
    return {
        "rows": len(records),
        "unique": unique,
        "dup_ids": len(ids) - unique,
        "blank_key": blank,
        "formula_cols": sorted(formula_cols),
    }


def read_sheet(ws) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    try:
        keys = common.sheet_keys(ws)
        records = list(common.iter_records(ws))
        formulas = common.formula_columns(ws)
    except ValueError as exc:
        raise DumpError(str(exc)) from exc
    id_key = next(iter(keys), None)
    return records, sheet_stats(records, id_key, formulas)


def select_sheets(wb, requested: list[str] | None) -> tuple[list[str], list[str]]:
    """(islenecek sayfalar, eksik veri sayfalari). Bilinmeyen --sheets adi kullanim hatasidir."""
    if requested is not None:
        unknown = [name for name in requested if name not in common.DATA_SHEETS]
        if unknown:
            raise DumpError(f"bilinmeyen veri sayfasi: {', '.join(unknown)}")
        wanted = requested
    else:
        wanted = list(common.DATA_SHEETS)
    present = [name for name in wanted if name in wb.sheetnames]
    return present, [name for name in wanted if name not in wb.sheetnames]


def write_jsonl(path: Path, records: list[dict[str, Any]]) -> None:
    lines = [json.dumps(r, ensure_ascii=False, separators=(",", ":"), default=str) for r in records]
    with path.open("w", encoding="utf-8", newline="\n") as stream:
        stream.write("".join(line + "\n" for line in lines))


def dump(xlsx: Path, requested: list[str] | None, out_dir: Path | None, summary_only: bool) -> dict[str, Any]:
    wb = open_workbook(xlsx)
    try:
        names, missing = select_sheets(wb, requested)
        sheets: dict[str, Any] = {}
        if out_dir is not None:
            out_dir.mkdir(parents=True, exist_ok=True)
        for name in names:
            records, stats = read_sheet(wb[name])
            sheets[name] = stats
            if out_dir is not None and not summary_only:
                write_jsonl(out_dir / f"{name}.jsonl", records)
    finally:
        wb.close()
    summary = {"sheets": sheets, "total_rows": sum(s["rows"] for s in sheets.values()), "missing_sheets": missing}
    if out_dir is not None:
        common.write_json(out_dir / "summary.json", summary)
    return summary


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Excel veri sayfalarini salt-okunur dokumler (K10 sayimlari).")
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("--out-dir", type=Path, default=None, help="<Sayfa>.jsonl + summary.json yazilacak klasor")
    parser.add_argument("--summary", action="store_true", help="--out-dir ile yalniz summary.json yaz")
    parser.add_argument("--sheets", default=None, help="virgulle ayrilmis sayfa adlari (varsayilan: tum veri sayfalari)")
    return parser


def main(argv: list[str] | None = None) -> int:
    common.setup_utf8()
    args = build_parser().parse_args(argv)
    requested = [s.strip() for s in args.sheets.split(",") if s.strip()] if args.sheets else None
    try:
        summary = dump(args.xlsx, requested, args.out_dir, args.summary)
    except DumpError as exc:
        common.emit({"error": str(exc)})
        return common.EXIT_USAGE
    common.emit(summary)
    return common.EXIT_OK


if __name__ == "__main__":
    sys.exit(main())
