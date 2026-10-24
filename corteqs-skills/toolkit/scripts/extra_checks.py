"""Ek kontroller: resmi v7 dogrulayicida olmayan proje kontrolleri.

Her kontrol, Excel'den okunmus {sayfa: [kayit]} sozlugu uzerinde saf bir fonksiyondur (girdiyi degistirmez).
Sonuclar 'ek kontrol' etiketiyle ayri raporlanir; resmi hata/uyari sayilarina asla karistirilmaz.
"""
from __future__ import annotations

import re
from collections import defaultdict
from pathlib import Path
from typing import Any, Callable, Iterable, Mapping
from urllib.parse import urlsplit

import openpyxl

from common import iter_records

LABEL = "ek kontrol"
Sheets = Mapping[str, list[dict[str, Any]]]

CONTENT_SHEETS = ("Mekânlar", "Yanıtlar", "Mahalleler", "Sözlük")
LOADED_SHEETS = CONTENT_SHEETS + ("Bağlantılar", "Kaynaklar")
ID_KEYS = ("entry_id", "term_id", "action_id", "source_id")
PLACEHOLDER_MARKERS = re.compile(r"\[\[|\]\]|\{\{|\}\}|lorem", re.IGNORECASE)  # 'TODO' ayrica, buyuk harf duyarli
DIGIT_QUESTION = re.compile(r"\d\?(?![\w%.-]*=)")  # '5621?' evet; URL sorgusu 'page2?id=7' hayir
# Tirnak icindeki metin (resmi baslik/alinti); duz tek tirnak yalniz kelime sinirinda: 'Eindhoven'a' kesme isaretidir
QUOTED = re.compile(r"\"[^\"\n]*\"|“[^”\n]*”|„[^“”\n]*[“”]|«[^»\n]*»|»[^«\n]*«|‘[^’\n]*’|(?<!\w)'[^'\n]*'(?!\w)")
CLAUSE_BREAK = re.compile(r"[,;:/()\[\]{}!?\n•|]|\.\s|\s[-–—]\s")
# Soru basligi bicimden taninir, uzunluktan degil (V4): cumlecikte NL/DE/EN soru sozcugu ya da yardimci fiil + ozne
# ('Kan ik', 'Gibt es'). Turkce 'Kan testi …', 'Hat 401?' ozne almadigi icin baslik sayilmaz; 'WHO' gibi kisaltma sayilmaz.
QUESTION_WORDS = frozenset({
    "wat", "wie", "waar", "wanneer", "hoe", "hoeveel", "welk", "welke", "waarom", "waarvoor", "waarmee",  # NL
    "was", "wer", "wo", "wann", "warum", "welche", "welcher", "welches", "wieviel", "wozu", "woher", "wohin",  # DE
    "what", "who", "where", "when", "how", "which", "why"})  # EN
AUX_VERBS = frozenset({
    "kan", "kun", "kunt", "moet", "mag", "wil", "wilt", "zal", "zou", "is", "zijn", "ben", "heb", "hebt", "heeft",  # NL
    "krijg", "krijgt", "word", "wordt", "ga", "gaat",
    "kann", "muss", "darf", "soll", "ist", "sind", "gibt", "habe", "hat", "bin", "wird", "bekomme",  # DE
    "can", "do", "does", "are", "will", "should", "may", "must"})  # EN
SUBJECTS = frozenset({
    "ik", "je", "jij", "u", "we", "wij", "ze", "zij", "hij", "men", "het", "er", "mijn", "uw", "de", "een", "dit", "deze",  # NL
    "ich", "du", "sie", "es", "man", "wir", "mein", "meine", "der", "die", "das", "ein", "eine",  # DE
    "i", "you", "it", "there", "they", "my", "the", "a", "an"})  # EN
QUOTE_CHARS = "\"'‘’“”„«»"
HOMEPAGE_PATH = re.compile(r"/?(?:[a-z]{2}(?:[-_][a-z]{2})?/?|(?:home|index)(?:\.html?)?)?", re.IGNORECASE)
OFFICIAL_WORDS = re.compile(r"resm[iî]|official", re.IGNORECASE)
HISTORY_GROUP = "Tarih ve mimari"
HISTORY_THRESHOLD = 0.35


def _row_id(record: Mapping[str, Any]) -> Any:
    return next((record[k] for k in ID_KEYS if record.get(k)), None)


def _item(sheet: str, record: Mapping[str, Any], **extra: Any) -> dict[str, Any]:
    return {"sheet": sheet, "row": record.get("_row"), "id": _row_id(record), **extra}


def _result(items: list[dict[str, Any]], **extra: Any) -> dict[str, Any]:
    return {"label": LABEL, "count": len(items), "items": items, **extra}


def _host(url: Any) -> str:
    try:
        return (urlsplit(str(url).strip()).hostname or "").lower()
    except ValueError:
        return ""


def _single_token_question(piece: str) -> bool:
    piece = piece.strip()
    return len(piece) > 1 and piece.endswith("?") and not re.search(r"\s", piece) and "://" not in piece


def _clause_words(text: str, mark: int) -> list[str]:
    """Son ayractan '?'e kadarki cumlecigin sozcukleri (kucuk harf, tirnaksiz); kisaltmalar ('WHO') bos sozcuk olur."""
    start = max((b.end() for b in CLAUSE_BREAK.finditer(text, 0, mark)), default=0)
    words = (w.strip(QUOTE_CHARS) for w in text[start:mark].split())
    return ["" if len(w) > 1 and w.isupper() else w.casefold() for w in words]


def _title_question(text: str, mark: int) -> bool:
    """'?' oncesi cumlecikte soru basligi bicimi: soru sozcugu ('wat', 'Wann') ya da yardimci fiil + ozne ('Kan ik')."""
    words = _clause_words(text, mark)
    return any(w in QUESTION_WORDS or (w in AUX_VERBS and nxt in SUBJECTS) for w, nxt in zip(words, words[1:] + [""]))


def is_placeholder(value: Any) -> bool:
    """'[[', ']]', '{{', '}}', 'TODO', 'lorem', rakamdan hemen sonra '?' (ör. '5621?') veya tek kelime + '?'.

    Soru cumleleri ('... yapmalıyım?') ve URL sorgu dizgeleri ('page2?id=7') isaretlenmez. '?' kurallari tirnak icindeki
    metne ("…", '…', “…”, „…“, «…», ‘…’) ve NL/DE/EN soru basligi bicimindeki cumlecige bakmaz (pilot: "Kan ik … 2025?",
    'Zorgtoeslag: wat verandert in 2026?'); uzunluk olcut degildir ('Kayıt ücreti şu an 85?' yer tutucudur). Acik isaretler
    ('{{', 'TODO' ...) tirnak icinde de yer tutucudur.
    """
    if not isinstance(value, str):
        return False
    if PLACEHOLDER_MARKERS.search(value) or "TODO" in value:
        return True
    text = QUOTED.sub(" ", value)
    if any(not _title_question(text, m.end() - 1) for m in DIGIT_QUESTION.finditer(text)):
        return True
    return any(_single_token_question(p) for p in re.split(r"[;\n]", text))


def placeholders(sheets: Sheets) -> dict[str, Any]:
    items = [_item(sheet, rec, key=key, value=str(value)[:120])
             for sheet in CONTENT_SHEETS for rec in sheets.get(sheet, [])
             for key, value in rec.items() if not key.startswith("_") and is_placeholder(value)]
    return _result(items)


def wikipedia_labeled_official(sheets: Sheets) -> dict[str, Any]:
    """Wikipedia adresini 'Resmî' diye etiketleyen baglantilar ve resmi diye tanitilan Wikipedia kaynaklari."""
    wiki = lambda rec: _host(rec.get("url")).endswith("wikipedia.org")  # noqa: E731
    links = [_item("Bağlantılar", rec, label=rec.get("label"), url=rec.get("url"))
             for rec in sheets.get("Bağlantılar", []) if wiki(rec) and OFFICIAL_WORDS.search(str(rec.get("label", "")))]
    sources = [_item("Kaynaklar", rec, publisher=rec.get("publisher"), url=rec.get("url"))
               for rec in sheets.get("Kaynaklar", []) if wiki(rec)
               and OFFICIAL_WORDS.search(f"{rec.get('publisher', '')} {rec.get('title', '')}")]
    return _result(links + sources, by_sheet={"Bağlantılar": len(links), "Kaynaklar": len(sources)})


def _verified(value: Any) -> bool:
    return value is True or value == "Evet"


def verified_without_snapshot(sheets: Sheets) -> dict[str, Any]:
    verified = [rec for rec in sheets.get("Kaynaklar", []) if _verified(rec.get("verified"))]
    items = [_item("Kaynaklar", rec, missing=[k for k in ("snapshot_file", "snapshot_sha256") if not rec.get(k)])
             for rec in verified if not (rec.get("snapshot_file") and rec.get("snapshot_sha256"))]
    return _result(items, verified_total=len(verified))


def shared_coordinates(sheets: Sheets) -> dict[str, Any]:
    groups: dict[tuple[float, float], set[str]] = defaultdict(set)
    for rec in sheets.get("Mekânlar", []):
        lat, lon = rec.get("latitude"), rec.get("longitude")
        if isinstance(lat, (int, float)) and isinstance(lon, (int, float)) and rec.get("entry_id"):
            groups[(round(float(lat), 7), round(float(lon), 7))].add(str(rec["entry_id"]))
    items = [{"sheet": "Mekânlar", "latitude": lat, "longitude": lon, "entry_ids": sorted(ids)}
             for (lat, lon), ids in sorted(groups.items()) if len(ids) > 1]
    return _result(items, entries=sum(len(i["entry_ids"]) for i in items))


def is_homepage(url: Any) -> bool:
    """Yalniz alan adi: yol bos, '/', dil koku ('/en', '/nl-nl/') veya '/home.html' / '/index.htm'; sorgu yok."""
    try:
        parts = urlsplit(str(url).strip())
    except ValueError:
        return False
    return bool(parts.hostname) and not parts.query and HOMEPAGE_PATH.fullmatch(parts.path) is not None


def _split_ids(value: Any) -> list[str]:
    return [p.strip() for p in str(value or "").split(";") if p.strip()]


def homepage_only_sources(sheets: Sheets) -> dict[str, Any]:
    urls = {rec.get("source_id"): rec.get("url") for rec in sheets.get("Kaynaklar", []) if rec.get("url")}
    items = []
    for rec in sheets.get("Yanıtlar", []):
        cited = [urls[s] for s in _split_ids(rec.get("source_ids")) if s in urls]
        if cited and all(is_homepage(u) for u in cited):
            items.append(_item("Yanıtlar", rec, urls=cited))
    return _result(items)


def history_share(sheets: Sheets) -> dict[str, Any]:
    places = sheets.get("Mekânlar", [])
    history = sum(1 for rec in places if rec.get("group") == HISTORY_GROUP)
    ratio = history / len(places) if places else 0.0
    flagged = ratio > HISTORY_THRESHOLD
    return {"label": LABEL, "count": int(flagged), "items": [], "history": history, "total": len(places),
            "ratio": round(ratio, 4), "threshold": HISTORY_THRESHOLD, "flagged": flagged,
            "threshold_note": f"öneri: '{HISTORY_GROUP}' payı > {HISTORY_THRESHOLD:.0%} ise çeşitlilik gözden geçirilmeli"}


CHECKS: dict[str, Callable[[Sheets], dict[str, Any]]] = {
    "placeholders": placeholders,
    "wikipedia_labeled_official": wikipedia_labeled_official,
    "verified_without_snapshot": verified_without_snapshot,
    "shared_coordinates": shared_coordinates,
    "homepage_only_sources": homepage_only_sources,
    "history_share": history_share,
}


def run_all(sheets: Sheets) -> dict[str, dict[str, Any]]:
    return {name: check(sheets) for name, check in CHECKS.items()}


def counts(results: Mapping[str, Mapping[str, Any]]) -> dict[str, int]:
    return {name: result["count"] for name, result in results.items()}


def load_sheets(workbook: Path, names: Iterable[str] = LOADED_SHEETS) -> dict[str, list[dict[str, Any]]]:
    """Excel'i salt okur (satir 7 anahtarlari); olmayan sayfa bos liste olur."""
    wb = openpyxl.load_workbook(workbook, data_only=False)
    try:
        return {name: list(iter_records(wb[name])) if name in wb.sheetnames else [] for name in names}
    finally:
        wb.close()
