"""build_xlsx hucre kurallari: sayfa profili (sablondan turetilir) ve kayit dogrulama/donusturme.

Profil sablondan okunur: satir 7 anahtarlari, formul sutunlari, veri dogrulama listeleri, kapasite, tarih sutunlari.
Alan turleri (id listesi, metin listesi, sayi, tarih) --tools ile resmi schema_definition.py'den, yoksa yerlesik kumeden.
"""
from __future__ import annotations

import datetime as _dt
import importlib.util
import json
import math
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils.cell import range_boundaries

import common

SCHEMA_RELPATH = Path("core") / "content_delivery" / "schema_definition.py"
FALLBACK_ID_LIST = frozenset({"source_ids", "stop_entry_ids", "stages"})
FALLBACK_TEXT_LIST = frozenset({"steps", "documents", "tags", "stops"})
FALLBACK_NUMBER = frozenset({"latitude", "longitude", "order", "web_width", "web_height", "source_position",
                             "priority", "directory_order"})
JSON_KEYS = frozenset({"visual_json", "compare_json", "source_evidence_json"})
JSON_VALUE_SHEETS = frozenset({"SayfaAlanları", "EkDetaylar"})
DATE_SUFFIXES = ("_on", "_date")
ID_SEPARATOR = ";"
TEXT_SEPARATOR = "\n"
DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
DATETIME_RE = re.compile(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}(:\d{2})?$")
SKIP = object()  # bos deger: hucreye yazilmaz


class RuleError(ValueError):
    """Tek bir hucre degerinin reddedilme nedeni."""


@dataclass(frozen=True)
class FieldTypes:
    """Sayfa -> anahtar -> tur ('id_list', 'text_list', 'number', 'integer', 'date', ...)."""
    by_sheet: dict[str, dict[str, str]] = field(default_factory=dict)
    source: str = "builtin"

    def kind(self, sheet: str, key: str) -> str | None:
        return self.by_sheet.get(sheet, {}).get(key)


def load_field_types(tools: Path | None) -> FieldTypes:
    """--tools verildiyse resmi semadan alan turlerini okur (modul yalniz tanim icerir; yazmaz)."""
    if tools is None:
        return FieldTypes()
    path = Path(tools) / SCHEMA_RELPATH
    if not path.is_file():
        raise FileNotFoundError(f"--tools altinda schema_definition.py yok: {path}")
    spec = importlib.util.spec_from_file_location("corteqs_schema_definition", path)
    module = importlib.util.module_from_spec(spec)
    previous = sys.dont_write_bytecode
    sys.dont_write_bytecode = True  # resmi araclar klasorune __pycache__ yazma
    try:
        spec.loader.exec_module(module)
    finally:
        sys.dont_write_bytecode = previous
    by_sheet: dict[str, dict[str, str]] = {}
    for table in getattr(module, "TABLES", {}).values():
        by_sheet[table["sheet"]] = {f["key"]: f.get("type", "string") for f in table.get("fields", [])}
    return FieldTypes(by_sheet=by_sheet, source=str(path))


@dataclass(frozen=True)
class SheetProfile:
    name: str
    keys: dict[str, int]
    formulas: frozenset[str]
    enums: dict[str, tuple[str, ...]]
    kinds: dict[str, str]
    capacity: int | None
    merge: bool
    identity: dict[str, Any]


def literal_list(formula1: str | None) -> tuple[str, ...] | None:
    """DV formula1 '"a,b,c"' -> ('a','b','c'); aralik referansi -> None (dogrulanamaz)."""
    if not formula1 or not (formula1.startswith('"') and formula1.endswith('"')):
        return None
    return tuple(formula1[1:-1].split(","))


def dv_enums(ws, keys: dict[str, int]) -> dict[str, tuple[str, ...]]:
    """Ilk veri satirini kapsayan liste dogrulamalari (sutun bazinda tekdüze varsayilir)."""
    enums: dict[str, tuple[str, ...]] = {}
    for key, col in keys.items():
        coord = ws.cell(common.FIRST_DATA_ROW, col).coordinate
        for dv in ws.data_validations.dataValidation:
            options = literal_list(dv.formula1) if dv.type == "list" else None
            if options is not None and coord in dv.sqref:
                enums[key] = options
    return enums


def sheet_capacity(ws, merge: bool) -> int | None:
    """Kapasite: tek satir sayfasi 1; Excel tablosu varsa tablo araliginin son satiri; yoksa DV araliginin son satiri;
    o da yoksa sablon max_row. Bilinmiyorsa None."""
    if merge:
        return 1
    table_ends = [range_boundaries(t.ref)[3] for t in ws.tables.values()]
    dv_ends = [rng.max_row for dv in ws.data_validations.dataValidation for rng in dv.sqref.ranges]
    end = max(table_ends or dv_ends or [ws.max_row])
    return end - common.KEY_ROW if end >= common.FIRST_DATA_ROW else None


def column_kind(ws, sheet: str, key: str, col: int, types: FieldTypes) -> str:
    schema_kind = types.kind(sheet, key)
    if schema_kind:
        return schema_kind
    if key in FALLBACK_ID_LIST:
        return "id_list"
    if key in FALLBACK_TEXT_LIST:
        return "text_list"
    if key in FALLBACK_NUMBER:
        return "number"
    if key.endswith(DATE_SUFFIXES) or is_date_format(ws.cell(common.FIRST_DATA_ROW, col).number_format):
        return "date"
    return "string"


def build_profile(ws, types: FieldTypes) -> SheetProfile:
    keys = common.sheet_keys(ws)
    merge = ws.title in common.SINGLE_ROW_SHEETS
    identity = {}
    if merge:
        identity = {k: ws.cell(common.FIRST_DATA_ROW, keys[k]).value for k in common.IDENTITY_KEYS.get(ws.title, ()) if k in keys}
    return SheetProfile(
        name=ws.title, keys=keys, formulas=frozenset(common.formula_columns(ws)), enums=dv_enums(ws, keys),
        kinds={k: column_kind(ws, ws.title, k, c, types) for k, c in keys.items()},
        capacity=sheet_capacity(ws, merge), merge=merge, identity=identity,
    )


def dedup_ids(value: str) -> tuple[str, int]:
    """';' listesindeki tekrarlari sirayi koruyarak atar. Tekrar yoksa metin aynen kalir."""
    parts = [p.strip() for p in value.split(ID_SEPARATOR)]
    unique = list(dict.fromkeys(p for p in parts if p))
    removed = sum(1 for p in parts if p) - len(unique)
    return (ID_SEPARATOR.join(unique), removed) if removed else (value, 0)


def to_number(value: Any) -> int | float:
    if isinstance(value, bool):
        raise RuleError("number expected, got boolean")
    if isinstance(value, (int, float)) and math.isfinite(value):
        return value
    if isinstance(value, str):
        text = value.strip()
        try:
            return int(text)
        except ValueError:
            pass
        try:
            number = float(text)
        except ValueError:
            number = math.nan
        if math.isfinite(number):
            return number
    raise RuleError(f"number expected, got {value!r}"[:200])


def to_date(value: Any) -> _dt.date | _dt.datetime:
    if isinstance(value, str) and DATE_RE.match(value.strip()):
        try:
            return _dt.date.fromisoformat(value.strip())
        except ValueError as exc:
            raise RuleError(f"invalid date {value!r}") from exc
    if isinstance(value, str) and DATETIME_RE.match(value.strip()):
        return _dt.datetime.fromisoformat(value.strip())
    raise RuleError(f"date expected as YYYY-MM-DD, got {value!r}"[:200])


def from_list(value: list, kind: str) -> tuple[str, int]:
    if kind not in ("id_list", "text_list") or not all(isinstance(v, (str, int, float)) for v in value):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":")), 0
    items = [str(v).strip() for v in value]
    if kind == "text_list":
        return TEXT_SEPARATOR.join(items), 0
    unique = list(dict.fromkeys(i for i in items if i))
    return ID_SEPARATOR.join(unique), len([i for i in items if i]) - len(unique)


def check_text(value: str) -> str:
    if value.startswith("="):
        raise RuleError("formula injection: text starts with '='")
    if len(value) > common.MAX_CELL_CHARS:
        raise RuleError(f"text too long: {len(value)} > {common.MAX_CELL_CHARS} chars")
    if ILLEGAL_CHARACTERS_RE.search(value):
        raise RuleError("illegal control character in text")
    return value


def is_json_cell(profile: SheetProfile, key: str, record: dict[str, Any]) -> bool:
    if key in JSON_KEYS:
        return True
    return profile.name in JSON_VALUE_SHEETS and key == "value" and record.get("value_type") == "JSON"


def convert(profile: SheetProfile, key: str, value: Any, record: dict[str, Any]) -> tuple[Any, int]:
    """Tek hucre degeri -> (yazilacak deger | SKIP, atilan tekrar sayisi). Hata: RuleError."""
    if value is None or value == "":
        return SKIP, 0
    if key in profile.formulas:
        raise RuleError("formula column is computed by the template; value not allowed")
    kind = profile.kinds.get(key, "string")
    removed = 0
    if kind in ("number", "integer"):
        return to_number(value), 0
    if isinstance(value, bool):
        value = "Evet" if value else "Hayır"
    elif isinstance(value, list):
        value, removed = from_list(value, kind)
    elif isinstance(value, dict):
        value = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    if kind == "date":
        return to_date(value), 0
    if isinstance(value, str):
        value = check_text(value)
        if is_json_cell(profile, key, record):
            try:
                json.loads(value)
            except ValueError as exc:
                raise RuleError(f"invalid JSON text: {exc}"[:200]) from exc
        if kind == "id_list":
            value, extra = dedup_ids(value)
            removed += extra
    return value, removed


@dataclass
class RecordResult:
    values: dict[str, Any] = field(default_factory=dict)
    errors: list[tuple[str | None, str]] = field(default_factory=list)
    deduped: dict[str, int] = field(default_factory=dict)
    defaults: list[str] = field(default_factory=list)


def check_enum_and_human(profile: SheetProfile, key: str, value: Any, allow_human: bool) -> None:
    options = profile.enums.get(key)
    if options is not None and str(value) not in options:
        raise RuleError(f"value {str(value)[:80]!r} not in list {','.join(options)}")
    if key in common.HUMAN_REVIEW_KEYS and value != common.HUMAN_PENDING_VALUE and not allow_human:
        raise RuleError(f"human review field: AI may only write {common.HUMAN_PENDING_VALUE!r} (use --allow-human-values)")


def check_identity(profile: SheetProfile, key: str, value: Any) -> None:
    current = profile.identity.get(key)
    if key not in profile.identity or current in (None, ""):
        return
    a, b = common.normalize_value(current), common.normalize_value(value)
    same = math.isclose(a, b, abs_tol=1e-9) if all(isinstance(v, (int, float)) for v in (a, b)) else str(a) == str(b)
    if not same:
        raise RuleError(f"identity key may not change: template {a!r} != {b!r}"[:200])


def apply_defaults(profile: SheetProfile, result: RecordResult, record: dict[str, Any]) -> None:
    if profile.merge:
        return
    for key in common.HUMAN_REVIEW_KEYS:
        absent = record.get(key) in (None, "")
        if key in profile.keys and absent and common.HUMAN_PENDING_VALUE in profile.enums.get(key, ()):
            result.values[key] = common.HUMAN_PENDING_VALUE
            result.defaults.append(key)


def check_record(profile: SheetProfile, record: dict[str, Any], allow_human: bool) -> RecordResult:
    result = RecordResult()
    for key, raw in record.items():
        if key == "_row":
            continue
        if key not in profile.keys:
            result.errors.append((key, f"unknown key (not in row {common.KEY_ROW} of {profile.name})"))
            continue
        try:
            value, removed = convert(profile, key, raw, record)
            if value is SKIP:
                continue
            check_enum_and_human(profile, key, value, allow_human)
            check_identity(profile, key, value)
        except RuleError as exc:
            result.errors.append((key, str(exc)))
            continue
        result.values[key] = value
        if removed:
            result.deduped[key] = removed
    apply_defaults(profile, result, record)
    return result
