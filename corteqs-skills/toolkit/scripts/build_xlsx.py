"""build_xlsx.py --template T.xlsx --data-dir D --out OUT.xlsx [--report R.json] [--tools <araclar/arac>]
                 [--overwrite] [--partial] [--allow-human-values]

Dogrulanmis JSONL verisini sehir sablonunun KOPYASINA yazar (sablon asla degistirilmez).
- D/<Sayfa>.jsonl (ASCII ad da olur: Yanitlar.jsonl). Dosya yoksa sayfa sablondaki gibi kalir; bilinmeyen dosya = hata.
- Yazim satir 7 anahtarina gore (sutun sirasina degil); `_row` yok sayilir. Teslimat/Sehir: satir 8 birlestirme (merge),
  kimlik anahtarlari degistirilemez. Formul sutunlari, DV listesi disi deger, '=' ile baslayan metin, >32767 karakter,
  bilinmeyen anahtar, kapasite asimi, insan inceleme alanina 'Bekliyor' disi deger = reddedilen satir.
- Reddedilen satir varsa hicbir sey yazilmaz (cikis 1); --partial ile gecerli satirlar yazilir (yine cikis 1).
- Yazim gecici dosyaya, sadakat denetimi (xlsx_fidelity) gecerse yeniden adlandirma. Fark = cikti yok, cikis 2.
stdout: {rows_written, rejected, deduped, defaults_applied, fidelity, out}. Cikis: 0 temiz / 1 red / 2 kullanim-yapi-sadakat.
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import time
import unicodedata
import uuid
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

import common
import xlsx_fidelity
import xlsx_rules


class BuildError(Exception):
    """Kullanim, yapi veya sadakat hatasi (cikis 2)."""


@dataclass(frozen=True)
class Options:
    template: Path
    data_dir: Path
    out: Path
    report: Path | None = None
    tools: Path | None = None
    overwrite: bool = False
    partial: bool = False
    allow_human: bool = False
    strict_parts: bool = False


@dataclass
class SheetPlan:
    rows: list[tuple[int, dict[str, Any]]] = field(default_factory=list)
    rejected: list[dict[str, Any]] = field(default_factory=list)
    rejected_rows: int = 0
    deduped: list[dict[str, Any]] = field(default_factory=list)
    defaults: int = 0


def ascii_name(name: str) -> str:
    text = unicodedata.normalize("NFKD", name.replace("ı", "i").replace("İ", "I"))
    return "".join(ch for ch in text if not unicodedata.combining(ch)).casefold()


def find_data_files(data_dir: Path, sheetnames: list[str]) -> dict[str, Path]:
    """<Sayfa>.jsonl -> sayfa. Ayni sayfaya iki dosya veya sablonda olmayan sayfa adi = kullanim hatasi."""
    index = {ascii_name(name): name for name in sheetnames}
    found: dict[str, Path] = {}
    unknown: list[str] = []
    for path in sorted(data_dir.glob("*.jsonl")):
        sheet = index.get(ascii_name(path.stem))
        if sheet is None:
            unknown.append(path.name)
        elif sheet in found:
            raise BuildError(f"ayni sayfa icin iki dosya: {found[sheet].name} ve {path.name}")
        else:
            found[sheet] = path
    if unknown:
        raise BuildError(f"sablonda karsiligi olmayan JSONL dosyasi: {', '.join(unknown)}")
    return found


def read_jsonl(path: Path) -> list[tuple[int, dict[str, Any] | None, str | None]]:
    """(satir_no, kayit | None, hata | None). Bos satirlar atlanir."""
    entries = []
    for number, line in enumerate(path.read_text(encoding="utf-8-sig").splitlines(), start=1):
        if not line.strip():
            continue
        try:
            record = json.loads(line)
        except ValueError as exc:
            entries.append((number, None, f"invalid JSON line: {exc}"[:200]))
            continue
        if not isinstance(record, dict):
            entries.append((number, None, "JSON line is not an object"))
            continue
        entries.append((number, record, None))
    return entries


def same_file(a: Path, b: Path) -> bool:
    if a.resolve() == b.resolve():
        return True
    try:
        return b.exists() and os.path.samefile(a, b)
    except OSError:
        return False


def check_paths(opts: Options) -> None:
    if not opts.template.is_file():
        raise BuildError(f"template bulunamadi: {opts.template}")
    if not opts.data_dir.is_dir():
        raise BuildError(f"--data-dir klasor degil: {opts.data_dir}")
    if same_file(opts.template, opts.out):
        raise BuildError("--out template ile ayni dosya; girdi asla uzerine yazilmaz")
    if opts.out.exists() and not opts.overwrite:
        raise BuildError(f"--out zaten var (uzerine yazmak icin --overwrite): {opts.out}")


def load_template(path: Path):
    try:
        return openpyxl.load_workbook(path)
    except (OSError, zipfile.BadZipFile, InvalidFileException, KeyError, ValueError) as exc:
        raise BuildError(f"template okunamadi: {path} ({type(exc).__name__}: {exc})") from exc


def rejection(sheet: str, line: int, record: dict | None, key: str | None, reason: str) -> dict[str, Any]:
    return {"sheet": sheet, "line": line, "_row": (record or {}).get("_row"), "key": key, "reason": reason}


def reject_all(sheet: str, entries, reason: str) -> SheetPlan:
    plan = SheetPlan()
    for line, record, error in entries:
        plan.rejected.append(rejection(sheet, line, record, None, error or reason))
        plan.rejected_rows += 1
    return plan


def add_valid(plan: SheetPlan, profile, line: int, record: dict, result) -> None:
    if profile.capacity is not None and len(plan.rows) >= profile.capacity:
        plan.rejected.append(rejection(profile.name, line, record, None,
                                       f"capacity exceeded: sheet holds {profile.capacity} data rows"))
        plan.rejected_rows += 1
        return
    row = common.FIRST_DATA_ROW + len(plan.rows)
    plan.rows.append((row, result.values))
    plan.deduped += [{"sheet": profile.name, "line": line, "row": row, "key": k, "removed": n} for k, n in result.deduped.items()]
    plan.defaults += len(result.defaults)


def plan_sheet(profile, entries, allow_human: bool) -> SheetPlan:
    plan = SheetPlan()
    for line, record, error in entries:
        if error:
            plan.rejected.append(rejection(profile.name, line, record, None, error))
            plan.rejected_rows += 1
            continue
        result = xlsx_rules.check_record(profile, record, allow_human)
        if result.errors:
            plan.rejected += [rejection(profile.name, line, record, key, reason) for key, reason in result.errors]
            plan.rejected_rows += 1
            continue
        add_valid(plan, profile, line, record, result)
    return plan


def plan_workbook(wb, files: dict[str, Path], types, allow_human: bool) -> tuple[dict[str, SheetPlan], dict[str, Any]]:
    plans: dict[str, SheetPlan] = {}
    capacities: dict[str, Any] = {}
    for sheet, path in files.items():
        entries = read_jsonl(path)
        ws = wb[sheet]
        if sheet not in common.DATA_SHEETS:
            plans[sheet] = reject_all(sheet, entries, "sheet not writable (formula/summary sheet)")
            continue
        try:
            profile = xlsx_rules.build_profile(ws, types)
        except ValueError as exc:
            raise BuildError(str(exc)) from exc
        if not profile.merge and next(common.iter_records(ws), None) is not None:
            raise BuildError(f"{sheet}: sablon sayfasinda zaten veri satiri var; bos sablon gerekir")
        capacities[sheet] = profile.capacity
        plans[sheet] = plan_sheet(profile, entries, allow_human)
    return plans, capacities


def write_plans(wb, plans: dict[str, SheetPlan]) -> dict[str, int]:
    written: dict[str, int] = {}
    for sheet, plan in plans.items():
        if not plan.rows:
            continue
        ws = wb[sheet]
        keys = common.sheet_keys(ws)
        for row, values in plan.rows:
            for key, value in values.items():
                ws.cell(row, keys[key]).value = value
        written[sheet] = len(plan.rows)
    return written


def save_checked(wb, template: Path, out: Path, accept_parts: bool) -> tuple[list[str], list[str]]:
    """Gecici dosyaya yazar, sadakat denetler; temizse --out'a tasir. Doner: (farklar, kabul edilen farklar)."""
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_name(f"{out.stem}.tmp-{uuid.uuid4().hex[:8]}.xlsx")
    try:
        wb.save(tmp)
        diffs, accepted = xlsx_fidelity.check(template, tmp, accept_parts)
        if not diffs:
            os.replace(tmp, out)
        return diffs, accepted
    finally:
        if tmp.exists():
            tmp.unlink()


def run_build(opts: Options) -> tuple[int, dict[str, Any], dict[str, Any]]:
    check_paths(opts)
    try:
        types = xlsx_rules.load_field_types(opts.tools)
    except (OSError, SyntaxError, ImportError, KeyError, TypeError) as exc:
        raise BuildError(f"--tools schema_definition okunamadi: {exc}") from exc
    wb = load_template(opts.template)
    files = find_data_files(opts.data_dir, wb.sheetnames)
    plans, capacities = plan_workbook(wb, files, types, opts.allow_human)
    rejected_rows = sum(p.rejected_rows for p in plans.values())
    summary: dict[str, Any] = {"rows_written": {}, "rejected": rejected_rows,
                               "deduped": sum(d["removed"] for p in plans.values() for d in p.deduped),
                               "defaults_applied": sum(p.defaults for p in plans.values()), "fidelity": "not_run", "fidelity_accepted": [], "out": None}
    details = {"files": {s: p.name for s, p in files.items()}, "capacities": capacities, "schema_source": types.source,
               "rejected": [r for p in plans.values() for r in p.rejected],
               "deduped": [d for p in plans.values() for d in p.deduped],
               "rows_planned": {s: len(p.rows) for s, p in plans.items()}}
    if rejected_rows and not opts.partial:
        return common.EXIT_FINDINGS, summary, details
    written = write_plans(wb, plans)
    diffs, accepted = save_checked(wb, opts.template, opts.out, not opts.strict_parts)
    summary.update(fidelity=diffs or "ok", fidelity_accepted=accepted)
    if diffs:
        return common.EXIT_USAGE, summary, details
    summary.update(rows_written=written, out=str(opts.out.resolve()))
    return (common.EXIT_FINDINGS if rejected_rows else common.EXIT_OK), summary, details


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Dogrulanmis JSONL verisini sehir sablonunun kopyasina yazar.")
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--data-dir", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--report", "--out-report", dest="report", type=Path, default=None)
    parser.add_argument("--tools", type=Path, default=None, help="araclar/arac koku (resmi schema_definition.py icin)")
    parser.add_argument("--overwrite", action="store_true", help="var olan --out dosyasinin uzerine yaz (template asla)")
    parser.add_argument("--partial", action="store_true", help="red varken de gecerli satirlari yaz")
    parser.add_argument("--allow-human-values", dest="allow_human", action="store_true",
                        help="insan inceleme alanlarina 'Bekliyor' disi degere izin ver (yalniz insan girdisi icin)")
    parser.add_argument("--strict-parts", action="store_true",
                        help="varsayilan olarak kabul edilen 3 openpyxl zip farkini (docProps ekleme, sharedStrings kalkma) hata say")
    return parser


def main(argv: list[str] | None = None) -> int:
    common.setup_utf8()
    args = build_parser().parse_args(argv)
    opts = Options(**{k: v for k, v in vars(args).items()})
    started = time.perf_counter()
    try:
        code, summary, details = run_build(opts)
    except BuildError as exc:
        code, summary, details = common.EXIT_USAGE, {"error": str(exc)}, {}
    if opts.report:
        common.write_json(opts.report, {"summary": summary, **details, "exit_code": code, "partial": opts.partial,
                                        "elapsed_s": round(time.perf_counter() - started, 3)})
    common.emit(summary)
    return code


if __name__ == "__main__":
    sys.exit(main())
