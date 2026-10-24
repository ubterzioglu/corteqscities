"""inventory.py <package_root> [--xlsx X] [--out dosya-envanteri.json]

Paketi salt-okunur gezer: dosya basina yol/boyut/sha256/uzanti, resimler icin GERCEK bicim+boyut (Pillow),
uzanti kusurlari (bad_ext, ext_mismatch) ve --xlsx verilirse Fotograflar / FotografKaynaklari capraz kontrolu.

stdout: tek satir JSON {files,total_mb,images,missing,sha_mismatch,file_ne_original,bad_ext,ext_mismatch,orphans,dup_sha,errors}.
--out: tum ayrintilar (dosya listesi + bulgular). Varsayilan: <cwd>/dosya-envanteri.json; paket icine dusuyorsa REDDEDILIR
(paket icine yalniz acik --out ile yazilir; o dosya bir sonraki taramada envantere girmez).
Cikis: 1 = missing/sha_mismatch/file_ne_original/bad_ext/ext_mismatch/errors (guvensiz yol, bos referans, bozuk resim dahil) > 0;
0 = temiz; 2 = kullanim/yapi hatasi. orphans ve dup_sha bilgi amaclidir, cikis kodunu etkilemez.
total_mb = bayt / 1.000.000 (ondalik MB); total_bytes ayrintida. missing/sha_mismatch tekil YOL sayisidir (ayrintida her referans).
bad_ext olan dosya ayrica ext_mismatch sayilmaz. Orphan: Excel'de (Fotograflar.file, FotografKaynaklari.original/web,
Kaynaklar.snapshot_file) hic anilmayan dosya; Teslimler/, xlsx'in kendisi ve *.json durum dosyalari haric (sayilari listelenir).
"""
from __future__ import annotations

import argparse
import os
import re
import sys
import zipfile
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException
from PIL import Image, UnidentifiedImageError

import common

IMAGE_EXTS = frozenset({"jpg", "jpeg", "png", "webp", "gif", "bmp", "tif", "tiff"})
FORMAT_EXTS = {"JPEG": {"jpg", "jpeg"}, "MPO": {"jpg", "jpeg"}, "TIFF": {"tif", "tiff"}}
GOOD_EXT = re.compile(r"[a-z0-9]*")
PHOTO_SHEETS = ("Fotoğraflar", "FotoğrafKaynakları")


class InventoryError(Exception):
    """Kullanim veya yapi hatasi (cikis kodu 2)."""


def ext_of(name: str) -> str:
    """Son noktadan sonrasi, kucuk harf, noktasiz ('x.JPG' -> 'jpg'; uzantisiz -> '')."""
    return Path(name).suffix[1:].lower()


def is_bad_ext(ext: str) -> bool:
    """Uzanti '.' ile baslayabilir; [a-z0-9] disinda karakter varsa kusurludur."""
    return GOOD_EXT.fullmatch(ext.lstrip(".").lower()) is None


def probe_image(path: Path) -> tuple[dict[str, Any] | None, str | None]:
    """(bicim bilgisi, hata). Pillow acamazsa (None, None); acip dogrulayamazsa (None, mesaj). Tek gecis: verify()."""
    try:
        with Image.open(path) as img:
            info = {"format": img.format, "width": img.width, "height": img.height}
            img.verify()
        return info, None
    except (UnidentifiedImageError, Image.DecompressionBombError):
        return None, None
    except (OSError, ValueError, SyntaxError) as exc:
        return None, f"{type(exc).__name__}: {exc}"


def describe_file(root: Path, path: Path) -> tuple[dict[str, Any], str | None]:
    ext = ext_of(path.name)
    record: dict[str, Any] = {"path": path.relative_to(root).as_posix(), "size": path.stat().st_size,
                              "sha256": common.sha256_file(path), "ext": ext, "bad_ext": is_bad_ext(ext)}
    info, error = probe_image(path)
    if info is not None:
        allowed = FORMAT_EXTS.get(info["format"], {str(info["format"]).lower()})
        record.update(info)
        record["ext_mismatch"] = (not record["bad_ext"]) and ext not in allowed
        return record, None
    record["ext_mismatch"] = False
    if error is None and ext in IMAGE_EXTS:
        error = "UnidentifiedImageError: resim olarak acilamadi"
    return record, error


def walk_files(root: Path, skip: set[Path]) -> list[Path]:
    found: list[Path] = []
    for folder, dirs, names in os.walk(root):
        dirs.sort()
        for name in sorted(names):
            path = Path(folder) / name
            if path.resolve() not in skip:
                found.append(path)
    return found


def scan(root: Path, skip: set[Path]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    files: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    for path in walk_files(root, skip):
        record, error = describe_file(root, path)
        files.append(record)
        if error:
            errors.append({"kind": "image_unreadable", "path": record["path"], "error": error})
    return files, errors


def folder_totals(files: list[dict[str, Any]]) -> dict[str, dict[str, int]]:
    totals: dict[str, dict[str, int]] = {}
    for f in files:
        parts = f["path"].split("/")
        top = parts[0] if len(parts) > 1 else "."
        entry = totals.setdefault(top, {"files": 0, "bytes": 0})
        entry["files"] += 1
        entry["bytes"] += f["size"]
    return dict(sorted(totals.items()))


def duplicate_groups(files: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_sha: dict[str, list[str]] = {}
    for f in files:
        by_sha.setdefault(f["sha256"], []).append(f["path"])
    return [{"sha256": sha, "files": paths} for sha, paths in sorted(by_sha.items()) if len(paths) > 1]


def load_sheets(xlsx: Path) -> dict[str, list[dict[str, Any]]]:
    try:
        wb = openpyxl.load_workbook(xlsx, data_only=False)
    except (OSError, zipfile.BadZipFile, InvalidFileException, KeyError, ValueError) as exc:
        raise InventoryError(f"xlsx okunamadi: {xlsx} ({type(exc).__name__}: {exc})") from exc
    try:
        absent = [name for name in PHOTO_SHEETS if name not in wb.sheetnames]
        if absent:
            raise InventoryError(f"xlsx'te sayfa yok: {', '.join(absent)}")
        sheets = {name: list(common.iter_records(wb[name])) for name in PHOTO_SHEETS}
        if "Kaynaklar" in wb.sheetnames:
            sheets["Kaynaklar"] = list(common.iter_records(wb["Kaynaklar"]))
        return sheets
    except ValueError as exc:
        raise InventoryError(str(exc)) from exc
    finally:
        wb.close()


class Checker:
    """Excel referanslarini paket dosyalariyla karsilastirir; bulgulari biriktirir."""

    def __init__(self, root: Path, files: list[dict[str, Any]]) -> None:
        self.root = root.resolve()
        self.sha_by_path = {f["path"]: f["sha256"] for f in files}
        self.referenced: set[str] = set()
        self.missing: list[dict[str, Any]] = []
        self.sha_mismatch: list[dict[str, Any]] = []
        self.errors: list[dict[str, Any]] = []

    def check(self, sheet: str, row: dict[str, Any], column: str, path_value: Any, sha_value: Any) -> None:
        ref = {"sheet": sheet, "row": row["_row"], "photo_id": row.get("photo_id"), "column": column}
        if not isinstance(path_value, str) or not path_value.strip():
            self.errors.append({"kind": "empty_reference", **ref})
            return
        try:
            target = common.safe_relpath(self.root, path_value)
        except ValueError as exc:
            self.errors.append({"kind": "unsafe_path", "path": path_value, "error": str(exc), **ref})
            return
        rel = target.relative_to(self.root).as_posix()
        self.referenced.add(rel)
        if not target.is_file():
            self.missing.append({"path": rel, **ref})
            return
        actual = self.sha_by_path.get(rel) or common.sha256_file(target)
        if sha_value is not None and str(sha_value).strip().lower() != actual:
            self.sha_mismatch.append({"path": rel, "expected": str(sha_value).strip().lower(), "actual": actual, **ref})

    def snapshot_refs(self, rows: list[dict[str, Any]]) -> None:
        for row in rows:
            try:
                value = row.get("snapshot_file")
                if value:
                    self.referenced.add(common.safe_relpath(self.root, value).relative_to(self.root).as_posix())
            except ValueError:
                continue


def norm_path(value: Any) -> str:
    return str(value).strip().replace("\\", "/") if value is not None else ""


def norm_sha(value: Any) -> str:
    return str(value).strip().lower() if value is not None else ""


def find_file_ne_original(photos: list[dict[str, Any]], provenance: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_id = {p["photo_id"]: p for p in provenance if "photo_id" in p}
    found = []
    for photo in photos:
        prov = by_id.get(photo.get("photo_id"))
        if prov is None:
            continue
        same_file = norm_path(photo.get("file")) == norm_path(prov.get("original_file"))
        same_sha = norm_sha(photo.get("sha256")) == norm_sha(prov.get("original_sha256"))
        if not (same_file and same_sha):
            found.append({"photo_id": photo.get("photo_id"), "row": photo["_row"], "file": photo.get("file"), "original_file": prov.get("original_file"),
                          "sha256": photo.get("sha256"), "original_sha256": prov.get("original_sha256")})
    return found


def cross_check(root: Path, files: list[dict[str, Any]], sheets: dict[str, list[dict[str, Any]]]) -> Checker:
    checker = Checker(root, files)
    for row in sheets["Fotoğraflar"]:
        checker.check("Fotoğraflar", row, "file", row.get("file"), row.get("sha256"))
    for row in sheets["FotoğrafKaynakları"]:
        checker.check("FotoğrafKaynakları", row, "original_file", row.get("original_file"), row.get("original_sha256"))
        checker.check("FotoğrafKaynakları", row, "web_file", row.get("web_file"), row.get("web_sha256"))
    checker.snapshot_refs(sheets.get("Kaynaklar", []))
    return checker


def find_orphans(files: list[dict[str, Any]], referenced: set[str], xlsx_rel: str | None) -> tuple[list[str], dict[str, int]]:
    excluded = {"teslimler": 0, "xlsx": 0, "json": 0}
    orphans = []
    for f in files:
        path = f["path"]
        if path in referenced:
            continue
        if path.split("/")[0] == "Teslimler" and "/" in path:
            excluded["teslimler"] += 1
        elif path == xlsx_rel:
            excluded["xlsx"] += 1
        elif f["ext"] == "json":
            excluded["json"] += 1
        else:
            orphans.append(path)
    return orphans, excluded


def build_report(root: Path, xlsx: Path | None, skip: set[Path]) -> dict[str, Any]:
    files, errors = scan(root, skip)
    findings: dict[str, Any] = {
        "bad_ext": [f for f in files if f["bad_ext"]],
        "ext_mismatch": [f for f in files if f["ext_mismatch"]],
        "dup_sha": duplicate_groups(files),
        "missing": [], "sha_mismatch": [], "file_ne_original": [], "orphans": [],
        "orphans_excluded": {"teslimler": 0, "xlsx": 0, "json": 0},
    }
    if xlsx is not None:
        sheets = load_sheets(xlsx)
        checker = cross_check(root, files, sheets)
        xlsx_rel = xlsx.resolve().relative_to(root.resolve()).as_posix() if root.resolve() in xlsx.resolve().parents else None
        orphans, excluded = find_orphans(files, checker.referenced, xlsx_rel)
        findings.update(missing=checker.missing, sha_mismatch=checker.sha_mismatch, file_ne_original=find_file_ne_original(sheets["Fotoğraflar"], sheets["FotoğrafKaynakları"]),
                        orphans=orphans, orphans_excluded=excluded)
        errors = errors + checker.errors
    findings["errors"] = errors
    total = sum(f["size"] for f in files)
    totals = {"files": len(files), "total_bytes": total, "total_mb": round(total / 1_000_000, 2)}
    summary = {
        "files": len(files), "total_mb": totals["total_mb"], "images": sum(1 for f in files if "format" in f),
        "missing": len({m["path"] for m in findings["missing"]}), "sha_mismatch": len({m["path"] for m in findings["sha_mismatch"]}),
        "file_ne_original": len(findings["file_ne_original"]), "bad_ext": len(findings["bad_ext"]),
        "ext_mismatch": len(findings["ext_mismatch"]), "orphans": len(findings["orphans"]), "dup_sha": len(findings["dup_sha"]),
        "errors": len(errors),
    }
    return {"root": root.name, "xlsx_checked": xlsx is not None, "summary": summary, "totals": totals,
            "folders": folder_totals(files), "findings": findings, "files": files}


EXIT_KEYS = ("missing", "sha_mismatch", "file_ne_original", "bad_ext", "ext_mismatch", "errors")


def resolve_out(root: Path, out: Path | None) -> Path:
    if out is not None:
        return out
    default = Path.cwd() / "dosya-envanteri.json"
    resolved_root = root.resolve()
    if resolved_root == default.resolve().parent or resolved_root in default.resolve().parents:
        raise InventoryError("varsayilan --out paket klasorune dusuyor; paket disinda bir --out verin")
    return default


def run(root: Path, xlsx: Path | None, out: Path | None) -> int:
    if not root.is_dir():
        raise InventoryError(f"paket klasoru yok: {root}")
    if xlsx is not None and not xlsx.is_file():
        raise InventoryError(f"xlsx yok: {xlsx}")
    out_path = resolve_out(root, out)
    report = build_report(root, xlsx, {out_path.resolve()})
    common.write_json(out_path, report)
    common.emit(report["summary"])
    return common.EXIT_FINDINGS if any(report["summary"][k] > 0 for k in EXIT_KEYS) else common.EXIT_OK


def main(argv: list[str] | None = None) -> int:
    common.setup_utf8()
    parser = argparse.ArgumentParser(description="Paket dosya envanteri + Excel capraz kontrolu (salt-okunur).")
    parser.add_argument("package_root", type=Path)
    parser.add_argument("--xlsx", type=Path, default=None)
    parser.add_argument("--out", type=Path, default=None, help="ayrinti JSON (varsayilan: <cwd>/dosya-envanteri.json, paket disinda olmali)")
    args = parser.parse_args(argv)
    try:
        return run(args.package_root, args.xlsx, args.out)
    except InventoryError as exc:
        common.emit({"error": str(exc)})
        return common.EXIT_USAGE


if __name__ == "__main__":
    sys.exit(main())
