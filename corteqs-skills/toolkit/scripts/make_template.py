"""make_template.py: Berlin örneğinden (ekler\\CITY-000_berlin_ornek_v001.xlsx) yeni şehir için BOŞ v7 şablonu türetir.

Atınç'ın o şehre ait şablonu varsa o kullanılır. Örnek salt okunur; iş zip/XML düzeyinde yapılır (openpyxl kaydetmez):
DV (x14 dahil), CF (extLst dahil), stiller, genişlik/yükseklik, dondurma ve tablo stilleri korunur. Ayrıntı: --help ve
city-pack/references/sablon_turetme.md. Sekmeler = base-schema.json tabloları + Özet + FotoğrafKaynakları (gerisi ve
yalnız onlara bağlı parçalar atılır); satır 8+ boşaltılır (formüller kalır), tablolar kapasiteye kırpılır, Özet notları
atılır, Teslimat/Şehir satır 8'e yalnız kimlik alanları yazılır (delivery_id küçük harf; NNN >= 900 = yerel deneme).
Öz denetim (yazmadan önce): sekmeler, satır 7 anahtarları, tablo/kapasite, boş veri satırları, örnek belirteci taraması
(Berlin, CITY-013, CITY-000 ...), medya/çizim/bağımsız parça yok, paket ([Content_Types].xml'de her ad tek kez, her
kayıt var olan parçayı gösterir; Excel ve .NET buna bakar), openpyxl ile açılış. Bulgu varsa çıktı yazılmaz.
Yollar her şeyden önce denetlenir (--out/--report örneği, araçları ya da birbirini ezemez); yol hatasında rapor yazılmaz.
stdout: tek satır JSON özet. Çıkış: 0 temiz / 1 öz-denetim bulgusu (çıktı yok) / 2 kullanım-yapı hatası.
Kod üç dosyada: make_template_zip (zip/XML, yeniden numaralama, paket denetimi, hücre yardımcıları),
make_template_sheets (sayfa, Özet, formül, çalışma kitabı) ve bu dosya (komut satırı, sözleşme, türetme, öz denetim).
"""
from __future__ import annotations

import argparse
import contextlib
import datetime as _dt
import io
import json
import os
import re
import sys
import time
import uuid
import warnings
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils.cell import range_boundaries

import common
from make_template_sheets import SUMMARY_SHEET, TITLE_CELL, Book, Sheet, Stats, refit_formulas, strip_notes, write_cached
from make_template_zip import (Package, TemplateError, cell_text, find_tokens, fold, get_cell, has_value, local, m,
                               package_findings, rels_name, renumber, row_map, set_value, split_ref)

PROVENANCE_SHEET = "FotoğrafKaynakları"  # resmî read_delivery_v7 bu adla okur; alanları provenance-fields.json
METADATA_TABLE, CITY_TABLE = "metadata", "city"
DEFAULT_EDITOR = "Claude Code city-pack"
SCOPE, STATUS, LOCAL_TRIAL_FROM = "Tam teslim", "Taslak", 900
FIXED_TOKENS = ("Berlin", "berlin", "CITY-013", "CITY-000")
NOTE = "{base} örneğinden türetildi (make_template); içerik üretilmedi."
LOCAL_TRIAL_NOTE = " YEREL DENEME; resmî katalog numarası değil."
SLUG_RE = re.compile(r"^[a-z0-9]+(?:-[a-z0-9]+)*$")
COUNTRY_CODE_RE = re.compile(r"^[A-Z]{2,3}$")
CITY_NO_RE = re.compile(r"city-(\d{3})", re.IGNORECASE)
# Örnekten okunurken beklenmedik yapı (docx, bozuk/kesik XML, UTF-8 olmayan parça, eksik parça): çıkış 2, hata dökümü yok
STRUCTURE_ERRORS = (ET.ParseError, ValueError, KeyError, IndexError, TypeError, AttributeError)


# --- sözleşme ve seçenekler --------------------------------------------------------------------------------------

@dataclass(frozen=True)
class TableSpec:
    key: str
    sheet: str
    capacity: int | None
    fields: tuple[dict[str, Any], ...]

    @property
    def keys(self) -> tuple[str, ...]:
        return tuple(f["key"] for f in self.fields)

    def field(self, key: str) -> dict[str, Any]:
        return next((f for f in self.fields if f["key"] == key), {})


@dataclass(frozen=True)
class Contract:
    source: str
    schema_version: str
    tables: dict[str, TableSpec]  # sayfa adı -> tablo
    provenance_keys: tuple[str, ...] | None

    def by_key(self, key: str) -> TableSpec:
        return next(t for t in self.tables.values() if t.key == key)

    @property
    def sheets(self) -> set[str]:
        return set(self.tables) | {SUMMARY_SHEET, PROVENANCE_SHEET}


def load_contract(tools: Path) -> Contract:
    path = tools / "base-schema.json"
    try:
        schema = json.loads(path.read_text(encoding="utf-8"))
        tables = {c["sheet"]: TableSpec(k, c["sheet"], c.get("capacity"), tuple(c.get("fields", []))) for k, c in schema["tables"].items()}
        prov_path = tools / "provenance-fields.json"
        prov = tuple(f[0] for f in json.loads(prov_path.read_text(encoding="utf-8"))) if prov_path.is_file() else None
    except (OSError, ValueError, KeyError, TypeError, IndexError) as exc:
        raise TemplateError(f"sözleşme okunamadı ({path}): {type(exc).__name__}: {exc}") from exc
    if not {METADATA_TABLE, CITY_TABLE} <= {t.key for t in tables.values()}:
        raise TemplateError(f"base-schema.json'da {METADATA_TABLE}/{CITY_TABLE} tablosu yok: {path}")
    return Contract(str(path), str(schema.get("schema_version", "")), tables, prov)


@dataclass(frozen=True)
class Options:
    base: Path
    out: Path
    city_id: str
    city_name: str
    country: str
    country_code: str
    region: str
    kind: str
    city_number: int
    updated_on: _dt.date
    editor: str


def text_arg(flag: str, value: str | None, limit: int | None) -> str:
    text = (value or "").strip()
    if not text or text.startswith("=") or ILLEGAL_CHARACTERS_RE.search(text) or (limit and len(text) > limit):
        rule = f", en çok {limit} karakter" if limit else ""
        raise TemplateError(f"{flag} boş olamaz, '=' ile başlayamaz, denetim karakteri içeremez{rule}: {value!r}")
    return text


def choice_arg(flag: str, value: str, spec: TableSpec, key: str) -> str:
    choices = spec.field(key).get("choices")
    if choices and value not in choices:
        raise TemplateError(f"{flag} şu değerlerden biri olmalı: {', '.join(choices)} (verilen: {value})")
    return value


def parse_day(value: str | None) -> _dt.date:
    if value is None:
        return _dt.date.today()
    try:
        day = _dt.date.fromisoformat(value) if re.fullmatch(r"\d{4}-\d{2}-\d{2}", value) else None
    except ValueError:
        day = None
    if day is None:
        raise TemplateError(f"--updated-on YYYY-MM-DD olmalı: {value}")
    if day > _dt.date.today():
        raise TemplateError(f"--updated-on gelecekte olamaz (resmî doğrulayıcı reddeder): {value}")
    return day


def check_options(args: argparse.Namespace, contract: Contract) -> Options:
    city = contract.by_key(CITY_TABLE)
    if not SLUG_RE.match(args.city_id or ""):
        raise TemplateError(f"--city-id küçük harf ASCII adres adı olmalı (ör. tilburg): {args.city_id}")
    if not COUNTRY_CODE_RE.match(args.country_code or ""):
        raise TemplateError(f"--country-code 2-3 büyük ASCII harf olmalı (ör. NL): {args.country_code}")
    if not re.fullmatch(r"\d{1,3}", (args.city_number or "").strip()) or not 1 <= int(args.city_number) <= 999:
        raise TemplateError(f"--city-number 1-999 arası tam sayı olmalı (katalog sıra no; bilinmiyorsa 900+): {args.city_number}")
    return Options(base=args.base, out=args.out, city_id=args.city_id,
                   city_name=text_arg("--city-name", args.city_name, city.field("name").get("max_length")),
                   country=text_arg("--country", args.country, city.field("country").get("max_length")),
                   country_code=args.country_code, region=choice_arg("--region", args.region, city, "region"),
                   kind=choice_arg("--kind", args.kind, city, "kind"), city_number=int(args.city_number),
                   updated_on=parse_day(args.updated_on), editor=text_arg("--editor", args.editor, 200))


def inside(path: Path, folder: Path) -> bool:
    target, root = os.path.normcase(path.resolve()), os.path.normcase(folder.resolve())
    return target == root or target.startswith(root.rstrip("\\/") + os.sep)


def check_target(flag: str, path: Path, suffix: str, inputs: dict[str, Path], read_only: list[Path]) -> None:
    """Yazılacak yol: girdi dosyası değil, salt okunur klasörde değil, doğru uzantılı, klasör değil, üst yolu klasör."""
    for name, other in inputs.items():
        if inside(path, other) or (path.exists() and other.exists() and os.path.samefile(path, other)):
            raise TemplateError(f"{flag}, {name} ile aynı dosya; girdi asla üzerine yazılmaz")
    if any(inside(path, folder) for folder in read_only):
        raise TemplateError(f"{flag} örneğin ya da --tools'un klasöründe/altında olamaz (salt okunur): {path}")
    if path.suffix.lower() != suffix:
        raise TemplateError(f"{flag} {suffix} olmalı: {path}")
    if path.exists() and not path.is_file():
        raise TemplateError(f"{flag} bir dosya değil (klasör?): {path}")
    parent = next((p for p in path.resolve().parents if p.exists()), None)
    if parent is None or not parent.is_dir():
        raise TemplateError(f"{flag} klasöre yazılamaz; üst yol klasör değil: {parent or path}")


def check_paths(args: argparse.Namespace) -> Path:
    """Her şeyden önce: geçmezse rapor da yazılmaz (yanlış yola hiçbir şey yazılmaz). Döner: --tools."""
    if not args.base.is_file():
        raise TemplateError(f"--base bulunamadı: {args.base}")
    tools = resolve_tools(args.tools)
    read_only = [args.base.resolve().parent, tools]
    check_target("--out", args.out, ".xlsx", {"--base": args.base}, read_only)
    if args.out.exists() and not args.overwrite:
        raise TemplateError(f"--out zaten var (üzerine yazmak için --overwrite): {args.out}")
    if args.report is not None:
        check_target("--report", args.report, ".json", {"--base": args.base, "--out": args.out}, read_only)
    return tools


def resolve_tools(value: Path | None) -> Path:
    tools = value or (Path(os.environ["CORTEQS_TOOLS"]) if os.environ.get("CORTEQS_TOOLS") else None)
    if tools is None:
        raise TemplateError("--tools (araclar/arac) veya CORTEQS_TOOLS gerekli (base-schema.json okunur)")
    if not (tools / "base-schema.json").is_file():
        raise TemplateError(f"--tools altında base-schema.json yok: {tools}")
    return tools


# --- türetme -----------------------------------------------------------------------------------------------------------

@dataclass(frozen=True)
class Derived:
    data: bytes
    keep: list[str]
    removed: list[str]
    tokens: list[str]
    expected: dict[tuple[str, str], Any]  # aracın yazdığı hücreler: belirteç taraması dışı, değeri ayrıca denetlenir


def sample_identity(base: Path, by_name: dict[str, Sheet], sst: list[str], contract: Contract) -> tuple[list[str], set[int], str]:
    """Örneğin belirteçleri (sabit + örnekten okunan), katalog numaraları ve şehir adı."""
    meta = by_name[contract.by_key(METADATA_TABLE).sheet].row8(sst)
    city = by_name[contract.by_key(CITY_TABLE).sheet].row8(sst)
    title = next((cell_text(c, sst) for c in by_name[SUMMARY_SHEET].doc.root.iter(m("c")) if c.get("r") == TITLE_CELL), "") or ""
    numbers = {int(n) for text in (meta.get("delivery_id", ""), title, base.stem) for n in CITY_NO_RE.findall(text)}
    tokens = [*FIXED_TOKENS, *[f"city-{n:03d}" for n in numbers], city.get("city_id", ""), city.get("name", ""),
              title.split(" · ")[0], *base.stem.split("_")[1:2]]
    return sorted({fold(t.strip()) for t in tokens if len(t.strip()) >= 3}), numbers, city.get("name") or "Berlin"


def identity_values(opts: Options, contract: Contract, base_name: str) -> dict[str, dict[str, Any]]:
    meta, city = contract.by_key(METADATA_TABLE), contract.by_key(CITY_TABLE)
    choice_arg("Teslimat.scope", SCOPE, meta, "scope")
    choice_arg("Teslimat.editorial_status", STATUS, meta, "editorial_status")
    note = NOTE.format(base=base_name) + (LOCAL_TRIAL_NOTE if opts.city_number >= LOCAL_TRIAL_FROM else "")
    return {meta.sheet: {"schema_version": contract.schema_version or "1.0", "delivery_id": f"city-{opts.city_number:03d}-v001",
                         "city_id": opts.city_id, "editor": opts.editor, "scope": SCOPE, "updated_on": opts.updated_on,
                         "editorial_status": STATUS, "notes": note},
            city.sheet: {"city_id": opts.city_id, "name": opts.city_name, "country": opts.country, "country_code": opts.country_code,
                         "region": opts.region, "kind": opts.kind, "source_position": opts.city_number}}


def transform(keep: list[Sheet], book: Book, contract: Contract, opts: Options, base_name: str, stats: Stats) -> None:
    by_name = {s.name: s for s in keep}
    plain_str = any(c.get("t") == "str" and c.find(m("f")) is None for s in keep for c in s.doc.root.iter(m("c")))
    text_type = "str" if plain_str else "inlineStr"  # örneğin metin hücresi biçimi (Berlin/Atınç: t="str")
    data_sheets = [s for s in keep if s.name != SUMMARY_SHEET]
    heights = sum((s.blank_heights() for s in data_sheets), Counter())
    for sheet in data_sheets:
        sheet.prepare(contract.tables.get(sheet.name), heights.most_common(1)[0][0] if heights else None, stats)
    summary = by_name[SUMMARY_SHEET]
    strip_notes(summary, book.sst, stats.tokens, stats)
    title_row = row_map(summary.doc.root).get(split_ref(TITLE_CELL)[1])
    if title_row is None:
        raise TemplateError(f"{SUMMARY_SHEET}: {TITLE_CELL} satırı yok")
    title = f"{opts.city_name} · {opts.country} · CITY-{opts.city_number:03d}"
    set_value(get_cell(title_row, TITLE_CELL), title, text_type)
    stats.written.append({"sheet": SUMMARY_SHEET, "key": "title", "cell": TITLE_CELL, "value": title})
    for sheet_name, values in identity_values(opts, contract, base_name).items():
        by_name[sheet_name].write_row8(contract.tables[sheet_name], values, text_type, stats)
    refit_formulas(keep, book, stats)
    write_cached(keep, stats)


def derive(opts: Options, contract: Contract, stats: Stats) -> Derived:
    pkg = Package.load(opts.base)
    book = Book.load(pkg)
    names = [s.name for s in book.sheets]
    if missing := sorted(contract.sheets - set(names)):
        raise TemplateError(f"örnekte sözleşme sekmesi yok: {', '.join(missing)}")
    keep = [s for s in book.sheets if s.name in contract.sheets]
    for sheet in keep:
        sheet.load(pkg, book.sst)
    stats.tokens, numbers, base_name = sample_identity(opts.base, {s.name: s for s in keep}, book.sst, contract)
    stats.tokens = sorted(set(stats.tokens) | {fold(n) for n in names if n not in contract.sheets})  # silinen sekmeye başvuru kalmaz
    if opts.city_number in numbers:
        raise TemplateError(f"--city-number {opts.city_number:03d} örnek şehrin katalog numarası; yeni şehrin numarasını verin")
    transform(keep, book, contract, opts, base_name, stats)
    book.compact_strings(keep, pkg)
    for sheet in keep:
        pkg.put(sheet.part, sheet.doc)
        if sheet.table_part:
            pkg.put(sheet.table_part, sheet.table_doc)
    book.drop_sheets(keep, stats)
    pkg.put(book.part, book.doc)
    pkg.put(rels_name(book.part), book.rels)
    app = pkg.xml("docProps/app.xml") if "docProps/app.xml" in pkg.data else None
    if app is not None and (titles := [c for c in app.root if local(c.tag) in ("HeadingPairs", "TitlesOfParts")]):
        for child in titles:  # Excel'in kaydettiği sekme adı listesi silinen sekmeleri de sayar; isteğe bağlı, Excel yeniden yazar
            app.root.remove(child)
        pkg.put("docProps/app.xml", app)
    stats.removed_parts = pkg.unreachable()
    pkg.drop(stats.removed_parts)
    stats.renamed = renumber(pkg, keep)
    return Derived(pkg.to_bytes(), [s.name for s in keep], [n for n in names if n not in contract.sheets], stats.tokens,
                   {(w["sheet"], w["cell"]): w["value"] for w in stats.written})


# --- öz denetim --------------------------------------------------------------------------------------------------------

def check_sheet(sheet: Sheet, contract: Contract, derived: Derived, sst: list[str]) -> list[str]:
    findings, stray = [], []
    spec = contract.tables.get(sheet.name)
    expected_keys = spec.keys if spec else (contract.provenance_keys if sheet.name == PROVENANCE_SHEET else None)
    if expected_keys is not None and tuple(sheet.keys) != tuple(expected_keys):
        findings.append(f"{sheet.name}: satır 7 anahtarları sözleşmeyle aynı değil")
    ref = sheet.table_doc.root.get("ref") if sheet.table_doc is not None else None
    if spec is not None and spec.capacity and (ref is None or range_boundaries(ref)[1:4:2] != (common.KEY_ROW, common.KEY_ROW + spec.capacity)):
        findings.append(f"{sheet.name}: tablo aralığı {ref} kapasiteyle ({spec.capacity}) uyumsuz")
    for cell in sheet.doc.root.iter(m("c")):
        ref, text = cell.get("r"), cell_text(cell, sst)
        if (sheet.name, ref) in derived.expected:
            expected = derived.expected[(sheet.name, ref)]
            if text not in (str(expected), f"{expected}T00:00:00.000Z"):  # tarih hücresi ISO 8601 (t="d")
                findings.append(f"{sheet.name}!{ref}: yazılan değer beklenenden farklı")
            continue
        if sheet.name != SUMMARY_SHEET and has_value(cell) and split_ref(ref)[1] >= common.FIRST_DATA_ROW:
            stray.append(ref)
        findings += [f"{sheet.name}!{ref}: örnek belirteci '{t}'" for t in find_tokens(text, derived.tokens)]
    if stray:
        findings.append(f"{sheet.name}: veri satırında {len(stray)} değer kaldı ({', '.join(stray[:5])})")
    outside = re.sub(r"<(\w+:)?sheetData\b.*</(\w+:)?sheetData>", "", ET.tostring(sheet.doc.root, encoding="unicode"), flags=re.S)
    return findings + [f"{sheet.name} (sayfa XML): örnek belirteci '{t}'" for t in find_tokens(outside, derived.tokens)]


def self_check(derived: Derived, contract: Contract) -> list[str]:
    pkg = Package.load(derived.data)
    book = Book.load(pkg)
    findings = package_findings(pkg)
    if media := [n for n in pkg.infos if n.startswith(("xl/media/", "xl/drawings/"))]:
        findings.append(f"medya/çizim parçası kaldı: {media[:5]}")
    if orphans := pkg.unreachable():
        findings.append(f"hiçbir ilişkiyle bağlanmayan parça: {orphans[:5]}")
    for name, data in [(n, d) for n, d in pkg.data.items() if n.endswith((".xml", ".rels")) and "/worksheets/sheet" not in n]:
        findings += [f"{name}: örnek belirteci '{t}'" for t in find_tokens(data.decode("utf-8", "ignore"), derived.tokens)]
    if (names := [s.name for s in book.sheets]) != derived.keep or set(names) != contract.sheets:
        findings.append(f"sekme listesi beklenenden farklı: {names}")
    for sheet in book.sheets:
        sheet.load(pkg, book.sst)
        findings += check_sheet(sheet, contract, derived, book.sst)
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")  # openpyxl x14 uzantılarını okumaz; uyarı, hata değil
            if openpyxl.load_workbook(io.BytesIO(derived.data)).sheetnames != derived.keep:
                findings.append("openpyxl sekme listesi farklı")
    except Exception as exc:  # noqa: BLE001 - açılamayan çıktı her nedenle bulgudur
        findings.append(f"openpyxl çıktıyı açamadı: {type(exc).__name__}: {exc}")
    return findings


# --- çalıştırma --------------------------------------------------------------------------------------------------------

def open_items(opts: Options, contract: Contract, stats: Stats) -> list[str]:
    items = []
    if priority := contract.by_key(CITY_TABLE).field("priority"):
        need = "zorunlu " if priority.get("required") else ""
        items.append(f"Şehir.priority: sözleşmede {need}tam sayı (≥{priority.get('minimum', 0)}), anlamı tanımlı değil; boş "
                     "bırakıldı. Değer katalogdan/Atınç'tan alınmalı.")
    if opts.city_number >= LOCAL_TRIAL_FROM:
        items.append(f"Şehir.source_position={opts.city_number}: yerel deneme numarası (katalogda yok); resmî numara gelince "
                     "şablon --city-number ile yeniden türetilmeli. Teslimat.notes bunu yazar.")
    if opts.out.name != (wanted := f"CITY-{opts.city_number:03d}_{opts.city_id}_v001.xlsx"):
        items.append(f"Çıktı adı {wanted} değil; city-pack intake CITY-NNN'i dosya adından okur.")
    if stats.uncached:
        items.append(f"Bilinmeyen formül deseni, önbellek değeri kaldırıldı: {', '.join(stats.uncached[:5])}")
    return items


def report(opts: Options, contract: Contract, derived: Derived, stats: Stats, findings: list[str], written: bool) -> tuple[dict, dict]:
    by_sheet: dict[str, list[str]] = {}
    for entry in [w for w in stats.written if w["sheet"] != SUMMARY_SHEET]:
        by_sheet.setdefault(entry["sheet"], []).append(entry["key"])
    summary = {"tool": "make_template", "out": str(opts.out.resolve()) if written else None, "sheets": len(derived.keep),
               "removed_sheets": derived.removed, "removed_parts": len(stats.removed_parts),
               "removed_media": sum(n.startswith("xl/media/") for n in stats.removed_parts), "tables_resized": stats.resized,
               "formulas_refit": len(stats.refits), "ozet_removed": stats.notes_removed, "written": by_sheet,
               "self_check": findings or "ok", "open_items": open_items(opts, contract, stats)}
    city = contract.by_key(CITY_TABLE)
    details = {"base": str(opts.base), "contract": contract.source, "forbidden_tokens": derived.tokens,
               "cleared_cells": dict(stats.cleared), "rows_dropped": dict(stats.dropped_rows), "heights_reset": dict(stats.heights),
               "formula_refits": stats.refits, "uncached_formulas": stats.uncached, "written_cells": stats.written,
               "city_required_blank": [f["key"] for f in city.fields if f.get("required") and f["key"] not in by_sheet.get(city.sheet, [])],
               "removed_part_names": stats.removed_parts, "renamed_parts": stats.renamed, "defined_names_removed": stats.names_removed}
    return summary, details


def write_output(opts: Options, data: bytes, base_sha: str) -> str:
    """Geçici dosyaya yazıp yerine taşır (yarım şablon kalmaz). Döner: çıktının SHA-256'sı."""
    tmp = opts.out.with_name(f"{opts.out.stem}.tmp-{uuid.uuid4().hex[:8]}.xlsx")
    try:
        opts.out.parent.mkdir(parents=True, exist_ok=True)
        tmp.write_bytes(data)
        os.replace(tmp, opts.out)
        base_changed, out_sha = common.sha256_file(opts.base) != base_sha, common.sha256_file(opts.out)
    except OSError as exc:
        raise TemplateError(f"--out yazılamadı: {type(exc).__name__}: {exc}") from exc
    finally:
        with contextlib.suppress(OSError):
            tmp.unlink(missing_ok=True)
    if base_changed:
        raise TemplateError("--base çalışma sırasında değişti; çıktı güvenilmez")
    return out_sha


def run(args: argparse.Namespace, tools: Path) -> tuple[int, dict, dict]:
    contract = load_contract(tools)
    opts = check_options(args, contract)
    stats = Stats()
    try:
        base_sha = common.sha256_file(opts.base)
        derived = derive(opts, contract, stats)
    except (OSError, *STRUCTURE_ERRORS) as exc:
        raise TemplateError(f"örnek okunamadı ya da beklenen xlsx yapısında değil: {type(exc).__name__}: {exc}") from exc
    try:
        findings = self_check(derived, contract)
    except (TemplateError, *STRUCTURE_ERRORS) as exc:  # kendi çıktımızı okuyamamak bulgudur: dosya yazılmaz
        findings = [f"öz denetim çıktıyı okuyamadı: {type(exc).__name__}: {exc}"]
    summary, details = report(opts, contract, derived, stats, findings, not findings)
    details["base_sha256"] = base_sha
    if findings:
        return common.EXIT_FINDINGS, summary, details
    details["out_sha256"] = write_output(opts, derived.data, base_sha)
    return common.EXIT_OK, summary, details


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Berlin örneğinden (CITY-000_berlin_ornek_v001.xlsx) yeni şehir için boş v7 şablonu türetir. "
                    "Atınç'ın o şehre ait şablonu varsa onu kullanın.",
        epilog="Örnek: make_template.py --base ROOT\\ekler\\CITY-000_berlin_ornek_v001.xlsx --out "
               "ROOT\\uretim\\tilburg\\_is\\sablon_kaynak\\CITY-901_tilburg_v001.xlsx --city-id tilburg --city-name Tilburg "
               "--country Hollanda --country-code NL --region EU --city-number 901 --tools ROOT\\araclar\\arac")
    parser.add_argument("--base", type=Path, required=True, help="Berlin örneği (salt okunur; aynı klasöre yazılmaz)")
    parser.add_argument("--out", type=Path, required=True, help="yeni şablon; adı CITY-NNN_<city_id>_v001.xlsx olmalı")
    for flag, help_text in (("--city-id", "küçük harf ASCII adres adı (ör. tilburg)"), ("--city-name", "şehir adı (ör. Tilburg)"),
                            ("--country", "Türkçe ülke adı (ör. Hollanda)"), ("--country-code", "2-3 büyük harf (ör. NL)"),
                            ("--region", "dünya bölgesi; sözleşmeden: EU, AS, AF, NA, SA, OC"),
                            ("--city-number", "katalog sıra no 1-999 (CITY-NNN); bilinmiyorsa 900+ = yerel deneme")):
        parser.add_argument(flag, required=True, help=help_text)
    parser.add_argument("--updated-on", default=None, help="Teslimat.updated_on YYYY-MM-DD (varsayılan bugün; gelecek olamaz)")
    parser.add_argument("--editor", default=DEFAULT_EDITOR, help=f"Teslimat.editor (varsayılan: {DEFAULT_EDITOR})")
    parser.add_argument("--kind", default="city", help="Şehir.kind (varsayılan city)")
    parser.add_argument("--tools", type=Path, default=None, help="araclar/arac (base-schema.json; yoksa CORTEQS_TOOLS)")
    parser.add_argument("--report", "--out-report", dest="report", type=Path, default=None,
                        help="ayrıntılı rapor (.json; --base/--out olamaz, örnek ya da --tools klasöründe olamaz; "
                             "yol hatasında yazılmaz)")
    parser.add_argument("--overwrite", action="store_true", help="var olan --out üzerine yaz (--base asla)")
    return parser


def main(argv: list[str] | None = None) -> int:
    common.setup_utf8()
    args = build_parser().parse_args(argv)
    started, paths_ok = time.perf_counter(), False
    try:
        tools = check_paths(args)
        paths_ok = True
        code, summary, details = run(args, tools)
    except (TemplateError, OSError) as exc:  # OSError: dosya sistemi (geçersiz ad, erişim); çıkış 1 'bulgu' ile karışmaz
        error = str(exc) if isinstance(exc, TemplateError) else f"dosya sistemi hatası: {type(exc).__name__}: {exc}"
        code, summary, details = common.EXIT_USAGE, {"tool": "make_template", "out": None, "error": error}, {}
    if args.report and paths_ok:  # yol hatasında rapor yazılmaz: yanlış yol örneği ya da şablonu ezebilir
        try:
            common.write_json(args.report, {"summary": summary, **details, "exit_code": code,
                                            "elapsed_s": round(time.perf_counter() - started, 3)})
        except OSError as exc:
            code, summary = common.EXIT_USAGE, {**summary, "error": f"--report yazılamadı: {type(exc).__name__}: {exc}"}
    common.emit(summary)
    return code


if __name__ == "__main__":
    sys.exit(main())
