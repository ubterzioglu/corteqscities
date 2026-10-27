"""Iddia defterini ve bagimsiz denetim sonucunu teslim klasorune yazar (A3/A4 kaniti pakette olsun).

Sevilla dis denetimi (25.10.2026): 2.434 iddialik defter ve denetim sonucu uretim klasorunde kaldi; denetci
iddia-kaynak paydasini kuramadi, A3/A4 (14 puan) sifir aldi. Bu arac son Excel'de bulunan kayitlarin
- yazar defterini (`claims\\*.jsonl`: iddia, deger, donem, kaynak, birebir alinti + kaynagin paketteki kopyasi),
- bagimsiz denetim satirlarini (`denetim\\<parca>[-ek|-dN]_iddia.jsonl`, tur adiyla) ve
- kayit basina son karari (`*_karar.jsonl`, en son tur gecerli)
`Teslimler\\` altina CSV + ozet JSON olarak yazar. Silinen kartin iddiasi deftere girmez (sayisi ozette).
Excel'i yalniz okur; bu bir onay degildir (insan inceleme alanlari Bekliyor kalir).

Cikis: 0 temiz; 1 kopyasi olmayan kaynaga dayanan iddia ya da 'düzelt' kararli ama degisiklik kaydi olmayan kart var;
2 girdi hatasi.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from pathlib import Path
from typing import Any, Iterable

import openpyxl

import common

LEDGER_NAME = "iddia-defteri-v7.csv"
AUDIT_NAME = "bagimsiz-denetim-v7.csv"
SUMMARY_NAME = "iddia-defteri-ozet-v7.json"
RECORD_SHEETS = {"Yanıtlar": "entry_id", "Mekânlar": "entry_id", "Mahalleler": "entry_id", "Rotalar": "route_id",
                 "Sözlük": "term_id", "Adımlar": "step_id", "SayfaAlanları": "field_id", "Şehir": "city_id"}
LEDGER_FIELDS = ("claim_id", "entry_id", "sayfa", "satir", "alan", "iddia", "deger", "birim", "donem", "source_id",
                 "url", "kaynak_kopyasi", "kaynak_kopyasi_sha256", "alinti", "parca")
AUDIT_FIELDS = ("tur", "claim_id", "entry_id", "sayfa", "satir", "iddia", "sonuc", "agirlik", "resmi_url",
                "kontrol_tarihi", "alinti", "dogrusu", "belirsizlik")
# iddia sahibi merge_package.CLAIM_OWNER_KEYS sirasiyla: sozluk satirlari term_id, rota route_id, sayfa alani field_id
# tasir (Sevilla 26.10.2026: yalniz entry_id okununca 60 terimin ve 12 rotanin iddialari defterden dustu, B4 0/60)
CLAIM_OWNER_KEYS = ("entry_id", "term_id", "field_id", "route_id")
ROUND = re.compile(r"^(?P<part>.+?)(?:-(?P<kind>ek|d)(?P<n>\d*))?$")


class DefterError(Exception):
    pass


def round_rank(stem: str) -> tuple[int, int]:
    """Tur sirasi: ilk tur < -ek/-ekN (tamamlama) < -dN (duzeltme turlari)."""
    match = ROUND.match(stem)
    kind, n = match.group("kind"), int(match.group("n") or 0)
    return (0, 0) if not kind else ((1, n) if kind == "ek" else (2, n))


def _part_and_round(path: Path, suffix: str) -> tuple[str, str]:
    stem = path.name[: -len(suffix)]
    return ROUND.match(stem).group("part"), stem


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    rows = []
    for number, line in enumerate(path.read_text(encoding="utf-8-sig").splitlines(), start=1):
        if line.strip():
            try:
                row = json.loads(line)
            except json.JSONDecodeError as exc:
                raise DefterError(f"{path.name}:{number} JSON degil ({exc})") from exc
            if not isinstance(row, dict):
                raise DefterError(f"{path.name}:{number} JSON nesnesi degil")
            rows.append(row)
    return rows


def load_excel(xlsx: Path) -> tuple[dict[str, tuple[str, int]], dict[str, list[str]], dict[str, dict[str, Any]]]:
    """({kayit_id: (sayfa, satir)}, {sayfa: [kayit_id]}, {source_id: kaynak satiri})."""
    try:
        wb = openpyxl.load_workbook(xlsx, read_only=False)
    except Exception as exc:  # openpyxl bozuk dosyada farkli hatalar atar
        raise DefterError(f"xlsx okunamadi: {xlsx} ({type(exc).__name__}: {exc})") from exc
    try:
        where, by_sheet = {}, {}
        for sheet, key in RECORD_SHEETS.items():
            if sheet not in wb.sheetnames:
                continue
            for record in common.iter_records(wb[sheet]):
                record_id = str(record.get(key) or record.get("entry_id") or "")
                if record_id:
                    where.setdefault(record_id, (sheet, record["_row"]))
                    by_sheet.setdefault(sheet, []).append(record_id)
        sources = {str(r["source_id"]): r for r in common.iter_records(wb["Kaynaklar"]) if r.get("source_id")} \
            if "Kaynaklar" in wb.sheetnames else {}
    finally:
        wb.close()
    return where, by_sheet, sources


def ledger_rows(claims_dir: Path, where: dict[str, tuple[str, int]], sources: dict[str, dict[str, Any]]
                ) -> tuple[list[dict[str, Any]], int]:
    rows, dropped = [], 0
    for path in sorted(claims_dir.glob("*.jsonl")):
        for claim in read_jsonl(path):
            entry = str(next((claim[k] for k in CLAIM_OWNER_KEYS if claim.get(k)), ""))
            if entry not in where:
                dropped += 1
                continue
            source = sources.get(str(claim.get("source_id") or ""), {})
            rows.append({"claim_id": claim.get("claim_id", ""), "entry_id": entry, "sayfa": where[entry][0],
                         "satir": where[entry][1], "alan": claim.get("field", ""), "iddia": claim.get("claim", ""),
                         "deger": claim.get("value", ""), "birim": claim.get("unit", ""), "donem": claim.get("period", ""),
                         "source_id": claim.get("source_id", ""), "url": source.get("url") or claim.get("url", ""),
                         "kaynak_kopyasi": source.get("snapshot_file", ""),
                         "kaynak_kopyasi_sha256": source.get("snapshot_sha256", ""), "alinti": claim.get("quote", ""),
                         "parca": path.stem})
    return rows, dropped


def _ordered(dirs: list[Path], suffix: str) -> list[tuple[str, Path]]:
    """Klasor sirasi (once ulke/bolge paketi, en son sehir), sonra tur, sonra parca: sonraki karar oncekini ezer.

    Tur parcadan once gelir: ayni kart iki parcada olabilir (Sevilla: 'mekan1' duzelt, 'mekan-d2' temiz).
    """
    named = [((index, *_part_and_round(path, suffix)), path)
             for index, folder in enumerate(dirs) if folder.is_dir() for path in folder.glob(f"*{suffix}")]
    return [(stem, path) for (_i, _part, stem), path in sorted(named, key=lambda x: (x[0][0], round_rank(x[0][2]), x[0][1]))]


class Resolver:
    """Paket kimligi ('es:slug') sehir kimligine ('sevilla:slug') slug uzerinden baglanir."""

    def __init__(self, where: dict[str, tuple[str, int]]) -> None:
        self.where = where
        self.by_slug = {key.split(":", 1)[-1]: key for key in where}

    def __call__(self, entry: Any) -> str | None:
        entry = str(entry or "")
        return entry if entry in self.where else self.by_slug.get(entry.split(":", 1)[-1]) if ":" in entry else None


def _entries(paths: Iterable[Path], resolve: Resolver, keep: Any = lambda item: True) -> set[str]:
    return {resolve(item.get("entry_id")) for path in paths for item in read_jsonl(path) if keep(item)} - {None}


def _rank(dirs: list[Path], path: Path, suffix: str) -> tuple[int, tuple[int, int]]:
    """Dosyanin sirasi: (klasor sirasi, tur sirasi); _ordered ile ayni sira."""
    index = next((i for i, folder in enumerate(dirs) if path.parent == folder), 0)
    return index, round_rank(path.name[: -len(suffix)])


def card_status(dirs: list[Path], verdict: dict[str, str], resolve: Resolver,
                verdict_rank: dict[str, tuple] | None = None) -> tuple[dict[str, str], list[str]]:
    """Son durum: temiz | düzeltildi (düzelt + editör degisiklik kaydi) | düzelt_kayitsiz; ve dogrulayicinin sorunlu buldugu kayitlar.

    Dogrulayici isareti yalniz kaydin son kararindan eski olmayan turdaysa sayilir: sonraki bagimsiz denetim (yeniden
    denetim, son tam denetim) eski isareti gecersiz kilar (Sevilla 26.10.2026: ilk turun 8 isareti d5-d10 denetimlerinden
    sonra da A1-A5'i dusuruyordu)."""
    verdict_rank = verdict_rank or {}
    changed = _entries((p for _s, p in _ordered(dirs, "_degisiklik.jsonl")), resolve)
    bad = {"hatali_duzeltme", "eksik_kalan"}
    flagged: set[str] = set()
    for _stem, path in _ordered(dirs, "_dogrulama.jsonl"):
        rank = _rank(dirs, path, "_dogrulama.jsonl")
        for item in read_jsonl(path):
            entry = resolve(item.get("entry_id"))
            if entry and item.get("sonuc") in bad and rank >= verdict_rank.get(entry, rank):
                flagged.add(entry)
    status = {entry: ("temiz" if karar == "temiz" else ("düzeltildi" if entry in changed else "düzelt_kayitsiz"))
              for entry, karar in verdict.items()}
    return status, sorted(entry for entry in flagged if entry in verdict)


def audit_rows(dirs: list[Path], resolve: Resolver) -> tuple[list[dict[str, Any]], dict[str, str]]:
    rows, verdict, _rank_of = audit_rows_ranked(dirs, resolve)
    return rows, verdict


def audit_rows_ranked(dirs: list[Path], resolve: Resolver) -> tuple[list[dict[str, Any]], dict[str, str], dict[str, tuple]]:
    """audit_rows + kayit basina son kararin sirasi (card_status dogrulama isaretini buna gore eler)."""
    rows, verdict, ranks, where = [], {}, {}, resolve.where
    for stem, path in _ordered(dirs, "_iddia.jsonl"):
        for item in read_jsonl(path):
            entry = resolve(item.get("entry_id"))
            if entry:
                rows.append({"tur": stem, "claim_id": item.get("claim_id", ""), "entry_id": entry,
                             "sayfa": where[entry][0], "satir": where[entry][1], "iddia": item.get("iddia", ""),
                             "sonuc": item.get("sonuç", ""), "agirlik": item.get("ağırlık", ""),
                             "resmi_url": item.get("resmî_URL", ""), "kontrol_tarihi": item.get("kontrol_tarihi", ""),
                             "alinti": item.get("alıntı", ""), "dogrusu": item.get("doğrusu", ""),
                             "belirsizlik": item.get("belirsizlik", "")})
    for _stem, path in _ordered(dirs, "_karar.jsonl"):
        rank = _rank(dirs, path, "_karar.jsonl")
        for item in read_jsonl(path):
            entry = resolve(item.get("entry_id"))
            if entry:
                verdict[entry] = str(item.get("karar", ""))  # sirali okunur: son tur gecerli
                ranks[entry] = rank
    return rows, verdict, ranks


def _missing_by_sheet(by_sheet: dict[str, list[str]], present: set[str]) -> dict[str, list[str]]:
    missing = {sheet: [i for i in ids if i not in present] for sheet, ids in by_sheet.items()}
    return {sheet: ids for sheet, ids in missing.items() if ids}


def write_csv(path: Path, fields: tuple[str, ...], rows: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as stream:  # BOM: Excel Turkce karakteri dogru acar
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def build(xlsx: Path, claims_dir: Path, audit_dirs: list[Path], out_dir: Path) -> dict[str, Any]:
    where, by_sheet, sources = load_excel(xlsx)
    resolve = Resolver(where)
    ledger, dropped = ledger_rows(claims_dir, where, sources)
    audit, verdict, verdict_rank = audit_rows_ranked(audit_dirs, resolve)
    status, check_flags = card_status(audit_dirs, verdict, resolve, verdict_rank)
    with_claims = {row["entry_id"] for row in ledger}
    audited = {row["entry_id"] for row in audit} | set(verdict)
    no_copy = [row["claim_id"] for row in ledger if not (row["kaynak_kopyasi"] and row["kaynak_kopyasi_sha256"])]
    out_dir.mkdir(parents=True, exist_ok=True)
    write_csv(out_dir / LEDGER_NAME, LEDGER_FIELDS, ledger)
    write_csv(out_dir / AUDIT_NAME, AUDIT_FIELDS, audit)
    summary = {"tool": "iddia_defteri", "workbook": xlsx.name, "kayit": len(where), "iddia": len(ledger),
               "kayit_iddiali": len(with_claims), "silinen_kayit_iddiasi": dropped,
               "denetim_satiri": len(audit), "kayit_denetlenen": len(audited),
               "son_durum_sayim": {k: list(status.values()).count(k) for k in sorted(set(status.values()))},
               "son_durum": status, "denetim_karari": verdict, "dogrulama_sorunlu": check_flags,
               "duzelt_kayitsiz": sorted(e for e, s in status.items() if s == "düzelt_kayitsiz"),
               "iddiasiz_kayit": _missing_by_sheet(by_sheet, with_claims),
               "denetlenmemis_kayit": _missing_by_sheet(by_sheet, audited), "kopyasiz_iddia": no_copy,
               "not": "Yazar defteri ve bagimsiz AI denetimi; insan onayi degildir. denetim_karari editörden "
                      "öncedir; son_durum editörün degisiklik kaydini da sayar. Ülke/bölge paketinden gelen kartların "
                      "denetimi paketin klasörü --audit-dir ile verilmediyse burada denetlenmemis görünür.",
               "denetim_klasorleri": [str(d) for d in audit_dirs]}
    common.write_json(out_dir / SUMMARY_NAME, summary)
    return summary


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Iddia defteri + bagimsiz denetim sonucu -> Teslimler")
    parser.add_argument("--xlsx", type=Path, required=True, help="teslim Excel'i (kayitlar buradan)")
    parser.add_argument("--claims-dir", type=Path, required=True, help="WORK\\_is\\claims")
    parser.add_argument("--audit-dir", type=Path, required=True, action="append",
                        help="denetim klasoru; tekrarlanabilir: once ulke/bolge paketi (ulke\\ES\\_is\\denetim), en son "
                             "WORK\\_is\\denetim (sonraki karar oncekini ezer)")
    parser.add_argument("--out-dir", type=Path, required=True, help="PKG\\Teslimler")
    return parser


def main(argv: list[str] | None = None) -> int:
    common.setup_utf8()
    args = build_parser().parse_args(argv)
    try:
        if not args.xlsx.is_file():
            raise DefterError(f"xlsx yok: {args.xlsx}")
        if not args.claims_dir.is_dir():
            raise DefterError(f"iddia klasoru yok: {args.claims_dir}")
        summary = build(args.xlsx, args.claims_dir, args.audit_dir, args.out_dir)
    except DefterError as exc:
        common.emit({"tool": "iddia_defteri", "error": str(exc)})
        return common.EXIT_USAGE
    common.emit({key: summary[key] for key in ("tool", "kayit", "iddia", "kayit_iddiali", "silinen_kayit_iddiasi",
                                               "denetim_satiri", "kayit_denetlenen", "son_durum_sayim")}
                | {"kopyasiz_iddia": len(summary["kopyasiz_iddia"]), "duzelt_kayitsiz": len(summary["duzelt_kayitsiz"]),
                   "dogrulama_sorunlu": len(summary["dogrulama_sorunlu"]), "out": str(args.out_dir)})
    return common.EXIT_FINDINGS if summary["kopyasiz_iddia"] or summary["duzelt_kayitsiz"] else common.EXIT_OK


if __name__ == "__main__":
    sys.exit(main())
