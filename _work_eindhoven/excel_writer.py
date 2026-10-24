# Excel yazıcı: şablonu bozmadan veriyi 8. satırdan itibaren yazar
import json, os, datetime, hashlib
import openpyxl

ROOT = r"C:\temp_private\corteqs\0000_newera\eindhoven_qwen38_max"
XLSX = os.path.join(ROOT, "CITY-039_eindhoven_v001.xlsx")
TODAY = datetime.date(2026, 10, 20)

data = json.load(open('excel_data.json', encoding='utf-8'))
photos = json.load(open('photos_rows.json', encoding='utf-8'))

wb = openpyxl.load_workbook(XLSX)

def to_date(v):
    if v is None:
        return None
    if isinstance(v, datetime.date):
        return v
    try:
        return datetime.date.fromisoformat(str(v)[:10])
    except Exception:
        return None

def write_rows(sheet_name, rows, keys, date_cols=(), num_cols=()):
    ws = wb[sheet_name]
    r = 8
    for row in rows:
        for c, k in enumerate(keys, start=1):
            v = row.get(k)
            if k in date_cols:
                v = to_date(v)
            if v is not None and k in num_cols:
                try:
                    v = float(v)
                except (TypeError, ValueError):
                    v = None
            ws.cell(row=r, column=c, value=v)
        r += 1
    return r - 8

# Konular
n = write_rows('Konular', data['topics'],
               ['topic_id', 'city_id', 'slug', 'title', 'description', 'icon', 'order'])
print('Konular', n)

# Kaynaklar
n = write_rows('Kaynaklar', data['sources'],
               ['source_id', 'city_id', 'title', 'publisher', 'url', 'checked_on', 'topic', 'summary',
                'verified', 'snapshot_file', 'snapshot_sha256'], date_cols={'checked_on'})
print('Kaynaklar', n)

# Yanıtlar
n = write_rows('Yanıtlar', data['answers'],
               ['entry_id', 'city_id', 'slug', 'title', 'subtitle', 'summary', 'keywords', 'stages',
                'source_ids', 'reviewed_on', 'order', 'featured', 'editorial_status', 'topic_id',
                'steps', 'documents', 'timing', 'cost', 'note', 'visual_json'], date_cols={'reviewed_on'})
print('Yanıtlar', n)

# Mekânlar
n = write_rows('Mekânlar', data['places'],
               ['entry_id', 'city_id', 'slug', 'title', 'subtitle', 'summary', 'keywords', 'stages',
                'source_ids', 'reviewed_on', 'order', 'featured', 'editorial_status', 'group', 'district',
                'setting', 'location', 'latitude', 'longitude', 'visit_status', 'note', 'collection',
                'media_key'], date_cols={'reviewed_on'}, num_cols={'latitude', 'longitude'})
print('Mekânlar', n)

# Mahalleler
n = write_rows('Mahalleler', data['districts'],
               ['entry_id', 'city_id', 'slug', 'title', 'subtitle', 'summary', 'keywords', 'stages',
                'source_ids', 'reviewed_on', 'order', 'featured', 'editorial_status', 'location',
                'tags', 'compare_json', 'note', 'media_key'], date_cols={'reviewed_on'})
print('Mahalleler', n)

# Bağlantılar
n = write_rows('Bağlantılar', data['actions'],
               ['action_id', 'city_id', 'entry_id', 'label', 'url', 'checked_on', 'order'],
               date_cols={'checked_on'})
print('Bağlantılar', n)

# Rotalar
n = write_rows('Rotalar', data['routes'],
               ['route_id', 'city_id', 'title', 'tag', 'stops', 'stop_entry_ids', 'entry_id', 'note',
                'reviewed_on', 'source_ids', 'order', 'editorial_status'], date_cols={'reviewed_on'})
print('Rotalar', n)

# Sözlük
n = write_rows('Sözlük', data['glossary'],
               ['term_id', 'city_id', 'term', 'definition', 'source_ids', 'reviewed_on', 'order',
                'editorial_status'], date_cols={'reviewed_on'})
print('Sözlük', n)

# Adımlar
n = write_rows('Adımlar', data['steps'],
               ['step_id', 'city_id', 'entry_id', 'title', 'when', 'description', 'order'])
print('Adımlar', n)

# SayfaAlanları
n = write_rows('SayfaAlanları', data['fields'],
               ['field_id', 'city_id', 'key', 'label', 'value_type', 'value', 'reviewed_on'],
               date_cols={'reviewed_on'})
print('SayfaAlanları', n)

# Fotoğraflar (+ order galeri sırası)
order_per_entry = {}
photo_rows = []
kapak_order = 0
for p in photos:
    eid = p['entry_id'] or '__kapak__'
    if eid == '__kapak__':
        o = kapak_order; kapak_order += 1
    else:
        o = order_per_entry.get(eid, 0)
        order_per_entry[eid] = o + 1
    photo_rows.append({
        'photo_id': p['photo_id'], 'city_id': 'eindhoven', 'role': p['role'],
        'entry_id': p['entry_id'], 'file': p['file'], 'sha256': p['sha256'],
        'caption': p['caption'], 'author': p['author'], 'license': p['license'],
        'license_url': p['license_url'], 'source_url': p['source_url'],
        'capture_date': to_date(p['capture_date']), 'source_checked_on': to_date(p['source_checked_on']),
        'modifications': p['modifications'], 'evidence_id': p['evidence_id'],
        'completion_credit': 'Evet', 'order': o, 'editorial_status': 'İnceleme',
    })
ws = wb['Fotoğraflar']
keys = ['photo_id', 'city_id', 'role', 'entry_id', 'file', 'sha256', 'caption', 'author', 'license',
        'license_url', 'source_url', 'capture_date', 'source_checked_on', 'modifications',
        'evidence_id', 'completion_credit', 'order', 'editorial_status']
r = 8
for row in photo_rows:
    for c, k in enumerate(keys, start=1):
        ws.cell(row=r, column=c, value=row[k])
    r += 1
print('Fotoğraflar', len(photo_rows))

# Kanıtlar
ev_rows = []
for p in photos:
    src_ev_path = os.path.join(ROOT, p['src_ev_rel'].replace('/', os.sep))
    api_json = None
    try:
        ev = json.load(open(src_ev_path, encoding='utf-8'))
        api_json = json.dumps(ev.get('api_imageinfo', {}), ensure_ascii=False)[:30000]
    except Exception:
        api_json = None
    ev_rows.append({
        'evidence_id': p['evidence_id'], 'city_id': 'eindhoven', 'photo_id': p['photo_id'],
        'source_url': p['source_url'], 'source_title': p['commons_title'], 'author': p['author'],
        'license': p['license'], 'license_url': p['license_url'], 'usage_terms': p['usage_terms'],
        'attribution': p['attribution'], 'restrictions': p['restrictions'],
        'checked_on': TODAY, 'metadata_review_status': 'Tamam',
        'pixel_review_status': 'Bekliyor', 'pixel_review_file': p['file'],
        'pixel_review_note': p['pixel_note'], 'rights_status': 'Bekliyor',
        'rights_basis': f"{p['license']} — ücretsiz ticari kullanıma uygun (API lisans kaydı); nihai hak onayı insan incelemesine bırakıldı.",
        'source_snapshot_file': p['src_ev_rel'], 'source_snapshot_sha256': p['src_ev_sha'],
        'license_snapshot_file': p['lic_ev_rel'], 'license_snapshot_sha256': p['lic_ev_sha'],
        'source_evidence_json': api_json,
    })
n = write_rows('Kanıtlar', ev_rows,
               ['evidence_id', 'city_id', 'photo_id', 'source_url', 'source_title', 'author', 'license',
                'license_url', 'usage_terms', 'attribution', 'restrictions', 'checked_on',
                'metadata_review_status', 'pixel_review_status', 'pixel_review_file', 'pixel_review_note',
                'rights_status', 'rights_basis', 'source_snapshot_file', 'source_snapshot_sha256',
                'license_snapshot_file', 'license_snapshot_sha256', 'source_evidence_json'],
               date_cols={'checked_on'})
print('Kanıtlar', n)

# FotoğrafKaynakları
fk_rows = []
for p in photos:
    fk_rows.append({
        'photo_id': p['photo_id'], 'author_url': None, 'attribution_url': p['source_url'],
        'original_file_url': p['orig_url'], 'downloaded_file_url': p['dl_url'],
        'downloaded_on': TODAY, 'original_file': p['orig_rel'], 'original_sha256': p['orig_sha'],
        'web_file': p['file'], 'web_sha256': p['sha256'], 'web_width': p['web_w'], 'web_height': p['web_h'],
        'derivation': 'Commons API thumburl (1600px türev) indirildi; Pillow ile WebP (kalite 82, ≤1600px) üretildi. Kırpma yok.',
        'provenance_status': 'Doğrulandı',
        'provenance_note': 'Özgün tam çözünürlük dosyası ayrıca indirilmedi; saklanan indirme kopyası Commons 1600px türevidir. API imageinfo kaydı, lisans kaydı ve SHA-256 zinciri Lisans kanıtları klasöründedir.',
    })
n = write_rows('FotoğrafKaynakları', fk_rows,
               ['photo_id', 'author_url', 'attribution_url', 'original_file_url', 'downloaded_file_url',
                'downloaded_on', 'original_file', 'original_sha256', 'web_file', 'web_sha256',
                'web_width', 'web_height', 'derivation', 'provenance_status', 'provenance_note'],
               date_cols={'downloaded_on'})
print('FotoğrafKaynakları', n)

# Şehir satırı (kimlik alanları korunur, rehber alanları doldurulur)
ws = wb['Şehir']
c = data['city']
colmap = {'title': 7, 'intro': 8, 'currency': 9, 'currency_symbol': 10, 'currency_label': 11,
          'language': 12, 'timezone': 13, 'reviewed_on': 22, 'tagline': 26,
          'forecast_url': 28, 'provider_url': 29, 'forecast_label': 30, 'hero_photo_id': 31,
          'editorial_status': 33}
for k, col in colmap.items():
    v = c.get(k)
    if k == 'reviewed_on':
        v = to_date(v)
    ws.cell(row=8, column=col, value=v)
print('Şehir güncellendi')

# Teslimat satırı
ws = wb['Teslimat']
t = data['teslimat']
ws.cell(row=8, column=4, value=t['editor'])
ws.cell(row=8, column=5, value=t['reviewer'])
ws.cell(row=8, column=6, value=t['scope'])
ws.cell(row=8, column=7, value=to_date(t['updated_on']))
ws.cell(row=8, column=8, value=t['editorial_status'])
ws.cell(row=8, column=9, value=t['notes'])
print('Teslimat güncellendi')

wb.save(XLSX)
print('SAVED', XLSX)
