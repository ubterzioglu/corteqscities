# QA: yapısal + içerik + görsel denetim (çıktı dosyası üzerinde)
import json, os, hashlib, datetime
import openpyxl
from PIL import Image

ROOT = r"C:\temp_private\corteqs\0000_newera\eindhoven_qwen38_max"
XLSX = os.path.join(ROOT, "CITY-039_eindhoven_v001.xlsx")
ALLOWED_LICENSE = ("CC0", "CC BY", "Public domain", "PD")

wb = openpyxl.load_workbook(XLSX, data_only=False)

def rows_of(sheet, keys):
    ws = wb[sheet]
    out = []
    for row in ws.iter_rows(min_row=8):
        vals = [c.value for c in row][:len(keys)]
        if vals[0] is None or all(v is None for v in vals):
            continue
        out.append(dict(zip(keys, vals)))
    return out

K = {
 'Konular': ['topic_id','city_id','slug','title','description','icon','order','used_count'],
 'Yanıtlar': ['entry_id','city_id','slug','title','subtitle','summary','keywords','stages','source_ids','reviewed_on','order','featured','editorial_status','topic_id','steps','documents','timing','cost','note','visual_json'],
 'Mekânlar': ['entry_id','city_id','slug','title','subtitle','summary','keywords','stages','source_ids','reviewed_on','order','featured','editorial_status','group','district','setting','location','latitude','longitude','visit_status','note','collection','media_key'],
 'Mahalleler': ['entry_id','city_id','slug','title','subtitle','summary','keywords','stages','source_ids','reviewed_on','order','featured','editorial_status','location','tags','compare_json','note','media_key'],
 'Bağlantılar': ['action_id','city_id','entry_id','label','url','checked_on','order'],
 'Rotalar': ['route_id','city_id','title','tag','stops','stop_entry_ids','entry_id','note','reviewed_on','source_ids','order','editorial_status'],
 'Sözlük': ['term_id','city_id','term','definition','source_ids','reviewed_on','order','editorial_status'],
 'Adımlar': ['step_id','city_id','entry_id','title','when','description','order'],
 'Kaynaklar': ['source_id','city_id','title','publisher','url','checked_on','topic','summary','verified','snapshot_file','snapshot_sha256'],
 'Fotoğraflar': ['photo_id','city_id','role','entry_id','file','sha256','caption','author','license','license_url','source_url','capture_date','source_checked_on','modifications','evidence_id','completion_credit','order','editorial_status'],
 'Kanıtlar': ['evidence_id','city_id','photo_id','source_url','source_title','author','license','license_url','usage_terms','attribution','restrictions','checked_on','metadata_review_status','pixel_review_status','pixel_review_file','pixel_review_note','rights_status','rights_basis','source_snapshot_file','source_snapshot_sha256','license_snapshot_file','license_snapshot_sha256','source_evidence_json'],
 'SayfaAlanları': ['field_id','city_id','key','label','value_type','value','reviewed_on'],
 'FotoğrafKaynakları': ['photo_id','author_url','attribution_url','original_file_url','downloaded_file_url','downloaded_on','original_file','original_sha256','web_file','web_sha256','web_width','web_height','derivation','provenance_status','provenance_note'],
}
T = {k: rows_of(k, v) for k, v in K.items()}

errors, warns = [], []
def err(m): errors.append(m)
def warn(m): warns.append(m)

# ---- kimlik setleri
source_ids = {r['source_id'] for r in T['Kaynaklar']}
topic_ids = {r['topic_id'] for r in T['Konular']}
answer_ids = {r['entry_id'] for r in T['Yanıtlar']}
place_ids = {r['entry_id'] for r in T['Mekânlar']}
district_ids = {r['entry_id'] for r in T['Mahalleler']}
photo_ids = {r['photo_id'] for r in T['Fotoğraflar']}

# ---- tekrarlanan kimlikler
for sheet, key in [('Yanıtlar','entry_id'),('Mekânlar','entry_id'),('Mahalleler','entry_id'),
                   ('Fotoğraflar','photo_id'),('Kanıtlar','evidence_id'),('Kaynaklar','source_id'),
                   ('Konular','topic_id'),('Rotalar','route_id'),('Sözlük','term_id'),
                   ('Adımlar','step_id'),('Bağlantılar','action_id'),('SayfaAlanları','field_id')]:
    seen = set()
    for r in T[sheet]:
        v = r[key]
        if v in seen:
            err(f"{sheet}: tekrarlanan {key} {v}")
        seen.add(v)

# ---- sayılar
counts = {
 'answers': len(T['Yanıtlar']), 'places': len(T['Mekânlar']), 'districts': len(T['Mahalleler']),
 'topics': len(T['Konular']), 'routes': len(T['Rotalar']), 'glossary': len(T['Sözlük']),
 'sources': len(T['Kaynaklar']), 'photos': len(T['Fotoğraflar']), 'evidence': len(T['Kanıtlar']),
 'actions': len(T['Bağlantılar']), 'steps': len(T['Adımlar']), 'fields': len(T['SayfaAlanları']),
}
MIN = {'answers': 212, 'places': 92, 'districts': 12, 'topics': 12, 'routes': 9, 'glossary': 48}
for k, m in MIN.items():
    if counts[k] < m:
        err(f"SAYI: {k} {counts[k]} < minimum {m}")

# ---- referans bütünlüğü
def check_srcs(val, where):
    if not val:
        err(f"{where}: source_ids boş")
        return
    for s in str(val).split(';'):
        s = s.strip()
        if s and s not in source_ids:
            err(f"{where}: bilinmeyen kaynak {s}")

for r in T['Yanıtlar']:
    check_srcs(r['source_ids'], f"Yanıt {r['entry_id']}")
    if r['topic_id'] not in topic_ids:
        err(f"Yanıt {r['entry_id']}: geçersiz topic {r['topic_id']}")
    for f in ['title','summary','steps','subtitle','keywords','stages']:
        if not r[f]:
            err(f"Yanıt {r['entry_id']}: {f} boş")
    if r['summary'] and len(str(r['summary'])) < 60:
        warn(f"Yanıt {r['entry_id']}: özet çok kısa")
    if len(str(r['steps']).split('\n')) < 2:
        err(f"Yanıt {r['entry_id']}: adım sayısı < 2")
    st = set(x.strip() for x in str(r['stages']).split(';') if x.strip())
    if not st <= {'before','planning','new','settled'}:
        err(f"Yanıt {r['entry_id']}: geçersiz stages {r['stages']}")
    if r['editorial_status'] != 'İnceleme':
        warn(f"Yanıt {r['entry_id']}: editorial_status {r['editorial_status']}")
for r in T['Mekânlar']:
    check_srcs(r['source_ids'], f"Mekân {r['entry_id']}")
    for f in ['title','summary','group','location','note','media_key']:
        if not r[f]:
            err(f"Mekân {r['entry_id']}: {f} boş")
    if r['group'] not in {'Müzeler','Sanat ve sergiler','Sahne ve müzik','Spor','Doğa ve açık hava','Çocuklarla keşif','Yeme ve içme','Tarih ve mimari'}:
        err(f"Mekân {r['entry_id']}: geçersiz grup {r['group']}")
for r in T['Mahalleler']:
    try:
        cj = json.loads(r['compare_json'])
        assert isinstance(cj, dict) and len(cj) >= 8
    except Exception:
        err(f"Mahalle {r['entry_id']}: compare_json geçersiz")
for r in T['Sözlük']:
    check_srcs(r['source_ids'], f"Terim {r['term_id']}")
for r in T['Rotalar']:
    ids = [x for x in str(r['stop_entry_ids']).split(';') if x]
    if len(ids) < 3:
        err(f"Rota {r['route_id']}: durak < 3")
    for x in ids:
        if x not in place_ids:
            err(f"Rota {r['route_id']}: durak yok {x}")
    if r['entry_id'] not in place_ids:
        err(f"Rota {r['route_id']}: ana kart yok {r['entry_id']}")
    check_srcs(r['source_ids'], f"Rota {r['route_id']}")
for r in T['Adımlar']:
    if r['entry_id'] not in answer_ids:
        err(f"Adım {r['step_id']}: entry yanıt değil {r['entry_id']}")
for r in T['Bağlantılar']:
    eid = r['entry_id']
    if eid not in answer_ids | place_ids | district_ids:
        err(f"Bağlantı {r['action_id']}: entry yok {eid}")
    if not str(r['url']).startswith('http'):
        err(f"Bağlantı {r['action_id']}: url geçersiz")
for r in T['Kaynaklar']:
    if not str(r['url']).startswith('http'):
        err(f"Kaynak {r['source_id']}: url geçersiz")
    if r['verified'] != 'Evet':
        warn(f"Kaynak {r['source_id']}: verified={r['verified']}")

# ---- fotoğraflar
def sha_file(p):
    h = hashlib.sha256()
    with open(p, 'rb') as f:
        for c in iter(lambda: f.read(1 << 20), b''):
            h.update(c)
    return h.hexdigest()

venue_shas, hero_shas = set(), set()
by_photo = {r['photo_id']: r for r in T['Fotoğraflar']}
for r in T['Fotoğraflar']:
    fp = os.path.join(ROOT, str(r['file']).replace('/', os.sep))
    if not os.path.exists(fp):
        err(f"Foto {r['photo_id']}: dosya yok {r['file']}"); continue
    h = sha_file(fp)
    if h != r['sha256']:
        err(f"Foto {r['photo_id']}: sha256 uyuşmuyor")
    try:
        im = Image.open(fp); im.verify()
        im = Image.open(fp)
        if max(im.size) > 1600:
            warn(f"Foto {r['photo_id']}: kenar >1600 {im.size}")
    except Exception as e:
        err(f"Foto {r['photo_id']}: açılamadı {e}")
    lic = str(r['license'] or '')
    if not any(lic.startswith(a) for a in ALLOWED_LICENSE):
        err(f"Foto {r['photo_id']}: lisans uygun değil: {lic}")
    if lic.startswith('CC BY') and not r['author']:
        err(f"Foto {r['photo_id']}: CC BY ama author boş")
    if 'Unknown' in str(r['author'] or ''):
        warn(f"Foto {r['photo_id']}: author belirsiz: {r['author']}")
    if r['role'] == 'Mekân':
        if r['entry_id'] not in place_ids | district_ids:
            err(f"Foto {r['photo_id']}: entry_id mekân/mahalle değil")
        venue_shas.add(r['sha256'])
    elif r['role'] == 'Kapak':
        hero_shas.add(r['sha256'])
    else:
        err(f"Foto {r['photo_id']}: rol geçersiz {r['role']}")
counts['unique_venue_photos'] = len(venue_shas)
counts['unique_hero_photos'] = len(hero_shas)
if len(venue_shas) < 187:
    err(f"SAYI: tekil mekân fotoğrafı {len(venue_shas)} < 187")
if len(hero_shas) < 6:
    err(f"SAYI: tekil kapak {len(hero_shas)} < 6")

# her mekân ve mahallede en az 1 fotoğraf
photo_entries = {r['entry_id'] for r in T['Fotoğraflar'] if r['role'] == 'Mekân'}
for pid in place_ids | district_ids:
    if pid not in photo_entries:
        err(f"FOTO EKSİK: {pid}")

# ---- kanıtlar
ev_photo_ids = {r['photo_id'] for r in T['Kanıtlar']}
for r in T['Fotoğraflar']:
    if r['photo_id'] not in ev_photo_ids:
        err(f"Kanıt yok: {r['photo_id']}")
for r in T['Kanıtlar']:
    if r['photo_id'] not in by_photo:
        err(f"Kanıt {r['evidence_id']}: photo yok")
    for f, sha in [('source_snapshot_file','source_snapshot_sha256'), ('license_snapshot_file','license_snapshot_sha256')]:
        p = os.path.join(ROOT, str(r[f]).replace('/', os.sep))
        if not r[f] or not os.path.exists(p):
            err(f"Kanıt {r['evidence_id']}: {f} yok")
        elif sha_file(p) != r[sha]:
            err(f"Kanıt {r['evidence_id']}: {sha} uyuşmuyor")
    if r['pixel_review_status'] != 'Bekliyor' or r['rights_status'] != 'Bekliyor':
        warn(f"Kanıt {r['evidence_id']}: insan onayı alanları Bekliyor değil")

# ---- FotoğrafKaynakları
fk_ids = {r['photo_id'] for r in T['FotoğrafKaynakları']}
for r in T['Fotoğraflar']:
    if r['photo_id'] not in fk_ids:
        err(f"FotoğrafKaynakları eksik: {r['photo_id']}")
for r in T['FotoğrafKaynakları']:
    op = os.path.join(ROOT, str(r['original_file']).replace('/', os.sep))
    if not os.path.exists(op):
        err(f"FK {r['photo_id']}: original_file yok")
    elif sha_file(op) != r['original_sha256']:
        err(f"FK {r['photo_id']}: original_sha uyuşmuyor")
    if r['web_sha256'] != by_photo.get(r['photo_id'], {}).get('sha256'):
        err(f"FK {r['photo_id']}: web_sha Fotoğraflar ile uyuşmuyor")

# ---- SayfaAlanları
need_keys = {'facts','timeline','emergency','quick_links','budget_fields','place_groups',
 'city_overview','registration_process','residence_rules','finding_home','utilities_setup',
 'transport_system','healthcare_system','work_search','education_system','family_life',
 'safety_support','community_life','discovery_intro','photo_rights','place_disclaimer',
 'glossary_intro','route_intro'}
have = {r['key'] for r in T['SayfaAlanları']}
for k in need_keys - have:
    err(f"SayfaAlanı eksik: {k}")
for r in T['SayfaAlanları']:
    if r['value_type'] == 'JSON':
        try:
            json.loads(r['value'])
        except Exception:
            err(f"SayfaAlanı {r['key']}: JSON geçersiz")

# ---- Şehir satırı
ws = wb['Şehir']
city = [c.value for c in ws[8]]
checks = {'title': 6, 'intro': 7, 'currency': 8, 'language': 11, 'timezone': 12,
          'forecast_url': 27, 'provider_url': 28, 'tagline': 25, 'hero_photo_id': 30}
for name, idx in checks.items():
    if not city[idx]:
        err(f"Şehir.{name} boş")
if city[30] and city[30] not in photo_ids:
    err(f"Şehir.hero_photo_id foto değil: {city[30]}")
if city[27] and 'eindhoven' not in str(city[27]).lower():
    warn(f"Şehir.forecast_url şehir içermiyor: {city[27]}")

report = {'counts': counts, 'errors': errors, 'warnings': warns}
json.dump(report, open('qa_report.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print(json.dumps(counts, indent=1))
print('ERRORS:', len(errors))
for e in errors[:40]:
    print('  E', e)
print('WARNINGS:', len(warns))
for w in warns[:20]:
    print('  W', w)
