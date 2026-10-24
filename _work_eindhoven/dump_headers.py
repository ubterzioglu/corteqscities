import openpyxl, json

wb = openpyxl.load_workbook(r'C:\temp_private\corteqs\0000_newera\ekler_eindhoven\CITY-039_eindhoven_v001.xlsx')
out = open('template_headers.txt', 'w', encoding='utf-8')
for ws in wb.worksheets:
    out.write(f"===== {ws.title} (max_row={ws.max_row}, max_col={ws.max_column})\n")
    r6 = [c.value for c in ws[6]]
    r7 = [c.value for c in ws[7]]
    out.write("row6: " + json.dumps(r6, ensure_ascii=False) + "\n")
    out.write("row7: " + json.dumps(r7, ensure_ascii=False) + "\n")
    r8 = [c.value for c in ws[8]]
    out.write("row8: " + json.dumps(r8, ensure_ascii=False, default=str)[:2000] + "\n\n")
out.close()

wb2 = openpyxl.load_workbook(r'C:\temp_private\corteqs\0000_newera\ekler_eindhoven\CITY-000_berlin_ornek_v001.xlsx', read_only=True)
out2 = open('berlin_model.txt', 'w', encoding='utf-8')
def dumprows(sheet, maxrows, maxcols=30, filt=None):
    ws = wb2[sheet]
    out2.write(f"\n===== {sheet}\n")
    r6 = [c.value for c in next(ws.iter_rows(min_row=6, max_row=6, max_col=maxcols))]
    r7 = [c.value for c in next(ws.iter_rows(min_row=7, max_row=7, max_col=maxcols))]
    out2.write("row6: " + json.dumps(r6, ensure_ascii=False) + "\n")
    out2.write("row7: " + json.dumps(r7, ensure_ascii=False) + "\n")
    n = 0
    for row in ws.iter_rows(min_row=8, max_col=maxcols):
        vals = [c.value for c in row]
        if all(v is None for v in vals):
            continue
        if filt and not filt(vals):
            continue
        out2.write(json.dumps(vals, ensure_ascii=False, default=str)[:3000] + "\n")
        n += 1
        if n >= maxrows:
            break
dumprows('SayfaAlanları', 60, 8)
dumprows('Şehir', 2, 40)
dumprows('Teslimat', 2, 40)
dumprows('Mahalleler', 2, 20)
dumprows('Rotalar', 2, 15)
dumprows('Adımlar', 6, 8)
dumprows('Bağlantılar', 6, 8)
dumprows('Sözlük', 4, 10)
dumprows('Özet', 0, 3)
out2.close()
print('done')
